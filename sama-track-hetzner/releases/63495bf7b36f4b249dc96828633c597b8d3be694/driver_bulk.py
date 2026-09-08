from io import BytesIO, StringIO
import csv
from fastapi import Request, HTTPException
from fastapi.responses import HTMLResponse


def register_driver_bulk(app, db, request_user, audit):
    def require_admin(request: Request):
        user=request_user(request)
        if not user or user.get("role") != "ADMIN":
            raise HTTPException(403,"Sadece ADMIN toplu şoför güncellemesi yapabilir.")
        return user

    def norm(v):
        s=str(v or "").strip().upper()
        for a,b in {"İ":"I","Ş":"S","Ğ":"G","Ü":"U","Ö":"O","Ç":"C"}.items():
            s=s.replace(a,b)
        for ch in " ._-/\\()[]:":
            s=s.replace(ch,"")
        return s

    def active_value(v):
        s=norm(v)
        return 0 if s in ("0","PASIF","INACTIVE","HAYIR","NO") else 1

    def parse_rows(raw: bytes, filename: str):
        ext=(filename.rsplit('.',1)[-1] if '.' in filename else '').lower()
        matrix=[]
        if ext == 'csv':
            txt=raw.decode('utf-8-sig',errors='replace')
            sample=txt[:4096]
            try:
                dialect=csv.Sniffer().sniff(sample,delimiters=';,\t,')
                matrix=list(csv.reader(StringIO(txt),dialect))
            except Exception:
                matrix=list(csv.reader(StringIO(txt),delimiter=';'))
        else:
            try:
                from openpyxl import load_workbook
                wb=load_workbook(BytesIO(raw),data_only=True,read_only=True)
                ws=wb[wb.sheetnames[0]]
                matrix=[list(r) for r in ws.iter_rows(values_only=True)]
            except Exception as e:
                raise HTTPException(400,f"Excel açılamadı: {e}")

        hi=-1; hm={}
        for i,row in enumerate(matrix[:20]):
            m={norm(v):j for j,v in enumerate(row) if str(v or '').strip()}
            if any(k in m for k in ('SOFORADI','DRIVERNAME','SOFOR','DRIVER')):
                hi=i; hm=m; break
        if hi < 0:
            raise HTTPException(400,"ŞOFÖR ADI / DRIVER NAME başlığı bulunamadı.")

        def pick(row,names):
            for name in names:
                k=norm(name)
                if k in hm and hm[k] < len(row):
                    return row[hm[k]]
            return ''

        out=[]
        for row in matrix[hi+1:]:
            if not any(str(v or '').strip() for v in row):
                continue
            name=str(pick(row,['ŞOFÖR ADI','DRIVER NAME','ŞOFÖR','DRIVER']) or '').strip()
            if not name:
                continue
            out.append({
                'name':name,
                'phone':str(pick(row,['TELEFON','PHONE','MOBILE']) or '').strip(),
                'd_no':str(pick(row,['D.NO','DNO','DRIVER NO','DRIVERNO']) or '').strip(),
                'is_active':active_value(pick(row,['DURUM','STATUS','ACTIVE','AKTİF']))
            })
        if not out:
            raise HTTPException(400,"Dosyada şoför kaydı bulunamadı.")
        return out

    def preview(rows,deactivate_missing=False):
        c=db()
        existing=[dict(r) for r in c.execute("SELECT id,name,phone,d_no,is_active FROM fleet_drivers")]
        c.close()
        by_dno={str(x.get('d_no') or '').strip().upper():x for x in existing if str(x.get('d_no') or '').strip()}
        by_name={str(x.get('name') or '').strip().upper():x for x in existing if str(x.get('name') or '').strip()}
        seen=set(); data=[]; counts={'new':0,'update':0,'same':0,'deactivate':0}
        for x in rows:
            old=(by_dno.get(x['d_no'].upper()) if x['d_no'] else None) or by_name.get(x['name'].upper())
            if old:
                seen.add(old['id'])
                changed=(str(old.get('name') or '').strip()!=x['name'] or str(old.get('phone') or '').strip()!=x['phone'] or str(old.get('d_no') or '').strip()!=x['d_no'] or int(old.get('is_active') or 0)!=x['is_active'])
                result='GÜNCELLENECEK' if changed else 'AYNI'
                counts['update' if changed else 'same']+=1
            else:
                result='YENİ'; counts['new']+=1
            data.append({**x,'result':result})
        if deactivate_missing:
            for old in existing:
                if int(old.get('is_active') or 0)==1 and old['id'] not in seen:
                    counts['deactivate']+=1
                    data.append({'name':old.get('name',''),'phone':old.get('phone',''),'d_no':old.get('d_no',''),'is_active':0,'result':'PASİFE ALINACAK'})
        return counts,data

    @app.get('/driver-bulk',response_class=HTMLResponse)
    def driver_bulk_page(request: Request):
        require_admin(request)
        return HTMLResponse('''<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>SAMA - Şoför Toplu Güncelleme</title><style>body{font-family:Arial;background:#f5f7fb;margin:0;padding:24px;color:#172033}.box{max-width:1100px;margin:auto;background:white;border-radius:16px;padding:22px;box-shadow:0 8px 28px #0001}.row{display:flex;gap:12px;align-items:center;flex-wrap:wrap}.btn{border:0;border-radius:10px;padding:12px 18px;font-weight:800;cursor:pointer}.primary{background:#155eef;color:white}.danger{background:#b42318;color:white}.muted{color:#667085;font-size:13px}.cards{display:flex;gap:10px;flex-wrap:wrap;margin:16px 0}.card{padding:12px 16px;border:1px solid #e4e7ec;border-radius:12px;min-width:140px}.card b{font-size:24px;display:block}table{width:100%;border-collapse:collapse;font-size:13px}th,td{padding:9px;border-bottom:1px solid #eee;text-align:left}th{position:sticky;top:0;background:#f9fafb}.table{max-height:500px;overflow:auto;border:1px solid #eee;border-radius:12px;margin-top:14px}</style></head><body><div class="box"><h2>Şoför Toplu Güncelleme</h2><p class="muted">Excel/CSV başlıkları: ŞOFÖR ADI, TELEFON, D.NO, DURUM. Önce önizleme yapılır, sonra onayla uygulanır.</p><div class="row"><input id="f" type="file" accept=".xlsx,.xls,.csv"><label><input id="deact" type="checkbox"> Listede olmayan aktif şoförleri pasife al</label><button class="btn primary" onclick="preview()">ÖNİZLE</button><button id="apply" class="btn danger" onclick="applyNow()" disabled>UYGULA</button></div><div id="summary" class="cards"></div><div id="msg"></div><div id="table"></div></div><script>let last=null;async function send(mode){const file=f.files[0];if(!file)throw new Error('Dosya seçin.');const raw=await file.arrayBuffer();const r=await fetch('/api/driver-bulk/'+mode+'?deactivate_missing='+(deact.checked?'1':'0'),{method:'POST',headers:{'Content-Type':'application/octet-stream','X-Filename':encodeURIComponent(file.name)},body:raw});const d=await r.json();if(!r.ok)throw new Error(d.detail||'Hata');return d}function cards(c){summary.innerHTML=`<div class="card">YENİ<b>${c.new||0}</b></div><div class="card">GÜNCELLENECEK<b>${c.update||0}</b></div><div class="card">AYNI<b>${c.same||0}</b></div><div class="card">PASİF<b>${c.deactivate||0}</b></div>`}async function preview(){try{msg.innerHTML='Okunuyor...';last=await send('preview');cards(last.counts||{});let h='<div class="table"><table><thead><tr><th>SONUÇ</th><th>ŞOFÖR</th><th>TELEFON</th><th>D.NO</th><th>DURUM</th></tr></thead><tbody>';for(const x of (last.rows||[])){h+=`<tr><td><b>${x.result||''}</b></td><td>${x.name||''}</td><td>${x.phone||''}</td><td>${x.d_no||''}</td><td>${x.is_active?'AKTİF':'PASİF'}</td></tr>`}h+='</tbody></table></div>';table.innerHTML=h;apply.disabled=false;msg.innerHTML='';}catch(e){msg.innerHTML='<b style="color:#b42318">'+e.message+'</b>';apply.disabled=true}}async function applyNow(){if(!last)return;if(!confirm('Önizlemedeki değişiklikler uygulanacak. Devam edilsin mi?'))return;try{const d=await send('apply');msg.innerHTML=`<b style="color:#067647">Tamamlandı. Yeni: ${d.added||0} | Güncellenen: ${d.updated||0} | Pasif: ${d.deactivated||0}</b>`;apply.disabled=true;}catch(e){msg.innerHTML='<b style="color:#b42318">'+e.message+'</b>'}}</script></body></html>''')

    @app.post('/api/driver-bulk/preview')
    async def driver_bulk_preview(request: Request, deactivate_missing: int=0):
        require_admin(request)
        raw=await request.body()
        filename=request.headers.get('X-Filename','drivers.xlsx')
        try:
            from urllib.parse import unquote
            filename=unquote(filename)
        except Exception:
            pass
        rows=parse_rows(raw,filename)
        counts,data=preview(rows,bool(deactivate_missing))
        return {'ok':True,'counts':counts,'rows':data[:3000]}

    @app.post('/api/driver-bulk/apply')
    async def driver_bulk_apply(request: Request, deactivate_missing: int=0):
        user=require_admin(request)
        raw=await request.body()
        filename=request.headers.get('X-Filename','drivers.xlsx')
        try:
            from urllib.parse import unquote
            filename=unquote(filename)
        except Exception:
            pass
        rows=parse_rows(raw,filename)
        c=db()
        existing=[dict(r) for r in c.execute("SELECT id,name,phone,d_no,is_active FROM fleet_drivers")]
        by_dno={str(x.get('d_no') or '').strip().upper():x for x in existing if str(x.get('d_no') or '').strip()}
        by_name={str(x.get('name') or '').strip().upper():x for x in existing if str(x.get('name') or '').strip()}
        seen=set(); added=updated=deactivated=0
        try:
            for x in rows:
                old=(by_dno.get(x['d_no'].upper()) if x['d_no'] else None) or by_name.get(x['name'].upper())
                if old:
                    seen.add(old['id'])
                    c.execute("UPDATE fleet_drivers SET name=?,phone=?,d_no=?,is_active=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(x['name'],x['phone'],x['d_no'],x['is_active'],old['id']))
                    updated+=1
                else:
                    cur=c.execute("INSERT INTO fleet_drivers(name,phone,d_no,is_active,note) VALUES(?,?,?,?,?)",(x['name'],x['phone'],x['d_no'],x['is_active'],''))
                    seen.add(cur.lastrowid); added+=1
                drow=None
                if x['d_no']:
                    drow=c.execute("SELECT id FROM drivers WHERE UPPER(TRIM(COALESCE(d_no,'')))=UPPER(TRIM(?)) LIMIT 1",(x['d_no'],)).fetchone()
                if not drow:
                    drow=c.execute("SELECT id FROM drivers WHERE UPPER(TRIM(name))=UPPER(TRIM(?)) LIMIT 1",(x['name'],)).fetchone()
                if drow:
                    c.execute("UPDATE drivers SET name=?,phone=?,d_no=?,active=? WHERE id=?",(x['name'],x['phone'],x['d_no'],x['is_active'],drow['id']))
                else:
                    c.execute("INSERT INTO drivers(name,phone,d_no,active) VALUES(?,?,?,?)",(x['name'],x['phone'],x['d_no'],x['is_active']))
            if deactivate_missing:
                for old in existing:
                    if int(old.get('is_active') or 0)==1 and old['id'] not in seen:
                        c.execute("UPDATE fleet_drivers SET is_active=0,updated_at=CURRENT_TIMESTAMP WHERE id=?",(old['id'],))
                        if str(old.get('d_no') or '').strip():
                            c.execute("UPDATE drivers SET active=0 WHERE UPPER(TRIM(COALESCE(d_no,'')))=UPPER(TRIM(?))",(old.get('d_no'),))
                        else:
                            c.execute("UPDATE drivers SET active=0 WHERE UPPER(TRIM(name))=UPPER(TRIM(?))",(old.get('name'),))
                        deactivated+=1
            c.commit()
        except Exception:
            c.rollback(); c.close(); raise
        c.close()
        audit('DRIVER_BULK_UPDATE','FILO',f"Yeni {added}, güncellenen {updated}, pasif {deactivated}")
        return {'ok':True,'added':added,'updated':updated,'deactivated':deactivated,'by':user.get('username','ADMIN')}
