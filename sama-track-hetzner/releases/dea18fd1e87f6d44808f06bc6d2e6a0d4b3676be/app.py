
from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import HTMLResponse, StreamingResponse, JSONResponse, Response
from pydantic import BaseModel
from typing import Optional
from pathlib import Path
import os
import shutil
import sqlite3
import uvicorn
from io import BytesIO
from datetime import datetime, date
import requests
import re
import html as html_lib
import hashlib
import secrets
import hmac
import contextvars
from datetime import timedelta

BASE = Path(__file__).resolve().parent

# Railway persistent SQLite storage.
# Prefer /data whenever the Railway volume is mounted there.
DATA_DIR = Path("/data")
ENV_VOLUME = os.environ.get("RAILWAY_VOLUME_MOUNT_PATH", "").strip()

if DATA_DIR.exists():
    VOLUME_DIR = DATA_DIR
elif ENV_VOLUME:
    VOLUME_DIR = Path(ENV_VOLUME).resolve()
else:
    VOLUME_DIR = BASE

DB = VOLUME_DIR / "sevkiyat.db"
SEED_DB = BASE / "sevkiyat.db"

def prepare_database_file():
    """Use persistent storage when /data or Railway volume is available."""
    VOLUME_DIR.mkdir(parents=True, exist_ok=True)

    # First deployment after attaching the volume:
    # copy the existing DB only once, then keep using the persistent DB.
    if DB != SEED_DB and not DB.exists() and SEED_DB.exists():
        shutil.copy2(SEED_DB, DB)
        print(f"[SAMA] Existing database copied to persistent volume: {DB}")

    print(f"[SAMA] SQLite database: {DB}")
    print(f"[SAMA] Data dir exists: {DATA_DIR.exists()}")
    print(f"[SAMA] Railway volume env: {ENV_VOLUME or 'not-set'}")
    print(f"[SAMA] Persistent volume selected: {DB.parent != BASE}")

app = FastAPI(title="SAMA TRACK V79 DRIVER VEHICLE")

AUTH_COOKIE="sama_session"
AUTH_SESSION_HOURS=12
CURRENT_AUTH_USER=contextvars.ContextVar("sama_auth_user",default=None)

PERMISSION_LABELS={
  "dashboard.view":"Ana Sayfa Gör",
  "shipment.view":"Sevkiyat Gör",
  "shipment.create":"Yeni Sevkiyat Ekle",
  "shipment.edit":"Sevkiyat / Çıkış / Giriş Düzenle",
  "shipment.delete":"Sevkiyat Sil",
  "excel.import":"Excel / OneDrive İçe Aktar",
  "operation.view":"Operasyon / Gemi / Yükleme Sırası Gör",
  "operation.edit":"Operasyon / Gemi / Yükleme Sırası Düzenle",
  "fleet.view":"Filo Gör",
  "fleet.edit":"Filo Düzenle",
  "driver.view":"Şoförler Gör",
  "driver.edit":"Şoförler Düzenle",
  "maintenance.view":"Bakım Gör",
  "maintenance.edit":"Bakım Düzenle",
  "report.view":"Raporlar Gör",
  "report.export":"Rapor / Excel Dışa Aktar",
  "audit.view":"İşlem Geçmişi Gör",
  "cash.expense.edit":"Günlük Kasa Harcama Düzelt",
  "users.manage":"Kullanıcılar & Yetkiler Yönet",
}

ROLE_DEFAULTS={
  "ADMIN":set(PERMISSION_LABELS),
  "OPERATOR":{
    "dashboard.view","shipment.view","shipment.create","shipment.edit",
    "operation.view","operation.edit","fleet.view","driver.view",
    "maintenance.view","maintenance.edit","report.view"
  },
  "VIEWER":{
    "dashboard.view","shipment.view","operation.view","fleet.view",
    "driver.view","maintenance.view","report.view"
  },
  "DRIVER":{
    "shipment.view","shipment.create","shipment.edit"
  },
}

def _normalize_scna_value(value):
    if value is None:
        return ""
    s=str(value).strip().upper()
    if s.endswith(".0"):
        try:
            s=str(int(float(s)))
        except Exception:
            pass
    return s

def _validate_scna_value(value):
    s=_normalize_scna_value(value)
    if not s:
        return False,"SCNA boş bırakılamaz."
    if s.isdigit():
        if len(s)<5:
            return False,"SCNA eksik görünüyor. En az 5 hane olmalı."
    else:
        compact=re.sub(r"[^A-Z0-9]","",s)
        if len(compact)<5:
            return False,"SCNA eksik görünüyor. En az 5 karakter olmalı."
    return True,s

def _password_hash(password: str, salt_hex: str | None=None):
    if salt_hex:
        salt=bytes.fromhex(salt_hex)
    else:
        salt=secrets.token_bytes(16)
        salt_hex=salt.hex()
    digest=hashlib.pbkdf2_hmac("sha256",password.encode("utf-8"),salt,210000)
    return salt_hex,digest.hex()

def _password_ok(password: str, salt_hex: str, expected: str):
    _,actual=_password_hash(password,salt_hex)
    return hmac.compare_digest(actual,expected)

def _user_permissions(c,user_id:int,role:str):
    base=set(ROLE_DEFAULTS.get((role or "VIEWER").upper(),set()))
    overrides=c.execute(
      "SELECT permission,allowed FROM auth_user_permissions WHERE user_id=?",
      (user_id,)
    ).fetchall()
    for r in overrides:
        if int(r["allowed"]):
            base.add(r["permission"])
        else:
            base.discard(r["permission"])
    return base

def _session_user(token: str | None):
    if not token:
        return None
    c=db()
    row=c.execute("""
      SELECT u.id,u.username,u.full_name,u.role,u.is_active,u.must_change_password,
             u.driver_id,u.current_plate
      FROM auth_sessions s
      JOIN auth_users u ON u.id=s.user_id
      WHERE s.token=? AND u.is_active=1 AND datetime(s.expires_at)>datetime('now')
    """,(token,)).fetchone()
    if not row:
        c.close();return None
    x=dict(row)
    x["permissions"]=sorted(_user_permissions(c,x["id"],x["role"]))
    c.close()
    return x

def _required_permission(path:str,method:str):
    m=method.upper()
    if path.startswith("/api/auth/"):
        return None
    if path in ("/api/health",):
        return None

    if path.startswith("/api/import-excel") or path.startswith("/api/public-onedrive"):
        return "excel.import"
    if path.startswith("/api/resolve/"):
        return "shipment.create"
    if path.startswith("/api/export-excel"):
        return "report.export"

    if path.startswith("/api/fleet/vehicles"):
        return "fleet.view" if m=="GET" else "fleet.edit"
    if path.startswith("/api/fleet/drivers"):
        return "driver.view" if m=="GET" else "driver.edit"
    if path.startswith("/api/maintenance"):
        return "maintenance.view" if m=="GET" else "maintenance.edit"

    if (path.startswith("/api/fleet-operations") or
        path.startswith("/api/integrated-fleet-status") or
        path.startswith("/api/operation-center") or
        path.startswith("/api/live-operations") or
        path.startswith("/api/route-standards") or
        path.startswith("/api/route-learning") or
        path.startswith("/api/operation-performance") or path.startswith("/api/vessel-operations")):
        return "operation.view" if m=="GET" else "operation.edit"

    if (path.startswith("/api/performance") or path.startswith("/api/finance") or
        path.startswith("/api/daily-manager-summary") or path.startswith("/api/anomalies") or
        path.startswith("/api/alerts") or path.startswith("/api/route-fuel-alerts")):
        return "report.view"

    if path.startswith("/api/audit"):
        return "audit.view"

    if path.startswith("/api/cash-control/expense/") and m=="PATCH":
        return "cash.expense.edit"

    if path.startswith("/api/trips") or path.startswith("/api/data-quality") or \
       path.startswith("/api/autocomplete") or path.startswith("/api/lookups") or \
       path.startswith("/api/stats") or path.startswith("/api/cash-diagnostic"):
        if m=="GET":
            return "shipment.view"
        if m=="POST" and path=="/api/trips":
            return "shipment.create"
        if m=="DELETE":
            return "shipment.delete"
        return "shipment.edit"

    return None

@app.middleware("http")
async def sama_auth_middleware(request: Request, call_next):
    path=request.url.path
    if path=="/" or path.startswith("/api/auth/") or path=="/api/health":
        return await call_next(request)

    if path.startswith("/api/"):
        user=_session_user(request.cookies.get(AUTH_COOKIE))
        if not user:
            return JSONResponse({"detail":"Oturum gerekli."},status_code=401)
        request.state.auth_user=user
        tok=CURRENT_AUTH_USER.set(user)
        try:
            if int(user.get("must_change_password") or 0)==1:
                return JSONResponse(
                    {"detail":"Şifrenizi değiştirmeden diğer işlemleri kullanamazsınız."},
                    status_code=403
                )
            perm=_required_permission(path,request.method)
            if perm and perm not in set(user.get("permissions") or []) and user.get("role")!="ADMIN":
                return JSONResponse(
                    {"detail":f"Bu işlem için yetkiniz yok: {PERMISSION_LABELS.get(perm,perm)}"},
                    status_code=403
                )
            return await call_next(request)
        finally:
            CURRENT_AUTH_USER.reset(tok)

    return await call_next(request)


def db():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    return c

def addcol(c, table, col, definition):
    cols = [r["name"] for r in c.execute(f"PRAGMA table_info({table})")]
    if col not in cols:
        c.execute(f"ALTER TABLE {table} ADD COLUMN {col} {definition}")

def init_db():
    c = db()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS drivers(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      name TEXT NOT NULL,
      phone TEXT DEFAULT '',
      d_no TEXT DEFAULT '',
      active INTEGER DEFAULT 1
    );

    CREATE TABLE IF NOT EXISTS vehicles(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      plate TEXT NOT NULL UNIQUE,
      driver_id INTEGER,
      active INTEGER DEFAULT 1
    );

    CREATE TABLE IF NOT EXISTS areas(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      name TEXT NOT NULL UNIQUE,
      km_go REAL DEFAULT 0,
      km_back REAL DEFAULT 0,
      active INTEGER DEFAULT 1
    );

    CREATE TABLE IF NOT EXISTS cargo_categories(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      name TEXT NOT NULL UNIQUE,
      active INTEGER DEFAULT 1
    );

    CREATE TABLE IF NOT EXISTS customers(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      name TEXT NOT NULL UNIQUE,
      active INTEGER DEFAULT 1
    );

    CREATE TABLE IF NOT EXISTS trips(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      scna TEXT NOT NULL UNIQUE,
      plate TEXT NOT NULL,
      driver_id INTEGER,
      area_id INTEGER,
      trip_date TEXT,

      net_kg REAL DEFAULT 0,
      freight_rate REAL DEFAULT 0,
      freight_basis TEXT DEFAULT 'TON',
      cargo_category_id INTEGER,
      cargo_type TEXT DEFAULT 'BULK',
      customer_id INTEGER,

      exit_km REAL DEFAULT 0,
      exit_cash REAL DEFAULT 0,

      tank_start_liters REAL DEFAULT 0,
      dock_fee REAL DEFAULT 0,
      port_fee REAL DEFAULT 0,
      sonar REAL DEFAULT 0,

      exit_official_fuel_liters REAL DEFAULT 0,
      exit_official_fuel_total REAL DEFAULT 0,

      exit_commercial_fuel_liters REAL DEFAULT 0,
      exit_commercial_fuel_total REAL DEFAULT 0,

      exit_baghdad_fuel_liters REAL DEFAULT 0,
      exit_baghdad_fuel_total REAL DEFAULT 0,

      exit_allowance REAL DEFAULT 0,
      exit_premium REAL DEFAULT 0,
      exit_other REAL DEFAULT 0,
      exit_note TEXT DEFAULT '',
      exit_done INTEGER DEFAULT 0,
      exit_at DATETIME,

      entry_km REAL DEFAULT 0,
      tank_end_liters REAL DEFAULT 0,

      entry_extra_expense_1 REAL DEFAULT 0,
      entry_extra_expense_2 REAL DEFAULT 0,
      entry_extra_expense_3 REAL DEFAULT 0,

      entry_collection REAL DEFAULT 0,
      entry_cash_handed REAL DEFAULT 0,
      excel_price_k REAL,
      excel_freight_au REAL,
      excel_amount REAL,
      excel_remain REAL,
      delivery_time TEXT DEFAULT '',
      entry_note TEXT DEFAULT '',
      entry_done INTEGER DEFAULT 0,
      entry_at DATETIME,

      status TEXT DEFAULT 'Bekliyor',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE IF NOT EXISTS fuel_purchases(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      scna TEXT NOT NULL,
      liters REAL DEFAULT 0,
      total REAL DEFAULT 0,
      note TEXT DEFAULT '',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE IF NOT EXISTS audit_log(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      action TEXT,
      scna TEXT,
      detail TEXT,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE IF NOT EXISTS auth_users(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      username TEXT NOT NULL UNIQUE COLLATE NOCASE,
      full_name TEXT DEFAULT '',
      role TEXT NOT NULL DEFAULT 'VIEWER',
      password_salt TEXT NOT NULL,
      password_hash TEXT NOT NULL,
      is_active INTEGER DEFAULT 1,
      must_change_password INTEGER DEFAULT 0,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE IF NOT EXISTS auth_user_permissions(
      user_id INTEGER NOT NULL,
      permission TEXT NOT NULL,
      allowed INTEGER NOT NULL DEFAULT 1,
      PRIMARY KEY(user_id,permission)
    );

    CREATE TABLE IF NOT EXISTS auth_sessions(
      token TEXT PRIMARY KEY,
      user_id INTEGER NOT NULL,
      expires_at DATETIME NOT NULL,
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP
    );


    CREATE TABLE IF NOT EXISTS route_standards(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      area_id INTEGER UNIQUE,
      expected_hours REAL DEFAULT 48,
      expected_km REAL DEFAULT 0,
      expected_l100 REAL DEFAULT 40,
      tolerance_percent REAL DEFAULT 10,
      updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE IF NOT EXISTS fleet_vehicles(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      plate TEXT NOT NULL UNIQUE,
      brand TEXT DEFAULT '',
      model TEXT DEFAULT '',
      vehicle_type TEXT DEFAULT 'DAMPER',
      is_active INTEGER DEFAULT 1,
      garage_state TEXT DEFAULT 'GARAGE',
      note TEXT DEFAULT '',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE IF NOT EXISTS fleet_drivers(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      name TEXT NOT NULL,
      phone TEXT DEFAULT '',
      d_no TEXT DEFAULT '',
      is_active INTEGER DEFAULT 1,
      note TEXT DEFAULT '',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE IF NOT EXISTS vehicle_operations(
      plate TEXT PRIMARY KEY,
      load_state TEXT DEFAULT 'BOS',
      queue_no INTEGER DEFAULT 0,
      vessel TEXT DEFAULT '',
      operation_note TEXT DEFAULT '',
      updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
    );

    CREATE TABLE IF NOT EXISTS maintenance(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      plate TEXT NOT NULL,
      last_service_km REAL DEFAULT 0,
      next_service_km REAL DEFAULT 0,
      oil_change_km REAL DEFAULT 0,
      tire_note TEXT DEFAULT '',
      brake_note TEXT DEFAULT '',
      engine_note TEXT DEFAULT '',
      note TEXT DEFAULT '',
      updated_at DATETIME DEFAULT CURRENT_TIMESTAMP
    );
    """)

    # Migrations from older versions
    addcol(c, "trips", "net_kg", "REAL DEFAULT 0")
    addcol(c, "trips", "freight_rate", "REAL DEFAULT 0")
    addcol(c, "trips", "freight_basis", "TEXT DEFAULT 'TON'")
    addcol(c, "trips", "entry_collection", "REAL DEFAULT 0")
    addcol(c, "trips", "entry_cash_handed", "REAL DEFAULT 0")
    addcol(c, "trips", "excel_price_k", "REAL")
    addcol(c, "trips", "excel_freight_au", "REAL")
    addcol(c, "trips", "excel_amount", "REAL")
    addcol(c, "trips", "excel_remain", "REAL")
    addcol(c, "trips", "entry_extra_expense_1", "REAL DEFAULT 0")
    addcol(c, "trips", "entry_extra_expense_2", "REAL DEFAULT 0")
    addcol(c, "trips", "entry_extra_expense_3", "REAL DEFAULT 0")

    addcol(c, "trips", "tank_start_liters", "REAL DEFAULT 0")
    addcol(c, "trips", "tank_end_liters", "REAL DEFAULT 0")

    addcol(c, "trips", "exit_official_fuel_liters", "REAL DEFAULT 0")
    addcol(c, "trips", "exit_official_fuel_total", "REAL DEFAULT 0")
    addcol(c, "trips", "exit_commercial_fuel_liters", "REAL DEFAULT 0")
    addcol(c, "trips", "exit_commercial_fuel_total", "REAL DEFAULT 0")
    addcol(c, "trips", "exit_baghdad_fuel_liters", "REAL DEFAULT 0")
    addcol(c, "trips", "exit_baghdad_fuel_total", "REAL DEFAULT 0")
    addcol(c, "trips", "cargo_category_id", "INTEGER")
    addcol(c, "trips", "cargo_type", "TEXT DEFAULT 'BULK'")
    addcol(c, "trips", "customer_id", "INTEGER")
    addcol(c, "trips", "dock_fee", "REAL DEFAULT 0")
    addcol(c, "trips", "port_fee", "REAL DEFAULT 0")
    addcol(c, "trips", "sonar", "REAL DEFAULT 0")
    addcol(c, "trips", "delivery_time", "TEXT DEFAULT ''")
    addcol(c, "trips", "manual_lock", "INTEGER DEFAULT 0")
    addcol(c, "trips", "manual_lock_at", "TEXT DEFAULT ''")
    addcol(c, "trips", "manual_lock_reason", "TEXT DEFAULT ''")
    addcol(c, "vehicle_operations", "vessel_state", "TEXT DEFAULT ''")
    addcol(c, "vehicle_operations", "gps_state", "TEXT DEFAULT ''")
    c.execute("""UPDATE vehicle_operations
                 SET vessel_state=load_state,load_state='BOS',queue_no=0
                 WHERE load_state IN ('CALISIYOR','CIKTI') AND COALESCE(vessel_state,'')=''""")

    addcol(c, "audit_log", "username", "TEXT DEFAULT ''")
    addcol(c, "audit_log", "user_id", "INTEGER")
    addcol(c, "audit_log", "entity_type", "TEXT DEFAULT ''")
    addcol(c, "audit_log", "old_value", "TEXT DEFAULT ''")
    addcol(c, "audit_log", "new_value", "TEXT DEFAULT ''")
    # V79 - kullanıcı hesabını gerçek şoför ve o an kullandığı araçla eşleştir.
    addcol(c, "auth_users", "driver_id", "INTEGER")
    addcol(c, "auth_users", "current_plate", "TEXT DEFAULT ''")
    # V83 - aciklamali OTHER ve 7 gun geri alinabilir soft delete
    addcol(c, "trips", "exit_other_note", "TEXT DEFAULT ''")
    addcol(c, "trips", "is_deleted", "INTEGER DEFAULT 0")
    addcol(c, "trips", "deleted_at", "TEXT DEFAULT ''")
    addcol(c, "trips", "deleted_by", "TEXT DEFAULT ''")


    if c.execute("SELECT COUNT(*) c FROM drivers").fetchone()["c"] == 0:
        c.executemany(
            "INSERT INTO drivers(name,phone,d_no) VALUES(?,?,?)",
            [
                ("Ali Hassan", "+964700000001", "D001"),
                ("Ahmed Kareem", "+964700000002", "D002"),
                ("Hassan Ali", "+964700000003", "D003"),
                ("Omar Saad", "+964700000004", "D004"),
            ]
        )

    if c.execute("SELECT COUNT(*) c FROM vehicles").fetchone()["c"] == 0:
        ids = {r["name"]: r["id"] for r in c.execute("SELECT id,name FROM drivers")}
        c.executemany(
            "INSERT INTO vehicles(plate,driver_id) VALUES(?,?)",
            [
                ("21L15441", ids["Ali Hassan"]),
                ("22L28424", ids["Ahmed Kareem"]),
                ("22K29376", ids["Hassan Ali"]),
                ("22B34933", ids["Omar Saad"]),
            ]
        )

    if c.execute("SELECT COUNT(*) c FROM areas").fetchone()["c"] == 0:
        c.executemany(
            "INSERT INTO areas(name,km_go,km_back) VALUES(?,?,?)",
            [
                ("Basra", 70, 70),
                ("Baghdad", 600, 600),
                ("Najaf", 520, 520),
                ("Erbil", 950, 950),
                ("Karbala", 570, 570),
                ("Nasiriya", 250, 250),
            ]
        )

    if c.execute("SELECT COUNT(*) c FROM cargo_categories").fetchone()["c"] == 0:
        c.executemany(
            "INSERT INTO cargo_categories(name) VALUES(?)",
            [
                ("Mısır",),
                ("Buğday",),
                ("Arpa",),
                ("Pirinç",),
                ("Un",),
                ("Gübre",),
                ("Diğer",),
            ]
        )

    if c.execute("SELECT COUNT(*) c FROM customers").fetchone()["c"] == 0:
        c.executemany(
            "INSERT INTO customers(name) VALUES(?)",
            [
                ("SAMA ALMANAR",),
                ("FREIGHT",),
                ("DİĞER",),
            ]
        )

    for a in c.execute("SELECT id,name,km_go,km_back FROM areas").fetchall():
        if not c.execute("SELECT 1 FROM route_standards WHERE area_id=?",(a["id"],)).fetchone():
            total_km=float(a["km_go"] or 0)+float(a["km_back"] or 0)
            name=str(a["name"] or "").upper()
            expected_hours=48
            if "BASRA" in name: expected_hours=12
            elif "NASIRI" in name: expected_hours=24
            elif "BAGHDAD" in name or "BAGDAT" in name: expected_hours=36
            elif "NAJAF" in name or "KARBALA" in name: expected_hours=48
            elif "ERBIL" in name: expected_hours=72
            c.execute(
                "INSERT INTO route_standards(area_id,expected_hours,expected_km,expected_l100,tolerance_percent) VALUES(?,?,?,?,?)",
                (a["id"],expected_hours,total_km,40,10)
            )

    c.execute("""
      INSERT OR IGNORE INTO fleet_vehicles(plate)
      SELECT DISTINCT UPPER(TRIM(plate))
      FROM trips
      WHERE TRIM(COALESCE(plate,''))<>''
    """)

    c.execute("""
      INSERT INTO fleet_drivers(name,phone,d_no)
      SELECT d.name,COALESCE(d.phone,''),COALESCE(d.d_no,'')
      FROM drivers d
      WHERE TRIM(COALESCE(d.name,''))<>''
        AND NOT EXISTS(
          SELECT 1 FROM fleet_drivers fd
          WHERE UPPER(TRIM(fd.name))=UPPER(TRIM(d.name))
        )
    """)

    c.commit()
    c.close()

@app.on_event("startup")
def startup():
    prepare_database_file()
    init_db()
    try:
        fixed=_repair_v4_exit_cash_misimports()
        if fixed:
            print(f'SAMA_EXCEL_EXIT_CASH_FIX_V5 repaired {len(fixed)} rows')
    except Exception as e:
        print('SAMA_EXCEL_EXIT_CASH_FIX_V5 repair warning:',e)

class AuthSetupIn(BaseModel):
    username: str
    full_name: str = ""
    password: str

class AuthLoginIn(BaseModel):
    username: str
    password: str

class AuthPasswordIn(BaseModel):
    current_password: str
    new_password: str

class AuthUserIn(BaseModel):
    username: str
    full_name: str = ""
    role: str = "VIEWER"
    password: str = ""
    is_active: int = 1
    must_change_password: int = 0
    driver_id: Optional[int] = None

class AuthUserUpdateIn(BaseModel):
    full_name: str = ""
    role: str = "VIEWER"
    is_active: int = 1
    must_change_password: int = 0
    driver_id: Optional[int] = None

class AuthResetPasswordIn(BaseModel):
    new_password: str
    must_change_password: int = 1

class AuthPermissionsIn(BaseModel):
    permissions: dict[str,bool] = {}

class DriverVehicleSelectIn(BaseModel):
    plate: str

class ResolveNameIn(BaseModel):
    name: str

class TripCreate(BaseModel):
    scna: str
    plate: str
    driver_id: Optional[int] = None
    area_id: Optional[int] = None
    cargo_category_id: Optional[int] = None
    cargo_type: str = "BULK"
    customer_id: Optional[int] = None
    trip_date: str = ""
    net_kg: float = 0
    freight_rate: float = 0
    freight_basis: str = "TON"
    exit_km: float = 0
    tank_start_liters: float = 0
    exit_premium: float = 0
    dock_fee: float = 0
    port_fee: float = 0
    sonar: float = 0

class ExitIn(BaseModel):
    exit_km: float = 0
    exit_cash: float = 0

    tank_start_liters: float = 0

    exit_official_fuel_liters: float = 0
    exit_official_fuel_total: float = 0

    exit_commercial_fuel_liters: float = 0
    exit_commercial_fuel_total: float = 0

    exit_baghdad_fuel_liters: float = 0
    exit_baghdad_fuel_total: float = 0

    exit_allowance: float = 0
    exit_other: float = 0
    exit_other_note: str = ""
    exit_note: str = ""

class EntryIn(BaseModel):
    entry_km: float = 0
    tank_end_liters: float = 0

    entry_extra_expense_1: float = 0
    entry_extra_expense_2: float = 0
    entry_extra_expense_3: float = 0

    entry_collection: float = 0
    entry_cash_handed: float = 0
    entry_note: str = ""

class FuelIn(BaseModel):
    liters: float
    total: float
    note: str = ""

class StatusIn(BaseModel):
    status: str

class CashExpenseEditIn(BaseModel):
    document_no: str = ""
    amount: float = 0
    note: str = ""
    expense_date: str = ""
    currency: str = "IQD"

class RouteStandardIn(BaseModel):
    area_id: int
    expected_hours: float = 48
    expected_km: float = 0
    expected_l100: float = 40
    tolerance_percent: float = 10

class MaintenanceIn(BaseModel):
    plate: str
    last_service_km: float = 0
    next_service_km: float = 0
    oil_change_km: float = 0
    tire_note: str = ""
    brake_note: str = ""
    engine_note: str = ""
    note: str = ""

class FleetVehicleIn(BaseModel):
    plate: str
    brand: str = ""
    model: str = ""
    vehicle_type: str = "DAMPER"
    is_active: int = 1
    garage_state: str = "GARAGE"
    note: str = ""

class FleetDriverIn(BaseModel):
    name: str
    phone: str = ""
    d_no: str = ""
    is_active: int = 1
    note: str = ""

class VehicleOperationIn(BaseModel):
    plate: str
    load_state: str = "BOS"
    queue_no: int = 0
    vessel: str = ""
    operation_note: str = ""

class BulkFixIn(BaseModel):
    scnas: list[str] = []
    field: str
    value: str = ""

class TripEditIn(BaseModel):
    plate: str = ""
    driver_id: Optional[int] = None
    area_id: Optional[int] = None
    cargo_category_id: Optional[int] = None
    cargo_type: str = "BULK"
    customer_id: Optional[int] = None
    trip_date: str = ""
    net_kg: float = 0
    freight_rate: float = 0
    freight_basis: str = "KG"
    exit_km: float = 0
    tank_start_liters: float = 0
    exit_premium: float = 0
    dock_fee: float = 0
    port_fee: float = 0
    sonar: float = 0
    exit_allowance: float = 0
    exit_other: float = 0
    entry_km: float = 0
    tank_end_liters: float = 0
    entry_extra_expense_1: float = 0
    entry_extra_expense_2: float = 0
    entry_extra_expense_3: float = 0
    entry_collection: float = 0
    entry_cash_handed: float = 0
    status: str = "Bekliyor"
    note: str = ""

def audit(action, scna, detail="", entity_type="", old_value="", new_value=""):
    user=CURRENT_AUTH_USER.get()
    c = db()
    c.execute(
        """INSERT INTO audit_log(
             action,scna,detail,username,user_id,entity_type,old_value,new_value
           ) VALUES(?,?,?,?,?,?,?,?)""",
        (
          action,scna,detail,
          (user or {}).get("username",""),
          (user or {}).get("id"),
          entity_type or "",
          old_value or "",
          new_value or ""
        )
    )
    c.commit()
    c.close()

def _audit_friendly_value(field, value):
    if value in (None,""):
        return ""
    lookup={
      "driver_id":("drivers","name"),
      "customer_id":("customers","name"),
      "area_id":("areas","name"),
      "cargo_category_id":("cargo_categories","name"),
    }
    if field not in lookup:
        return str(value)
    table,col=lookup[field]
    try:
        c=db()
        r=c.execute(f"SELECT {col} value FROM {table} WHERE id=?",(value,)).fetchone()
        c.close()
        return str(r["value"]) if r else str(value)
    except Exception:
        return str(value)

def audit_trip_changes(scna, before, after):
    if not before or not after:
        return
    fields={
      "plate":"Plaka",
      "driver_id":"Şoför",
      "customer_id":"Müşteri",
      "area_id":"Bölge",
      "cargo_category_id":"Mal",
      "cargo_type":"Tip",
      "net_kg":"KG",
      "freight_rate":"Navlun",
      "trip_date":"Çıkış Tarihi",
      "entry_at":"Giriş Tarihi",
      "delivery_time":"Teslim/Giriş Tarihi",
      "status":"Durum",
      "exit_km":"Çıkış KM",
      "entry_km":"Giriş KM"
    }
    for key,label in fields.items():
        ov=before.get(key)
        nv=after.get(key)
        if str(ov or "") != str(nv or ""):
            audit(
              "EDIT",scna,f"{label} değiştirildi","SEVKIYAT",
              _audit_friendly_value(key,ov),
              _audit_friendly_value(key,nv)
            )

def tq():
    return """
    SELECT
      t.*,
      d.name driver_name,
      d.phone,
      d.d_no,
      a.name area_name,
      a.km_go planned_km_go,
      a.km_back planned_km_back,
      cc.name cargo_name,
      cu.name customer_name,
      t.cargo_type cargo_type_name,

      COALESCE((SELECT vo.vessel FROM vehicle_operations vo
                WHERE UPPER(TRIM(vo.plate))=UPPER(TRIM(t.plate)) LIMIT 1),'') vessel_name,
      COALESCE((SELECT vo.load_state FROM vehicle_operations vo
                WHERE UPPER(TRIM(vo.plate))=UPPER(TRIM(t.plate)) LIMIT 1),'') vessel_operation_state,

      COALESCE(NULLIF(t.entry_at,''), NULLIF(t.delivery_time,'')) entry_display_at,
      CASE
        WHEN COALESCE(NULLIF(t.entry_at,''), NULLIF(t.delivery_time,'')) IS NOT NULL THEN 'Tamamlandı'
        ELSE t.status
      END status_display,

      COALESCE((SELECT SUM(f.total) FROM fuel_purchases f WHERE f.scna=t.scna),0) road_fuel_total,
      COALESCE((SELECT SUM(f.liters) FROM fuel_purchases f WHERE f.scna=t.scna),0) road_fuel_liters,

      CASE
        WHEN COALESCE((SELECT SUM(f.liters) FROM fuel_purchases f WHERE f.scna=t.scna),0) > 0
        THEN COALESCE((SELECT SUM(f.total) FROM fuel_purchases f WHERE f.scna=t.scna),0) /
             COALESCE((SELECT SUM(f.liters) FROM fuel_purchases f WHERE f.scna=t.scna),0)
        ELSE 0
      END road_fuel_avg_price,

      CASE
        WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
        ELSE (t.net_kg/1000.0)*t.freight_rate
      END freight_total,

      CASE
        WHEN t.entry_done=1 AND t.exit_km>0 AND t.entry_km>=t.exit_km THEN t.entry_km-t.exit_km
        ELSE 0
      END actual_km,

      CASE
        WHEN t.entry_done=1 AND t.delivery_time<>'' AND t.trip_date<>'' THEN
          ROUND((JULIANDAY(t.delivery_time)-JULIANDAY(t.trip_date))*24,1)
        ELSE 0
      END trip_duration_hours,

      (
        t.tank_start_liters +
        t.exit_official_fuel_liters +
        t.exit_commercial_fuel_liters +
        t.exit_baghdad_fuel_liters +
        COALESCE((SELECT SUM(f.liters) FROM fuel_purchases f WHERE f.scna=t.scna),0)
      ) total_fuel_available,

      CASE
        WHEN t.entry_done=1 AND t.exit_km>0 AND t.entry_km>=t.exit_km THEN
          MAX(
            (
              t.tank_start_liters +
              t.exit_official_fuel_liters +
              t.exit_commercial_fuel_liters +
              t.exit_baghdad_fuel_liters +
              COALESCE((SELECT SUM(f.liters) FROM fuel_purchases f WHERE f.scna=t.scna),0)
            ) - t.tank_end_liters,
            0
          )
        ELSE 0
      END fuel_consumed_liters,

      CASE
        WHEN t.entry_done=1 AND (t.entry_km-t.exit_km)>0 THEN
          (
            MAX(
              (
                t.tank_start_liters +
                t.exit_official_fuel_liters +
                t.exit_commercial_fuel_liters +
                t.exit_baghdad_fuel_liters +
                COALESCE((SELECT SUM(f.liters) FROM fuel_purchases f WHERE f.scna=t.scna),0)
              ) - t.tank_end_liters,
              0
            ) / (t.entry_km-t.exit_km)
          ) * 100
        ELSE 0
      END liters_per_100km,

      CASE
        WHEN t.entry_done=1 AND t.exit_km>0 AND t.entry_km>t.exit_km AND
             MAX(
               (
                 t.tank_start_liters +
                 t.exit_official_fuel_liters +
                 t.exit_commercial_fuel_liters +
                 t.exit_baghdad_fuel_liters +
                 COALESCE((SELECT SUM(f.liters) FROM fuel_purchases f WHERE f.scna=t.scna),0)
               ) - t.tank_end_liters,
               0
             ) > 0
        THEN
          (t.entry_km-t.exit_km) /
          MAX(
            (
              t.tank_start_liters +
              t.exit_official_fuel_liters +
              t.exit_commercial_fuel_liters +
              t.exit_baghdad_fuel_liters +
              COALESCE((SELECT SUM(f.liters) FROM fuel_purchases f WHERE f.scna=t.scna),0)
            ) - t.tank_end_liters,
            0
          )
        ELSE 0
      END km_per_liter,

      (
        t.exit_official_fuel_total +
        t.exit_commercial_fuel_total +
        t.exit_baghdad_fuel_total +
        COALESCE((SELECT SUM(f.total) FROM fuel_purchases f WHERE f.scna=t.scna),0)
      ) fuel_expense_total,

      (
        t.exit_official_fuel_total +
        t.exit_commercial_fuel_total +
        t.exit_baghdad_fuel_total +
        COALESCE((SELECT SUM(f.total) FROM fuel_purchases f WHERE f.scna=t.scna),0) +
        t.exit_allowance +
        t.exit_premium +
        t.dock_fee +
        t.port_fee +
        t.sonar +
        t.exit_other +
        t.entry_extra_expense_1 +
        t.entry_extra_expense_2 +
        t.entry_extra_expense_3
      ) total_expense,

      (
        CASE
          WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
          ELSE (t.net_kg/1000.0)*t.freight_rate
        END
      ) -
      (
        t.exit_official_fuel_total +
        t.exit_commercial_fuel_total +
        t.exit_baghdad_fuel_total +
        COALESCE((SELECT SUM(f.total) FROM fuel_purchases f WHERE f.scna=t.scna),0) +
        t.exit_allowance +
        t.exit_premium +
        t.dock_fee +
        t.port_fee +
        t.sonar +
        t.exit_other +
        t.entry_extra_expense_1 +
        t.entry_extra_expense_2 +
        t.entry_extra_expense_3
      ) trip_profit,

      (
        CASE
          WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
          ELSE (t.net_kg/1000.0)*t.freight_rate
        END
      ) - t.entry_collection amount_due,

      (
        t.exit_official_fuel_total +
        t.exit_commercial_fuel_total +
        t.exit_baghdad_fuel_total +
        COALESCE((SELECT SUM(f.total) FROM fuel_purchases f WHERE f.scna=t.scna),0)
      ) fuel_to_deduct_from_collection,

      (
        CASE
          WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
          ELSE (t.net_kg/1000.0)*t.freight_rate
        END
      ) -
      (
        COALESCE((SELECT SUM(f.total) FROM fuel_purchases f WHERE f.scna=t.scna),0) +
        t.exit_other +
        t.entry_extra_expense_1 +
        t.entry_extra_expense_2 +
        t.entry_extra_expense_3
      ) expected_cash_handover,

      CASE
        WHEN t.excel_remain IS NOT NULL THEN
          CASE
            WHEN COALESCE(t.entry_collection,0)>0
            THEN t.entry_collection - (
              CASE
                WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
                ELSE (t.net_kg/1000.0)*t.freight_rate
              END
            )
            ELSE 0
          END
        WHEN COALESCE(t.entry_cash_handed,0)=0
         AND COALESCE(t.entry_collection,0)>0
         AND t.entry_done=1
        THEN t.entry_collection - (
              CASE
                WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
                ELSE (t.net_kg/1000.0)*t.freight_rate
              END
            )
        ELSE
          t.entry_cash_handed -
          (
            (
              CASE
                WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
                ELSE (t.net_kg/1000.0)*t.freight_rate
              END
            ) -
            (
              COALESCE((SELECT SUM(f.total) FROM fuel_purchases f WHERE f.scna=t.scna),0) +
              t.entry_extra_expense_1 +
              t.entry_extra_expense_2 +
              t.entry_extra_expense_3
            )
          )
      END driver_cash_diff,

      CASE
        WHEN t.entry_done=0 THEN 'AÇIK'
        WHEN ABS(CASE
            WHEN t.excel_remain IS NOT NULL THEN
              CASE
                WHEN COALESCE(t.entry_collection,0)>0
                THEN t.entry_collection - (
              CASE
                WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
                ELSE (t.net_kg/1000.0)*t.freight_rate
              END
            )
                ELSE 0
              END
            WHEN COALESCE(t.entry_cash_handed,0)=0
             AND COALESCE(t.entry_collection,0)>0
             AND t.entry_done=1
            THEN t.entry_collection - (
              CASE
                WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
                ELSE (t.net_kg/1000.0)*t.freight_rate
              END
            )
            ELSE
              t.entry_cash_handed -
              (
                (
                  CASE
                    WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
                    ELSE (t.net_kg/1000.0)*t.freight_rate
                  END
                ) -
                (
                  COALESCE((SELECT SUM(f.total) FROM fuel_purchases f WHERE f.scna=t.scna),0) +
                  t.entry_extra_expense_1 +
                  t.entry_extra_expense_2 +
                  t.entry_extra_expense_3
                )
              )
          END) < 0.01 THEN 'HESAP TAMAM'
        WHEN (CASE
            WHEN t.excel_remain IS NOT NULL THEN
              CASE
                WHEN COALESCE(t.entry_collection,0)>0
                THEN t.entry_collection - (
              CASE
                WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
                ELSE (t.net_kg/1000.0)*t.freight_rate
              END
            )
                ELSE 0
              END
            WHEN COALESCE(t.entry_cash_handed,0)=0
             AND COALESCE(t.entry_collection,0)>0
             AND t.entry_done=1
            THEN t.entry_collection - (
              CASE
                WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
                ELSE (t.net_kg/1000.0)*t.freight_rate
              END
            )
            ELSE
              t.entry_cash_handed -
              (
                (
                  CASE
                    WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
                    ELSE (t.net_kg/1000.0)*t.freight_rate
                  END
                ) -
                (
                  COALESCE((SELECT SUM(f.total) FROM fuel_purchases f WHERE f.scna=t.scna),0) +
                  t.entry_extra_expense_1 +
                  t.entry_extra_expense_2 +
                  t.entry_extra_expense_3
                )
              )
          END) < 0 THEN 'EKSİK PARA'
        ELSE 'FAZLA PARA'
      END driver_cash_status

    FROM trips t
    LEFT JOIN drivers d ON d.id=t.driver_id
    LEFT JOIN areas a ON a.id=t.area_id
    LEFT JOIN cargo_categories cc ON cc.id=t.cargo_category_id
    LEFT JOIN customers cu ON cu.id=t.customer_id
    """


def _norm_header(v):
    if v is None:
        return ""
    s = str(v).strip().upper()
    repl = {
        "İ":"I","Ş":"S","Ğ":"G","Ü":"U","Ö":"O","Ç":"C",
        "ı":"I","ş":"S","ğ":"G","ü":"U","ö":"O","ç":"C"
    }
    for a,b in repl.items():
        s=s.replace(a,b)
    for ch in [" ", "_", "-", ".", "/", "\\", "(", ")", "[", "]", ":"]:
        s=s.replace(ch,"")
    return s

def _excel_num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int,float)):
        return float(v)
    s=str(v).strip().replace(" ","")
    if not s:
        return 0.0
    # 1.500.000 / 469.770 gibi IQD formatları
    if "." in s and "," not in s:
        parts=s.split(".")
        if len(parts)>1 and all(len(p)==3 for p in parts[1:]):
            s="".join(parts)
    elif "," in s and "." not in s:
        parts=s.split(",")
        if len(parts)>1 and all(len(p)==3 for p in parts[1:]):
            s="".join(parts)
        else:
            s=s.replace(",",".")
    elif "." in s and "," in s:
        if s.rfind(",") > s.rfind("."):
            s=s.replace(".","").replace(",",".")
        else:
            s=s.replace(",","")
    try:
        return float(s)
    except:
        return 0.0


def _excel_optional_num(v):
    if v is None:
        return None
    s=str(v).strip()
    if s in ("", "-", "—", ".", "#VALUE!", "#DEĞER!", "#N/A", "#SAYI/0!"):
        return None
    return _excel_num(v)

def _excel_date(v):
    if v is None or v == "":
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    # Excel serial date: örn. 45999
    if isinstance(v, (int,float)) and 20000 < float(v) < 80000:
        try:
            from openpyxl.utils.datetime import from_excel
            return from_excel(v).strftime("%Y-%m-%d")
        except:
            pass
    s=str(v).strip()
    for fmt in ("%d.%m.%Y","%d/%m/%Y","%Y-%m-%d","%d-%m-%Y"):
        try:
            return datetime.strptime(s,fmt).strftime("%Y-%m-%d")
        except:
            pass
    return s[:10]

def _excel_datetime(v):
    if v is None or v == "":
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return datetime(v.year,v.month,v.day).strftime("%Y-%m-%d 00:00:00")
    if isinstance(v, (int,float)) and 20000 < float(v) < 80000:
        try:
            from openpyxl.utils.datetime import from_excel
            return from_excel(v).strftime("%Y-%m-%d %H:%M:%S")
        except:
            pass

    s=str(v).strip()
    if s in ("", ".", "-", "—", "#VALUE!", "#DEĞER!", "#N/A", "#SAYI/0!"):
        return ""

    for fmt in (
        "%d.%m.%Y","%d/%m/%Y","%Y-%m-%d",
        "%d-%m-%Y","%d.%m.%Y %H:%M","%Y-%m-%d %H:%M:%S"
    ):
        try:
            return datetime.strptime(s,fmt).strftime("%Y-%m-%d %H:%M:%S")
        except:
            pass
    return s

def _first(row, headers, aliases, default=None):
    for alias in aliases:
        k=_norm_header(alias)
        if k in headers:
            val=row[headers[k]]
            if val is not None and str(val).strip()!="":
                return val
    return default

@app.post("/api/import-excel")
async def import_excel(request: Request, filename: str = "", update_existing: int = 0, preview_only: int = 0, skip_warnings: int = 0):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(
            500,
            "Excel aktarımı için openpyxl gerekli. Komut: pip install openpyxl"
        )

    raw = await request.body()
    if not raw:
        raise HTTPException(400, "Excel dosyası boş")

    try:
        wb = load_workbook(BytesIO(raw), data_only=True, read_only=True)
    except Exception as e:
        raise HTTPException(400, f"Excel dosyası açılamadı: {e}")

    # Öncelik gerçek operasyon sayfasında.
    ws = None
    for name in wb.sheetnames:
        if _norm_header(name) in ("YUKLEMELISTESI","YUKLEMELIST","LOADINGLIST"):
            ws = wb[name]
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]

    # read_only + iter_rows ile hızlı oku. Bu dosyada 8.797 x 62 hücre var;
    # ws.cell() ile tek tek okumak aşırı yavaştı.
    row_iter = ws.iter_rows(values_only=True)

    header_row = None
    header_map = {}
    buffered_rows = []

    for rno, vals in enumerate(row_iter, start=1):
        vals = tuple(vals)
        hm={_norm_header(v): i for i,v in enumerate(vals) if v is not None}
        score=0
        for candidate in ("SCNA","SERINO","SERIALNO","PLATE","PLAKA","DRIVER","DRIVERNAME","WEIGHT","FRIGHT"):
            if _norm_header(candidate) in hm:
                score += 1
        if score >= 2 or "SCNA" in hm:
            header_row=rno
            header_map=hm
            break
        if rno >= 30:
            break

    if not header_row:
        raise HTTPException(400, "Başlık satırı bulunamadı. SCNA/Plate gibi başlıklar gerekli.")

    # REMAIN kolonunu başlıktan bul. Sabit AX varsayımına güvenmiyoruz.
    remain_col = None
    for candidate in ("REMAIN","KALAN","BALANCE","REMAINING","REMAINDER"):
        key=_norm_header(candidate)
        if key in header_map:
            remain_col=header_map[key]
            break

    c=db()
    added=0
    updated=0
    skipped=0
    warning_skipped=0
    manual_protected_skipped=0
    remain_updated=0
    remain_missing=0
    delivery_date_source_count=0
    delivery_date_written_count=0
    invalid_scna_skipped=0
    excel_duplicate_count=0
    excel_duplicate_records=[]
    seen_excel_scna=set()
    errors=[]
    failed_records=[]
    preview_records=[]
    problem_records=[]
    preview_counts={"clean":0,"warning":0,"error":0,"duplicate":0,"update":0}
    created={"drivers":0,"vehicles":0,"customers":0,"areas":0,"cargo":0}

    def get_or_create_driver(name, phone="", dno=""):
        if not name:
            return None
        row=c.execute("SELECT id FROM drivers WHERE UPPER(name)=UPPER(?)",(name,)).fetchone()
        if row:
            return row["id"]
        cur=c.execute("INSERT INTO drivers(name,phone,d_no) VALUES(?,?,?)",(name,phone or "",dno or ""))
        created["drivers"]+=1
        return cur.lastrowid

    def get_or_create_vehicle(plate, driver_id=None):
        # Şoför araç sabiti değildir; her seferde trips.driver_id ile tutulur.
        if not plate:
            return None
        row=c.execute("SELECT id FROM vehicles WHERE UPPER(plate)=UPPER(?)",(plate,)).fetchone()
        if row:
            return row["id"]
        cur=c.execute("INSERT INTO vehicles(plate,driver_id) VALUES(?,NULL)",(plate.upper(),))
        created["vehicles"]+=1
        return cur.lastrowid

    def get_or_create_named(table, name, key):
        if not name:
            return None
        row=c.execute(f"SELECT id FROM {table} WHERE UPPER(name)=UPPER(?)",(name,)).fetchone()
        if row:
            return row["id"]
        cur=c.execute(f"INSERT INTO {table}(name) VALUES(?)",(name,))
        created[key]+=1
        return cur.lastrowid

    def get_or_create_area(name, km_go=0, km_back=0):
        if not name:
            return None
        row=c.execute("SELECT id,km_go,km_back FROM areas WHERE UPPER(name)=UPPER(?)",(name,)).fetchone()
        if row:
            if (not row["km_go"] and km_go) or (not row["km_back"] and km_back):
                c.execute("UPDATE areas SET km_go=CASE WHEN COALESCE(km_go,0)=0 THEN ? ELSE km_go END, km_back=CASE WHEN COALESCE(km_back,0)=0 THEN ? ELSE km_back END WHERE id=?",(km_go,km_back,row["id"]))
            return row["id"]
        cur=c.execute("INSERT INTO areas(name,km_go,km_back) VALUES(?,?,?)",(name,km_go or 0,km_back or 0))
        created["areas"]+=1
        return cur.lastrowid

    try:
        empty_streak = 0
        processed = 0

        # Header bulunduğunda row_iter zaten bir sonraki satırdan devam ediyor.
        for rno, row in enumerate(row_iter, start=header_row+1):
            row = tuple(row)
            processed += 1

            # Uzun boş kuyruk varsa Excel'in stil uygulanmış ama veri olmayan satırlarını bırak.
            if not any(v is not None and str(v).strip() != "" for v in row):
                empty_streak += 1
                if empty_streak >= 100:
                    break
                continue
            empty_streak = 0

            scna=_first(row,header_map,["SCNA","Serial No","Seri No","SNCA"])
            plate=_first(row,header_map,["Plate","Plaka","Plate - الرقم"])
            if not scna and not plate:
                continue

            scna=str(scna or "").strip().upper()
            invalid_scna_values={"",".","-","—","NONE","NAN","#VALUE!","#N/A"}
            if scna in invalid_scna_values:
                invalid_scna_skipped+=1
                # Kullanıcıyı 5.000 adet nokta satırıyla boğma; ilk örnekleri raporla.
                if plate and len(failed_records)<50:
                    failed_records.append({"row":rno,"scna":scna,"reason":"Geçersiz/boş SCNA; satır aktarılmadı"})
                continue

            if scna in seen_excel_scna:
                excel_duplicate_count+=1
                rec={"row":rno,"scna":scna,"reason":"Aynı SCNA Excel içinde birden fazla kez geçiyor; ikinci kayıt aktarılmadı"}
                excel_duplicate_records.append(rec)
                errors.append(f"Satır {rno} / SCNA {scna}: Excel içinde mükerrer")
                continue
            seen_excel_scna.add(scna)

            existing = c.execute(
                "SELECT id,COALESCE(manual_lock,0) manual_lock FROM trips WHERE UPPER(scna)=UPPER(?)",
                (scna,)
            ).fetchone()
            if existing and int(existing["manual_lock"] or 0)==1:
                protected_delivery=_excel_datetime(row[2] if len(row)>2 else "")
                protected_entry_km=_excel_num(row[29] if len(row)>29 else 0)

                protected_row=c.execute(
                    "SELECT entry_at,delivery_time,entry_km FROM trips WHERE id=?",
                    (existing["id"],)
                ).fetchone()

                if protected_delivery and protected_row and not (
                    (protected_row["entry_at"] or "").strip() or
                    (protected_row["delivery_time"] or "").strip()
                ):
                    c.execute("""
                      UPDATE trips SET
                        delivery_time=?,
                        entry_at=?,
                        entry_km=CASE
                          WHEN COALESCE(entry_km,0)<=0 AND ?>0 THEN ?
                          ELSE entry_km
                        END,
                        entry_done=1,
                        exit_done=1,
                        status='Tamamlandı',
                        updated_at=CURRENT_TIMESTAMP
                      WHERE id=?
                    """,(
                      protected_delivery,protected_delivery,
                      protected_entry_km,protected_entry_km,
                      existing["id"]
                    ))

                manual_protected_skipped += 1
                rec={
                    "row":rno,"scna":scna,"plate":str(plate or "").strip().upper(),
                    "driver":"","customer":"","area":"","cargo":"",
                    "kg":0,"price":0,"freight_rate":0,"amount":0,"collection":0,"excel_remain":None,
                    "delivery_time":"","exit_km":0,"entry_km":0,
                    "status":"MANUEL KORUMALI",
                    "messages":["Program içinde elle değiştirilmiş kayıt. Excel bu SCNA'nın üzerine yazmadı."]
                }
                if len(preview_records)<100: preview_records.append(rec)
                if len(problem_records)<1000: problem_records.append(rec)
                continue

            # Tek senkronizasyon mantığı:
            # - Yeni SCNA: ekle
            # - Mevcut SCNA: güncelle
            # - Manuel korumalı SCNA: dokunma
            # "Mükerrer diye atla" mantığı kaldırıldı; bu mantık giriş tarihlerini kaçırıyordu.

            try:
                c.execute("SAVEPOINT excel_row")
                # Excel'de G=Driver Name, X=Driver (PARA). X sütununu isim sanma.
                driver_name=str(_first(row,header_map,["Driver Name","DRIVERNAME","Şoför","Sofor"],"") or "").strip()
                if not driver_name and len(row) > 6:
                    driver_name=str(row[6] or "").strip()
                phone = str(row[7] or "").strip() if len(row) > 7 else str(_first(row,header_map,["Phone","Telefon"],"") or "").strip()
                dno = str(row[5] or "").strip() if len(row) > 5 else str(_first(row,header_map,["D.No","DNo","Driver No"],"") or "").strip()
                driver_id=get_or_create_driver(driver_name,phone,dno)

                # I = Plate
                if len(row) > 8 and row[8] not in (None,""):
                    plate=str(row[8]).strip()
                plate=str(plate or "").strip().upper()
                if not plate:
                    msg=f"Satır {rno}: Plaka boş"
                    errors.append(msg)
                    failed_records.append({"row":rno,"scna":scna,"reason":"Plaka boş"})
                    continue
                get_or_create_vehicle(plate,driver_id)

                # YUKLEME LISTESI gerçek kolon düzeni:
                # F=D.No, G=Driver Name, H=Phone, I=Plate, J=FRIGHT(Net KG),
                # K=PRICE(KG başı), L=Customer, M=CARGO, N=Cargo Type, O=Delivery Area.
                dno = str(row[5] or "").strip() if len(row) > 5 else str(_first(row,header_map,["D.No","DNo"],"") or "").strip()
                phone = str(row[7] or "").strip() if len(row) > 7 else str(_first(row,header_map,["Phone","Telefon"],"") or "").strip()

                customer = str(row[11] or "").strip() if len(row) > 11 else str(_first(row,header_map,["Customer","Müşteri","Musteri"],"") or "").strip()
                customer_id=get_or_create_named("customers",customer,"customers")

                area = str(row[14] or "").strip() if len(row) > 14 else str(_first(row,header_map,["Delivery Area","DeliveryArea","Teslimat Bölgesi","Teslim Yeri"],"") or "").strip()
                route_km_go=_excel_num(row[15] if len(row)>15 else 0)
                route_km_back=_excel_num(row[16] if len(row)>16 else 0)
                area_id=get_or_create_area(area,route_km_go,route_km_back)

                cargo = str(row[12] or "").strip() if len(row) > 12 else str(_first(row,header_map,["Cargo","Mal Cinsi","Mal","Product","Ürün"],"") or "").strip()
                cargo_id=get_or_create_named("cargo_categories",cargo,"cargo")

                cargo_type = str(row[13] or "").strip().upper() if len(row) > 13 else str(_first(row,header_map,["Cargo Type","CargoType","Yük Tipi","Yuk Tipi"],"") or "").strip().upper()
                # Excel'deki gerçek operasyon tiplerini koru. Sadece bariz yazım hatalarını düzelt.
                cargo_type_alias={"BLUK":"BULK","BULG":"BULK"}
                cargo_type=cargo_type_alias.get(cargo_type,cargo_type or "BULK")

                trip_date = _excel_date(row[1] if len(row) > 1 else _first(row,header_map,["Loading Date","LoadingDate","Date","Tarih"],""))
                delivery_time = _excel_datetime(
                    row[2] if len(row) > 2 else
                    _first(row,header_map,["Delivery Time","DeliveryTime","Teslim Tarihi","Teslim Zamanı"],"")
                )
                if delivery_time:
                    delivery_date_source_count += 1

                excel_exit_at=_excel_datetime(_first(row,header_map,["Exit Date","Exit Time","Çıkış Tarihi","Cikis Tarihi","Çıkış Saati"],""))
                excel_entry_at=_excel_datetime(_first(row,header_map,["Entry Date","Entry Time","Giriş Tarihi","Giris Tarihi","Giriş Saati"],""))

                net_kg = _excel_num(row[9] if len(row) > 9 else 0)       # J = FRIGHT = Net KG
                price_k = _excel_num(row[10] if len(row) > 10 else 0)      # K = PRICE
                freight_rate_au = _excel_num(row[46] if len(row) > 46 else 0) # AU = FREIGHT (final)
                freight_basis = "KG"

                # AV = AMOUNT = gerçek navlun toplamı; AW = COLLECTION = fiilen alınan para.
                freight_total = _excel_num(row[47] if len(row) > 47 else 0)

                # Finansal doğrulukta AV=AMOUNT ana kaynaktır. AU veya K bazı satırlarda
                # AMOUNT ile uyuşmuyor. Bu yüzden efektif navlun oranı mümkünse AMOUNT/KG.
                if freight_total and net_kg:
                    freight_rate = freight_total / net_kg
                else:
                    freight_rate = freight_rate_au or price_k

                exit_km=_excel_num(_first(row,header_map,["GIDIS KM","Exit KM","Çıkış KM","Cikis KM"],0))
                entry_km=_excel_num(_first(row,header_map,["GELIS KM","Entry KM","Giriş KM","Giris KM"],0))
                tank_start=_excel_num(_first(row,header_map,["DEPOSUNDAKI LITRE","Tank Start","Depodaki Mazot","Depo Mazot","Başlangıç Depo"],0))
                tank_end=_excel_num(_first(row,header_map,["DEPODA KALAN","Tank End","Depoda Kalan","Kalan Mazot"],0))
                if len(row) >= 45:
                    if not exit_km: exit_km=_excel_num(row[28])
                    if not entry_km: entry_km=_excel_num(row[29])
                    if not tank_start: tank_start=_excel_num(row[31])
                    if not tank_end: tank_end=_excel_num(row[44])

                # T=Premium, U=Harcırah, V=Port Fee, W=Dock Fee, Y=OTHER, Z=SONAR
                premium=_excel_num(row[19] if len(row)>19 else 0)
                allowance=_excel_num(row[20] if len(row)>20 else 0)
                port_fee=_excel_num(row[21] if len(row)>21 else 0)
                dock_fee=_excel_num(row[22] if len(row)>22 else 0)
                other=_excel_num(row[24] if len(row)>24 else 0)
                sonar=_excel_num(row[25] if len(row)>25 else 0)

                # Çıkış yakıtları: LT ve toplam tutar için farklı olası başlıklar.
                # Başlıklar tekrar ettiği için tam otomatik eşleme bazı eski Excel'lerde belirsiz olabilir.
                # Bilinen alternatif adları yine destekliyoruz.
                off_lt=_excel_num(_first(row,header_map,["RESMI LT","Official Fuel LT","Resmi Mazot LT"],0))
                off_total=_excel_num(_first(row,header_map,["RESMI TOPLAM","Official Fuel Total","Resmi Mazot Toplam"],0))
                com_lt=_excel_num(_first(row,header_map,["TICARI LT","Commercial Fuel LT","Ticari Mazot LT"],0))
                com_total=_excel_num(_first(row,header_map,["TICARI TOPLAM","Commercial Fuel Total","Ticari Mazot Toplam"],0))
                bag_lt=_excel_num(_first(row,header_map,["BAGDAT LT","Baghdad Fuel LT","Bağdat Mazot LT"],0))
                bag_total=_excel_num(_first(row,header_map,["BAGDAT TOPLAM","Baghdad Fuel Total","Bağdat Mazot Toplam"],0))

                # Kullanıcının mevcut YUKLEME LISTESI sabit kolonları:
                # AG=33 LITRE, AH=34 RESMI, AI=35 QID
                # AJ=36 LITRE, AK=37 TICARI, AL=38 IQD
                # AM=39 BAGDAT LITRE, AN=40 RESMI, AO=41 IQD
                # Python 0-index: 32..40
                if len(row) >= 41:
                    if not off_lt: off_lt=_excel_num(row[32])
                    if not off_total: off_total=_excel_num(row[34])
                    if not com_lt: com_lt=_excel_num(row[35])
                    if not com_total: com_total=_excel_num(row[37])
                    if not bag_lt: bag_lt=_excel_num(row[38])
                    if not bag_total: bag_total=_excel_num(row[40])

                # AW = COLLECTION = müşteriden fiilen alınan para.
                collection=_excel_num(row[48] if len(row)>48 else _first(row,header_map,["COLLECTION","Tahsilat","Müşteriden Alınan Para"],0))

                # SAMA_EXCEL_CASH_FIELDS_V4
                # Muhasebe için COLLECTION ile şoförün fiilen teslim ettiği para aynı şey değildir.
                # Bu yüzden yalnızca açıkça isimlendirilmiş Excel alanlarını entry_cash_handed olarak okuruz.
                entry_cash_handed=_excel_num(_first(row,header_map,[
                    "Şoförün Teslim Ettiği Para","Soforun Teslim Ettigi Para",
                    "Şoförün Verdiği Para","Soforun Verdigi Para",
                    "Entry Cash Handed","Cash Handed","Driver Cash Handed",
                    "Teslim Edilen Para","Kasa Teslim"
                ],0))

                # Çıkışta şoföre verilen toplam nakit. Yeni/ayrıntılı Excel başlıkları önceliklidir.
                exit_cash=_excel_num(_first(row,header_map,[
                    "Şoföre Verilen Avans","Sofore Verilen Avans",
                    "Şoföre Verilen Para","Sofore Verilen Para",
                    "Exit Cash","Driver Cash","Cash Given To Driver"
                ],0))

                # SAMA_EXCEL_EXIT_CASH_FIX_V5
                # X sütunu (Driver) toplam kasa çıkışı değildir. Eski dosyada sürücü masraf alanıdır;
                # bu yüzden artık exit_cash için kullanılmaz. Açık bir nakit kolonu yoksa, kullanıcı
                # kuralına göre şoföre verilen toplam nakit görünür çıkış giderlerinin toplamıdır.
                if not exit_cash:
                    exit_cash=(
                        off_total+com_total+bag_total+allowance+premium+other+
                        dock_fee+port_fee+sonar
                    )

                # REMAIN başlıktan okunur. Başlık bulunamazsa eski dosyalar için
                # yalnızca fallback olarak index 49 denenir.
                remain_raw=None
                if remain_col is not None and remain_col < len(row):
                    remain_raw=row[remain_col]
                elif len(row)>49:
                    remain_raw=row[49]

                excel_remain=_excel_optional_num(remain_raw)
                if excel_remain is None:
                    remain_missing+=1
                else:
                    remain_updated+=1

                # ---- EXCEL KONTROL MERKEZİ DOĞRULAMALARI ----
                row_errors=[]
                row_warnings=[]

                expected_amount=net_kg*freight_rate if net_kg and freight_rate else 0

                if price_k and freight_rate_au and abs(price_k-freight_rate_au)>0.0001:
                    row_warnings.append(
                        f"PRICE K ({price_k:,.6f}) ile FREIGHT AU ({freight_rate_au:,.6f}) farklı. "
                        "Finansal hesapta AMOUNT/KG kullanılacak."
                    )
                if excel_remain is not None and freight_total and collection:
                    calc_remain=collection-freight_total
                    if abs(excel_remain-calc_remain)>1:
                        row_warnings.append(
                            f"Excel REMAIN hatalı görünüyor: hücre {excel_remain:,.2f}, doğru COLLECTION-AMOUNT {calc_remain:,.2f}. "
                            "Program hesaplanan farkı kullanacak."
                        )

                if not driver_name:
                    row_warnings.append("Şoför adı boş.")
                if not customer:
                    row_warnings.append("Müşteri boş.")
                if not area:
                    row_warnings.append("Bölge boş.")
                if not cargo:
                    row_warnings.append("Mal cinsi boş.")

                if net_kg < 0:
                    row_errors.append("Net KG negatif.")
                elif net_kg > 80000:
                    row_warnings.append(f"Net KG olağandışı yüksek: {net_kg:,.0f}")

                if freight_rate < 0:
                    row_errors.append("PRICE negatif.")
                elif freight_rate > 100000:
                    row_warnings.append(f"PRICE olağandışı yüksek: {freight_rate:,.2f}")

                if freight_total and expected_amount and abs(freight_total-expected_amount) > 1:
                    row_warnings.append(
                        f"AMOUNT uyuşmuyor: Excel {freight_total:,.2f}, KG×PRICE {expected_amount:,.2f}"
                    )

                if delivery_time and entry_km <= 0:
                    row_warnings.append(
                        f"Giriş tarihi var ({delivery_time[:10]}) fakat GELİŞ KM boş. "
                        "Kayıt Tamamlandı olarak alınacak; KM bilgisi eksik kalacak."
                    )

                if entry_km > 0 and exit_km > 0 and entry_km < exit_km:
                    row_warnings.append(
                        f"Giriş KM ({entry_km:,.0f}) çıkış KM'den ({exit_km:,.0f}) küçük. "
                        "Kayıt aktarılacak fakat KM/anormallik hesabına alınmayacak."
                    )
                if entry_km > 0 and exit_km > 0 and (entry_km-exit_km) > 5000:
                    row_warnings.append(f"Sefer KM farkı çok yüksek: {entry_km-exit_km:,.0f} KM")

                fuel_liters=off_lt+com_lt+bag_lt
                if fuel_liters > 5000:
                    row_errors.append(f"Yakıt miktarı gerçekçi değil: {fuel_liters:,.0f} LT")
                elif fuel_liters > 1500:
                    row_warnings.append(f"Yakıt miktarı yüksek görünüyor: {fuel_liters:,.0f} LT")

                if tank_start > 2000 or tank_end > 2000:
                    row_warnings.append("Depo litre değeri 2.000 LT üzerinde.")

                if collection < 0:
                    row_errors.append("COLLECTION negatif.")
                if collection and expected_amount and collection > expected_amount*1.10:
                    row_warnings.append(
                        f"COLLECTION navlunun %10'dan fazla üzerinde: {collection:,.2f} > {expected_amount:,.2f}"
                    )

                preview_status="HATALI" if row_errors else (
                    "UYARI" if row_warnings else ("GÜNCELLENECEK" if existing else "TEMİZ")
                )
                messages=row_errors+row_warnings
                preview_rec={
                    "row":rno,"scna":scna,"plate":plate,"driver":driver_name,
                    "customer":customer,"area":area,"cargo":cargo,
                    "kg":net_kg,"price":price_k,"freight_rate":freight_rate,
                    "amount":freight_total or expected_amount,
                    "collection":collection,
                    "entry_cash_handed":entry_cash_handed,
                    "exit_cash":exit_cash,
                    "excel_remain":excel_remain,
                    "delivery_time":delivery_time,
                    "exit_km":exit_km,
                    "entry_km":entry_km,
                    "status":preview_status,"messages":messages
                }

                if row_errors:
                    preview_counts["error"]+=1
                    failed_records.append({
                        "row":rno,"scna":scna,"reason":" | ".join(row_errors)
                    })
                    errors.append(f"Satır {rno} / SCNA {scna}: {' | '.join(row_errors)}")
                    if len(preview_records)<100: preview_records.append(preview_rec)
                    if len(problem_records)<1000: problem_records.append(preview_rec)
                    c.execute("ROLLBACK TO excel_row")
                    c.execute("RELEASE excel_row")
                    continue

                if row_warnings:
                    preview_counts["warning"]+=1
                    if len(problem_records)<1000:
                        problem_records.append(preview_rec)
                    # UYARI sadece bilgi amaçlıdır. Satırın tamamını atlama.
                    # Özellikle "giriş tarihi var / giriş KM boş" gibi durumlarda
                    # dönüş tarihinin kaybolmasına izin verme.
                elif existing:
                    preview_counts["update"]+=1
                else:
                    preview_counts["clean"]+=1

                if len(preview_records)<100:
                    preview_records.append(preview_rec)

                # Durum Excel'deki gerçek hareket bilgisine göre belirlenir.
                # Delivery Time gerçek giriş tarihidir. GELİŞ KM boş olsa bile giriş tarihi varsa
                # kayıt tamamlanmış kabul edilir; KM eksikliği önizlemede ayrıca uyarılır.
                has_delivery_time = bool(delivery_time)
                has_entry_km = entry_km > 0 and (exit_km <= 0 or entry_km >= exit_km)

                entry_done = 1 if (has_delivery_time or has_entry_km) else 0
                exit_done = 1 if (
                    entry_done or exit_km > 0 or off_lt > 0 or com_lt > 0 or
                    bag_lt > 0 or allowance > 0 or other > 0
                ) else 0

                status="Tamamlandı" if entry_done else ("Yolda" if exit_done else "Bekliyor")

                if existing:
                    c.execute("""
                      UPDATE trips SET
                        plate=?,
                        driver_id=?,
                        area_id=?,
                        cargo_category_id=?,
                        cargo_type=?,
                        customer_id=?,
                        trip_date=?,
                        net_kg=?,
                        freight_rate=?,
                        freight_basis=?,
                        exit_km=?,
                        exit_cash=?,
                        tank_start_liters=?,
                        exit_official_fuel_liters=?,
                        exit_official_fuel_total=?,
                        exit_commercial_fuel_liters=?,
                        exit_commercial_fuel_total=?,
                        exit_baghdad_fuel_liters=?,
                        exit_baghdad_fuel_total=?,
                        exit_allowance=?,
                        exit_premium=?,
                        exit_other=?,
                        dock_fee=?,
                        port_fee=?,
                        sonar=?,
                        entry_km=?,
                        tank_end_liters=?,
                        entry_collection=?,
                        entry_cash_handed=CASE WHEN ?>0 THEN ? ELSE entry_cash_handed END,
                        excel_price_k=?,
                        excel_freight_au=?,
                        excel_amount=?,
                        excel_remain=?,
                        delivery_time=?,
                        exit_done=?,
                        entry_done=?,
                        status=?,
                        exit_at=?,
                        entry_at=?,
                        updated_at=CURRENT_TIMESTAMP
                      WHERE id=?
                    """,(
                        plate,driver_id,area_id,cargo_id,cargo_type,customer_id,trip_date,
                        net_kg,freight_rate,freight_basis,
                        exit_km,exit_cash,tank_start,off_lt,off_total,
                        com_lt,com_total,bag_lt,bag_total,
                        allowance,premium,other,dock_fee,port_fee,sonar,
                        entry_km,tank_end,collection,entry_cash_handed,entry_cash_handed,price_k,freight_rate_au,freight_total,excel_remain,delivery_time,
                        exit_done,entry_done,status,
                        excel_exit_at or (trip_date+" 00:00:00" if exit_done and trip_date else None),
                        delivery_time or excel_entry_at or None,
                        existing["id"]
                    ))
                    updated+=1
                else:
                    c.execute("""
                    INSERT INTO trips(
                      scna,plate,driver_id,area_id,cargo_category_id,cargo_type,customer_id,trip_date,
                      net_kg,freight_rate,freight_basis,
                      exit_km,exit_cash,tank_start_liters,exit_official_fuel_liters,exit_official_fuel_total,
                      exit_commercial_fuel_liters,exit_commercial_fuel_total,
                      exit_baghdad_fuel_liters,exit_baghdad_fuel_total,
                      exit_allowance,exit_premium,exit_other,dock_fee,port_fee,sonar,
                      entry_km,tank_end_liters,entry_collection,entry_cash_handed,excel_price_k,excel_freight_au,excel_amount,excel_remain,delivery_time,
                      exit_done,entry_done,status,exit_at,entry_at
                    ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,(
                    scna,plate,driver_id,area_id,cargo_id,cargo_type,customer_id,trip_date,
                    net_kg,freight_rate,freight_basis,
                    exit_km,exit_cash,tank_start,off_lt,off_total,
                    com_lt,com_total,bag_lt,bag_total,
                    allowance,premium,other,dock_fee,port_fee,sonar,
                    entry_km,tank_end,collection,entry_cash_handed,price_k,freight_rate_au,freight_total,excel_remain,delivery_time,
                    exit_done,entry_done,status,
                    excel_exit_at or (trip_date+" 00:00:00" if exit_done and trip_date else None),
                    delivery_time or excel_entry_at or None
                ))
                    added+=1

                c.execute("RELEASE excel_row")
                if delivery_time:
                    delivery_date_written_count += 1

            except Exception as e:
                try:
                    c.execute("ROLLBACK TO excel_row")
                    c.execute("RELEASE excel_row")
                except Exception:
                    pass
                reason=str(e)
                errors.append(f"Satır {rno} / SCNA {scna}: {reason}")
                failed_records.append({
                    "row":rno,
                    "scna":scna,
                    "reason":reason
                })

        if preview_only:
            c.rollback()
        else:
            c.commit()
    except Exception:
        c.rollback()
        raise
    finally:
        c.close()

    if not preview_only:
        audit(
            "EXCEL_IMPORT",
            filename or "EXCEL",
            f"{added} eklendi, {updated} güncellendi, {skipped} mükerrer atlandı, "
            f"{warning_skipped} uyarılı atlandı, {remain_updated} REMAIN okundu, "
            f"{remain_missing} REMAIN boş, {len(errors)} hata"
        )

    return {
        "ok": True,
        "sheet": ws.title,
        "header_row": header_row,
        "processed": processed,
        "added": added,
        "updated": updated,
        "skipped": skipped,
        "warning_skipped": warning_skipped,
        "manual_protected_skipped": manual_protected_skipped,
        "remain_header_found": remain_col is not None,
        "remain_column_index": (remain_col+1) if remain_col is not None else None,
        "remain_updated": remain_updated,
        "remain_missing": remain_missing,
        "delivery_date_source_count": delivery_date_source_count,
        "delivery_date_written_count": delivery_date_written_count,
        "delivery_date_not_written_count": max(delivery_date_source_count-delivery_date_written_count,0),
        "invalid_scna_skipped": invalid_scna_skipped,
        "excel_duplicate_count": excel_duplicate_count,
        "excel_duplicate_records": excel_duplicate_records[:100],
        "error_count": len(errors),
        "errors": errors[:100],
        "failed_records": failed_records[:500],
        "failed_scna": [x["scna"] for x in failed_records if x.get("scna")][:500],
        "preview_only": bool(preview_only),
        "preview_counts": preview_counts,
        "preview_records": preview_records,
        "problem_records": problem_records,
        "created": created
    }


def _safe_float(v):
    try:
        return float(v or 0)
    except:
        return 0.0

def _export_rows(q="", status=""):
    c=db()
    sql=tq()+" WHERE 1=1 "
    p=[]
    if q:
        like=f"%{q}%"
        sql += " AND (t.scna LIKE ? OR t.plate LIKE ? OR d.name LIKE ? OR a.name LIKE ? OR cu.name LIKE ? OR cc.name LIKE ?)"
        p += [like,like,like,like,like,like]
    if status and status != "Tümü":
        sql += " AND t.status=?"
        p.append(status)
    sql += " ORDER BY t.id ASC"
    rows=[dict(r) for r in c.execute(sql,p)]
    c.close()
    return rows

def _excel_headers():
    # Kullanıcının YUKLEME LISTESI düzeni, A:BG.
    # BC ve BD orijinal dosyada başlıksız olduğu için boş bırakılır.
    return [
        "SCNA","Loading\nDate","Delivery Time","AY","SEFER SURESI","D.No",
        "Driver Name","Phone","Plate","FRIGHT","PRICE","Customer","CARGO",
        "Cargo \nType","Delivery Area","KM-GO","KM-BACK","Total","0.35",
        "Premium","Harcırah","Port Fee","Dock Fee","Driver","OTHER","SONAR",
        "G.Total","DRIVER ADVANTAGE","GIDIS KM ","GELIS KM","FARK",
        "DEPOSUNDAKI \nLITRE","LITRE","RESMI","QID","LITRE","TICARI","IQD",
        "BAGDAT\nLITRE","RESMI","IQD","Genel\nToplam","TOPLAM\nLITRE",
        "KALMASI BEKLENEN","DEPODA \nKALAN","% TUKETIM ORANI","FREIGHT",
        "AMOUNT","COLLECTION ","REMAIN","GELİŞ","Maintanence \nCost IQD / Km",
        "Maintanence \nCost ","Gelir","","","GPS CONSUMPTION\nLITRE",
        "TOTAL KM \nGPS","GPS CONSUMPTION %"
    ]

def _build_export_workbook(rows):
    try:
        from openpyxl import Workbook, load_workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(
            500,
            "Excel dışa aktarma için openpyxl gerekli. Komut: pip install openpyxl"
        )

    # İstenirse APP.py yanına EXPORT_TEMPLATE.xlsx konabilir.
    # Bu durumda kullanıcının kendi Excel görünümü/stili korunur.
    template = BASE / "EXPORT_TEMPLATE.xlsx"
    if template.exists():
        wb = load_workbook(template)
        if "YUKLEME LISTESI" in wb.sheetnames:
            ws=wb["YUKLEME LISTESI"]
        else:
            ws=wb.create_sheet("YUKLEME LISTESI")

        # Mevcut veri satırlarını temizle; başlık/stil kalsın.
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row-1)
    else:
        wb=Workbook()
        ws=wb.active
        ws.title="YUKLEME LISTESI"

        # Gönderilen çalışma kitabındaki diğer ana sekmeleri de oluştur.
        for name in ["ANA FORM","MAZOT TAKIP","URUNLER","CARILER"]:
            if name not in wb.sheetnames:
                wb.create_sheet(name)

        headers=_excel_headers()
        for idx,h in enumerate(headers,1):
            ws.cell(1,idx,h)

        # Basit ama okunaklı varsayılan stil.
        fill=PatternFill("solid",fgColor="1F4E78")
        font=Font(color="FFFFFF",bold=True)
        thin=Side(style="thin",color="D9E2F3")
        for cell in ws[1]:
            cell.fill=fill
            cell.font=font
            cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
            cell.border=Border(bottom=thin)
        ws.row_dimensions[1].height=38
        ws.freeze_panes="A2"
        ws.auto_filter.ref=f"A1:BG1"

        widths={
            "A":14,"B":13,"C":18,"F":10,"G":26,"H":18,"I":14,"J":12,"K":10,
            "L":24,"M":18,"N":12,"O":24,"P":10,"Q":10,"T":12,"U":12,"V":12,
            "W":12,"X":12,"Y":12,"Z":12,"AC":12,"AD":12,"AF":14,"AG":10,
            "AH":10,"AI":12,"AJ":10,"AK":10,"AL":12,"AM":12,"AN":10,"AO":12,
            "AP":14,"AQ":12,"AS":12,"AT":14,"AU":10,"AV":14,"AW":14,"AX":14,
            "AZ":18,"BA":16,"BB":14,"BE":16,"BF":14,"BG":16
        }
        for col,w in widths.items():
            ws.column_dimensions[col].width=w

    # Eğer template başlıkları boşsa/eksikse başlıkları yine yaz.
    headers=_excel_headers()
    for idx,h in enumerate(headers,1):
        if ws.cell(1,idx).value in (None,""):
            ws.cell(1,idx,h)

    # CARILER / URUNLER sekmelerini veritabanından doldur.
    c=db()
    customers=[r["name"] for r in c.execute("SELECT name FROM customers ORDER BY name")]
    cargos=[r["name"] for r in c.execute("SELECT name FROM cargo_categories ORDER BY name")]
    c.close()

    if "CARILER" in wb.sheetnames:
        sh=wb["CARILER"]
        if sh.max_row>1: sh.delete_rows(2,sh.max_row-1)
        if sh["A1"].value in (None,""): sh["A1"]="CARILER"
        for i,v in enumerate(customers,2): sh.cell(i,1,v)

    if "URUNLER" in wb.sheetnames:
        sh=wb["URUNLER"]
        if sh.max_row>1: sh.delete_rows(2,sh.max_row-1)
        if sh["A1"].value in (None,""): sh["A1"]="URUNLER"
        for i,v in enumerate(cargos,2): sh.cell(i,1,v)

    # YUKLEME LISTESI veri satırları.
    for excel_row, x in enumerate(rows, start=2):
        kg=_safe_float(x.get("net_kg"))
        rate=_safe_float(x.get("freight_rate"))
        freight=_safe_float(x.get("freight_total"))
        collection=_safe_float(x.get("entry_collection"))

        km_go=_safe_float(x.get("planned_km_go"))
        km_back=_safe_float(x.get("planned_km_back"))
        planned_total=km_go+km_back

        exit_km=_safe_float(x.get("exit_km"))
        entry_km=_safe_float(x.get("entry_km"))
        actual_km=max(entry_km-exit_km,0) if entry_km and exit_km else 0

        off_lt=_safe_float(x.get("exit_official_fuel_liters"))
        off_total=_safe_float(x.get("exit_official_fuel_total"))
        off_price=off_total/off_lt if off_lt else 0

        com_lt=_safe_float(x.get("exit_commercial_fuel_liters"))
        com_total=_safe_float(x.get("exit_commercial_fuel_total"))
        com_price=com_total/com_lt if com_lt else 0

        bag_lt=_safe_float(x.get("exit_baghdad_fuel_liters"))
        bag_total=_safe_float(x.get("exit_baghdad_fuel_total"))
        bag_price=bag_total/bag_lt if bag_lt else 0

        road_lt=_safe_float(x.get("road_fuel_liters"))
        road_total=_safe_float(x.get("road_fuel_total"))

        tank_start=_safe_float(x.get("tank_start_liters"))
        tank_end=_safe_float(x.get("tank_end_liters"))
        total_liters=tank_start+off_lt+com_lt+bag_lt+road_lt
        consumed=_safe_float(x.get("fuel_consumed_liters"))
        consumption=_safe_float(x.get("liters_per_100km"))

        premium=_safe_float(x.get("exit_premium"))
        allowance=_safe_float(x.get("exit_allowance"))
        port_fee=_safe_float(x.get("port_fee"))
        dock_fee=_safe_float(x.get("dock_fee"))
        other=_safe_float(x.get("exit_other"))
        sonar=_safe_float(x.get("sonar"))
        fuel_total=off_total+com_total+bag_total+road_total
        gtotal=premium+allowance+port_fee+dock_fee+other+sonar+fuel_total

        loading=x.get("trip_date") or ""
        delivery=x.get("delivery_time") or x.get("entry_at") or ""

        # A:BG positional mapping.
        vals = [None]*59
        vals[0]=x.get("scna")
        vals[1]=loading
        vals[2]=delivery

        # AY / sefer süresi
        try:
            if loading:
                vals[3]=int(str(loading)[5:7])
        except:
            vals[3]=""
        vals[4]=_safe_float(x.get("trip_duration_hours"))/24 if x.get("trip_duration_hours") else ""

        vals[5]=x.get("d_no") or ""
        vals[6]=x.get("driver_name") or ""
        vals[7]=x.get("phone") or ""
        vals[8]=x.get("plate") or ""
        vals[9]=kg
        vals[10]=rate
        vals[11]=x.get("customer_name") or ""
        vals[12]=x.get("cargo_name") or ""
        vals[13]=x.get("cargo_type_name") or ""
        vals[14]=x.get("area_name") or ""
        vals[15]=km_go
        vals[16]=km_back
        vals[17]=planned_total
        vals[18]=planned_total*0.35 if planned_total else 0
        vals[19]=premium
        vals[20]=allowance
        vals[21]=port_fee
        vals[22]=dock_fee

        # X Driver: eski Excel'deki sürücü masraf alanı.
        # Programda ayrı alan olmadığı için 0 bırakılır.
        vals[23]=0
        vals[24]=other
        vals[25]=sonar
        vals[26]=gtotal
        vals[27]=""

        vals[28]=exit_km
        vals[29]=entry_km
        vals[30]=actual_km
        vals[31]=tank_start

        vals[32]=off_lt
        vals[33]=off_price
        vals[34]=off_total
        vals[35]=com_lt
        vals[36]=com_price
        vals[37]=com_total
        vals[38]=bag_lt
        vals[39]=bag_price
        vals[40]=bag_total

        vals[41]=fuel_total
        vals[42]=total_liters
        vals[43]=total_liters-consumed if total_liters else ""
        vals[44]=tank_end
        vals[45]=consumption

        vals[46]=rate
        vals[47]=freight
        vals[48]=collection
        vals[49]=collection-freight
        vals[50]=x.get("entry_at") or ""

        # AZ/BA bakım maliyeti DB'de yok, boş bırak.
        vals[51]=""
        vals[52]=""

        vals[53]=_safe_float(x.get("trip_profit"))
        vals[54]=""
        vals[55]=""
        vals[56]=""
        vals[57]=""
        vals[58]=""

        for col,val in enumerate(vals,1):
            ws.cell(excel_row,col,val)

    # Görsel/format ayarları.
    ws.freeze_panes="A2"
    ws.auto_filter.ref=f"A1:BG{max(ws.max_row,1)}"

    # Tarihler metin olarak gelir; para/sayı kolonlarını düzgün formatla.
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for idx,cell in enumerate(row,1):
            if idx in (10,11,16,17,18,19,20,21,22,23,24,25,26,27,29,30,31,32,
                       33,34,35,36,37,38,39,40,41,42,43,44,45,46,47,48,49,50,
                       52,53,54):
                cell.number_format='#,##0.00'

    return wb


def _request_user(request: Request):
    user=getattr(request.state,"auth_user",None) or _session_user(request.cookies.get(AUTH_COOKIE))
    return user

def _require_user_perm(request: Request,perm:str):
    user=_request_user(request)
    if not user:
        raise HTTPException(401,"Oturum gerekli.")
    if user.get("role")!="ADMIN" and perm not in set(user.get("permissions") or []):
        raise HTTPException(403,"Bu işlem için yetkiniz yok.")
    return user

@app.get("/api/auth/status")
def auth_status(request: Request):
    c=db()
    count=c.execute("SELECT COUNT(*) n FROM auth_users").fetchone()["n"]
    c.close()
    user=_session_user(request.cookies.get(AUTH_COOKIE))
    return {"setup_required":count==0,"authenticated":bool(user),"user":user}

@app.post("/api/auth/setup")
def auth_setup(x: AuthSetupIn):
    c=db()
    if c.execute("SELECT COUNT(*) n FROM auth_users").fetchone()["n"]>0:
        c.close();raise HTTPException(409,"İlk admin zaten oluşturulmuş.")
    username=x.username.strip()
    if len(username)<3 or len(x.password)<8:
        c.close();raise HTTPException(400,"Kullanıcı adı en az 3, şifre en az 8 karakter olmalı.")
    salt,pwh=_password_hash(x.password)
    c.execute("""INSERT INTO auth_users(username,full_name,role,password_salt,password_hash,is_active,must_change_password)
                 VALUES(?,?,?,?,?,1,0)""",(username,x.full_name.strip(),"ADMIN",salt,pwh))
    c.commit();c.close()
    return {"ok":True}

@app.post("/api/auth/login")
def auth_login(x: AuthLoginIn):
    c=db()
    row=c.execute("""SELECT * FROM auth_users WHERE username=? COLLATE NOCASE""",(x.username.strip(),)).fetchone()
    if not row or not int(row["is_active"]) or not _password_ok(x.password,row["password_salt"],row["password_hash"]):
        c.close();raise HTTPException(401,"Kullanıcı adı veya şifre hatalı.")
    token=secrets.token_urlsafe(32)
    expires=(datetime.now()+timedelta(hours=AUTH_SESSION_HOURS)).strftime("%Y-%m-%d %H:%M:%S")
    c.execute("DELETE FROM auth_sessions WHERE datetime(expires_at)<=datetime('now')")
    c.execute("INSERT INTO auth_sessions(token,user_id,expires_at) VALUES(?,?,?)",(token,row["id"],expires))
    c.commit()
    user=dict(row)
    user["permissions"]=sorted(_user_permissions(c,row["id"],row["role"]))
    c.close()
    resp=JSONResponse({"ok":True,"user":{
      "id":user["id"],"username":user["username"],"full_name":user["full_name"],
      "role":user["role"],"must_change_password":user["must_change_password"],
      "driver_id":user.get("driver_id"),"current_plate":user.get("current_plate") or "",
      "permissions":user["permissions"]
    }})
    resp.set_cookie(AUTH_COOKIE,token,max_age=AUTH_SESSION_HOURS*3600,httponly=True,samesite="lax")
    return resp

@app.post("/api/auth/logout")
def auth_logout(request: Request):
    token=request.cookies.get(AUTH_COOKIE)
    if token:
        c=db();c.execute("DELETE FROM auth_sessions WHERE token=?",(token,));c.commit();c.close()
    resp=JSONResponse({"ok":True})
    resp.delete_cookie(AUTH_COOKIE)
    return resp

@app.get("/api/auth/me")
def auth_me(request: Request):
    user=_session_user(request.cookies.get(AUTH_COOKIE))
    if not user:
        return {"authenticated":False}
    return {"authenticated":True,"user":user}

@app.post("/api/auth/change-password")
def auth_change_password(x: AuthPasswordIn, request: Request):
    user=_request_user(request)
    if not user:
        raise HTTPException(401,"Oturum gerekli.")
    if len(x.new_password)<8:
        raise HTTPException(400,"Yeni şifre en az 8 karakter olmalı.")
    c=db()
    row=c.execute("SELECT * FROM auth_users WHERE id=?",(user["id"],)).fetchone()
    if not row or not _password_ok(x.current_password,row["password_salt"],row["password_hash"]):
        c.close();raise HTTPException(400,"Mevcut şifre hatalı.")
    salt,pwh=_password_hash(x.new_password)
    c.execute("""UPDATE auth_users SET password_salt=?,password_hash=?,must_change_password=0,
                 updated_at=CURRENT_TIMESTAMP WHERE id=?""",(salt,pwh,user["id"]))
    c.commit();c.close()
    audit("PASSWORD_CHANGE","AUTH","Kullanıcı kendi şifresini değiştirdi")
    return {"ok":True}

@app.get("/api/auth/permission-catalog")
def auth_permission_catalog(request: Request):
    _require_user_perm(request,"users.manage")
    return {
      "permissions":[{"key":k,"label":v} for k,v in PERMISSION_LABELS.items()],
      "role_defaults":{k:sorted(v) for k,v in ROLE_DEFAULTS.items()}
    }

@app.get("/api/auth/users")
def auth_users(request: Request):
    _require_user_perm(request,"users.manage")
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT id,username,full_name,role,is_active,must_change_password,driver_id,current_plate,created_at,updated_at
      FROM auth_users ORDER BY is_active DESC,username
    """)]
    for x in rows:
        x["permissions"]=sorted(_user_permissions(c,x["id"],x["role"]))
    c.close()
    return rows

@app.post("/api/auth/users")
def auth_create_user(x: AuthUserIn, request: Request):
    _require_user_perm(request,"users.manage")
    role=x.role.strip().upper()
    if role not in ROLE_DEFAULTS:
        raise HTTPException(400,"Geçersiz rol.")
    if len(x.username.strip())<3 or len(x.password)<8:
        raise HTTPException(400,"Kullanıcı adı en az 3, şifre en az 8 karakter olmalı.")
    salt,pwh=_password_hash(x.password)
    c=db()
    try:
        cur=c.execute("""INSERT INTO auth_users(username,full_name,role,password_salt,password_hash,is_active,must_change_password,driver_id,current_plate)
                         VALUES(?,?,?,?,?,?,?,?,?)""",
          (x.username.strip(),x.full_name.strip(),role,salt,pwh,1 if x.is_active else 0,1 if x.must_change_password else 0,x.driver_id,""))
        c.commit()
    except sqlite3.IntegrityError:
        c.close();raise HTTPException(409,"Bu kullanıcı adı zaten mevcut.")
    c.close()
    audit("USER_CREATE","AUTH",f"Kullanıcı oluşturuldu: {x.username.strip()} / {role}","KULLANICI","",f"{x.username.strip()} | {role}")
    return {"ok":True}

@app.patch("/api/auth/users/{user_id}")
def auth_update_user(user_id:int,x:AuthUserUpdateIn,request:Request):
    current=_require_user_perm(request,"users.manage")
    role=x.role.strip().upper()
    if role not in ROLE_DEFAULTS:
        raise HTTPException(400,"Geçersiz rol.")
    if current["id"]==user_id and not x.is_active:
        raise HTTPException(400,"Kendi hesabınızı pasife alamazsınız.")
    c=db()
    c.execute("""UPDATE auth_users SET full_name=?,role=?,is_active=?,must_change_password=?,driver_id=?,
                 current_plate=CASE WHEN COALESCE(driver_id,-1)=COALESCE(?,-1) THEN current_plate ELSE '' END,
                 updated_at=CURRENT_TIMESTAMP WHERE id=?""",
      (x.full_name.strip(),role,1 if x.is_active else 0,1 if x.must_change_password else 0,x.driver_id,x.driver_id,user_id))
    c.commit();c.close()
    audit("USER_UPDATE","AUTH",f"Kullanıcı güncellendi: id={user_id}")
    return {"ok":True}

@app.post("/api/auth/users/{user_id}/reset-password")
def auth_reset_password(user_id:int,x:AuthResetPasswordIn,request:Request):
    _require_user_perm(request,"users.manage")
    if len(x.new_password)<8:
        raise HTTPException(400,"Yeni şifre en az 8 karakter olmalı.")
    salt,pwh=_password_hash(x.new_password)
    c=db();c.execute("""UPDATE auth_users SET password_salt=?,password_hash=?,must_change_password=?,
                        updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                     (salt,pwh,1 if x.must_change_password else 0,user_id))
    c.execute("DELETE FROM auth_sessions WHERE user_id=?",(user_id,))
    c.commit();c.close()
    audit("PASSWORD_RESET","AUTH",f"Admin şifre sıfırladı: id={user_id}","KULLANICI","Eski şifre hash","Yeni şifre hash")
    return {"ok":True}

@app.put("/api/auth/users/{user_id}/permissions")
def auth_set_permissions(user_id:int,x:AuthPermissionsIn,request:Request):
    _require_user_perm(request,"users.manage")
    c=db()
    row=c.execute("SELECT role FROM auth_users WHERE id=?",(user_id,)).fetchone()
    if not row:
        c.close();raise HTTPException(404,"Kullanıcı bulunamadı.")
    role=row["role"].upper()
    defaults=ROLE_DEFAULTS.get(role,set())
    c.execute("DELETE FROM auth_user_permissions WHERE user_id=?",(user_id,))
    for perm in PERMISSION_LABELS:
        desired=bool(x.permissions.get(perm,perm in defaults))
        default=perm in defaults
        if desired!=default:
            c.execute("INSERT INTO auth_user_permissions(user_id,permission,allowed) VALUES(?,?,?)",
                      (user_id,perm,1 if desired else 0))
    c.commit();c.close()
    audit("PERMISSIONS_UPDATE","AUTH",f"Kullanıcı yetkileri değişti: id={user_id}","KULLANICI","","Yetki matrisi güncellendi")
    return {"ok":True}


@app.get("/api/auth/driver-options")
def auth_driver_options(request: Request):
    _require_user_perm(request,"users.manage")
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT id,name,phone,d_no FROM drivers WHERE active=1 ORDER BY name
    """)]
    c.close()
    return rows

@app.get("/api/driver-session")
def driver_session(request: Request):
    user=_request_user(request)
    if not user:
        raise HTTPException(401,"Oturum gerekli.")
    if not user.get("driver_id"):
        return {"linked":False,"driver":None,"current_plate":"","vehicles":[]}
    c=db()
    driver=c.execute("SELECT id,name,phone,d_no FROM drivers WHERE id=?",(user["driver_id"],)).fetchone()
    vehicles=[dict(r) for r in c.execute("""
      SELECT DISTINCT plate FROM (
        SELECT UPPER(TRIM(plate)) plate FROM fleet_vehicles WHERE is_active=1
        UNION
        SELECT UPPER(TRIM(plate)) plate FROM vehicles WHERE active=1
        UNION
        SELECT UPPER(TRIM(plate)) plate FROM trips WHERE TRIM(COALESCE(plate,''))<>''
      ) WHERE TRIM(COALESCE(plate,''))<>'' ORDER BY plate
    """)]
    c.close()
    return {
      "linked":bool(driver),
      "driver":dict(driver) if driver else None,
      "current_plate":user.get("current_plate") or "",
      "vehicles":vehicles
    }

@app.post("/api/driver-session/select-vehicle")
def driver_select_vehicle(x: DriverVehicleSelectIn, request: Request):
    user=_request_user(request)
    if not user:
        raise HTTPException(401,"Oturum gerekli.")
    if not user.get("driver_id"):
        raise HTTPException(400,"Bu kullanıcı bir şoför kaydıyla eşleştirilmemiş.")
    plate=x.plate.strip().upper()
    if not plate:
        raise HTTPException(400,"Plaka seçin.")
    c=db()
    exists=c.execute("""
      SELECT 1 FROM (
        SELECT UPPER(TRIM(plate)) plate FROM fleet_vehicles WHERE is_active=1
        UNION SELECT UPPER(TRIM(plate)) plate FROM vehicles WHERE active=1
        UNION SELECT UPPER(TRIM(plate)) plate FROM trips WHERE TRIM(COALESCE(plate,''))<>''
      ) WHERE plate=? LIMIT 1
    """,(plate,)).fetchone()
    if not exists:
        c.close();raise HTTPException(404,"Aktif araç listesinde bu plaka bulunamadı.")
    c.execute("UPDATE auth_users SET current_plate=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(plate,user["id"]))
    c.commit();c.close()
    audit("DRIVER_VEHICLE","AUTH",f"Aktif araç seçildi: {plate}","ARAC","",plate)
    return {"ok":True,"current_plate":plate}

@app.post("/api/resolve/customer")
def resolve_customer(x: ResolveNameIn):
    name=x.name.strip()
    if not name:
        return {"id":None,"name":""}
    c=db()
    row=c.execute("SELECT id,name FROM customers WHERE UPPER(TRIM(name))=UPPER(TRIM(?))",(name,)).fetchone()
    if not row:
        cur=c.execute("INSERT INTO customers(name,active) VALUES(?,1)",(name,))
        c.commit()
        row=c.execute("SELECT id,name FROM customers WHERE id=?",(cur.lastrowid,)).fetchone()
    out=dict(row);c.close()
    return out

@app.post("/api/resolve/driver")
def resolve_driver(x: ResolveNameIn):
    name=x.name.strip()
    if not name:
        return {"id":None,"name":""}
    c=db()
    row=c.execute("SELECT id,name FROM drivers WHERE UPPER(TRIM(name))=UPPER(TRIM(?))",(name,)).fetchone()
    if not row:
        cur=c.execute("INSERT INTO drivers(name) VALUES(?)",(name,))
        c.commit()
        row=c.execute("SELECT id,name FROM drivers WHERE id=?",(cur.lastrowid,)).fetchone()
    out=dict(row);c.close()
    return out

@app.get("/api/export-excel")
def export_excel(mode: str="all", q: str="", status: str=""):
    try:
        from openpyxl import Workbook
    except ImportError:
        raise HTTPException(
            500,
            "Excel dışa aktarma için openpyxl gerekli. Komut: pip install openpyxl"
        )

    if mode == "all":
        rows=_export_rows("","")
        suffix="TUM_KAYITLAR"
    else:
        rows=_export_rows(q,status)
        suffix="FILTRELI"

    wb=_build_export_workbook(rows)
    bio=BytesIO()
    wb.save(bio)
    bio.seek(0)

    stamp=datetime.now().strftime("%Y%m%d_%H%M")
    filename=f"SAMA_YUKLEME_LISTESI_{suffix}_{stamp}.xlsx"

    headers={
        "Content-Disposition": f'attachment; filename="{filename}"'
    }
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )


@app.get("/api/scna-detail/{scna}")
def scna_detail(scna: str):
    c=db()
    trip=c.execute(tq()+" WHERE t.scna=?",(scna,)).fetchone()
    if not trip:
        c.close()
        raise HTTPException(404,"SCNA bulunamadı")

    fuels=[dict(r) for r in c.execute(
        "SELECT * FROM fuel_purchases WHERE scna=? ORDER BY id",
        (scna,)
    )]

    logs=[dict(r) for r in c.execute(
        "SELECT * FROM audit_log WHERE scna=? ORDER BY id DESC LIMIT 100",
        (scna,)
    )]

    c.close()
    return {
        "trip": dict(trip),
        "fuels": fuels,
        "logs": logs
    }



@app.get("/api/performance/vehicles")
def vehicle_performance():
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT
        t.plate,
        COUNT(*) trip_count,
        SUM(CASE WHEN t.entry_done=1 THEN 1 ELSE 0 END) completed_count,
        ROUND(SUM(t.net_kg)/1000.0,2) total_ton,
        ROUND(SUM(CASE WHEN t.freight_basis='KG'
                       THEN t.net_kg*t.freight_rate
                       ELSE (t.net_kg/1000.0)*t.freight_rate END),2) total_freight,
        ROUND(SUM(
          CASE WHEN t.entry_done=1 THEN MAX(t.entry_km-t.exit_km,0) ELSE 0 END
        ),2) total_km,
        ROUND(AVG(
          CASE
            WHEN t.entry_done=1
             AND (t.entry_km-t.exit_km)>0
             AND (
               t.tank_start_liters+
               t.exit_official_fuel_liters+
               t.exit_commercial_fuel_liters+
               t.exit_baghdad_fuel_liters+
               COALESCE(f.road_liters,0)-
               t.tank_end_liters
             )>0
            THEN (
              (
                t.tank_start_liters+
                t.exit_official_fuel_liters+
                t.exit_commercial_fuel_liters+
                t.exit_baghdad_fuel_liters+
                COALESCE(f.road_liters,0)-
                t.tank_end_liters
              )/(t.entry_km-t.exit_km)
            )*100
          END
        ),2) avg_l100,
        ROUND(SUM(
          (CASE WHEN t.freight_basis='KG'
             THEN t.net_kg*t.freight_rate
             ELSE (t.net_kg/1000.0)*t.freight_rate END)
          -
          (
            t.exit_official_fuel_total+
            t.exit_commercial_fuel_total+
            t.exit_baghdad_fuel_total+
            COALESCE(f.road_total,0)+
            t.exit_allowance+
            t.exit_premium+
            t.dock_fee+
            t.port_fee+
            t.sonar+
            t.exit_other+
            t.entry_extra_expense_1+
            t.entry_extra_expense_2+
            t.entry_extra_expense_3
          )
        ),2) total_profit
      FROM trips t
      LEFT JOIN (
        SELECT scna,SUM(liters) road_liters,SUM(total) road_total
        FROM fuel_purchases GROUP BY scna
      ) f ON f.scna=t.scna
      GROUP BY t.plate
      ORDER BY total_profit DESC
    """)]
    c.close()
    return rows

@app.get("/api/performance/drivers")
def driver_performance():
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT
        COALESCE(d.name,'') driver_name,
        COUNT(*) trip_count,
        SUM(CASE WHEN t.entry_done=1 THEN 1 ELSE 0 END) completed_count,
        ROUND(SUM(t.net_kg)/1000.0,2) total_ton,
        ROUND(SUM(CASE WHEN t.freight_basis='KG'
                       THEN t.net_kg*t.freight_rate
                       ELSE (t.net_kg/1000.0)*t.freight_rate END),2) total_freight,
        ROUND(SUM(
          CASE WHEN t.entry_done=1 THEN MAX(t.entry_km-t.exit_km,0) ELSE 0 END
        ),2) total_km,
        ROUND(AVG(
          CASE
            WHEN t.entry_done=1
             AND (t.entry_km-t.exit_km)>0
             AND (
               t.tank_start_liters+
               t.exit_official_fuel_liters+
               t.exit_commercial_fuel_liters+
               t.exit_baghdad_fuel_liters+
               COALESCE(f.road_liters,0)-
               t.tank_end_liters
             )>0
            THEN (
              (
                t.tank_start_liters+
                t.exit_official_fuel_liters+
                t.exit_commercial_fuel_liters+
                t.exit_baghdad_fuel_liters+
                COALESCE(f.road_liters,0)-
                t.tank_end_liters
              )/(t.entry_km-t.exit_km)
            )*100
          END
        ),2) avg_l100,
        ROUND(SUM(
          CASE WHEN t.entry_done=1 THEN
            COALESCE(
              t.excel_remain,
              CASE
                WHEN COALESCE(t.entry_cash_handed,0)=0
                 AND COALESCE(t.entry_collection,0)>0
                 AND t.entry_done=1
                THEN t.entry_collection -
                  (CASE WHEN t.freight_basis='KG'
                        THEN t.net_kg*t.freight_rate
                        ELSE (t.net_kg/1000.0)*t.freight_rate END)
                ELSE
                  t.entry_cash_handed -
                  (
                    (CASE WHEN t.freight_basis='KG'
                          THEN t.net_kg*t.freight_rate
                          ELSE (t.net_kg/1000.0)*t.freight_rate END) -
                    (
                      COALESCE(f.road_total,0)+
                      t.exit_other+
                      t.entry_extra_expense_1+
                      t.entry_extra_expense_2+
                      t.entry_extra_expense_3
                    )
                  )
              END
            )
          ELSE 0 END
        ),2) cash_diff
      FROM trips t
      LEFT JOIN drivers d ON d.id=t.driver_id
      LEFT JOIN (
        SELECT scna,SUM(liters) road_liters,SUM(total) road_total
        FROM fuel_purchases GROUP BY scna
      ) f ON f.scna=t.scna
      GROUP BY d.id,d.name
      ORDER BY completed_count DESC, trip_count DESC
    """)]
    c.close()
    return rows

@app.get("/api/alerts")
def alerts():
    c=db()

    delayed=[dict(r) for r in c.execute("""
      SELECT
        t.scna,t.plate,d.name driver_name,a.name area_name,t.exit_at,
        ROUND((JULIANDAY('now','localtime')-JULIANDAY(t.exit_at))*24,1) hours_on_road
      FROM trips t
      LEFT JOIN drivers d ON d.id=t.driver_id
      LEFT JOIN areas a ON a.id=t.area_id
      WHERE t.status='Yolda'
        AND t.exit_at IS NOT NULL
        AND (JULIANDAY('now','localtime')-JULIANDAY(t.exit_at))*24 > 48
      ORDER BY hours_on_road DESC
    """)]

    # Completed-trip per vehicle historical average.
    fuel_alerts=[dict(r) for r in c.execute("""
      WITH trip_cons AS (
        SELECT
          t.id,t.scna,t.plate,d.name driver_name,
          CASE
            WHEN t.entry_done=1
             AND (t.entry_km-t.exit_km)>0
             AND (
               t.tank_start_liters+
               t.exit_official_fuel_liters+
               t.exit_commercial_fuel_liters+
               t.exit_baghdad_fuel_liters+
               COALESCE(f.road_liters,0)-
               t.tank_end_liters
             )>0
            THEN (
              (
                t.tank_start_liters+
                t.exit_official_fuel_liters+
                t.exit_commercial_fuel_liters+
                t.exit_baghdad_fuel_liters+
                COALESCE(f.road_liters,0)-
                t.tank_end_liters
              )/(t.entry_km-t.exit_km)
            )*100
          END l100
        FROM trips t
        LEFT JOIN drivers d ON d.id=t.driver_id
        LEFT JOIN (
          SELECT scna,SUM(liters) road_liters
          FROM fuel_purchases GROUP BY scna
        ) f ON f.scna=t.scna
        WHERE t.entry_done=1
      ),
      vehicle_avg AS (
        SELECT plate,AVG(l100) avg_l100,COUNT(l100) n
        FROM trip_cons
        WHERE l100 IS NOT NULL AND l100>0
        GROUP BY plate
      )
      SELECT
        tc.scna,tc.plate,tc.driver_name,
        ROUND(tc.l100,2) trip_l100,
        ROUND(va.avg_l100,2) avg_l100,
        ROUND(((tc.l100/va.avg_l100)-1)*100,1) percent_over
      FROM trip_cons tc
      JOIN vehicle_avg va ON va.plate=tc.plate
      WHERE va.n>=3
        AND tc.l100 > va.avg_l100*1.20
      ORDER BY percent_over DESC
      LIMIT 100
    """)]

    c.close()
    return {
        "delay_threshold_hours":48,
        "fuel_threshold_percent":20,
        "delayed": delayed,
        "fuel_alerts": fuel_alerts
    }


@app.get("/api/live-operations")
def live_operations():
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT t.scna,t.plate,t.status,t.exit_at,t.entry_at,
             d.name driver_name,a.name area_name,
             rs.expected_hours,rs.expected_km,rs.expected_l100,rs.tolerance_percent,
             CASE WHEN t.status='Yolda' AND t.exit_at IS NOT NULL
                  THEN ROUND((JULIANDAY('now','localtime')-JULIANDAY(t.exit_at))*24,1)
                  ELSE 0 END hours_on_road,
             CASE WHEN t.status='Yolda' AND t.exit_at IS NOT NULL
                       AND rs.expected_hours IS NOT NULL
                       AND (JULIANDAY('now','localtime')-JULIANDAY(t.exit_at))*24 > rs.expected_hours
                  THEN 1 ELSE 0 END delayed
      FROM trips t
      LEFT JOIN drivers d ON d.id=t.driver_id
      LEFT JOIN areas a ON a.id=t.area_id
      LEFT JOIN route_standards rs ON rs.area_id=t.area_id
      WHERE t.status IN ('Bekliyor','Yolda')
      ORDER BY delayed DESC,t.exit_at ASC,t.id DESC
      LIMIT 500
    """)]
    s=dict(c.execute("""
      SELECT COUNT(*) total,
             SUM(CASE WHEN status='Yolda' THEN 1 ELSE 0 END) onroad,
             SUM(CASE WHEN status='Bekliyor' THEN 1 ELSE 0 END) waiting,
             SUM(CASE WHEN status='Tamamlandı' THEN 1 ELSE 0 END) completed
      FROM trips
    """).fetchone())
    s["delayed"]=sum(1 for r in rows if r["delayed"])
    s["maintenance_due"]=c.execute("""
      SELECT COUNT(*) c FROM maintenance m
      LEFT JOIN (
        SELECT plate,MAX(CASE WHEN entry_km>exit_km THEN entry_km ELSE exit_km END) current_km
        FROM trips GROUP BY plate
      ) k ON k.plate=m.plate
      WHERE m.next_service_km>0 AND COALESCE(k.current_km,0)>=m.next_service_km-2000
    """).fetchone()["c"]
    c.close()
    return {"summary":s,"rows":rows}

@app.get("/api/route-standards")
def route_standards():
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT rs.*,a.name area_name,a.km_go,a.km_back
      FROM route_standards rs
      LEFT JOIN areas a ON a.id=rs.area_id
      ORDER BY a.name
    """)]
    c.close()
    return rows

@app.post("/api/route-standards")
def save_route_standard(x: RouteStandardIn):
    c=db()
    c.execute("""
      INSERT INTO route_standards(area_id,expected_hours,expected_km,expected_l100,tolerance_percent,updated_at)
      VALUES(?,?,?,?,?,CURRENT_TIMESTAMP)
      ON CONFLICT(area_id) DO UPDATE SET
        expected_hours=excluded.expected_hours,
        expected_km=excluded.expected_km,
        expected_l100=excluded.expected_l100,
        tolerance_percent=excluded.tolerance_percent,
        updated_at=CURRENT_TIMESTAMP
    """,(x.area_id,x.expected_hours,x.expected_km,x.expected_l100,x.tolerance_percent))
    c.commit(); c.close()
    return {"ok":True}

@app.get("/api/maintenance")
def maintenance_list():
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT m.*,
             COALESCE(k.current_km,0) current_km,
             CASE WHEN m.next_service_km<=0 THEN 999999999
                  ELSE m.next_service_km-COALESCE(k.current_km,0) END km_left
      FROM maintenance m
      LEFT JOIN (
        SELECT plate,MAX(CASE WHEN entry_km>exit_km THEN entry_km ELSE exit_km END) current_km
        FROM trips GROUP BY plate
      ) k ON k.plate=m.plate
      ORDER BY km_left ASC,m.plate
    """)]
    c.close()
    return rows

@app.post("/api/maintenance")
def maintenance_save(x: MaintenanceIn):
    c=db()
    plate=x.plate.strip().upper()
    old=c.execute("SELECT id FROM maintenance WHERE plate=? ORDER BY id DESC LIMIT 1",(plate,)).fetchone()
    if old:
        c.execute("""
          UPDATE maintenance SET
            last_service_km=?,next_service_km=?,oil_change_km=?,
            tire_note=?,brake_note=?,engine_note=?,note=?,updated_at=CURRENT_TIMESTAMP
          WHERE id=?
        """,(x.last_service_km,x.next_service_km,x.oil_change_km,
             x.tire_note,x.brake_note,x.engine_note,x.note,old["id"]))
    else:
        c.execute("""
          INSERT INTO maintenance(
            plate,last_service_km,next_service_km,oil_change_km,
            tire_note,brake_note,engine_note,note
          ) VALUES(?,?,?,?,?,?,?,?)
        """,(plate,x.last_service_km,x.next_service_km,x.oil_change_km,
             x.tire_note,x.brake_note,x.engine_note,x.note))
    c.commit(); c.close()
    return {"ok":True}

@app.get("/api/route-fuel-alerts")
def route_fuel_alerts():
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT t.scna,t.plate,d.name driver_name,a.name area_name,
             rs.expected_l100,rs.tolerance_percent,
             ROUND(
               ((t.tank_start_liters+t.exit_official_fuel_liters+t.exit_commercial_fuel_liters+
                 t.exit_baghdad_fuel_liters+COALESCE(f.road_liters,0)-t.tank_end_liters)
                /(t.entry_km-t.exit_km))*100,2
             ) actual_l100
      FROM trips t
      LEFT JOIN drivers d ON d.id=t.driver_id
      LEFT JOIN areas a ON a.id=t.area_id
      LEFT JOIN route_standards rs ON rs.area_id=t.area_id
      LEFT JOIN (
        SELECT scna,SUM(liters) road_liters FROM fuel_purchases GROUP BY scna
      ) f ON f.scna=t.scna
      WHERE t.entry_done=1
        AND rs.expected_l100>0
        AND (t.entry_km-t.exit_km)>0
        AND (((t.tank_start_liters+t.exit_official_fuel_liters+t.exit_commercial_fuel_liters+
               t.exit_baghdad_fuel_liters+COALESCE(f.road_liters,0)-t.tank_end_liters)
              /(t.entry_km-t.exit_km))*100)
            > rs.expected_l100*(1+rs.tolerance_percent/100.0)
      ORDER BY actual_l100 DESC
      LIMIT 200
    """)]
    c.close()
    return rows


def _date_filter_sql(date_from="", date_to=""):
    where = ""
    params = []
    if date_from:
        where += " AND DATE(COALESCE(NULLIF(t.trip_date,''),t.created_at)) >= DATE(?)"
        params.append(date_from)
    if date_to:
        where += " AND DATE(COALESCE(NULLIF(t.trip_date,''),t.created_at)) <= DATE(?)"
        params.append(date_to)
    return where, params

@app.get("/api/finance-summary")
def finance_summary(date_from: str="", date_to: str=""):
    c=db()
    date_sql, params = _date_filter_sql(date_from,date_to)

    row=c.execute(f"""
      SELECT
        COUNT(*) trip_count,
        COALESCE(SUM(t.net_kg),0) total_kg,

        COALESCE(SUM(
          CASE WHEN t.freight_basis='KG'
               THEN t.net_kg*t.freight_rate
               ELSE (t.net_kg/1000.0)*t.freight_rate END
        ),0) total_freight,

        COALESCE(SUM(
          t.exit_official_fuel_total +
          t.exit_commercial_fuel_total +
          t.exit_baghdad_fuel_total +
          COALESCE(f.road_total,0)
        ),0) fuel_cost,

        COALESCE(SUM(
          t.exit_allowance +
          t.exit_premium +
          t.dock_fee +
          t.port_fee +
          t.sonar +
          t.exit_other +
          t.entry_extra_expense_1 +
          t.entry_extra_expense_2 +
          t.entry_extra_expense_3
        ),0) other_cost,

        COALESCE(SUM(
          CASE WHEN t.entry_done=1 THEN MAX(t.entry_km-t.exit_km,0) ELSE 0 END
        ),0) total_km,

        COALESCE(SUM(t.entry_collection),0) collected,

        COALESCE(SUM(
          CASE WHEN t.entry_done=1 THEN
            MAX(
              (CASE WHEN t.freight_basis='KG'
                    THEN t.net_kg*t.freight_rate
                    ELSE (t.net_kg/1000.0)*t.freight_rate END)
              - t.entry_collection,
              0
            )
          ELSE 0 END
        ),0) receivable

      FROM trips t
      LEFT JOIN (
        SELECT scna,SUM(total) road_total
        FROM fuel_purchases
        GROUP BY scna
      ) f ON f.scna=t.scna
      WHERE 1=1 {date_sql}
    """,params).fetchone()

    r=dict(row)

    total_cost=float(r["fuel_cost"] or 0)+float(r["other_cost"] or 0)
    total_profit=float(r["total_freight"] or 0)-total_cost
    total_ton=float(r["total_kg"] or 0)/1000.0
    total_km=float(r["total_km"] or 0)

    r["total_cost"]=total_cost
    r["total_profit"]=total_profit
    r["profit_per_km"]=total_profit/total_km if total_km>0 else 0
    r["profit_per_ton"]=total_profit/total_ton if total_ton>0 else 0
    r["margin_percent"]=(total_profit/float(r["total_freight"]))*100 if float(r["total_freight"] or 0)>0 else 0
    r["total_ton"]=total_ton

    c.close()
    return r

@app.get("/api/finance-by-customer")
def finance_by_customer(date_from: str="", date_to: str=""):
    c=db()
    date_sql, params = _date_filter_sql(date_from,date_to)
    rows=[dict(r) for r in c.execute(f"""
      SELECT
        COALESCE(cu.name,'Tanımsız') name,
        COUNT(*) trip_count,
        ROUND(SUM(t.net_kg)/1000.0,2) total_ton,

        ROUND(SUM(
          CASE WHEN t.freight_basis='KG'
               THEN t.net_kg*t.freight_rate
               ELSE (t.net_kg/1000.0)*t.freight_rate END
        ),2) freight,

        ROUND(SUM(
          t.exit_official_fuel_total +
          t.exit_commercial_fuel_total +
          t.exit_baghdad_fuel_total +
          COALESCE(f.road_total,0) +
          t.exit_allowance +
          t.exit_premium +
          t.dock_fee +
          t.port_fee +
          t.sonar +
          t.exit_other +
          t.entry_extra_expense_1 +
          t.entry_extra_expense_2 +
          t.entry_extra_expense_3
        ),2) cost,

        ROUND(SUM(
          (CASE WHEN t.freight_basis='KG'
                THEN t.net_kg*t.freight_rate
                ELSE (t.net_kg/1000.0)*t.freight_rate END)
          -
          (
            t.exit_official_fuel_total +
            t.exit_commercial_fuel_total +
            t.exit_baghdad_fuel_total +
            COALESCE(f.road_total,0) +
            t.exit_allowance +
            t.exit_premium +
            t.dock_fee +
            t.port_fee +
            t.sonar +
            t.exit_other +
            t.entry_extra_expense_1 +
            t.entry_extra_expense_2 +
            t.entry_extra_expense_3
          )
        ),2) profit

      FROM trips t
      LEFT JOIN customers cu ON cu.id=t.customer_id
      LEFT JOIN (
        SELECT scna,SUM(total) road_total
        FROM fuel_purchases GROUP BY scna
      ) f ON f.scna=t.scna
      WHERE 1=1 {date_sql}
      GROUP BY cu.id,cu.name
      ORDER BY profit DESC
    """,params)]
    c.close()
    return rows

@app.get("/api/finance-by-area")
def finance_by_area(date_from: str="", date_to: str=""):
    c=db()
    date_sql, params = _date_filter_sql(date_from,date_to)
    rows=[dict(r) for r in c.execute(f"""
      SELECT
        COALESCE(a.name,'Tanımsız') name,
        COUNT(*) trip_count,
        ROUND(SUM(t.net_kg)/1000.0,2) total_ton,

        ROUND(SUM(
          CASE WHEN t.freight_basis='KG'
               THEN t.net_kg*t.freight_rate
               ELSE (t.net_kg/1000.0)*t.freight_rate END
        ),2) freight,

        ROUND(SUM(
          CASE WHEN t.entry_done=1 THEN MAX(t.entry_km-t.exit_km,0) ELSE 0 END
        ),2) total_km,

        ROUND(SUM(
          t.exit_official_fuel_total +
          t.exit_commercial_fuel_total +
          t.exit_baghdad_fuel_total +
          COALESCE(f.road_total,0) +
          t.exit_allowance +
          t.exit_premium +
          t.dock_fee +
          t.port_fee +
          t.sonar +
          t.exit_other +
          t.entry_extra_expense_1 +
          t.entry_extra_expense_2 +
          t.entry_extra_expense_3
        ),2) cost,

        ROUND(SUM(
          (CASE WHEN t.freight_basis='KG'
                THEN t.net_kg*t.freight_rate
                ELSE (t.net_kg/1000.0)*t.freight_rate END)
          -
          (
            t.exit_official_fuel_total +
            t.exit_commercial_fuel_total +
            t.exit_baghdad_fuel_total +
            COALESCE(f.road_total,0) +
            t.exit_allowance +
            t.exit_premium +
            t.dock_fee +
            t.port_fee +
            t.sonar +
            t.exit_other +
            t.entry_extra_expense_1 +
            t.entry_extra_expense_2 +
            t.entry_extra_expense_3
          )
        ),2) profit

      FROM trips t
      LEFT JOIN areas a ON a.id=t.area_id
      LEFT JOIN (
        SELECT scna,SUM(total) road_total
        FROM fuel_purchases GROUP BY scna
      ) f ON f.scna=t.scna
      WHERE 1=1 {date_sql}
      GROUP BY a.id,a.name
      ORDER BY profit DESC
    """,params)]
    c.close()
    return rows

@app.get("/api/finance-by-cargo")
def finance_by_cargo(date_from: str="", date_to: str=""):
    c=db()
    date_sql, params = _date_filter_sql(date_from,date_to)
    rows=[dict(r) for r in c.execute(f"""
      SELECT
        COALESCE(cc.name,'Tanımsız') name,
        COUNT(*) trip_count,
        ROUND(SUM(t.net_kg)/1000.0,2) total_ton,

        ROUND(SUM(
          CASE WHEN t.freight_basis='KG'
               THEN t.net_kg*t.freight_rate
               ELSE (t.net_kg/1000.0)*t.freight_rate END
        ),2) freight,

        ROUND(SUM(
          t.exit_official_fuel_total +
          t.exit_commercial_fuel_total +
          t.exit_baghdad_fuel_total +
          COALESCE(f.road_total,0) +
          t.exit_allowance +
          t.exit_premium +
          t.dock_fee +
          t.port_fee +
          t.sonar +
          t.exit_other +
          t.entry_extra_expense_1 +
          t.entry_extra_expense_2 +
          t.entry_extra_expense_3
        ),2) cost,

        ROUND(SUM(
          (CASE WHEN t.freight_basis='KG'
                THEN t.net_kg*t.freight_rate
                ELSE (t.net_kg/1000.0)*t.freight_rate END)
          -
          (
            t.exit_official_fuel_total +
            t.exit_commercial_fuel_total +
            t.exit_baghdad_fuel_total +
            COALESCE(f.road_total,0) +
            t.exit_allowance +
            t.exit_premium +
            t.dock_fee +
            t.port_fee +
            t.sonar +
            t.exit_other +
            t.entry_extra_expense_1 +
            t.entry_extra_expense_2 +
            t.entry_extra_expense_3
          )
        ),2) profit

      FROM trips t
      LEFT JOIN cargo_categories cc ON cc.id=t.cargo_category_id
      LEFT JOIN (
        SELECT scna,SUM(total) road_total
        FROM fuel_purchases GROUP BY scna
      ) f ON f.scna=t.scna
      WHERE 1=1 {date_sql}
      GROUP BY cc.id,cc.name
      ORDER BY profit DESC
    """,params)]
    c.close()
    return rows

@app.get("/api/finance-monthly")
def finance_monthly():
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT
        SUBSTR(COALESCE(NULLIF(t.trip_date,''),t.created_at),1,7) month,
        COUNT(*) trip_count,
        ROUND(SUM(t.net_kg)/1000.0,2) total_ton,
        ROUND(SUM(
          CASE WHEN t.freight_basis='KG'
               THEN t.net_kg*t.freight_rate
               ELSE (t.net_kg/1000.0)*t.freight_rate END
        ),2) freight,
        ROUND(SUM(
          t.exit_official_fuel_total+t.exit_commercial_fuel_total+
          t.exit_baghdad_fuel_total+COALESCE(f.road_total,0)+
          t.exit_allowance+t.exit_premium+t.dock_fee+t.port_fee+t.sonar+
          t.exit_other+t.entry_extra_expense_1+t.entry_extra_expense_2+t.entry_extra_expense_3
        ),2) cost
      FROM trips t
      LEFT JOIN (
        SELECT scna,SUM(total) road_total
        FROM fuel_purchases GROUP BY scna
      ) f ON f.scna=t.scna
      GROUP BY month
      HAVING month IS NOT NULL AND month<>''
      ORDER BY month DESC
      LIMIT 24
    """)]
    for r in rows:
        r["profit"]=float(r["freight"] or 0)-float(r["cost"] or 0)
    c.close()
    return rows


@app.get("/api/vehicle-card/{plate}")
def vehicle_card(plate: str):
    c=db()

    vehicle=c.execute("""
      SELECT
        v.plate,
        d.name driver_name,
        d.phone,
        d.d_no
      FROM vehicles v
      LEFT JOIN trips lt ON lt.id=(
        SELECT t2.id
        FROM trips t2
        WHERE UPPER(t2.plate)=UPPER(v.plate)
        ORDER BY
          CASE WHEN t2.status='Yolda' THEN 0 ELSE 1 END,
          t2.id DESC
        LIMIT 1
      )
      LEFT JOIN drivers d ON d.id=lt.driver_id
      WHERE UPPER(v.plate)=UPPER(?)
      LIMIT 1
    """,(plate,)).fetchone()

    if not vehicle:
        c.close()
        raise HTTPException(404,"Araç bulunamadı")

    summary=c.execute("""
      SELECT
        COUNT(*) trip_count,
        SUM(CASE WHEN t.entry_done=1 THEN 1 ELSE 0 END) completed_count,
        ROUND(SUM(t.net_kg)/1000.0,2) total_ton,
        ROUND(SUM(CASE WHEN t.entry_done=1 THEN MAX(t.entry_km-t.exit_km,0) ELSE 0 END),2) total_km,
        ROUND(SUM(
          CASE WHEN t.freight_basis='KG'
               THEN t.net_kg*t.freight_rate
               ELSE (t.net_kg/1000.0)*t.freight_rate END
        ),2) total_freight,
        ROUND(AVG(
          CASE
            WHEN t.entry_done=1
             AND (t.entry_km-t.exit_km)>0
             AND (
               t.tank_start_liters+t.exit_official_fuel_liters+t.exit_commercial_fuel_liters+
               t.exit_baghdad_fuel_liters+COALESCE(f.road_liters,0)-t.tank_end_liters
             )>0
            THEN (
              (
                t.tank_start_liters+t.exit_official_fuel_liters+t.exit_commercial_fuel_liters+
                t.exit_baghdad_fuel_liters+COALESCE(f.road_liters,0)-t.tank_end_liters
              )/(t.entry_km-t.exit_km)
            )*100
          END
        ),2) avg_l100,
        ROUND(SUM(
          (CASE WHEN t.freight_basis='KG'
                THEN t.net_kg*t.freight_rate
                ELSE (t.net_kg/1000.0)*t.freight_rate END)
          -
          (
            t.exit_official_fuel_total+t.exit_commercial_fuel_total+t.exit_baghdad_fuel_total+
            COALESCE(f.road_total,0)+t.exit_allowance+t.exit_premium+t.dock_fee+t.port_fee+
            t.sonar+t.exit_other+t.entry_extra_expense_1+t.entry_extra_expense_2+t.entry_extra_expense_3
          )
        ),2) total_profit
      FROM trips t
      LEFT JOIN (
        SELECT scna,SUM(liters) road_liters,SUM(total) road_total
        FROM fuel_purchases GROUP BY scna
      ) f ON f.scna=t.scna
      WHERE UPPER(t.plate)=UPPER(?)
    """,(plate,)).fetchone()

    trips=[dict(r) for r in c.execute(tq()+"""
      WHERE UPPER(t.plate)=UPPER(?)
      ORDER BY t.id DESC
      LIMIT 20
    """,(plate,))]

    maint=c.execute("""
      SELECT * FROM maintenance
      WHERE UPPER(plate)=UPPER(?)
      ORDER BY id DESC LIMIT 1
    """,(plate,)).fetchone()

    c.close()

    return {
      "vehicle":dict(vehicle),
      "summary":dict(summary),
      "maintenance":dict(maint) if maint else None,
      "trips":trips
    }

@app.get("/api/driver-card/{driver_name}")
def driver_card(driver_name: str):
    c=db()

    driver=c.execute("""
      SELECT id,name,phone,d_no
      FROM drivers
      WHERE UPPER(name)=UPPER(?)
      LIMIT 1
    """,(driver_name,)).fetchone()

    if not driver:
        c.close()
        raise HTTPException(404,"Şoför bulunamadı")

    summary=c.execute("""
      SELECT
        COUNT(*) trip_count,
        SUM(CASE WHEN t.entry_done=1 THEN 1 ELSE 0 END) completed_count,
        ROUND(SUM(t.net_kg)/1000.0,2) total_ton,
        ROUND(SUM(CASE WHEN t.entry_done=1 THEN MAX(t.entry_km-t.exit_km,0) ELSE 0 END),2) total_km,
        ROUND(AVG(
          CASE
            WHEN t.entry_done=1
             AND (t.entry_km-t.exit_km)>0
             AND (
               t.tank_start_liters+t.exit_official_fuel_liters+t.exit_commercial_fuel_liters+
               t.exit_baghdad_fuel_liters+COALESCE(f.road_liters,0)-t.tank_end_liters
             )>0
            THEN (
              (
                t.tank_start_liters+t.exit_official_fuel_liters+t.exit_commercial_fuel_liters+
                t.exit_baghdad_fuel_liters+COALESCE(f.road_liters,0)-t.tank_end_liters
              )/(t.entry_km-t.exit_km)
            )*100
          END
        ),2) avg_l100,
        ROUND(SUM(
          CASE WHEN t.entry_done=1 THEN
            COALESCE(
              t.excel_remain,
              CASE
                WHEN COALESCE(t.entry_cash_handed,0)=0
                 AND COALESCE(t.entry_collection,0)>0
                 AND t.entry_done=1
                THEN t.entry_collection -
                  (CASE WHEN t.freight_basis='KG'
                        THEN t.net_kg*t.freight_rate
                        ELSE (t.net_kg/1000.0)*t.freight_rate END)
                ELSE
                  t.entry_cash_handed -
                  (
                    (CASE WHEN t.freight_basis='KG'
                          THEN t.net_kg*t.freight_rate
                          ELSE (t.net_kg/1000.0)*t.freight_rate END) -
                    (
                      COALESCE(f.road_total,0)+
                      t.exit_other+
                      t.entry_extra_expense_1+
                      t.entry_extra_expense_2+
                      t.entry_extra_expense_3
                    )
                  )
              END
            )
          ELSE 0 END
        ),2) cash_diff,
        SUM(CASE WHEN t.status='Yolda' THEN 1 ELSE 0 END) active_trips
      FROM trips t
      LEFT JOIN (
        SELECT scna,SUM(liters) road_liters,SUM(total) road_total
        FROM fuel_purchases GROUP BY scna
      ) f ON f.scna=t.scna
      WHERE t.driver_id=?
    """,(driver["id"],)).fetchone()

    trips=[dict(r) for r in c.execute(tq()+"""
      WHERE t.driver_id=?
      ORDER BY t.id DESC
      LIMIT 20
    """,(driver["id"],))]

    plates=[r["plate"] for r in c.execute("""
      SELECT DISTINCT plate
      FROM trips
      WHERE driver_id=?
      ORDER BY plate
    """,(driver["id"],))]

    c.close()

    return {
      "driver":dict(driver),
      "summary":dict(summary),
      "plates":plates,
      "trips":trips
    }

@app.get("/api/anomalies")
def anomalies():
    c=db()

    # 1) Yakıt anomalisi: aracın kendi geçmiş ortalamasından %20 fazla
    fuel=[dict(r) for r in c.execute("""
      WITH trip_cons AS (
        SELECT
          t.id,t.scna,t.plate,d.name driver_name,a.name area_name,
          CASE
            WHEN t.entry_done=1
             AND (t.entry_km-t.exit_km)>0
             AND (
               t.tank_start_liters+t.exit_official_fuel_liters+t.exit_commercial_fuel_liters+
               t.exit_baghdad_fuel_liters+COALESCE(f.road_liters,0)-t.tank_end_liters
             )>0
            THEN (
              (
                t.tank_start_liters+t.exit_official_fuel_liters+t.exit_commercial_fuel_liters+
                t.exit_baghdad_fuel_liters+COALESCE(f.road_liters,0)-t.tank_end_liters
              )/(t.entry_km-t.exit_km)
            )*100
          END l100
        FROM trips t
        LEFT JOIN drivers d ON d.id=t.driver_id
        LEFT JOIN areas a ON a.id=t.area_id
        LEFT JOIN (
          SELECT scna,SUM(liters) road_liters
          FROM fuel_purchases GROUP BY scna
        ) f ON f.scna=t.scna
        WHERE t.entry_done=1
      ),
      avg_cons AS (
        SELECT plate,AVG(l100) avg_l100,COUNT(l100) n
        FROM trip_cons
        WHERE l100 IS NOT NULL AND l100>0
        GROUP BY plate
      )
      SELECT
        tc.scna,tc.plate,tc.driver_name,tc.area_name,
        ROUND(tc.l100,2) value,
        ROUND(ac.avg_l100,2) baseline,
        ROUND(((tc.l100/ac.avg_l100)-1)*100,1) deviation
      FROM trip_cons tc
      JOIN avg_cons ac ON ac.plate=tc.plate
      WHERE ac.n>=3 AND tc.l100>ac.avg_l100*1.20
      ORDER BY deviation DESC
      LIMIT 100
    """)]

    # 2) KM anomalisi
    km=[dict(r) for r in c.execute("""
      SELECT
        t.scna,t.plate,d.name driver_name,a.name area_name,
        ROUND(MAX(t.entry_km-t.exit_km,0),2) value,
        ROUND(rs.expected_km,2) baseline,
        ROUND(((MAX(t.entry_km-t.exit_km,0)/rs.expected_km)-1)*100,1) deviation
      FROM trips t
      LEFT JOIN drivers d ON d.id=t.driver_id
      LEFT JOIN areas a ON a.id=t.area_id
      LEFT JOIN route_standards rs ON rs.area_id=t.area_id
      WHERE t.entry_done=1
        AND rs.expected_km>0
        AND MAX(t.entry_km-t.exit_km,0) > rs.expected_km*(1+rs.tolerance_percent/100.0)
      ORDER BY deviation DESC
      LIMIT 100
    """)]

    # Ortak şoför para farkı mantığı:
    # Excel kökenli kayıtta REMAIN hücresine güvenmiyoruz.
    # COLLECTION varsa fark = COLLECTION - AMOUNT.
    # Yeni program girişinde gerçek teslim edilen para hesabı kullanılır.
    cash=[dict(r) for r in c.execute("""
      WITH cash_calc AS (
        SELECT
          t.scna,t.plate,d.name driver_name,a.name area_name,
          CASE
            WHEN t.excel_remain IS NOT NULL THEN
              CASE
                WHEN COALESCE(t.entry_collection,0)>0 THEN
                  t.entry_collection -
                  (
                    CASE
                      WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
                      ELSE (t.net_kg/1000.0)*t.freight_rate
                    END
                  )
                ELSE 0
              END

            WHEN COALESCE(t.entry_cash_handed,0)=0
             AND COALESCE(t.entry_collection,0)>0
             AND t.entry_done=1 THEN
              t.entry_collection -
              (
                CASE
                  WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
                  ELSE (t.net_kg/1000.0)*t.freight_rate
                END
              )

            ELSE
              t.entry_cash_handed -
                  (
                    (CASE WHEN t.freight_basis='KG'
                          THEN t.net_kg*t.freight_rate
                          ELSE (t.net_kg/1000.0)*t.freight_rate END) -
                    (
                      COALESCE(f.road_total,0)+
                      t.exit_other+
                      t.entry_extra_expense_1+
                      t.entry_extra_expense_2+
                      t.entry_extra_expense_3
                    )
                  )
          END cash_diff,

          CASE
            WHEN t.excel_remain IS NOT NULL THEN 'EXCEL COLLECTION-AMOUNT'
            WHEN COALESCE(t.entry_cash_handed,0)=0
             AND COALESCE(t.entry_collection,0)>0
             AND t.entry_done=1 THEN 'ESKİ EXCEL HESABI'
            ELSE 'PROGRAM HESABI'
          END source

        FROM trips t
        LEFT JOIN drivers d ON d.id=t.driver_id
        LEFT JOIN areas a ON a.id=t.area_id
        LEFT JOIN (
          SELECT scna,SUM(total) road_total
          FROM fuel_purchases
          GROUP BY scna
        ) f ON f.scna=t.scna
        WHERE t.entry_done=1
      )
      SELECT
        scna,plate,driver_name,area_name,
        ROUND(cash_diff,2) value,
        0 baseline,
        0 deviation,
        source
      FROM cash_calc
      WHERE cash_diff < -0.01
      ORDER BY cash_diff ASC
      LIMIT 500
    """)]

    # 4) Son 5 tamamlanmış seferin 3+ tanesinde eksik para
    repeated=[dict(r) for r in c.execute("""
      WITH trip_cash AS (
        SELECT
          t.id,t.driver_id,d.name driver_name,t.scna,
          CASE
            WHEN t.excel_remain IS NOT NULL THEN
              CASE
                WHEN COALESCE(t.entry_collection,0)>0 THEN
                  t.entry_collection -
                  (
                    CASE
                      WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
                      ELSE (t.net_kg/1000.0)*t.freight_rate
                    END
                  )
                ELSE 0
              END

            WHEN COALESCE(t.entry_cash_handed,0)=0
             AND COALESCE(t.entry_collection,0)>0
             AND t.entry_done=1 THEN
              t.entry_collection -
              (
                CASE
                  WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate
                  ELSE (t.net_kg/1000.0)*t.freight_rate
                END
              )

            ELSE
              t.entry_cash_handed -
                  (
                    (CASE WHEN t.freight_basis='KG'
                          THEN t.net_kg*t.freight_rate
                          ELSE (t.net_kg/1000.0)*t.freight_rate END) -
                    (
                      COALESCE(f.road_total,0)+
                      t.exit_other+
                      t.entry_extra_expense_1+
                      t.entry_extra_expense_2+
                      t.entry_extra_expense_3
                    )
                  )
          END cash_diff

        FROM trips t
        LEFT JOIN drivers d ON d.id=t.driver_id
        LEFT JOIN (
          SELECT scna,SUM(total) road_total
          FROM fuel_purchases
          GROUP BY scna
        ) f ON f.scna=t.scna
        WHERE t.entry_done=1
          AND t.driver_id IS NOT NULL
      ),
      ranked AS (
        SELECT *,
          ROW_NUMBER() OVER(PARTITION BY driver_id ORDER BY id DESC) rn
        FROM trip_cash
      )
      SELECT
        driver_name,
        SUM(CASE WHEN cash_diff < -0.01 THEN 1 ELSE 0 END) short_count,
        COUNT(*) checked_count
      FROM ranked
      WHERE rn<=5
      GROUP BY driver_id,driver_name
      HAVING COUNT(*)>=3
         AND SUM(CASE WHEN cash_diff < -0.01 THEN 1 ELSE 0 END)>=3
      ORDER BY short_count DESC,driver_name
    """)]

    # Sayaçlar: limitten bağımsız gerçek toplamı göster.
    cash_count=c.execute("""
      WITH cash_calc AS (
        SELECT
          CASE
            WHEN t.excel_remain IS NOT NULL THEN
              CASE
                WHEN COALESCE(t.entry_collection,0)>0 THEN
                  t.entry_collection -
                  (CASE WHEN t.freight_basis='KG'
                        THEN t.net_kg*t.freight_rate
                        ELSE (t.net_kg/1000.0)*t.freight_rate END)
                ELSE 0
              END
            WHEN COALESCE(t.entry_cash_handed,0)=0
             AND COALESCE(t.entry_collection,0)>0
             AND t.entry_done=1 THEN
              t.entry_collection -
              (CASE WHEN t.freight_basis='KG'
                    THEN t.net_kg*t.freight_rate
                    ELSE (t.net_kg/1000.0)*t.freight_rate END)
            ELSE
              t.entry_cash_handed -
                  (
                    (CASE WHEN t.freight_basis='KG'
                          THEN t.net_kg*t.freight_rate
                          ELSE (t.net_kg/1000.0)*t.freight_rate END) -
                    (
                      COALESCE(f.road_total,0)+
                      t.exit_other+
                      t.entry_extra_expense_1+
                      t.entry_extra_expense_2+
                      t.entry_extra_expense_3
                    )
                  )
          END cash_diff
        FROM trips t
        LEFT JOIN (
          SELECT scna,SUM(total) road_total
          FROM fuel_purchases GROUP BY scna
        ) f ON f.scna=t.scna
        WHERE t.entry_done=1
      )
      SELECT COUNT(*) c
      FROM cash_calc
      WHERE cash_diff < -0.01
    """).fetchone()["c"]

    c.close()
    return {
      "fuel":fuel,
      "km":km,
      "cash":cash,
      "cash_total_count":cash_count,
      "repeated_driver_cash":repeated
    }


@app.get("/api/daily-manager-summary")
def daily_manager_summary(day: str=""):
    if not day:
        day=datetime.now().strftime("%Y-%m-%d")
    c=db()
    exit_row=c.execute("""
      SELECT COUNT(*) exited,
             COALESCE(SUM(net_kg),0)/1000.0 ton,
             COALESCE(SUM(CASE WHEN freight_basis='KG' THEN net_kg*freight_rate ELSE (net_kg/1000.0)*freight_rate END),0) freight
      FROM trips
      WHERE COALESCE(is_deleted,0)=0 AND exit_done=1 AND DATE(exit_at)=DATE(?)
    """,(day,)).fetchone()
    entry_row=c.execute("""
      SELECT COUNT(*) entered,
             COALESCE(SUM(entry_collection),0) collected,
             COALESCE(SUM(entry_extra_expense_1+entry_extra_expense_2+entry_extra_expense_3),0) entry_cost
      FROM trips
      WHERE COALESCE(is_deleted,0)=0 AND entry_done=1 AND DATE(entry_at)=DATE(?)
    """,(day,)).fetchone()
    exit_cost=c.execute("""
      SELECT COALESCE(SUM(
        t.exit_official_fuel_total+t.exit_commercial_fuel_total+t.exit_baghdad_fuel_total+
        COALESCE(f.road_total,0)+t.exit_allowance+t.exit_premium+t.dock_fee+t.port_fee+t.sonar+t.exit_other
      ),0) cost
      FROM trips t
      LEFT JOIN (SELECT scna,SUM(total) road_total FROM fuel_purchases GROUP BY scna) f ON f.scna=t.scna
      WHERE COALESCE(t.is_deleted,0)=0 AND t.exit_done=1 AND DATE(t.exit_at)=DATE(?)
    """,(day,)).fetchone()["cost"]
    total_cost=float(exit_cost or 0)+float(entry_row["entry_cost"] or 0)
    freight=float(exit_row["freight"] or 0)
    onroad=c.execute("SELECT COUNT(*) c FROM trips WHERE COALESCE(is_deleted,0)=0 AND status='Yolda'").fetchone()["c"]
    delayed=c.execute("""
      SELECT COUNT(*) c FROM trips t LEFT JOIN route_standards rs ON rs.area_id=t.area_id
      WHERE COALESCE(t.is_deleted,0)=0 AND t.status='Yolda' AND t.exit_at IS NOT NULL
        AND (JULIANDAY('now','localtime')-JULIANDAY(t.exit_at))*24 > COALESCE(NULLIF(rs.expected_hours,0),48)
    """).fetchone()["c"]
    driver_diff=c.execute("""
      SELECT COALESCE(SUM(
        CASE
          WHEN excel_remain IS NOT NULL AND COALESCE(entry_collection,0)>0 THEN
            entry_collection-(CASE WHEN freight_basis='KG' THEN net_kg*freight_rate ELSE (net_kg/1000.0)*freight_rate END)
          WHEN COALESCE(entry_cash_handed,0)=0 AND COALESCE(entry_collection,0)>0 THEN
            entry_collection-(CASE WHEN freight_basis='KG' THEN net_kg*freight_rate ELSE (net_kg/1000.0)*freight_rate END)
          ELSE entry_cash_handed-((CASE WHEN freight_basis='KG' THEN net_kg*freight_rate ELSE (net_kg/1000.0)*freight_rate END)-
               (exit_other+entry_extra_expense_1+entry_extra_expense_2+entry_extra_expense_3))
        END),0) v
      FROM trips WHERE COALESCE(is_deleted,0)=0 AND entry_done=1 AND DATE(entry_at)=DATE(?)
    """,(day,)).fetchone()["v"]
    rows=[dict(z) for z in c.execute("""
      SELECT t.scna,t.plate,d.name driver_name,a.name area_name,t.status,t.exit_at,t.entry_at,t.net_kg,
             CASE WHEN t.freight_basis='KG' THEN t.net_kg*t.freight_rate ELSE (t.net_kg/1000.0)*t.freight_rate END freight_total
      FROM trips t LEFT JOIN drivers d ON d.id=t.driver_id LEFT JOIN areas a ON a.id=t.area_id
      WHERE COALESCE(t.is_deleted,0)=0 AND (DATE(t.exit_at)=DATE(?) OR DATE(t.entry_at)=DATE(?))
      ORDER BY COALESCE(t.entry_at,t.exit_at) DESC
    """,(day,day))]
    c.close()
    return {"day":day,"exited":int(exit_row["exited"] or 0),"entered":int(entry_row["entered"] or 0),
            "onroad":int(onroad or 0),"delayed":int(delayed or 0),"ton":round(float(exit_row["ton"] or 0),2),
            "freight":round(freight,2),"collected":round(float(entry_row["collected"] or 0),2),
            "cost":round(total_cost,2),"profit":round(freight-total_cost,2),
            "driver_diff":round(float(driver_diff or 0),2),"rows":rows}


def _validate_trip_values(data):
    errors=[]
    warnings=[]

    kg=float(data.get("net_kg") or 0)
    rate=float(data.get("freight_rate") or 0)
    exit_km=float(data.get("exit_km") or 0)
    entry_km=float(data.get("entry_km") or 0)
    tank_start=float(data.get("tank_start_liters") or 0)
    tank_end=float(data.get("tank_end_liters") or 0)
    collection=float(data.get("entry_collection") or 0)

    if kg < 0:
        errors.append("Net KG negatif olamaz.")
    if kg > 80000:
        warnings.append(f"Net KG çok yüksek görünüyor: {kg:,.0f}")

    if rate < 0:
        errors.append("Navlun birim fiyatı negatif olamaz.")
    if rate > 1000000:
        warnings.append(f"Navlun birim fiyatı çok yüksek görünüyor: {rate:,.0f}")

    if exit_km < 0 or entry_km < 0:
        errors.append("KM değerleri negatif olamaz.")
    if entry_km > 0 and exit_km > 0 and entry_km < exit_km:
        errors.append("Giriş KM, çıkış KM'den küçük olamaz.")
    if entry_km > 0 and exit_km > 0 and (entry_km-exit_km) > 5000:
        warnings.append(f"Sefer KM farkı çok yüksek: {entry_km-exit_km:,.0f} KM")

    for label,val in [
        ("Çıkış depo mazotu",tank_start),
        ("Dönüş depo mazotu",tank_end),
    ]:
        if val < 0:
            errors.append(f"{label} negatif olamaz.")
        if val > 2000:
            warnings.append(f"{label} çok yüksek görünüyor: {val:,.0f} LT")

    if collection < 0:
        errors.append("Müşteriden alınan para negatif olamaz.")

    return errors,warnings

@app.get("/api/trips/{scna}/edit-data")
def trip_edit_data(scna: str):
    c=db()
    row=c.execute("""
      SELECT *
      FROM trips
      WHERE UPPER(TRIM(scna))=UPPER(TRIM(?))
    """,(scna,)).fetchone()
    if not row:
        c.close()
        raise HTTPException(404,"SCNA bulunamadı")
    c.close()
    return dict(row)

@app.post("/api/trips/{scna}/validate")
def validate_trip(scna: str, x: TripEditIn):
    data=x.model_dump()
    errors,warnings=_validate_trip_values(data)
    return {"ok":len(errors)==0,"errors":errors,"warnings":warnings}

@app.put("/api/trips/{scna}")
def update_trip(scna: str, x: TripEditIn):
    data=x.model_dump()
    errors,warnings=_validate_trip_values(data)
    if errors:
        raise HTTPException(400,{"errors":errors,"warnings":warnings})

    c=db()
    old=c.execute("SELECT * FROM trips WHERE scna=?",(scna,)).fetchone()
    if not old:
        c.close()
        raise HTTPException(404,"SCNA bulunamadı")
    before=dict(old)

    status=x.status if x.status in ("Bekliyor","Yolda","Tamamlandı") else "Bekliyor"
    exit_done=1 if status in ("Yolda","Tamamlandı") else 0
    entry_done=1 if status=="Tamamlandı" else 0

    c.execute("""
      UPDATE trips SET
        plate=?,
        driver_id=?,
        area_id=?,
        cargo_category_id=?,
        cargo_type=?,
        customer_id=?,
        trip_date=?,
        net_kg=?,
        freight_rate=?,
        freight_basis=?,
        exit_km=?,
        tank_start_liters=?,
        exit_premium=?,
        dock_fee=?,
        port_fee=?,
        sonar=?,
        exit_allowance=?,
        exit_other=?,
        entry_km=?,
        tank_end_liters=?,
        entry_extra_expense_1=?,
        entry_extra_expense_2=?,
        entry_extra_expense_3=?,
        entry_collection=?,
        entry_cash_handed=?,
        excel_remain=NULL,
        entry_note=?,
        manual_lock=1,
        manual_lock_at=DATETIME('now','localtime'),
        manual_lock_reason='SCNA Düzenleme Merkezi',
        status=?,
        exit_done=?,
        entry_done=?,
        exit_at=CASE
          WHEN ?=1 THEN COALESCE(exit_at,DATETIME('now','localtime'))
          ELSE NULL
        END,
        entry_at=CASE
          WHEN ?=1 THEN COALESCE(entry_at,DATETIME('now','localtime'))
          ELSE NULL
        END,
        updated_at=CURRENT_TIMESTAMP
      WHERE scna=?
    """,(
      x.plate.strip().upper(),x.driver_id,x.area_id,x.cargo_category_id,
      x.cargo_type.upper(),x.customer_id,x.trip_date,
      x.net_kg,x.freight_rate,x.freight_basis.upper(),
      x.exit_km,x.tank_start_liters,x.exit_premium,x.dock_fee,x.port_fee,x.sonar,
      x.exit_allowance,x.exit_other,x.entry_km,x.tank_end_liters,
      x.entry_extra_expense_1,x.entry_extra_expense_2,x.entry_extra_expense_3,
      x.entry_collection,x.entry_cash_handed,x.note,status,exit_done,entry_done,
      exit_done,entry_done,scna
    ))
    c.commit()
    after_row=c.execute("SELECT * FROM trips WHERE scna=?",(scna,)).fetchone()
    after=dict(after_row) if after_row else None
    c.close()

    audit_trip_changes(scna,before,after)
    audit("EDIT_SUMMARY",scna,f"SCNA düzenlendi. Uyarı: {' | '.join(warnings) if warnings else 'Yok'}","SEVKIYAT")
    return {"ok":True,"warnings":warnings}


@app.post("/api/trips/{scna}/manual-lock")
def set_manual_lock(scna: str):
    c=db()
    row=c.execute("SELECT id FROM trips WHERE UPPER(scna)=UPPER(?)",(scna,)).fetchone()
    if not row:
        c.close()
        raise HTTPException(404,"SCNA bulunamadı")
    c.execute("""
      UPDATE trips SET manual_lock=1,
        manual_lock_at=DATETIME('now','localtime'),
        manual_lock_reason='Kullanıcı koruması',
        updated_at=CURRENT_TIMESTAMP
      WHERE UPPER(scna)=UPPER(?)
    """,(scna,))
    c.commit(); c.close()
    audit("MANUAL_LOCK",scna,"Excel güncelleme koruması açıldı")
    return {"ok":True,"manual_lock":1}

@app.post("/api/trips/{scna}/manual-unlock")
def clear_manual_lock(scna: str):
    c=db()
    row=c.execute("SELECT id FROM trips WHERE UPPER(scna)=UPPER(?)",(scna,)).fetchone()
    if not row:
        c.close()
        raise HTTPException(404,"SCNA bulunamadı")
    c.execute("""
      UPDATE trips SET manual_lock=0,
        manual_lock_at='',
        manual_lock_reason='',
        updated_at=CURRENT_TIMESTAMP
      WHERE UPPER(scna)=UPPER(?)
    """,(scna,))
    c.commit(); c.close()
    audit("MANUAL_UNLOCK",scna,"Excel güncelleme koruması kaldırıldı")
    return {"ok":True,"manual_lock":0}

@app.post("/api/trips/{scna}/undo-entry")
def undo_entry(scna: str):
    c=db()
    row=c.execute("SELECT * FROM trips WHERE scna=?",(scna,)).fetchone()
    if not row:
        c.close()
        raise HTTPException(404,"SCNA bulunamadı")

    c.execute("""
      UPDATE trips SET
        entry_km=0,
        tank_end_liters=0,
        entry_extra_expense_1=0,
        entry_extra_expense_2=0,
        entry_extra_expense_3=0,
        entry_collection=0,
        entry_cash_handed=0,
        excel_remain=NULL,
        entry_note='',
        entry_done=0,
        entry_at=NULL,
        delivery_time='',
        status=CASE WHEN exit_done=1 THEN 'Yolda' ELSE 'Bekliyor' END,
        updated_at=CURRENT_TIMESTAMP
      WHERE scna=?
    """,(scna,))
    c.commit(); c.close()
    audit("UNDO_ENTRY",scna,"Giriş işlemi geri alındı")
    return {"ok":True}

@app.post("/api/trips/{scna}/undo-exit")
def undo_exit(scna: str):
    c=db()
    row=c.execute("SELECT * FROM trips WHERE scna=?",(scna,)).fetchone()
    if not row:
        c.close()
        raise HTTPException(404,"SCNA bulunamadı")

    c.execute("DELETE FROM fuel_purchases WHERE scna=?",(scna,))
    c.execute("""
      UPDATE trips SET
        exit_km=0,
        exit_cash=0,
        tank_start_liters=0,
        exit_official_fuel_liters=0,
        exit_official_fuel_total=0,
        exit_commercial_fuel_liters=0,
        exit_commercial_fuel_total=0,
        exit_baghdad_fuel_liters=0,
        exit_baghdad_fuel_total=0,
        exit_allowance=0,
        exit_other=0,
        exit_done=0,
        exit_at=NULL,

        entry_km=0,
        tank_end_liters=0,
        entry_extra_expense_1=0,
        entry_extra_expense_2=0,
        entry_extra_expense_3=0,
        entry_collection=0,
        entry_cash_handed=0,
        excel_remain=NULL,
        entry_note='',
        entry_done=0,
        entry_at=NULL,
        delivery_time='',

        status='Bekliyor',
        updated_at=CURRENT_TIMESTAMP
      WHERE scna=?
    """,(scna,))
    c.commit(); c.close()
    audit("UNDO_EXIT",scna,"Çıkış ve buna bağlı giriş/yol mazotu işlemleri geri alındı")
    return {"ok":True}


@app.post("/api/export-import-errors")
async def export_import_errors(request: Request):
    try:
        from openpyxl import Workbook
    except ImportError:
        raise HTTPException(500,"openpyxl gerekli")

    payload=await request.json()
    rows=payload.get("failed_records",[]) if isinstance(payload,dict) else []

    wb=Workbook()
    ws=wb.active
    ws.title="HATALI KAYITLAR"
    ws.append(["SATIR","SCNA","HATA NEDENİ"])

    for x in rows:
        ws.append([
            x.get("row",""),
            x.get("scna",""),
            x.get("reason","")
        ])

    ws.column_dimensions["A"].width=12
    ws.column_dimensions["B"].width=18
    ws.column_dimensions["C"].width=80

    bio=BytesIO()
    wb.save(bio)
    bio.seek(0)

    stamp=datetime.now().strftime("%Y%m%d_%H%M%S")
    return StreamingResponse(
        bio,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":f'attachment; filename="EXCEL_AKTARIM_HATALARI_{stamp}.xlsx"'}
    )


@app.delete("/api/trips/{scna}")
def soft_delete_trip(scna: str, request: Request):
    user=_request_user(request)
    if not user or user.get("role")!="ADMIN":
        raise HTTPException(403,"Sadece ADMIN sevkiyat silebilir.")
    c=db()
    row=c.execute("SELECT id FROM trips WHERE UPPER(TRIM(scna))=UPPER(TRIM(?)) AND COALESCE(is_deleted,0)=0",(scna,)).fetchone()
    if not row:
        c.close(); raise HTTPException(404,"Aktif sevkiyat bulunamadı.")
    c.execute("""UPDATE trips SET is_deleted=1,deleted_at=DATETIME('now','localtime'),deleted_by=?,updated_at=CURRENT_TIMESTAMP
                 WHERE id=?""",(user.get("username") or "ADMIN",row["id"]))
    c.commit();c.close()
    audit("SOFT_DELETE",scna,"Sevkiyat 7 gün geri alınabilir şekilde silindi","SEVKIYAT")
    return {"ok":True,"recoverable_days":7}

@app.get("/api/trips-deleted")
def deleted_trips(request: Request):
    user=_request_user(request)
    if not user or user.get("role")!="ADMIN":
        raise HTTPException(403,"Sadece ADMIN görebilir.")
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT scna,plate,deleted_at,deleted_by,
             CAST(7-(JULIANDAY('now','localtime')-JULIANDAY(deleted_at)) AS INTEGER) days_left
      FROM trips
      WHERE COALESCE(is_deleted,0)=1
      ORDER BY deleted_at DESC
      LIMIT 500
    """)]
    c.close(); return rows

@app.post("/api/trips/{scna}/restore")
def restore_deleted_trip(scna: str, request: Request):
    user=_request_user(request)
    if not user or user.get("role")!="ADMIN":
        raise HTTPException(403,"Sadece ADMIN geri alabilir.")
    c=db()
    row=c.execute("""SELECT id,deleted_at FROM trips
                     WHERE UPPER(TRIM(scna))=UPPER(TRIM(?)) AND COALESCE(is_deleted,0)=1""",(scna,)).fetchone()
    if not row:
        c.close(); raise HTTPException(404,"Silinen sevkiyat bulunamadı.")
    age=c.execute("SELECT JULIANDAY('now','localtime')-JULIANDAY(?) d",(row["deleted_at"],)).fetchone()["d"]
    if age is None or float(age)>7:
        c.close(); raise HTTPException(410,"7 günlük geri alma süresi dolmuş.")
    c.execute("UPDATE trips SET is_deleted=0,deleted_at='',deleted_by='',updated_at=CURRENT_TIMESTAMP WHERE id=?",(row["id"],))
    c.commit();c.close()
    audit("RESTORE",scna,"Silinen sevkiyat geri alındı","SEVKIYAT")
    return {"ok":True}

@app.get("/api/data-quality")
def data_quality():
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT
        t.scna,t.plate,d.name driver_name,cu.name customer_name,a.name area_name,
        cc.name cargo_name,t.status,t.trip_date,t.exit_at,t.entry_at,
        t.exit_km,t.entry_km,t.net_kg,t.freight_rate,t.entry_collection,
        CASE WHEN (t.entry_at IS NOT NULL AND t.entry_at<>'' AND t.entry_km<=0) THEN 1 ELSE 0 END entry_date_no_km,
        CASE WHEN (t.exit_km>0 AND (t.exit_at IS NULL OR t.exit_at='')) THEN 1 ELSE 0 END exit_km_no_date,
        CASE WHEN ((t.entry_at IS NOT NULL AND t.entry_at<>'') AND t.entry_done=0) THEN 1 ELSE 0 END entry_date_status_mismatch,
        CASE WHEN (t.freight_rate>0 AND t.net_kg>0 AND COALESCE(t.entry_collection,0)=0) THEN 1 ELSE 0 END collection_missing,
        CASE WHEN t.driver_id IS NULL THEN 1 ELSE 0 END driver_missing,
        CASE WHEN t.customer_id IS NULL THEN 1 ELSE 0 END customer_missing,
        CASE WHEN t.area_id IS NULL THEN 1 ELSE 0 END area_missing,
        CASE WHEN t.cargo_category_id IS NULL THEN 1 ELSE 0 END cargo_missing
      FROM trips t
      LEFT JOIN drivers d ON d.id=t.driver_id
      LEFT JOIN customers cu ON cu.id=t.customer_id
      LEFT JOIN areas a ON a.id=t.area_id
      LEFT JOIN cargo_categories cc ON cc.id=t.cargo_category_id
      ORDER BY t.id DESC
    """)]
    problems=[]
    counts={k:0 for k in [
      "entry_date_no_km","exit_km_no_date","entry_date_status_mismatch",
      "collection_missing","driver_missing","customer_missing","area_missing","cargo_missing"
    ]}
    for x in rows:
        msgs=[]
        mapping=[
          ("entry_date_no_km","Giriş tarihi var, giriş KM yok"),
          ("exit_km_no_date","Çıkış KM var, çıkış tarihi yok"),
          ("entry_date_status_mismatch","Giriş tarihi var ama kayıt tamamlanmış değil"),
          ("collection_missing","Navlun var, COLLECTION boş/0"),
          ("driver_missing","Şoför boş"),
          ("customer_missing","Müşteri boş"),
          ("area_missing","Bölge boş"),
          ("cargo_missing","Mal cinsi boş"),
        ]
        for key,msg in mapping:
            if x[key]:
                counts[key]+=1
                msgs.append(msg)
        if msgs:
            x["problems"]=msgs
            problems.append(x)
    c.close()
    return {"counts":counts,"total":len(problems),"rows":problems[:2000]}

@app.post("/api/bulk-fix")
def bulk_fix(x: BulkFixIn):
    allowed={"status","customer_id","area_id","cargo_category_id","driver_id","cargo_type","trip_date"}
    if x.field not in allowed:
        raise HTTPException(400,"Bu alan toplu düzeltmeye açık değil.")
    if not x.scnas:
        raise HTTPException(400,"SCNA seçilmedi.")
    c=db()
    done=0
    failed=[]
    for scna in x.scnas:
        try:
            val=x.value
            if x.field in ("customer_id","area_id","cargo_category_id","driver_id"):
                val=int(val) if str(val).strip() else None
            if x.field=="status":
                status=val if val in ("Bekliyor","Yolda","Tamamlandı") else "Bekliyor"
                exit_done=1 if status in ("Yolda","Tamamlandı") else 0
                entry_done=1 if status=="Tamamlandı" else 0
                c.execute("""
                  UPDATE trips SET status=?,exit_done=?,entry_done=?,
                    manual_lock=1,manual_lock_at=DATETIME('now','localtime'),manual_lock_reason='Toplu durum düzeltme',
                    exit_at=CASE WHEN ?=1 THEN COALESCE(exit_at,DATETIME('now','localtime')) ELSE NULL END,
                    entry_at=CASE WHEN ?=1 THEN COALESCE(entry_at,DATETIME('now','localtime')) ELSE NULL END,
                    updated_at=CURRENT_TIMESTAMP
                  WHERE scna=?
                """,(status,exit_done,entry_done,exit_done,entry_done,scna))
            else:
                c.execute(
                    f"UPDATE trips SET {x.field}=?,manual_lock=1,manual_lock_at=DATETIME('now','localtime'),manual_lock_reason='Toplu düzeltme',updated_at=CURRENT_TIMESTAMP WHERE scna=?",
                    (val,scna)
                )
            done+=1
        except Exception as e:
            failed.append({"scna":scna,"reason":str(e)})
    c.commit(); c.close()
    audit("BULK_FIX","TOPLU",f"{x.field} alanı {done} kayıtta güncellendi")
    return {"ok":True,"updated":done,"failed":failed}

@app.post("/api/compare-excel-db")
async def compare_excel_db(request: Request):
    raw=await request.body()
    if not raw:
        raise HTTPException(400,"Excel verisi boş.")
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(raw),data_only=False,read_only=True)
    sheet=None
    for cand in ("YUKLEME LISTESI","YÜKLEME LİSTESİ","YUKLEME LİSTESİ"):
        if cand in wb.sheetnames:
            sheet=wb[cand]; break
    if sheet is None:
        sheet=wb[wb.sheetnames[0]]

    c=db()
    db_rows={str(r["scna"]).strip().upper():dict(r) for r in c.execute(tq())}
    c.close()

    excel_rows={}
    compare_rows=[]
    def sval(v): return str(v or "").strip()

    for rno,row in enumerate(sheet.iter_rows(min_row=2,values_only=True),start=2):
        if not row: continue
        scna=sval(row[0]).upper() if len(row)>0 else ""
        if not scna or scna=="SCNA": continue
        e={
          "row":rno,"scna":scna,
          "delivery_time":_excel_datetime(row[2] if len(row)>2 else ""),
          "driver":sval(row[6] if len(row)>6 else ""),
          "plate":sval(row[8] if len(row)>8 else "").upper(),
          "kg":_excel_num(row[9] if len(row)>9 else 0),
          "price_k":_excel_num(row[10] if len(row)>10 else 0),
          "freight_rate":_excel_num(row[46] if len(row)>46 else 0),
          "amount":_excel_num(row[47] if len(row)>47 else 0),
          "customer":sval(row[11] if len(row)>11 else ""),
          "cargo":sval(row[12] if len(row)>12 else ""),
          "cargo_type":sval(row[13] if len(row)>13 else "").upper(),
          "area":sval(row[14] if len(row)>14 else ""),
          "exit_km":_excel_num(row[28] if len(row)>28 else 0),
          "entry_km":_excel_num(row[29] if len(row)>29 else 0),
          "collection":_excel_num(row[48] if len(row)>48 else 0),
        }
        excel_rows[scna]=e
        d=db_rows.get(scna)
        if not d:
            compare_rows.append({"scna":scna,"row":rno,"result":"EXCEL_VAR_DB_YOK","differences":["Excel'de var, veritabanında yok"]})
            continue
        diffs=[]
        for label,a,b in [
          ("Plaka",e["plate"],sval(d.get("plate")).upper()),
          ("Şoför",e["driver"],sval(d.get("driver_name"))),
          ("Müşteri",e["customer"],sval(d.get("customer_name"))),
          ("Bölge",e["area"],sval(d.get("area_name"))),
          ("Mal",e["cargo"],sval(d.get("cargo_name"))),
          ("Yük Tipi",e["cargo_type"],sval(d.get("cargo_type_name")).upper())
        ]:
            if a.strip().upper()!=b.strip().upper():
                diffs.append(f"{label}: Excel='{a}' / DB='{b}'")
        for label,a,b in [
          ("KG",e["kg"],float(d.get("net_kg") or 0)),
          ("FREIGHT AU",e["freight_rate"] or (e["amount"]/e["kg"] if e["kg"] else 0),float(d.get("freight_rate") or 0)),
          ("Çıkış KM",e["exit_km"],float(d.get("exit_km") or 0)),
          ("Giriş KM",e["entry_km"],float(d.get("entry_km") or 0)),
          ("COLLECTION",e["collection"],float(d.get("entry_collection") or 0))
        ]:
            if abs(float(a or 0)-float(b or 0))>0.01:
                diffs.append(f"{label}: Excel={a} / DB={b}")
        if bool(e["delivery_time"])!=bool(sval(d.get("entry_at"))):
            diffs.append(f"Giriş tarihi: Excel='{e['delivery_time']}' / DB='{sval(d.get('entry_at'))}'")
        if diffs:
            compare_rows.append({"scna":scna,"row":rno,"result":"FARKLI","differences":diffs})

    for scna in db_rows:
        if scna not in excel_rows:
            compare_rows.append({"scna":scna,"row":"","result":"DB_VAR_EXCEL_YOK","differences":["Veritabanında var, Excel'de yok"]})

    counts={
      "excel_only":sum(x["result"]=="EXCEL_VAR_DB_YOK" for x in compare_rows),
      "db_only":sum(x["result"]=="DB_VAR_EXCEL_YOK" for x in compare_rows),
      "different":sum(x["result"]=="FARKLI" for x in compare_rows)
    }
    return {"sheet":sheet.title,"counts":counts,"total":len(compare_rows),"rows":compare_rows[:3000]}


@app.get("/api/operation-center")
def operation_center():
    c=db()

    # Aktif operasyonlar: rota standardı ve geçmiş gerçek rota ortalamaları birlikte
    rows=[dict(r) for r in c.execute("""
      WITH road AS (
        SELECT scna,SUM(liters) road_liters
        FROM fuel_purchases GROUP BY scna
      ),
      route_hist AS (
        SELECT
          t.area_id,
          COUNT(*) completed_trips,
          AVG(CASE WHEN t.entry_done=1 AND t.entry_km>t.exit_km
                   THEN t.entry_km-t.exit_km END) avg_km,
          AVG(CASE WHEN t.entry_done=1 AND t.exit_at IS NOT NULL AND t.entry_at IS NOT NULL
                   THEN (JULIANDAY(t.entry_at)-JULIANDAY(t.exit_at))*24 END) avg_hours,
          AVG(CASE WHEN t.entry_done=1 AND t.entry_km>t.exit_km AND
                        (t.tank_start_liters+t.exit_official_fuel_liters+
                         t.exit_commercial_fuel_liters+t.exit_baghdad_fuel_liters+
                         COALESCE(r.road_liters,0)-t.tank_end_liters)>0
                   THEN (
                     (t.tank_start_liters+t.exit_official_fuel_liters+
                      t.exit_commercial_fuel_liters+t.exit_baghdad_fuel_liters+
                      COALESCE(r.road_liters,0)-t.tank_end_liters)
                     /(t.entry_km-t.exit_km)
                   )*100 END) avg_l100
        FROM trips t
        LEFT JOIN road r ON r.scna=t.scna
        WHERE t.entry_done=1
        GROUP BY t.area_id
      )
      SELECT
        t.scna,t.plate,t.status,t.exit_at,t.trip_date,t.net_kg,
        d.name driver_name,a.name area_name,
        rs.expected_hours,rs.expected_km,rs.expected_l100,rs.tolerance_percent,
        rh.completed_trips,
        ROUND(rh.avg_km,1) route_avg_km,
        ROUND(rh.avg_hours,1) route_avg_hours,
        ROUND(rh.avg_l100,2) route_avg_l100,
        CASE WHEN t.status='Yolda' AND t.exit_at IS NOT NULL
             THEN ROUND((JULIANDAY('now','localtime')-JULIANDAY(t.exit_at))*24,1)
             ELSE 0 END hours_on_road,
        CASE WHEN t.status='Yolda' AND t.exit_at IS NOT NULL AND
                  COALESCE(NULLIF(rh.avg_hours,0),rs.expected_hours,48)>0 AND
                  (JULIANDAY('now','localtime')-JULIANDAY(t.exit_at))*24 >
                  COALESCE(NULLIF(rh.avg_hours,0),rs.expected_hours,48)
             THEN 1 ELSE 0 END delayed
      FROM trips t
      LEFT JOIN drivers d ON d.id=t.driver_id
      LEFT JOIN areas a ON a.id=t.area_id
      LEFT JOIN route_standards rs ON rs.area_id=t.area_id
      LEFT JOIN route_hist rh ON rh.area_id=t.area_id
      WHERE t.status IN ('Bekliyor','Yolda')
      ORDER BY delayed DESC,t.exit_at ASC,t.id DESC
      LIMIT 500
    """)]

    summary=dict(c.execute("""
      SELECT
        COUNT(*) total,
        SUM(CASE WHEN status='Yolda' THEN 1 ELSE 0 END) onroad,
        SUM(CASE WHEN status='Bekliyor' THEN 1 ELSE 0 END) waiting,
        SUM(CASE WHEN DATE(exit_at)=DATE('now','localtime') THEN 1 ELSE 0 END) today_exit,
        SUM(CASE WHEN DATE(entry_at)=DATE('now','localtime') THEN 1 ELSE 0 END) today_entry
      FROM trips
    """).fetchone())
    summary["delayed"]=sum(1 for x in rows if x["delayed"])

    c.close()
    return {"summary":summary,"rows":rows}

@app.get("/api/route-learning")
def route_learning():
    c=db()
    rows=[dict(r) for r in c.execute("""
      WITH road AS (
        SELECT scna,SUM(liters) road_liters,SUM(total) road_total
        FROM fuel_purchases GROUP BY scna
      )
      SELECT
        COALESCE(a.name,'Tanımsız') area_name,
        COUNT(*) trip_count,
        ROUND(AVG(CASE WHEN t.entry_done=1 AND t.entry_km>t.exit_km
                       THEN t.entry_km-t.exit_km END),1) avg_km,
        ROUND(AVG(CASE WHEN t.entry_done=1 AND t.exit_at IS NOT NULL AND t.entry_at IS NOT NULL
                       THEN (JULIANDAY(t.entry_at)-JULIANDAY(t.exit_at))*24 END),1) avg_hours,
        ROUND(AVG(CASE WHEN t.entry_done=1 AND t.entry_km>t.exit_km AND
                            (t.tank_start_liters+t.exit_official_fuel_liters+
                             t.exit_commercial_fuel_liters+t.exit_baghdad_fuel_liters+
                             COALESCE(r.road_liters,0)-t.tank_end_liters)>0
                       THEN (
                         (t.tank_start_liters+t.exit_official_fuel_liters+
                          t.exit_commercial_fuel_liters+t.exit_baghdad_fuel_liters+
                          COALESCE(r.road_liters,0)-t.tank_end_liters)
                         /(t.entry_km-t.exit_km)
                       )*100 END),2) avg_l100,
        ROUND(AVG(
          t.exit_official_fuel_total+t.exit_commercial_fuel_total+t.exit_baghdad_fuel_total+
          COALESCE(r.road_total,0)+t.exit_allowance+t.exit_premium+t.dock_fee+t.port_fee+
          t.sonar+t.exit_other+t.entry_extra_expense_1+t.entry_extra_expense_2+t.entry_extra_expense_3
        ),2) avg_cost
      FROM trips t
      LEFT JOIN areas a ON a.id=t.area_id
      LEFT JOIN road r ON r.scna=t.scna
      WHERE t.entry_done=1
      GROUP BY t.area_id,a.name
      HAVING COUNT(*)>=1
      ORDER BY trip_count DESC,a.name
    """)]
    c.close()
    return rows

@app.get("/api/operation-performance")
def operation_performance():
    c=db()

    vehicles=[dict(r) for r in c.execute("""
      WITH road AS (
        SELECT scna,SUM(liters) road_liters,SUM(total) road_total
        FROM fuel_purchases GROUP BY scna
      )
      SELECT
        t.plate name,
        COUNT(*) trip_count,
        ROUND(SUM(CASE WHEN t.entry_done=1 THEN MAX(t.entry_km-t.exit_km,0) ELSE 0 END),1) total_km,
        ROUND(AVG(CASE WHEN t.entry_done=1 AND t.entry_km>t.exit_km AND
                            (t.tank_start_liters+t.exit_official_fuel_liters+
                             t.exit_commercial_fuel_liters+t.exit_baghdad_fuel_liters+
                             COALESCE(r.road_liters,0)-t.tank_end_liters)>0
                       THEN ((t.tank_start_liters+t.exit_official_fuel_liters+
                              t.exit_commercial_fuel_liters+t.exit_baghdad_fuel_liters+
                              COALESCE(r.road_liters,0)-t.tank_end_liters)
                             /(t.entry_km-t.exit_km))*100 END),2) avg_l100,
        ROUND(AVG(CASE WHEN t.entry_done=1 AND t.exit_at IS NOT NULL AND t.entry_at IS NOT NULL
                       THEN (JULIANDAY(t.entry_at)-JULIANDAY(t.exit_at))*24 END),1) avg_hours
      FROM trips t
      LEFT JOIN road r ON r.scna=t.scna
      GROUP BY t.plate
      ORDER BY trip_count DESC
      LIMIT 30
    """)]

    drivers=[dict(r) for r in c.execute("""
      WITH road AS (
        SELECT scna,SUM(liters) road_liters
        FROM fuel_purchases GROUP BY scna
      )
      SELECT
        COALESCE(d.name,'Tanımsız') name,
        COUNT(*) trip_count,
        ROUND(SUM(CASE WHEN t.entry_done=1 THEN MAX(t.entry_km-t.exit_km,0) ELSE 0 END),1) total_km,
        ROUND(AVG(CASE WHEN t.entry_done=1 AND t.entry_km>t.exit_km AND
                            (t.tank_start_liters+t.exit_official_fuel_liters+
                             t.exit_commercial_fuel_liters+t.exit_baghdad_fuel_liters+
                             COALESCE(r.road_liters,0)-t.tank_end_liters)>0
                       THEN ((t.tank_start_liters+t.exit_official_fuel_liters+
                              t.exit_commercial_fuel_liters+t.exit_baghdad_fuel_liters+
                              COALESCE(r.road_liters,0)-t.tank_end_liters)
                             /(t.entry_km-t.exit_km))*100 END),2) avg_l100,
        ROUND(AVG(CASE WHEN t.entry_done=1 AND t.exit_at IS NOT NULL AND t.entry_at IS NOT NULL
                       THEN (JULIANDAY(t.entry_at)-JULIANDAY(t.exit_at))*24 END),1) avg_hours
      FROM trips t
      LEFT JOIN drivers d ON d.id=t.driver_id
      LEFT JOIN road r ON r.scna=t.scna
      GROUP BY t.driver_id,d.name
      ORDER BY trip_count DESC
      LIMIT 30
    """)]

    c.close()
    return {"vehicles":vehicles,"drivers":drivers}


@app.post("/api/repair-remain-from-excel")
async def repair_remain_from_excel(request: Request):
    raw=await request.body()
    if not raw:
        raise HTTPException(400,"Excel verisi boş.")

    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(500,"openpyxl gerekli")

    wb=load_workbook(BytesIO(raw),data_only=True,read_only=True)

    ws=None
    for name in wb.sheetnames:
        if _norm_header(name) in ("YUKLEMELISTESI","YUKLEMELIST","LOADINGLIST"):
            ws=wb[name]
            break
    if ws is None:
        ws=wb[wb.sheetnames[0]]

    row_iter=ws.iter_rows(values_only=True)
    header_row=None
    header_map={}

    for rno,vals in enumerate(row_iter,start=1):
        vals=tuple(vals)
        hm={_norm_header(v):i for i,v in enumerate(vals) if v is not None}
        if "SCNA" in hm or ("PLATE" in hm and ("DRIVER" in hm or "DRIVERNAME" in hm)):
            header_row=rno
            header_map=hm
            break
        if rno>=30:
            break

    if not header_row:
        raise HTTPException(400,"Başlık satırı bulunamadı.")

    scna_col=header_map.get("SCNA")
    if scna_col is None:
        scna_col=header_map.get("SNCA",0)

    remain_col=None
    for candidate in ("REMAIN","KALAN","BALANCE","REMAINING","REMAINDER"):
        key=_norm_header(candidate)
        if key in header_map:
            remain_col=header_map[key]
            break

    # Başlık gerçekten bulunmadıysa eski format fallback.
    if remain_col is None:
        remain_col=49

    c=db()
    updated=0
    blank=0
    db_not_found=0
    parsed=0
    samples=[]

    for rno,row in enumerate(row_iter,start=header_row+1):
        row=tuple(row)
        if scna_col>=len(row):
            continue

        scna=str(row[scna_col] or "").strip().upper()
        if not scna:
            continue

        raw_remain=row[remain_col] if remain_col<len(row) else None
        remain=_excel_optional_num(raw_remain)

        if remain is None:
            blank+=1
            continue

        parsed+=1
        cur=c.execute(
            "UPDATE trips SET excel_remain=?,updated_at=CURRENT_TIMESTAMP WHERE UPPER(scna)=UPPER(?)",
            (remain,scna)
        )
        if cur.rowcount:
            updated+=1
            if len(samples)<25:
                samples.append({"scna":scna,"remain":remain,"row":rno})
        else:
            db_not_found+=1

    c.commit()
    c.close()

    audit("REMAIN_REPAIR","EXCEL",f"{updated} kayıt REMAIN ile güncellendi")

    return {
        "ok":True,
        "sheet":ws.title,
        "header_row":header_row,
        "remain_header_found": any(_norm_header(x) in header_map for x in ("REMAIN","KALAN","BALANCE","REMAINING","REMAINDER")),
        "remain_column_index":remain_col+1,
        "parsed":parsed,
        "updated":updated,
        "blank":blank,
        "db_not_found":db_not_found,
        "samples":samples
    }


@app.get("/api/cash-diagnostic/{scna}")
def cash_diagnostic(scna: str):
    c=db()
    row=c.execute("""
      SELECT
        t.scna,t.plate,t.entry_done,t.freight_basis,t.net_kg,t.freight_rate,
        t.entry_collection,t.entry_cash_handed,t.excel_price_k,t.excel_freight_au,t.excel_amount,t.excel_remain raw_excel_remain,
        CASE WHEN t.freight_basis='KG'
             THEN t.net_kg*t.freight_rate
             ELSE (t.net_kg/1000.0)*t.freight_rate END freight_total,
        CASE
          WHEN t.excel_remain IS NOT NULL THEN
            CASE WHEN COALESCE(t.entry_collection,0)>0
                 THEN t.entry_collection -
                   (CASE WHEN t.freight_basis='KG'
                         THEN t.net_kg*t.freight_rate
                         ELSE (t.net_kg/1000.0)*t.freight_rate END)
                 ELSE 0 END
          WHEN COALESCE(t.entry_cash_handed,0)=0 AND COALESCE(t.entry_collection,0)>0 AND t.entry_done=1
          THEN t.entry_collection -
            (CASE WHEN t.freight_basis='KG'
                  THEN t.net_kg*t.freight_rate
                  ELSE (t.net_kg/1000.0)*t.freight_rate END)
          ELSE t.entry_cash_handed -
            (
              (CASE WHEN t.freight_basis='KG'
                    THEN t.net_kg*t.freight_rate
                    ELSE (t.net_kg/1000.0)*t.freight_rate END) -
              (
                COALESCE((SELECT SUM(total) FROM fuel_purchases f WHERE f.scna=t.scna),0)+
                t.entry_extra_expense_1+t.entry_extra_expense_2+t.entry_extra_expense_3
              )
            )
        END final_driver_cash_diff,
        CASE
          WHEN t.excel_remain IS NOT NULL THEN 'EXCEL COLLECTION-AMOUNT'
          WHEN COALESCE(t.entry_cash_handed,0)=0 AND COALESCE(t.entry_collection,0)>0 AND t.entry_done=1
          THEN 'ESKİ EXCEL HESABI'
          ELSE 'PROGRAM HESABI'
        END calculation_source
      FROM trips t
      WHERE UPPER(t.scna)=UPPER(?)
      LIMIT 1
    """,(scna,)).fetchone()
    c.close()
    if not row:
        raise HTTPException(404,"SCNA bulunamadı")
    return dict(row)


@app.post("/api/repair-freight-from-excel")
async def repair_freight_from_excel(request: Request):
    raw=await request.body()
    if not raw:
        raise HTTPException(400,"Excel verisi boş.")

    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(raw),data_only=True,read_only=True)

    ws=None
    for name in wb.sheetnames:
        if _norm_header(name) in ("YUKLEMELISTESI","YUKLEMELIST","LOADINGLIST"):
            ws=wb[name]
            break
    if ws is None:
        ws=wb[wb.sheetnames[0]]

    it=ws.iter_rows(values_only=True)
    header_row=0
    scna_col=0

    for rno,row in enumerate(it,start=1):
        hm={_norm_header(v):i for i,v in enumerate(row) if v is not None}
        if "SCNA" in hm:
            header_row=rno
            scna_col=hm["SCNA"]
            break
        if rno>=30:
            break

    if not header_row:
        raise HTTPException(400,"SCNA başlığı bulunamadı.")

    c=db()
    updated=0
    unchanged=0
    missing=0
    db_not_found=0
    samples=[]

    for rno,row in enumerate(it,start=header_row+1):
        row=tuple(row)
        if scna_col>=len(row):
            continue
        scna=str(row[scna_col] or "").strip().upper()
        if not scna:
            continue

        kg=_excel_num(row[9] if len(row)>9 else 0)
        freight_au=_excel_num(row[46] if len(row)>46 else 0)
        amount=_excel_num(row[47] if len(row)>47 else 0)

        rate=freight_au or (amount/kg if amount and kg else 0)
        if not rate:
            missing+=1
            continue

        oldrow=c.execute("SELECT freight_rate FROM trips WHERE UPPER(scna)=UPPER(?)",(scna,)).fetchone()
        if not oldrow:
            db_not_found+=1
            continue

        old_rate=float(oldrow["freight_rate"] or 0)
        if abs(old_rate-rate)<=0.000001:
            unchanged+=1
            continue

        c.execute(
            "UPDATE trips SET freight_rate=?,freight_basis='KG',updated_at=CURRENT_TIMESTAMP WHERE UPPER(scna)=UPPER(?)",
            (rate,scna)
        )
        updated+=1

        if len(samples)<25:
            samples.append({
                "scna":scna,
                "old_rate":old_rate,
                "new_rate":rate,
                "amount":amount
            })

    c.commit(); c.close()
    audit("FREIGHT_REPAIR","EXCEL",f"{updated} kayıt AU/FREIGHT ile güncellendi")

    return {
        "ok":True,"sheet":ws.title,"updated":updated,"unchanged":unchanged,
        "missing":missing,"db_not_found":db_not_found,"samples":samples
    }


@app.post("/api/check-entry-dates-from-excel")
async def check_entry_dates_from_excel(request: Request):
    raw=await request.body()
    if not raw:
        raise HTTPException(400,"Excel verisi boş.")

    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(raw),data_only=True,read_only=True)

    ws=None
    for name in wb.sheetnames:
        if _norm_header(name) in ("YUKLEMELISTESI","YUKLEMELIST","LOADINGLIST"):
            ws=wb[name]
            break
    if ws is None:
        ws=wb[wb.sheetnames[0]]

    it=ws.iter_rows(values_only=True)
    header_row=0
    scna_col=0
    for rno,row in enumerate(it,start=1):
        hm={_norm_header(v):i for i,v in enumerate(row) if v is not None}
        if "SCNA" in hm:
            header_row=rno
            scna_col=hm["SCNA"]
            break
        if rno>=30:
            break
    if not header_row:
        raise HTTPException(400,"SCNA başlığı bulunamadı.")

    c=db()
    excel_with_date=0
    matched=0
    missing_in_db=[]
    date_mismatch=[]

    for rno,row in enumerate(it,start=header_row+1):
        row=tuple(row)
        if scna_col>=len(row):
            continue
        raw_scna=row[scna_col]
        if raw_scna is None:
            continue
        scna=str(raw_scna).strip().upper()
        if scna in ("", ".", "-", "—", "NAN", "NONE", "NULL"):
            continue
        if scna.endswith(".0"):
            try:
                scna=str(int(float(scna)))
            except:
                pass
        excel_date=_excel_datetime(row[2] if len(row)>2 else "")
        if not excel_date:
            continue

        excel_with_date+=1
        dbrow=c.execute(
            "SELECT entry_at,delivery_time,status,manual_lock FROM trips WHERE UPPER(scna)=UPPER(?)",
            (scna,)
        ).fetchone()

        if not dbrow:
            if len(missing_in_db)<200:
                missing_in_db.append({"row":rno,"scna":scna,"excel_date":excel_date})
            continue

        dbdate=dbrow["entry_at"] or dbrow["delivery_time"] or ""
        if not dbdate:
            if len(date_mismatch)<500:
                date_mismatch.append({
                    "row":rno,"scna":scna,"excel_date":excel_date,
                    "db_date":"","status":dbrow["status"],
                    "manual_lock":int(dbrow["manual_lock"] or 0)
                })
        else:
            matched+=1

    c.close()
    return {
        "ok":True,
        "excel_with_entry_date":excel_with_date,
        "db_with_entry_date":matched,
        "missing_date_count":len(date_mismatch),
        "missing_db_count":len(missing_in_db),
        "missing_dates":date_mismatch[:200],
        "missing_db":missing_in_db[:100]
    }


@app.post("/api/repair-entry-dates-from-excel")
async def repair_entry_dates_from_excel(request: Request):
    raw=await request.body()
    if not raw:
        raise HTTPException(400,"Excel verisi boş.")

    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(raw),data_only=True,read_only=True)

    ws=None
    for name in wb.sheetnames:
        if _norm_header(name) in ("YUKLEMELISTESI","YUKLEMELIST","LOADINGLIST"):
            ws=wb[name]
            break
    if ws is None:
        ws=wb[wb.sheetnames[0]]

    it=ws.iter_rows(values_only=True)
    header_row=0
    scna_col=0

    for rno,row in enumerate(it,start=1):
        hm={_norm_header(v):i for i,v in enumerate(row) if v is not None}
        if "SCNA" in hm:
            header_row=rno
            scna_col=hm["SCNA"]
            break
        if rno>=30:
            break

    if not header_row:
        raise HTTPException(400,"SCNA başlığı bulunamadı.")

    c=db()
    excel_dates=0
    already_ok=0
    repaired=0
    db_missing=0
    no_date=0
    repaired_rows=[]

    for rno,row in enumerate(it,start=header_row+1):
        row=tuple(row)
        if scna_col>=len(row):
            continue

        raw_scna=row[scna_col]
        if raw_scna is None:
            continue
        scna=str(raw_scna).strip().upper()
        if scna in ("", ".", "-", "—", "NAN", "NONE", "NULL"):
            continue
        if scna.endswith(".0"):
            try:
                scna=str(int(float(scna)))
            except:
                pass

        delivery=_excel_datetime(row[2] if len(row)>2 else "")
        if not delivery:
            no_date+=1
            continue

        excel_dates+=1
        entry_km=_excel_num(row[29] if len(row)>29 else 0)

        dbrow=c.execute("""
          SELECT id,entry_at,delivery_time,entry_km,status,manual_lock
          FROM trips
          WHERE UPPER(scna)=UPPER(?)
          LIMIT 1
        """,(scna,)).fetchone()

        if not dbrow:
            db_missing+=1
            continue

        current_date=(dbrow["entry_at"] or dbrow["delivery_time"] or "").strip()

        if current_date:
            already_ok+=1
            continue

        # Sadece eksik giriş bilgilerini doldur.
        # Plaka, şoför, müşteri, para, navlun, mazot vb. hiçbir alana dokunmaz.
        c.execute("""
          UPDATE trips SET
            delivery_time=?,
            entry_at=?,
            entry_km=CASE
              WHEN COALESCE(entry_km,0)<=0 AND ?>0 THEN ?
              ELSE entry_km
            END,
            entry_done=1,
            exit_done=1,
            status='Tamamlandı',
            updated_at=CURRENT_TIMESTAMP
          WHERE id=?
        """,(delivery,delivery,entry_km,entry_km,dbrow["id"]))

        repaired+=1
        if len(repaired_rows)<300:
            repaired_rows.append({
                "row":rno,
                "scna":scna,
                "entry_date":delivery,
                "entry_km":entry_km,
                "was_manual_lock":int(dbrow["manual_lock"] or 0)
            })

    c.commit()

    # Repair sonrası ikinci doğrulama.
    verify_missing=[]
    verify_it=ws.iter_rows(values_only=True)
    for _ in range(header_row):
        try:
            next(verify_it)
        except StopIteration:
            break

    for rno,row in enumerate(verify_it,start=header_row+1):
        row=tuple(row)
        if scna_col>=len(row):
            continue
        raw_scna=row[scna_col]
        if raw_scna is None:
            continue
        scna=str(raw_scna).strip().upper()
        if scna in ("", ".", "-", "—", "NAN", "NONE", "NULL"):
            continue
        if scna.endswith(".0"):
            try:
                scna=str(int(float(scna)))
            except:
                pass
        if not scna:
            continue
        delivery=_excel_datetime(row[2] if len(row)>2 else "")
        if not delivery:
            continue

        chk=c.execute("""
          SELECT entry_at,delivery_time
          FROM trips
          WHERE UPPER(scna)=UPPER(?)
          LIMIT 1
        """,(scna,)).fetchone()

        if chk and not (chk["entry_at"] or chk["delivery_time"]):
            if len(verify_missing)<300:
                verify_missing.append({
                    "row":rno,
                    "scna":scna,
                    "excel_date":delivery
                })

    c.close()
    audit("ENTRY_DATE_REPAIR","EXCEL",f"{repaired} eksik giriş tarihi onarıldı")

    return {
        "ok":True,
        "sheet":ws.title,
        "excel_dates":excel_dates,
        "already_ok":already_ok,
        "repaired":repaired,
        "db_missing":db_missing,
        "no_date":no_date,
        "verify_missing_count":len(verify_missing),
        "verify_missing":verify_missing,
        "repaired_rows":repaired_rows
    }


def _looks_like_excel_bytes(content: bytes, content_type: str=""):
    if not content:
        return False
    # xlsx/xlsm are ZIP containers and normally begin PK.
    if content[:2] == b"PK":
        return True
    ct=(content_type or "").lower()
    return (
        "spreadsheetml" in ct or
        "application/vnd.ms-excel" in ct or
        "application/octet-stream" in ct
    ) and not content.lstrip().lower().startswith(b"<!doctype html")

def _extract_onedrive_download_candidates(text: str):
    if not text:
        return []
    decoded=html_lib.unescape(text)

    # JSON escaped URLs
    decoded=decoded.replace("\\u0026","&").replace("\\/","/")

    patterns=[
        r'https://[^"\']+(?:download|Download)[^"\']*',
        r'https://[^"\']+/_layouts/15/download\.aspx[^"\']*',
        r'https://[^"\']+\.files\.1drv\.com/[^"\']+',
        r'https://[^"\']+\.sharepoint\.com/[^"\']+(?:download|Download)[^"\']*',
    ]

    out=[]
    seen=set()
    for p in patterns:
        for m in re.findall(p,decoded,re.I):
            u=m.rstrip("\\")
            if u not in seen:
                seen.add(u)
                out.append(u)
    return out[:20]

def _public_onedrive_download(share_url: str):
    share_url=(share_url or "").strip()
    if not share_url.startswith(("https://1drv.ms/","https://onedrive.live.com/")):
        raise HTTPException(400,"Geçerli bir OneDrive paylaşım linki gir.")

    session=requests.Session()
    session.headers.update({
        "User-Agent":"Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/142 Safari/537.36",
        "Accept":"*/*"
    })

    attempts=[]

    # Microsoft sometimes honors download=1 directly on sharing links.
    variants=[]
    sep="&" if "?" in share_url else "?"
    variants.append(share_url + sep + "download=1")
    variants.append(share_url)

    for url in variants:
        try:
            r=session.get(url,allow_redirects=True,timeout=45)
            attempts.append({
                "url":url,
                "status":r.status_code,
                "final_url":r.url,
                "content_type":r.headers.get("content-type",""),
                "size":len(r.content)
            })

            if r.ok and _looks_like_excel_bytes(r.content,r.headers.get("content-type","")):
                return r.content,{
                    "method":"DIRECT_SHARE_LINK",
                    "final_url":r.url,
                    "attempts":attempts
                }

            # If Microsoft returns a viewer HTML page, inspect it for actual download URLs.
            ct=(r.headers.get("content-type") or "").lower()
            if r.ok and ("text/html" in ct or r.content.lstrip().lower().startswith(b"<!doctype html")):
                text=r.text
                candidates=_extract_onedrive_download_candidates(text)

                # Also look through redirection history URLs.
                for h in list(r.history)+[r]:
                    if h.url and ("download" in h.url.lower() or "files.1drv.com" in h.url.lower()):
                        candidates.insert(0,h.url)

                seen=set()
                for candidate in candidates:
                    if candidate in seen:
                        continue
                    seen.add(candidate)
                    try:
                        dr=session.get(candidate,allow_redirects=True,timeout=60)
                        attempts.append({
                            "url":"HTML_CANDIDATE",
                            "status":dr.status_code,
                            "final_url":dr.url,
                            "content_type":dr.headers.get("content-type",""),
                            "size":len(dr.content)
                        })
                        if dr.ok and _looks_like_excel_bytes(dr.content,dr.headers.get("content-type","")):
                            return dr.content,{
                                "method":"HTML_DOWNLOAD_URL",
                                "final_url":dr.url,
                                "attempts":attempts
                            }
                    except Exception as e:
                        attempts.append({"url":"HTML_CANDIDATE","error":str(e)[:200]})

        except Exception as e:
            attempts.append({"url":url,"error":str(e)[:300]})

    raise HTTPException(
        400,
        {
            "message":"OneDrive paylaşım linkinden Excel dosyası doğrudan indirilemedi.",
            "hint":"Dosyanın paylaşım ayarında 'Bağlantıya sahip herkes görüntüleyebilir' açık olmalı. Microsoft bazı kişisel OneDrive bağlantılarında anonim doğrudan indirmeyi engelleyebilir.",
            "attempts":attempts[-10:]
        }
    )

@app.post("/api/public-onedrive/preview")
def public_onedrive_preview(payload: dict):
    share_url=str(payload.get("share_url") or "").strip()
    content,info=_public_onedrive_download(share_url)

    # Excel'i gerçekten açarak doğrula.
    try:
        from openpyxl import load_workbook
        wb=load_workbook(BytesIO(content),data_only=True,read_only=True)
        sheets=wb.sheetnames
        sheet=None
        for name in sheets:
            if _norm_header(name) in ("YUKLEMELISTESI","YUKLEMELIST","LOADINGLIST"):
                sheet=name
                break
        if sheet is None:
            sheet=sheets[0] if sheets else ""
    except Exception as e:
        raise HTTPException(400,f"Bağlantıdan veri geldi ama geçerli Excel olarak açılamadı: {e}")

    return {
        "ok":True,
        "size":len(content),
        "sheet":sheet,
        "sheets":sheets,
        "method":info.get("method"),
        "final_url":info.get("final_url"),
        "attempts":info.get("attempts",[])
    }

@app.post("/api/public-onedrive/file")
def public_onedrive_file(payload: dict):
    share_url=str(payload.get("share_url") or "").strip()
    content,info=_public_onedrive_download(share_url)
    return StreamingResponse(
        BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":'attachment; filename="WEB_EXCEL.xlsx"',
            "X-OneDrive-Method":str(info.get("method") or "")
        }
    )



@app.get("/api/fleet/vehicles")
def fleet_vehicles():
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT fv.*,
        (
          SELECT d.name FROM trips t
          LEFT JOIN drivers d ON d.id=t.driver_id
          WHERE UPPER(TRIM(t.plate))=UPPER(TRIM(fv.plate))
          ORDER BY t.id DESC LIMIT 1
        ) current_driver
      FROM fleet_vehicles fv
      ORDER BY fv.is_active DESC,fv.plate
    """)]
    c.close()
    return {
      "total":len(rows),
      "active":sum(1 for x in rows if x["is_active"]),
      "passive":sum(1 for x in rows if not x["is_active"]),
      "rows":rows
    }

@app.post("/api/fleet/vehicles")
def save_fleet_vehicle(x: FleetVehicleIn):
    plate=x.plate.strip().upper()
    if not plate: raise HTTPException(400,"Plaka gerekli.")
    c=db()
    c.execute("""
      INSERT INTO fleet_vehicles(plate,brand,model,vehicle_type,is_active,garage_state,note,updated_at)
      VALUES(?,?,?,?,?,?,?,CURRENT_TIMESTAMP)
      ON CONFLICT(plate) DO UPDATE SET
        brand=excluded.brand,model=excluded.model,vehicle_type=excluded.vehicle_type,
        is_active=excluded.is_active,garage_state=excluded.garage_state,note=excluded.note,
        updated_at=CURRENT_TIMESTAMP
    """,(plate,x.brand.strip(),x.model.strip(),x.vehicle_type.strip().upper(),
         1 if x.is_active else 0,x.garage_state.strip().upper(),x.note.strip()))
    c.commit();c.close()
    audit("FLEET_VEHICLE",plate,"Araç kaydı güncellendi")
    return {"ok":True}

@app.patch("/api/fleet/vehicles/{plate}/active")
def set_fleet_vehicle_active(plate: str, x: dict):
    active=1 if x.get("is_active") else 0
    c=db()
    c.execute("UPDATE fleet_vehicles SET is_active=?,updated_at=CURRENT_TIMESTAMP WHERE UPPER(plate)=UPPER(?)",(active,plate))
    c.commit();c.close()
    audit("FLEET_VEHICLE_ACTIVE",plate,f"Aktif={active}")
    return {"ok":True}

@app.get("/api/fleet/drivers")
def fleet_drivers():
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT fd.*,
        (
          SELECT t.plate FROM trips t
          LEFT JOIN drivers d ON d.id=t.driver_id
          WHERE UPPER(TRIM(d.name))=UPPER(TRIM(fd.name))
          ORDER BY t.id DESC LIMIT 1
        ) last_plate
      FROM fleet_drivers fd
      ORDER BY fd.is_active DESC,fd.name
    """)]
    c.close()
    return {
      "total":len(rows),
      "active":sum(1 for x in rows if x["is_active"]),
      "passive":sum(1 for x in rows if not x["is_active"]),
      "rows":rows
    }

@app.post("/api/fleet/drivers")
def save_fleet_driver(x: FleetDriverIn):
    name=x.name.strip()
    if not name: raise HTTPException(400,"Şoför adı gerekli.")
    c=db()
    row=c.execute("SELECT id FROM fleet_drivers WHERE UPPER(TRIM(name))=UPPER(TRIM(?))",(name,)).fetchone()
    if row:
        c.execute("""UPDATE fleet_drivers SET phone=?,d_no=?,is_active=?,note=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                  (x.phone.strip(),x.d_no.strip(),1 if x.is_active else 0,x.note.strip(),row["id"]))
    else:
        c.execute("""INSERT INTO fleet_drivers(name,phone,d_no,is_active,note) VALUES(?,?,?,?,?)""",
                  (name,x.phone.strip(),x.d_no.strip(),1 if x.is_active else 0,x.note.strip()))
    c.commit();c.close()
    audit("FLEET_DRIVER",name,"Şoför kaydı güncellendi")
    return {"ok":True}

@app.patch("/api/fleet/drivers/{driver_id}/active")
def set_fleet_driver_active(driver_id: int, x: dict):
    active=1 if x.get("is_active") else 0
    c=db();c.execute("UPDATE fleet_drivers SET is_active=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(active,driver_id))
    c.commit();c.close()
    return {"ok":True}

from driver_bulk import register_driver_bulk
register_driver_bulk(app, db, _request_user, audit)

@app.get("/api/fleet-operations")
def fleet_operations():
    c=db()
    rows=[dict(r) for r in c.execute("""
      WITH plates AS (
        SELECT UPPER(TRIM(plate)) plate FROM fleet_vehicles WHERE is_active=1
      ),
      last_trip AS (
        SELECT t.* FROM trips t JOIN (
          SELECT UPPER(TRIM(plate)) plate,MAX(id) max_id FROM trips
          WHERE TRIM(COALESCE(plate,''))<>'' GROUP BY UPPER(TRIM(plate))
        ) z ON z.max_id=t.id
      )
      SELECT p.plate,
        COALESCE(v.load_state,CASE WHEN COALESCE(NULLIF(lt.entry_at,''),NULLIF(lt.delivery_time,'')) IS NULL AND lt.exit_done=1 THEN 'DOLU' ELSE 'BOS' END) load_state,
        COALESCE(v.queue_no,0) queue_no,COALESCE(v.vessel,'') vessel,COALESCE(v.operation_note,'') operation_note,
        COALESCE(v.vessel_state,'') vessel_state,
        lt.scna,COALESCE(d.name,'') driver,COALESCE(a.name,'') area,lt.exit_at,
        COALESCE(NULLIF(lt.entry_at,''),NULLIF(lt.delivery_time,'')) entry_at
      FROM plates p
      LEFT JOIN last_trip lt ON UPPER(TRIM(lt.plate))=p.plate
      LEFT JOIN drivers d ON d.id=lt.driver_id LEFT JOIN areas a ON a.id=lt.area_id
      LEFT JOIN vehicle_operations v ON UPPER(TRIM(v.plate))=p.plate
      ORDER BY CASE COALESCE(v.load_state,'') WHEN 'SIRA' THEN 1 WHEN 'YUKLEMEDE' THEN 2 WHEN 'DOLU' THEN 3 WHEN 'BOS' THEN 4 ELSE 5 END,
        CASE WHEN COALESCE(v.queue_no,0)>0 THEN v.queue_no ELSE 999999 END,p.plate
    """)]
    c.close()
    counts={"TOPLAM":len(rows),"BOS":0,"DOLU":0,"SIRA":0,"YUKLEMEDE":0}
    for x in rows:
        s=x.get("load_state") or "BOS";counts[s]=counts.get(s,0)+1
    return {"counts":counts,"rows":rows}

@app.post("/api/fleet-operations")
def save_fleet_operation(x: VehicleOperationIn):
    plate=x.plate.strip().upper();state=x.load_state.strip().upper()
    if not plate: raise HTTPException(400,"Plaka gerekli.")
    if state not in ("BOS","DOLU","SIRA","YUKLEMEDE"): raise HTTPException(400,"Geçersiz durum.")
    q=max(int(x.queue_no or 0),0)
    if state!="SIRA":q=0
    c=db();c.execute("""INSERT INTO vehicle_operations(plate,load_state,queue_no,vessel,operation_note,updated_at)
      VALUES(?,?,?,?,?,CURRENT_TIMESTAMP) ON CONFLICT(plate) DO UPDATE SET
      load_state=excluded.load_state,queue_no=excluded.queue_no,vessel=excluded.vessel,
      operation_note=excluded.operation_note,updated_at=CURRENT_TIMESTAMP""",
      (plate,state,q,x.vessel.strip(),x.operation_note.strip()))
    c.commit();c.close();audit("FLEET_OPERATION",plate,f"{state} / sıra {q} / gemi {x.vessel.strip()}")
    return {"ok":True}


@app.post("/api/vessel-operations")
def save_vessel_operation(x: VehicleOperationIn):
    plate=x.plate.strip().upper()
    state=x.load_state.strip().upper()
    if not plate:
        raise HTTPException(400,"Plaka gerekli.")
    if state not in ("CALISIYOR","CIKTI"):
        raise HTTPException(400,"Geçersiz gemi operasyon durumu.")
    c=db()
    c.execute("""
      INSERT INTO vehicle_operations(plate,load_state,queue_no,vessel,operation_note,vessel_state,updated_at)
      VALUES(?, 'BOS',0,?,?,?,CURRENT_TIMESTAMP)
      ON CONFLICT(plate) DO UPDATE SET
        vessel=excluded.vessel,
        operation_note=excluded.operation_note,
        vessel_state=excluded.vessel_state,
        updated_at=CURRENT_TIMESTAMP
    """,(plate,x.vessel.strip(),x.operation_note.strip(),state))
    c.commit();c.close()
    audit("VESSEL_OPERATION",plate,f"{state} / gemi {x.vessel.strip()}")
    return {"ok":True}

@app.post("/api/fleet/gps-state")
async def set_fleet_gps_state(request: Request):
    user=_request_user(request)
    if not user or user.get("role") not in ("ADMIN","OPERATOR"):
        raise HTTPException(403,"Bu işlem için ADMIN veya OPERATOR yetkisi gerekli.")
    x=await request.json()
    plate=str(x.get("plate") or "").strip().upper()
    state=str(x.get("state") or "").strip().upper()
    if not plate:
        raise HTTPException(400,"Plaka gerekli.")
    if state not in ("DONUYOR",""):
        raise HTTPException(400,"Geçersiz GPS durumu.")
    c=db()
    c.execute("""INSERT INTO vehicle_operations(plate,load_state,queue_no,vessel,operation_note,vessel_state,gps_state,updated_at)
                 VALUES(?, 'BOS',0,'','','',?,CURRENT_TIMESTAMP)
                 ON CONFLICT(plate) DO UPDATE SET gps_state=excluded.gps_state,updated_at=CURRENT_TIMESTAMP""",(plate,state))
    c.commit();c.close()
    audit("GPS_STATE",plate,state or "TEMIZLENDI","FILO")
    return {"ok":True,"plate":plate,"gps_state":state}

@app.post("/api/fleet-operations/resequence")
def resequence_fleet_operations():
    c=db();rows=c.execute("""SELECT plate FROM vehicle_operations WHERE load_state='SIRA'
      ORDER BY CASE WHEN queue_no>0 THEN queue_no ELSE 999999 END,updated_at,plate""").fetchall()
    for i,r in enumerate(rows,1):c.execute("UPDATE vehicle_operations SET queue_no=? WHERE plate=?",(i,r["plate"]))
    c.commit();c.close();return {"ok":True,"count":len(rows)}


# SAMA_CASH_EXPENSE_EDIT_PERMISSION_V1
@app.patch("/api/cash-control/expense/{expense_id}")
def edit_cash_control_expense(expense_id:int, x:CashExpenseEditIn, request:Request):
    # Middleware already enforces cash.expense.edit. Keep a second explicit guard for clarity/safety.
    user=_request_user(request)
    if not user:
        raise HTTPException(401,"Oturum gerekli.")
    if user.get("role")!="ADMIN" and "cash.expense.edit" not in set(user.get("permissions") or []):
        raise HTTPException(403,"Günlük Kasa harcaması düzeltme yetkiniz yok.")

    doc=str(x.document_no or "").strip()
    note=str(x.note or "").strip()
    cur=str(x.currency or "IQD").strip().upper()
    day=str(x.expense_date or "").strip() or datetime.now().strftime("%Y-%m-%d")
    amount=float(x.amount or 0)
    if not doc:
        raise HTTPException(400,"Fatura/Fiş No gerekli.")
    if amount<=0:
        raise HTTPException(400,"Tutar 0’dan büyük olmalı.")
    if cur not in ("IQD","USD","TRY"):
        raise HTTPException(400,"Geçersiz para birimi.")
    try:
        datetime.strptime(day[:10],"%Y-%m-%d")
        day=day[:10]
    except Exception:
        raise HTTPException(400,"Tarih YYYY-MM-DD formatında olmalı.")

    c=db()
    old=c.execute("SELECT * FROM cash_daily_expenses WHERE id=?",(expense_id,)).fetchone()
    if not old:
        c.close(); raise HTTPException(404,"Harcama kaydı bulunamadı.")
    oldd=dict(old)
    c.execute("""UPDATE cash_daily_expenses
                 SET document_no=?,amount=?,note=?,expense_date=?,currency=?
                 WHERE id=?""",(doc,amount,note,day,cur,expense_id))
    c.commit(); c.close()
    audit("CASH_EXPENSE_EDIT",str(expense_id),
          f"Günlük kasa harcaması düzeltildi: {doc} / {amount:.2f} {cur}",
          "GUNLUK_KASA",
          f"{oldd.get('document_no','')} | {float(oldd.get('amount') or 0):.2f} {oldd.get('currency','')} | {oldd.get('expense_date','')} | {oldd.get('note','')}",
          f"{doc} | {amount:.2f} {cur} | {day} | {note}")
    return {"ok":True,"id":expense_id}

@app.get("/api/health")
def health():
    return {
        "ok": True,
        "service": "SAMA TRACK",
        "db_file": str(DB),
        "db_exists": DB.exists(),
        "persistent_volume": DB.parent != BASE,
        "data_dir_exists": DATA_DIR.exists(),
        "railway_volume_env": ENV_VOLUME,
    }


@app.get("/api/integrated-fleet-status")
def integrated_fleet_status():
    c=db()
    _ensure_maintenance_ops(c)
    rows=[dict(r) for r in c.execute("""
      WITH active_fleet AS (
        SELECT UPPER(TRIM(plate)) plate,brand,model,vehicle_type,garage_state,note
        FROM fleet_vehicles
        WHERE is_active=1
      ),
      latest_trip AS (
        SELECT t.*
        FROM trips t
        WHERE t.id=(
          SELECT t2.id
          FROM trips t2
          WHERE UPPER(TRIM(t2.plate))=UPPER(TRIM(t.plate))
            AND TRIM(COALESCE(t2.plate,''))<>''
            AND COALESCE(t2.is_deleted,0)=0
          ORDER BY
            CASE
              WHEN COALESCE(t2.status,'')='Yolda' AND COALESCE(t2.exit_done,0)=1 AND COALESCE(t2.entry_done,0)=0 THEN 0
              WHEN COALESCE(t2.status,'')='Bekliyor' AND COALESCE(t2.entry_done,0)=0 THEN 1
              ELSE 2
            END,
            t2.id DESC
          LIMIT 1
        )
      )
      SELECT
        f.plate,f.brand,f.model,f.vehicle_type,f.garage_state,f.note fleet_note,
        lt.scna,
        COALESCE(d.name,'') driver,
        COALESCE(a.name,'') area,
        lt.status trip_status,
        lt.exit_done,lt.entry_done,lt.exit_at,
        COALESCE(NULLIF(lt.entry_at,''),NULLIF(lt.delivery_time,'')) entry_at,
        COALESCE(vo.vessel,'') vessel,
        COALESCE(vo.vessel_state,'') vessel_state,
        COALESCE(vo.gps_state,'') gps_state,
        COALESCE(vo.operation_note,'') operation_note,
        CASE
          WHEN UPPER(COALESCE(f.garage_state,''))='BAKIM' THEN 'BAKIM'
          WHEN COALESCE(vo.vessel_state,'')='CALISIYOR' AND TRIM(COALESCE(vo.vessel,''))<>'' THEN 'GEMIDE'
          WHEN COALESCE(lt.entry_done,0)=1 THEN 'BOSTA'
          WHEN UPPER(COALESCE(vo.gps_state,''))='DONUYOR' THEN 'DONUYOR'
          WHEN lt.exit_done=1 AND COALESCE(lt.entry_done,0)=0 THEN 'YOLDA'
          WHEN COALESCE(lt.status,'')='Bekliyor' THEN 'BEKLEMEDE'
          ELSE 'BOSTA'
        END operation_state
      FROM active_fleet f
      LEFT JOIN latest_trip lt ON UPPER(TRIM(lt.plate))=f.plate
      LEFT JOIN drivers d ON d.id=lt.driver_id
      LEFT JOIN areas a ON a.id=lt.area_id
      LEFT JOIN vehicle_operations vo ON UPPER(TRIM(vo.plate))=f.plate
      ORDER BY f.plate
    """)]
    active_maintenance={str(r['plate'] or '').strip().upper() for r in c.execute("SELECT plate FROM maintenance_operations WHERE is_active=1")}
    for x in rows:
        if str(x.get('plate') or '').strip().upper() in active_maintenance:
            x['operation_state']='BAKIM'
    c.close()
    counts={"TOPLAM":len(rows),"GEMIDE":0,"BAKIM":0,"YOLDA":0,"DONUYOR":0,"BEKLEMEDE":0,"BOSTA":0}
    for x in rows:
        s=x.get("operation_state") or "BOSTA"
        counts[s]=counts.get(s,0)+1
    return {"counts":counts,"rows":rows}



@app.get("/api/autocomplete/plate-last-driver")
def autocomplete_plate_last_driver(plate: str = ""):
    plate=(plate or "").strip().upper()
    if not plate:
        return {"found":False}
    c=db()
    row=c.execute("""
      SELECT
        t.scna,
        t.plate,
        d.id driver_id,
        COALESCE(d.name,'') driver_name,
        COALESCE(d.phone,'') phone,
        COALESCE(d.d_no,'') d_no
      FROM trips t
      LEFT JOIN drivers d ON d.id=t.driver_id
      WHERE UPPER(TRIM(t.plate))=?
      ORDER BY t.id DESC
      LIMIT 1
    """,(plate,)).fetchone()
    c.close()
    if not row:
        return {"found":False}
    x=dict(row)
    x["found"]=bool((x.get("driver_name") or "").strip())
    return x

@app.get("/api/autocomplete/plates")
def autocomplete_plates(q: str = ""):
    q=(q or "").strip().upper()
    if len(q)<3:
        return []
    like=f"%{q}%"
    c=db()
    rows=[dict(r) for r in c.execute("""
      WITH all_plates AS (
        SELECT UPPER(TRIM(plate)) plate,
               COALESCE(brand,'') brand,
               COALESCE(model,'') model,
               COALESCE(vehicle_type,'') vehicle_type,
               is_active
        FROM fleet_vehicles
        WHERE TRIM(COALESCE(plate,''))<>''
        UNION
        SELECT UPPER(TRIM(plate)) plate,
               '' brand,'' model,'' vehicle_type,
               1 is_active
        FROM trips
        WHERE TRIM(COALESCE(plate,''))<>''
      )
      SELECT plate,
             MAX(brand) brand,
             MAX(model) model,
             MAX(vehicle_type) vehicle_type,
             MAX(is_active) is_active
      FROM all_plates
      WHERE plate LIKE ?
      GROUP BY plate
      ORDER BY
        CASE WHEN plate LIKE ? THEN 0 ELSE 1 END,
        plate
      LIMIT 40
    """,(like,q+"%"))]
    c.close()
    return rows


@app.get("/api/autocomplete/customers")
def autocomplete_customers(q: str = ""):
    q=(q or "").strip()
    if len(q)<3:
        return []
    like=f"%{q}%"
    c=db()
    out=[]
    seen=set()

    # Önce trips tablosundaki gerçek müşteri alanını şema üzerinden bul.
    tcols={r["name"] for r in c.execute("PRAGMA table_info(trips)").fetchall()}
    customer_col=next((x for x in ("customer","customer_name","client","client_name") if x in tcols),None)
    if customer_col:
        sql=f"""
          SELECT DISTINCT TRIM(COALESCE({customer_col},'')) name
          FROM trips
          WHERE TRIM(COALESCE({customer_col},''))<>''
            AND UPPER(TRIM({customer_col})) LIKE UPPER(?)
          ORDER BY name
          LIMIT 100
        """
        for r in c.execute(sql,(like,)).fetchall():
            name=(r["name"] or "").strip()
            k=name.upper()
            if name and k not in seen:
                seen.add(k);out.append({"name":name})

    # Müşteri master tablosu varsa onu da tara.
    tables={r["name"] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()}
    for table in ("customers","clients"):
        if table not in tables:
            continue
        cols={r["name"] for r in c.execute(f"PRAGMA table_info({table})").fetchall()}
        name_col=next((x for x in ("name","customer_name","client_name","title") if x in cols),None)
        if not name_col:
            continue
        sql=f"""
          SELECT DISTINCT TRIM(COALESCE({name_col},'')) name
          FROM {table}
          WHERE TRIM(COALESCE({name_col},''))<>''
            AND UPPER(TRIM({name_col})) LIKE UPPER(?)
          ORDER BY name
          LIMIT 100
        """
        for r in c.execute(sql,(like,)).fetchall():
            name=(r["name"] or "").strip()
            k=name.upper()
            if name and k not in seen:
                seen.add(k);out.append({"name":name})

    c.close()
    return out[:50]

@app.get("/api/autocomplete/drivers")
def autocomplete_drivers(q: str = ""):
    q=(q or "").strip()
    if len(q)<3:
        return []
    like=f"%{q}%"
    c=db()

    # drivers tablosunun gerçek kolonlarını kontrol et; opsiyonel alan yüzünden arama çökmesin.
    cols={r["name"] for r in c.execute("PRAGMA table_info(drivers)").fetchall()}
    phone_expr="COALESCE(d.phone,'')" if "phone" in cols else "''"
    dno_expr="COALESCE(d.d_no,'')" if "d_no" in cols else "''"

    sql=f"""
      SELECT d.id,
             TRIM(COALESCE(d.name,'')) name,
             {phone_expr} phone,
             {dno_expr} d_no
      FROM drivers d
      WHERE TRIM(COALESCE(d.name,''))<>''
        AND UPPER(TRIM(d.name)) LIKE UPPER(?)
      ORDER BY
        CASE WHEN UPPER(TRIM(d.name)) LIKE UPPER(?) THEN 0 ELSE 1 END,
        d.name
      LIMIT 50
    """
    rows=[dict(r) for r in c.execute(sql,(like,q+"%"))]

    # Filo şoför tablosu varsa onu da tara.
    tables={r["name"] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()}
    if "fleet_drivers" in tables:
        fcols={r["name"] for r in c.execute("PRAGMA table_info(fleet_drivers)").fetchall()}
        fphone="COALESCE(phone,'')" if "phone" in fcols else "''"
        fdno="COALESCE(d_no,'')" if "d_no" in fcols else "''"
        frows=[dict(r) for r in c.execute(f"""
          SELECT id,TRIM(COALESCE(name,'')) name,{fphone} phone,{fdno} d_no
          FROM fleet_drivers
          WHERE TRIM(COALESCE(name,''))<>''
            AND UPPER(TRIM(name)) LIKE UPPER(?)
          LIMIT 50
        """,(like,))]
        rows.extend(frows)

    c.close()

    seen=set();out=[]
    for x in rows:
        name=(x.get("name") or "").strip()
        key=name.upper()
        if not name or key in seen: continue
        seen.add(key);out.append(x)
        if len(out)>=50: break
    return out

@app.get("/api/lookups")
def lookups():
    c = db()
    vehicles = [dict(r) for r in c.execute("""
      SELECT
        v.*,
        lt.driver_id current_driver_id,
        d.name driver_name,
        d.phone,
        d.d_no
      FROM vehicles v
      LEFT JOIN trips lt ON lt.id=(
        SELECT t2.id
        FROM trips t2
        WHERE UPPER(t2.plate)=UPPER(v.plate)
        ORDER BY
          CASE WHEN t2.status='Yolda' THEN 0 ELSE 1 END,
          t2.id DESC
        LIMIT 1
      )
      LEFT JOIN drivers d ON d.id=lt.driver_id
      WHERE v.active=1
      ORDER BY v.plate
    """)]
    areas = [dict(r) for r in c.execute(
        "SELECT * FROM areas WHERE active=1 ORDER BY name"
    )]
    cargo_categories = [dict(r) for r in c.execute(
        "SELECT * FROM cargo_categories WHERE active=1 ORDER BY name"
    )]
    customers = [dict(r) for r in c.execute(
        "SELECT * FROM customers WHERE active=1 ORDER BY name"
    )]
    c.close()
    return {
        "vehicles": vehicles,
        "areas": areas,
        "cargo_categories": cargo_categories,
        "customers": customers
    }

@app.get("/api/trips")
def trips(q: str = "", status: str = "", customer: str = "", area: str = "", cargo: str = "", limit: int = 50, offset: int = 0):
    # Binlerce kaydı tarayıcıya tek seferde göndermiyoruz.
    limit = max(1, min(int(limit or 50), 200))
    offset = max(0, int(offset or 0))
    c = db()
    sql = tq() + " WHERE 1=1 AND COALESCE(t.is_deleted,0)=0 "
    p = []
    if q:
        like = f"%{q}%"
        sql += " AND (t.scna LIKE ? OR t.plate LIKE ? OR d.name LIKE ? OR a.name LIKE ? OR cu.name LIKE ? OR cc.name LIKE ?)"
        p += [like, like, like, like, like, like]
    if status and status != "Tümü":
        sql += " AND t.status=?"
        p.append(status)
    if customer:
        sql += " AND cu.name=?"
        p.append(customer)
    if area:
        sql += " AND a.name=?"
        p.append(area)
    if cargo:
        sql += " AND cc.name=?"
        p.append(cargo)
    sql += " ORDER BY t.id DESC LIMIT ? OFFSET ?"
    p += [limit, offset]
    rows = [dict(r) for r in c.execute(sql, p)]
    c.close()
    return rows

@app.get("/api/trips-count")
def trips_count(q: str = "", status: str = "", customer: str = "", area: str = "", cargo: str = ""):
    c = db()
    sql = """SELECT COUNT(*) c FROM trips t
             LEFT JOIN drivers d ON d.id=t.driver_id
             LEFT JOIN areas a ON a.id=t.area_id
             LEFT JOIN cargo_categories cc ON cc.id=t.cargo_category_id
             LEFT JOIN customers cu ON cu.id=t.customer_id
             WHERE 1=1 AND COALESCE(t.is_deleted,0)=0"""
    p=[]
    if q:
        like=f"%{q}%"
        sql += " AND (t.scna LIKE ? OR t.plate LIKE ? OR d.name LIKE ? OR a.name LIKE ? OR cu.name LIKE ? OR cc.name LIKE ?)"
        p += [like,like,like,like,like,like]
    if status and status != "Tümü":
        if status=="Tamamlandı":
            sql += " AND (t.status='Tamamlandı' OR NULLIF(t.entry_at,'') IS NOT NULL OR NULLIF(t.delivery_time,'') IS NOT NULL)"
        elif status=="Yolda":
            sql += " AND t.status='Yolda' AND NULLIF(t.entry_at,'') IS NULL AND NULLIF(t.delivery_time,'') IS NULL"
        else:
            sql += " AND t.status=?"
            p.append(status)
    if customer:
        sql += " AND cu.name=?"
        p.append(customer)
    if area:
        sql += " AND a.name=?"
        p.append(area)
    if cargo:
        sql += " AND cc.name=?"
        p.append(cargo)
    n=c.execute(sql,p).fetchone()["c"]
    c.close()
    return {"count":n}

@app.get("/api/stats")
def stats():
    c=db()
    sql="""SELECT
      COUNT(*) total,
      COALESCE(SUM(CASE WHEN status_display='Bekliyor' THEN 1 ELSE 0 END),0) waiting,
      COALESCE(SUM(CASE WHEN status_display='Yolda' THEN 1 ELSE 0 END),0) onroad,
      COALESCE(SUM(CASE WHEN status_display='Tamamlandı' THEN 1 ELSE 0 END),0) completed,
      COALESCE(SUM(freight_total),0) freight_total,
      COALESCE(SUM(CASE WHEN entry_done=1 THEN fuel_consumed_liters ELSE 0 END),0) fuel_consumed,
      COALESCE(SUM(CASE WHEN entry_done=1 THEN MAX(amount_due,0) ELSE 0 END),0) due_total,
      COALESCE(SUM(CASE WHEN entry_done=1 THEN driver_cash_diff ELSE 0 END),0) driver_diff
      FROM ("""+tq()+""") x"""
    r=c.execute(sql).fetchone()
    c.close()
    return dict(r)

@app.get("/api/trips/{scna}")
def trip(scna: str):
    c = db()
    r = c.execute(tq() + " WHERE UPPER(TRIM(t.scna))=UPPER(TRIM(?))", (scna,)).fetchone()
    c.close()
    if not r:
        raise HTTPException(404, "Kayıt bulunamadı")
    return dict(r)

@app.patch("/api/trips/{scna}/status")
def set_trip_status(scna: str, x: StatusIn):
    status=x.status.strip()
    if status not in ("Bekliyor","Yolda","Tamamlandı"):
        raise HTTPException(400,"Geçersiz durum")
    c=db()
    row=c.execute("SELECT id FROM trips WHERE scna=?",(scna,)).fetchone()
    if not row:
        c.close(); raise HTTPException(404,"Kayıt bulunamadı")
    if status=="Bekliyor":
        c.execute("""
          UPDATE trips SET
            status=?,exit_done=0,entry_done=0,
            exit_at=NULL,entry_at=NULL,
            updated_at=CURRENT_TIMESTAMP
          WHERE scna=?
        """,(status,scna))
    elif status=="Yolda":
        c.execute("""
          UPDATE trips SET
            status=?,exit_done=1,entry_done=0,
            exit_at=COALESCE(exit_at,DATETIME('now','localtime')),
            entry_at=NULL,
            updated_at=CURRENT_TIMESTAMP
          WHERE scna=?
        """,(status,scna))
    else:
        c.execute("""
          UPDATE trips SET
            status=?,exit_done=1,entry_done=1,
            exit_at=COALESCE(exit_at,DATETIME('now','localtime')),
            entry_at=COALESCE(entry_at,DATETIME('now','localtime')),
            updated_at=CURRENT_TIMESTAMP
          WHERE scna=?
        """,(status,scna))
    c.commit(); c.close()
    audit("STATUS",scna,f"Durum: {status}")
    return {"ok":True,"status":status}

@app.post("/api/trips")
def create_trip(x: TripCreate):
    ok,scna_norm=_validate_scna_value(x.scna)
    if not ok:
        raise HTTPException(400,scna_norm)
    x.scna=scna_norm
    basis = x.freight_basis.upper()
    if basis not in ("KG", "TON"):
        raise HTTPException(400, "Navlun birimi KG veya TON olmalı")

    c = db()
    try:
        c.execute("""
          INSERT INTO trips(
            scna,plate,driver_id,area_id,cargo_category_id,cargo_type,customer_id,trip_date,
            net_kg,freight_rate,freight_basis,exit_km,tank_start_liters,
            exit_premium,dock_fee,port_fee,sonar,status
          )
          VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?, 'Bekliyor')
        """, (
            x.scna.strip().upper(),
            x.plate.strip().upper(),
            x.driver_id,
            x.area_id,
            x.cargo_category_id,
            x.cargo_type.upper(),
            x.customer_id,
            x.trip_date,
            x.net_kg,
            x.freight_rate,
            basis,
            x.exit_km,
            x.tank_start_liters,
            x.exit_premium,
            x.dock_fee,
            x.port_fee,
            x.sonar
        ))
        c.commit()
    except sqlite3.IntegrityError:
        c.close()
        raise HTTPException(409, "Bu SCNA zaten mevcut")

    c.close()
    audit("CREATE", x.scna.upper(), "Sevkiyat oluşturuldu")
    return {"ok": True}

@app.patch("/api/trips/{scna}/exit")
def exit_trip(scna: str, x: ExitIn):
    if float(x.exit_other or 0)>0 and not (x.exit_other_note or "").strip():
        raise HTTPException(400,"OTHER tutarı girildiyse açıklama zorunludur.")
    c = db()
    if not c.execute("SELECT 1 FROM trips WHERE scna=?", (scna,)).fetchone():
        c.close()
        raise HTTPException(404, "Kayıt bulunamadı")

    c.execute("""
      UPDATE trips SET
        exit_cash=?,

        exit_official_fuel_liters=?,
        exit_official_fuel_total=?,

        exit_commercial_fuel_liters=?,
        exit_commercial_fuel_total=?,

        exit_baghdad_fuel_liters=?,
        exit_baghdad_fuel_total=?,

        exit_allowance=?,
        exit_other=?,
        exit_other_note=?,
        exit_note=?,

        exit_done=1,
        exit_at=DATETIME('now','localtime'),
        status='Yolda',
        updated_at=CURRENT_TIMESTAMP
      WHERE scna=?
    """, (
        x.exit_cash,

        x.exit_official_fuel_liters,
        x.exit_official_fuel_total,

        x.exit_commercial_fuel_liters,
        x.exit_commercial_fuel_total,

        x.exit_baghdad_fuel_liters,
        x.exit_baghdad_fuel_total,

        x.exit_allowance,
        x.exit_other,
        x.exit_other_note.strip(),
        x.exit_note,

        scna
    ))

    c.commit()
    c.close()

    audit("EXIT", scna, "Çıkış, depo ve mazot litre/tutar bilgileri kaydedildi")
    return {"ok": True}

@app.patch("/api/trips/{scna}/entry")
def entry_trip(scna: str, x: EntryIn):
    c = db()
    row = c.execute("SELECT * FROM trips WHERE scna=?", (scna,)).fetchone()

    if not row:
        c.close()
        raise HTTPException(404, "Kayıt bulunamadı")

    if not row["exit_done"]:
        c.close()
        raise HTTPException(409, "Önce çıkış işlemi yapılmalı")

    if x.entry_km < row["exit_km"]:
        c.close()
        raise HTTPException(409, "Giriş KM, çıkış KM'den küçük olamaz")

    c.execute("""
      UPDATE trips SET
        entry_km=?,
        tank_end_liters=?,

        entry_extra_expense_1=?,
        entry_extra_expense_2=?,
        entry_extra_expense_3=?,

        entry_collection=?,
        entry_cash_handed=?,
        excel_remain=NULL,
        delivery_time=DATETIME('now','localtime'),
        entry_note=?,

        entry_done=1,
        entry_at=DATETIME('now','localtime'),
        status='Tamamlandı',
        updated_at=CURRENT_TIMESTAMP
      WHERE scna=?
    """, (
        x.entry_km,
        x.tank_end_liters,

        x.entry_extra_expense_1,
        x.entry_extra_expense_2,
        x.entry_extra_expense_3,

        x.entry_collection,
        x.entry_cash_handed,
        x.entry_note,

        scna
    ))

    c.commit()
    c.close()

    audit("ENTRY", scna, "Giriş, kalan depo, ek gider ve para hesabı kaydedildi")
    return {"ok": True}

@app.get("/api/trips/{scna}/fuel")
def fuel_list(scna: str):
    c = db()
    rows = [dict(r) for r in c.execute(
        "SELECT * FROM fuel_purchases WHERE scna=? ORDER BY id",
        (scna,)
    )]
    c.close()
    return rows

@app.post("/api/trips/{scna}/fuel")
def fuel_add(scna: str, x: FuelIn):
    if x.liters <= 0:
        raise HTTPException(400, "Litre 0'dan büyük olmalı")

    unit_price = x.total / x.liters if x.liters else 0

    c = db()
    if not c.execute("SELECT 1 FROM trips WHERE scna=?", (scna,)).fetchone():
        c.close()
        raise HTTPException(404, "Sevkiyat bulunamadı")

    c.execute("""
      INSERT INTO fuel_purchases(scna,liters,total,note)
      VALUES(?,?,?,?)
    """, (scna, x.liters, x.total, x.note))

    c.commit()
    c.close()

    audit("FUEL", scna, f"{x.liters} LT / {x.total} toplam / {unit_price:.2f} birim fiyat")
    return {"ok": True, "unit_price": unit_price}

@app.delete("/api/fuel/{id}")
def fuel_delete(id: int):
    c = db()
    row = c.execute(
        "SELECT scna FROM fuel_purchases WHERE id=?",
        (id,)
    ).fetchone()

    if not row:
        c.close()
        raise HTTPException(404, "Mazot kaydı bulunamadı")

    scna = row["scna"]
    c.execute("DELETE FROM fuel_purchases WHERE id=?", (id,))
    c.commit()
    c.close()

    audit("FUEL_DELETE", scna, "Mazot kaydı silindi")
    return {"ok": True}

@app.get("/api/audit")
def logs(
    q: str = "",
    username: str = "",
    action: str = "",
    scna: str = "",
    date_from: str = "",
    date_to: str = "",
    limit: int = 500
):
    c=db()
    sql="""SELECT id,action,scna,detail,username,user_id,entity_type,old_value,new_value,created_at
           FROM audit_log WHERE 1=1"""
    p=[]
    if q.strip():
        like=f"%{q.strip()}%"
        sql+=" AND (action LIKE ? OR scna LIKE ? OR detail LIKE ? OR username LIKE ? OR old_value LIKE ? OR new_value LIKE ?)"
        p += [like,like,like,like,like,like]
    if username.strip():
        sql+=" AND UPPER(username)=UPPER(?)";p.append(username.strip())
    if action.strip():
        sql+=" AND action=?";p.append(action.strip())
    if scna.strip():
        sql+=" AND UPPER(scna) LIKE UPPER(?)";p.append(f"%{scna.strip()}%")
    if date_from.strip():
        sql+=" AND datetime(created_at)>=datetime(?)";p.append(date_from.strip()+" 00:00:00")
    if date_to.strip():
        sql+=" AND datetime(created_at)<=datetime(?)";p.append(date_to.strip()+" 23:59:59")
    sql+=" ORDER BY id DESC LIMIT ?";p.append(max(1,min(int(limit or 500),2000)))
    rows=[dict(r) for r in c.execute(sql,p)]
    users=[r["username"] for r in c.execute(
      "SELECT DISTINCT username FROM audit_log WHERE TRIM(COALESCE(username,''))<>'' ORDER BY username"
    ).fetchall()]
    actions=[r["action"] for r in c.execute(
      "SELECT DISTINCT action FROM audit_log WHERE TRIM(COALESCE(action,''))<>'' ORDER BY action"
    ).fetchall()]
    c.close()
    return {"rows":rows,"users":users,"actions":actions}

# ===================== SAMA DRIVER PWA =====================
class DriverTripCreateIn(BaseModel):
    scna: str
    area_id: Optional[int] = None
    cargo_category_id: Optional[int] = None
    customer_id: Optional[int] = None
    trip_date: str = ""
    net_kg: float = 0
    freight_rate: float = 0
    freight_basis: str = "TON"

class DriverFuelAddIn(BaseModel):
    liters: float = 0
    total: float = 0
    note: str = ""


def _driver_user(request: Request):
    user=_request_user(request)
    if not user:
        raise HTTPException(401,"Oturum gerekli.")
    if not user.get("driver_id"):
        raise HTTPException(403,"Bu kullanıcı bir şoför kaydıyla eşleştirilmemiş.")
    return user


def _driver_owns_trip(c, user, scna: str):
    row=c.execute("SELECT * FROM trips WHERE UPPER(scna)=UPPER(?)",(scna.strip(),)).fetchone()
    if not row:
        raise HTTPException(404,"Sefer bulunamadı.")
    if int(row["driver_id"] or 0)!=int(user["driver_id"] or 0):
        raise HTTPException(403,"Bu sefer size ait değil.")
    return row


@app.get("/api/driver/lookups")
def driver_mobile_lookups(request: Request):
    _driver_user(request)
    c=db()
    out={
      "areas":[dict(r) for r in c.execute("SELECT id,name FROM areas WHERE active=1 ORDER BY name")],
      "cargo":[dict(r) for r in c.execute("SELECT id,name FROM cargo_categories WHERE active=1 ORDER BY name")],
      "customers":[dict(r) for r in c.execute("SELECT id,name FROM customers WHERE active=1 ORDER BY name")]
    }
    c.close()
    return out


@app.get("/api/driver/my-trips")
def driver_my_trips(request: Request, limit:int=30):
    user=_driver_user(request)
    limit=max(1,min(int(limit or 30),100))
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT t.scna,t.plate,t.trip_date,t.status,t.exit_done,t.entry_done,t.net_kg,
             COALESCE(a.name,'') area_name,COALESCE(cc.name,'') cargo_name,
             COALESCE(cu.name,'') customer_name,
             COALESCE((SELECT SUM(f.liters) FROM fuel_purchases f WHERE f.scna=t.scna),0) road_fuel_liters
      FROM trips t
      LEFT JOIN areas a ON a.id=t.area_id
      LEFT JOIN cargo_categories cc ON cc.id=t.cargo_category_id
      LEFT JOIN customers cu ON cu.id=t.customer_id
      WHERE t.driver_id=?
      ORDER BY COALESCE(t.trip_date,t.created_at) DESC,t.id DESC
      LIMIT ?
    """,(user["driver_id"],limit))]
    c.close()
    return rows


@app.post("/api/driver/my-trips")
def driver_create_trip(x: DriverTripCreateIn, request: Request):
    user=_driver_user(request)
    current_plate=(user.get("current_plate") or "").strip().upper()
    if not current_plate:
        raise HTTPException(400,"Önce aktif araç/plaka seçin.")
    ok,scna_or_error=_validate_scna_value(x.scna)
    if not ok:
        raise HTTPException(400,scna_or_error)
    scna=scna_or_error
    c=db()
    try:
        if c.execute("SELECT 1 FROM trips WHERE UPPER(scna)=UPPER(?)",(scna,)).fetchone():
            raise HTTPException(409,"Bu SCNA zaten kayıtlı.")
        trip_date=(x.trip_date or "").strip() or datetime.now().strftime("%Y-%m-%d")
        basis=(x.freight_basis or "TON").strip().upper()
        if basis not in ("TON","KG"):
            basis="TON"
        c.execute("""
          INSERT INTO trips(scna,plate,driver_id,area_id,cargo_category_id,customer_id,trip_date,
                            net_kg,freight_rate,freight_basis,status,created_at,updated_at)
          VALUES(?,?,?,?,?,?,?,?,?,?,'Bekliyor',CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)
        """,(scna,current_plate,user["driver_id"],x.area_id,x.cargo_category_id,x.customer_id,
             trip_date,float(x.net_kg or 0),float(x.freight_rate or 0),basis))
        c.commit()
    finally:
        c.close()
    audit("DRIVER_TRIP_CREATE",scna,f"Şoför uygulamasından sefer oluşturuldu: {current_plate}","SEVKIYAT","",current_plate)
    return {"ok":True,"scna":scna,"plate":current_plate}


@app.post("/api/driver/my-trips/{scna}/fuel")
def driver_add_fuel(scna:str, x:DriverFuelAddIn, request:Request):
    user=_driver_user(request)
    if float(x.liters or 0)<=0:
        raise HTTPException(400,"Yakıt litresi 0'dan büyük olmalı.")
    if float(x.total or 0)<0:
        raise HTTPException(400,"Yakıt toplamı negatif olamaz.")
    c=db()
    _driver_owns_trip(c,user,scna)
    c.execute("INSERT INTO fuel_purchases(scna,liters,total,note) VALUES(?,?,?,?)",
              (scna.strip().upper(),float(x.liters or 0),float(x.total or 0),(x.note or '').strip()))
    c.commit();c.close()
    audit("DRIVER_FUEL",scna,f"Şoför uygulamasından yakıt: {x.liters} LT","YAKIT","",str(x.liters))
    return {"ok":True}


DRIVER_MANIFEST={
  "name":"SAMA DRIVER",
  "short_name":"SAMA DRIVER",
  "description":"SAMA sürücü ve sefer uygulaması",
  "start_url":"/driver",
  "scope":"/",
  "display":"standalone",
  "orientation":"portrait-primary",
  "background_color":"#07111f",
  "theme_color":"#0b63ce",
  "icons":[{"src":"/driver-icon.svg","sizes":"any","type":"image/svg+xml","purpose":"any maskable"}]
}

@app.get("/driver.webmanifest")
def driver_manifest():
    import json
    return Response(content=json.dumps(DRIVER_MANIFEST,ensure_ascii=False),media_type="application/manifest+json")

@app.get("/driver-icon.svg")
def driver_icon():
    svg='''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 512 512"><rect width="512" height="512" rx="105" fill="#07111f"/><rect x="62" y="128" width="388" height="235" rx="44" fill="#0b63ce"/><path d="M125 315h262l-30-91H155z" fill="#fff" opacity=".96"/><circle cx="154" cy="366" r="45" fill="#e9f2ff"/><circle cx="358" cy="366" r="45" fill="#e9f2ff"/><text x="256" y="208" text-anchor="middle" font-family="Arial,sans-serif" font-weight="700" font-size="64" fill="#fff">SAMA</text></svg>'''
    return Response(content=svg,media_type="image/svg+xml")

@app.get("/driver-sw.js")
def driver_service_worker():
    js="""const CACHE='sama-driver-v80';
const CORE=['/driver','/driver.webmanifest','/driver-icon.svg'];
self.addEventListener('install',e=>{e.waitUntil(caches.open(CACHE).then(c=>c.addAll(CORE)));self.skipWaiting();});
self.addEventListener('activate',e=>{e.waitUntil(caches.keys().then(keys=>Promise.all(keys.filter(k=>k!==CACHE).map(k=>caches.delete(k)))));self.clients.claim();});
self.addEventListener('fetch',e=>{const u=new URL(e.request.url);if(e.request.method!=='GET'||u.pathname.startsWith('/api/'))return;e.respondWith(fetch(e.request).then(r=>{const copy=r.clone();caches.open(CACHE).then(c=>c.put(e.request,copy));return r;}).catch(()=>caches.match(e.request).then(r=>r||caches.match('/driver'))));});"""
    return Response(content=js,media_type="application/javascript",headers={"Service-Worker-Allowed":"/","Cache-Control":"no-cache"})

DRIVER_HTML=r'''<!doctype html>
<html lang="tr"><head>
<meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1,viewport-fit=cover">
<meta name="theme-color" content="#0b63ce"><meta name="apple-mobile-web-app-capable" content="yes"><meta name="apple-mobile-web-app-title" content="SAMA DRIVER">
<link rel="manifest" href="/driver.webmanifest"><link rel="icon" href="/driver-icon.svg"><title>SAMA DRIVER</title>
<style>
:root{--bg:#07111f;--card:#101d2f;--card2:#172943;--text:#fff;--muted:#9cafc6;--blue:#1677ff;--green:#18ad6b;--line:#2a3e59}
*{box-sizing:border-box}html,body{margin:0;background:var(--bg);color:var(--text);font-family:system-ui,-apple-system,Segoe UI,Roboto,Arial,sans-serif;min-height:100%}
body{padding:env(safe-area-inset-top) 0 env(safe-area-inset-bottom)}button,input,select{font:inherit}.app{max-width:560px;margin:auto;min-height:100vh;padding:14px}.hidden{display:none!important}
.brand{text-align:center;padding:8px 0 14px}.brand .logo{display:none}.brand h1{font-size:28px;margin:0;font-weight:950}.brand small{display:none}.card{background:var(--card);border:1px solid var(--line);border-radius:22px;padding:16px;margin-bottom:13px}
#main>.card:not(.install){text-align:center}.drivername{font-size:20px;font-weight:900}.plate{font-size:35px;font-weight:950;letter-spacing:2px;margin:9px 0}.muted{color:var(--muted)}
.grid{display:grid;grid-template-columns:1fr 1fr;gap:12px}.action{min-height:145px;border:1px solid var(--line);border-radius:24px;background:var(--card2);color:#fff;font-weight:950;font-size:20px;text-align:center}
.action:first-child{background:#123d72}.action:nth-child(2){background:#0e5038}.action br{display:block}.action::first-line{font-size:56px}
.btn{width:100%;border:0;border-radius:16px;padding:16px;font-weight:900;color:#fff;background:var(--blue);margin-top:12px;font-size:18px}.btn.secondary{background:var(--card2);border:1px solid var(--line)}.btn.green{background:var(--green)}
.inp{width:100%;border:1px solid var(--line);background:#091626;color:#fff;border-radius:16px;padding:16px;font-size:18px;outline:none}label{display:block;color:var(--muted);font-size:14px;margin:12px 0 6px}
.topbar{display:flex;justify-content:space-between;align-items:center;gap:8px}.trip{padding:13px;border-radius:15px;background:var(--card2);margin-top:9px;text-align:left}.triphead{display:flex;justify-content:space-between}.scna{font-size:21px;font-weight:950}.tripmeta{font-size:13px;color:var(--muted);margin-top:4px}.badge{font-size:11px}
.msg{padding:11px;border-radius:12px;margin-top:10px}.err{background:#4b1d25}.ok{background:#123c2d}.modal{position:fixed;inset:0;background:#000c;display:grid;align-items:end;z-index:20}.sheet{background:var(--card);border-radius:26px 26px 0 0;padding:18px;max-height:90vh;overflow:auto}.sheet h2{text-align:center;font-size:24px}.x{float:right;background:none;border:0;color:#fff;font-size:30px}.install{text-align:center}
@media(min-width:520px){.modal{align-items:center}.sheet{width:min(520px,94vw);margin:auto;border-radius:26px}}
</style></head><body><div class="app"><div class="brand"><div class="logo">SD</div><div><h1>SAMA DRIVER</h1><small></small></div></div>
<section id="loginCard" class="card hidden"><h2>Giriş</h2><div class="muted">Şoför kullanıcı adınız ve şifrenizle giriş yapın.</div><label>Kullanıcı adı</label><input id="loginUser" class="inp" autocomplete="username"><label>Şifre</label><input id="loginPass" type="password" class="inp" autocomplete="current-password"><button class="btn" onclick="login()">GİRİŞ YAP</button><div id="loginMsg"></div></section>
<section id="main" class="hidden"><div id="installCard" class="card install hidden"><b>SAMA DRIVER'ı telefona kur</b><div class="muted" style="margin-top:4px">Ana ekranda normal uygulama gibi açılır.</div><button class="btn" onclick="installApp()">UYGULAMAYI KUR</button></div><div class="card"><div class="topbar"><div><div id="driverName" class="drivername">-</div><div id="driverNo" class="muted">-</div></div><button onclick="logout()" style="background:none;border:1px solid var(--line);color:#fff;border-radius:12px;padding:9px 11px">Çıkış</button></div><hr style="border:0;border-top:1px solid var(--line);margin:15px 0"><div class="muted">AKTİF ARAÇ</div><div id="plate" class="plate">PLAKA SEÇİLMEDİ</div><button class="btn secondary" onclick="openVehicle()">🚛 ARAÇ</button></div><div class="grid"><button class="action" onclick="openNewTrip()">📦<br>SEFER</button><button class="action" onclick="openFuel()">⛽<br>YAKIT</button></div><div class="card" style="margin-top:14px"><div class="topbar"><b>📋 GEÇMİŞ</b><button onclick="loadTrips()" style="background:none;border:0;color:#8ec1ff">↻</button></div><div id="tripList" style="margin-top:8px"></div></div></section></div>
<div id="vehicleModal" class="modal hidden"><div class="sheet"><button class="x" onclick="closeModal('vehicleModal')">×</button><h2>Araç Seç</h2><input id="plateSearch" class="inp" placeholder="Plaka ara..." oninput="renderVehicles()"><div id="vehicleList" style="margin-top:10px"></div></div></div>
<div id="tripModal" class="modal hidden"><div class="sheet"><button class="x" onclick="closeModal('tripModal')">×</button><h2>Yeni SCNA</h2><label>SCNA</label><input id="nScna" class="inp" placeholder="Örn: 93540"><label>Bölge</label><select id="nArea" class="inp"></select><label>Mal</label><select id="nCargo" class="inp"></select><label>Müşteri</label><select id="nCustomer" class="inp"></select><label>Net KG</label><input id="nKg" type="number" inputmode="decimal" class="inp"><label>Çıkış Tarihi</label><input id="nDate" type="date" class="inp"><button class="btn green" onclick="createTrip()">SEFERİ KAYDET</button><div id="tripMsg"></div></div></div>
<div id="fuelModal" class="modal hidden"><div class="sheet"><button class="x" onclick="closeModal('fuelModal')">×</button><h2>Yakıt Ekle</h2><label>SCNA</label><select id="fScna" class="inp"></select><label>Litre</label><input id="fLiters" type="number" inputmode="decimal" class="inp"><label>Toplam Tutar</label><input id="fTotal" type="number" inputmode="decimal" class="inp"><label>Not</label><input id="fNote" class="inp" placeholder="İsteğe bağlı"><button class="btn green" onclick="addFuel()">YAKITI KAYDET</button><div id="fuelMsg"></div></div></div>
<script>/* SAMA_LOGIN_JS_API_COLLISION_FIX_V1 */

let SESSION=null,VEHICLES=[],TRIPS=[],LOOKUPS=null,installPrompt=null;const $=id=>document.getElementById(id);async function authApi(url,opt={}){const r=await fetch(url,{credentials:'same-origin',headers:{'Content-Type':'application/json',...(opt.headers||{})},...opt});let d=null;try{d=await r.json()}catch{}if(!r.ok)throw new Error((d&&d.detail)||'İşlem başarısız');return d}function show(id,on=true){$(id).classList.toggle('hidden',!on)}function msg(id,text,ok=false){$(id).innerHTML=text?`<div class="msg ${ok?'ok':'err'}">${text}</div>`:''}function closeModal(id){show(id,false)}async function boot(){try{const st=await authApi('/api/auth/status');if(!st.authenticated){show('loginCard');show('main',false);return}await afterLogin()}catch(e){show('loginCard');msg('loginMsg',e.message)}}async function login(){msg('loginMsg','');try{await authApi('/api/auth/login',{method:'POST',body:JSON.stringify({username:$('loginUser').value,password:$('loginPass').value})});await afterLogin()}catch(e){msg('loginMsg',e.message)}}async function logout(){try{await authApi('/api/auth/logout',{method:'POST'})}catch{}location.reload()}async function afterLogin(){SESSION=await authApi('/api/driver-session');if(!SESSION.linked){show('loginCard');show('main',false);msg('loginMsg','Bu kullanıcı bir şoför kaydıyla eşleştirilmemiş. Admin eşleştirmesi gerekiyor.');return}show('loginCard',false);show('main');$('driverName').textContent=SESSION.driver.name||'-';$('driverNo').textContent=(SESSION.driver.d_no||'')+(SESSION.driver.phone?' • '+SESSION.driver.phone:'');$('plate').textContent=SESSION.current_plate||'PLAKA SEÇİLMEDİ';VEHICLES=SESSION.vehicles||[];await Promise.all([loadTrips(),loadLookups()]);if(!SESSION.current_plate)openVehicle()}function openVehicle(){show('vehicleModal');$('plateSearch').value='';renderVehicles()}function renderVehicles(){const q=$('plateSearch').value.trim().toUpperCase();const arr=VEHICLES.filter(x=>(x.plate||'').includes(q)).slice(0,80);$('vehicleList').innerHTML=arr.map(x=>`<button class="btn secondary" onclick="selectVehicle('${x.plate}')">${x.plate}</button>`).join('')||'<div class="muted">Araç bulunamadı.</div>'}async function selectVehicle(p){try{await authApi('/api/driver-session/select-vehicle',{method:'POST',body:JSON.stringify({plate:p})});SESSION.current_plate=p;$('plate').textContent=p;closeModal('vehicleModal')}catch(e){alert(e.message)}}async function loadLookups(){try{LOOKUPS=await authApi('/api/driver/lookups');fillSelect('nArea',LOOKUPS.areas,'Bölge seç');fillSelect('nCargo',LOOKUPS.cargo,'Mal seç');fillSelect('nCustomer',LOOKUPS.customers,'Müşteri seç')}catch(e){console.error(e)}}function fillSelect(id,arr,placeholder){$(id).innerHTML=`<option value="">${placeholder}</option>`+(arr||[]).map(x=>`<option value="${x.id}">${x.name}</option>`).join('')}async function loadTrips(){try{TRIPS=await authApi('/api/driver/my-trips?limit=40');$('tripList').innerHTML=TRIPS.length?TRIPS.map(t=>`<div class="trip"><div class="triphead"><span class="scna">${t.scna}</span><span class="badge">${t.status||'-'}</span></div><div class="tripmeta">${t.plate||''} • ${t.area_name||'-'} • ${t.trip_date||'-'}</div><div class="tripmeta">${t.customer_name||''}${t.cargo_name?' • '+t.cargo_name:''}${t.road_fuel_liters?' • Yakıt '+Number(t.road_fuel_liters).toLocaleString('tr-TR')+' LT':''}</div></div>`).join(''):'<div class="muted" style="padding:12px 0">Henüz sefer kaydı yok.</div>';$('fScna').innerHTML=TRIPS.filter(t=>!t.entry_done).map(t=>`<option value="${t.scna}">${t.scna} • ${t.plate}</option>`).join('')}catch(e){$('tripList').innerHTML=`<div class="msg err">${e.message}</div>`}}function openNewTrip(){if(!SESSION?.current_plate){openVehicle();return}$('nScna').value='';$('nKg').value='';$('nDate').value=new Date().toISOString().slice(0,10);msg('tripMsg','');show('tripModal')}async function createTrip(){msg('tripMsg','');try{const body={scna:$('nScna').value,area_id:numOrNull($('nArea').value),cargo_category_id:numOrNull($('nCargo').value),customer_id:numOrNull($('nCustomer').value),trip_date:$('nDate').value,net_kg:Number($('nKg').value||0),freight_rate:0,freight_basis:'TON'};const r=await authApi('/api/driver/my-trips',{method:'POST',body:JSON.stringify(body)});msg('tripMsg',`SCNA ${r.scna} kaydedildi.`,true);await loadTrips();setTimeout(()=>closeModal('tripModal'),500)}catch(e){msg('tripMsg',e.message)}}function numOrNull(v){return v?Number(v):null}function openFuel(){if(!TRIPS.filter(t=>!t.entry_done).length){alert('Yakıt eklemek için açık sefer bulunamadı.');return}$('fLiters').value='';$('fTotal').value='';$('fNote').value='';msg('fuelMsg','');show('fuelModal')}async function addFuel(){msg('fuelMsg','');try{const scna=$('fScna').value;if(!scna)throw new Error('SCNA seçin.');await authApi(`/api/driver/my-trips/${encodeURIComponent(scna)}/fuel`,{method:'POST',body:JSON.stringify({liters:Number($('fLiters').value||0),total:Number($('fTotal').value||0),note:$('fNote').value})});msg('fuelMsg','Yakıt kaydedildi.',true);await loadTrips();setTimeout(()=>closeModal('fuelModal'),500)}catch(e){msg('fuelMsg',e.message)}}window.addEventListener('beforeinstallprompt',e=>{e.preventDefault();installPrompt=e;show('installCard')});async function installApp(){if(!installPrompt)return;installPrompt.prompt();await installPrompt.userChoice;installPrompt=null;show('installCard',false)}window.addEventListener('appinstalled',()=>show('installCard',false));if('serviceWorker' in navigator){window.addEventListener('load',()=>navigator.serviceWorker.register('/driver-sw.js').catch(console.error))}boot();
</script></body></html>'''

@app.get("/driver",response_class=HTMLResponse)
def driver_home():
    return DRIVER_HTML

# =================== /SAMA DRIVER PWA =====================


# SAMA_MAINTENANCE_OPS_V3
class MaintenanceOperationIn(BaseModel):
    plate: str
    reason: str = ""
    description: str = ""
    estimated_hours: float = 0
    work_done: str = ""


def _ensure_maintenance_ops(c):
    c.execute("""CREATE TABLE IF NOT EXISTS maintenance_operations(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      plate TEXT NOT NULL,
      previous_state TEXT DEFAULT '',
      reason TEXT DEFAULT '',
      description TEXT DEFAULT '',
      estimated_hours REAL DEFAULT 0,
      estimated_finish_at DATETIME,
      work_done TEXT DEFAULT '',
      started_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      finished_at DATETIME,
      started_by TEXT DEFAULT '',
      finished_by TEXT DEFAULT '',
      is_active INTEGER DEFAULT 1
    )""")
    cols={r['name'] for r in c.execute('PRAGMA table_info(maintenance_operations)')}
    for name,definition in [('description',"TEXT DEFAULT ''"),('estimated_hours','REAL DEFAULT 0'),('estimated_finish_at','DATETIME')]:
        if name not in cols:
            c.execute(f'ALTER TABLE maintenance_operations ADD COLUMN {name} {definition}')
    c.commit()

@app.get("/api/maintenance-operations")
def maintenance_operations_list(active_only:int=0):
    c=db(); _ensure_maintenance_ops(c)
    sql="SELECT * FROM maintenance_operations"
    if active_only: sql += " WHERE is_active=1"
    sql += " ORDER BY is_active DESC,id DESC LIMIT 1000"
    rows=[dict(r) for r in c.execute(sql)]
    c.close(); return rows

@app.post("/api/maintenance-operations/start")
def maintenance_operations_start(x:MaintenanceOperationIn):
    plate=(x.plate or '').strip().upper(); reason=(x.reason or '').strip(); description=(x.description or '').strip()
    hours=max(0,float(x.estimated_hours or 0))
    if not plate: raise HTTPException(400,'Plaka zorunlu.')
    if not reason: raise HTTPException(400,'Bakım nedeni zorunludur.')
    c=db(); _ensure_maintenance_ops(c)
    if c.execute("SELECT 1 FROM maintenance_operations WHERE UPPER(plate)=UPPER(?) AND is_active=1",(plate,)).fetchone():
        c.close(); raise HTTPException(409,'Bu araç zaten bakımda.')
    prev='BOSTA'
    tr=c.execute("""SELECT status FROM trips WHERE UPPER(TRIM(plate))=UPPER(TRIM(?)) AND COALESCE(is_deleted,0)=0
      ORDER BY CASE WHEN status='Yolda' AND COALESCE(entry_done,0)=0 THEN 0 ELSE 1 END,id DESC LIMIT 1""",(plate,)).fetchone()
    vo=c.execute("SELECT load_state FROM vehicle_operations WHERE UPPER(TRIM(plate))=UPPER(TRIM(?))",(plate,)).fetchone()
    if tr and tr['status']=='Yolda': prev='YOLDA'
    elif vo and (vo['load_state'] or '') in ('SIRA','YUKLEMEDE','DOLU','BOS'): prev=vo['load_state'] or 'BOSTA'
    user=CURRENT_AUTH_USER.get() or {}; username=user.get('username','') if isinstance(user,dict) else ''
    est=(datetime.now()+timedelta(hours=hours)).strftime('%Y-%m-%d %H:%M:%S') if hours>0 else None
    c.execute("""INSERT INTO maintenance_operations(plate,previous_state,reason,description,estimated_hours,estimated_finish_at,started_at,started_by,is_active)
      VALUES(?,?,?,?,?,?,DATETIME('now','localtime'),?,1)""",(plate,prev,reason,description,hours,est,username))
    c.commit(); c.close()
    audit('MAINTENANCE_START',plate,f'Bakıma alındı. Önceki durum: {prev}. Neden: {reason}. Tahmini süre: {hours:g} saat','ARAC',prev,'BAKIM')
    return {'ok':True,'plate':plate,'previous_state':prev,'estimated_finish_at':est}

@app.post("/api/maintenance-operations/{op_id}/update")
def maintenance_operations_update(op_id:int,x:MaintenanceOperationIn):
    plate=(x.plate or '').strip().upper()
    reason=(x.reason or '').strip()
    description=(x.description or '').strip()
    hours=max(0,float(x.estimated_hours or 0))
    if not plate: raise HTTPException(400,'Plaka zorunlu.')
    if not reason: raise HTTPException(400,'Bakım nedeni zorunludur.')
    c=db(); _ensure_maintenance_ops(c)
    row=c.execute("SELECT * FROM maintenance_operations WHERE id=?",(op_id,)).fetchone()
    if not row: c.close(); raise HTTPException(404,'Bakım kaydı bulunamadı.')
    dup=c.execute("SELECT 1 FROM maintenance_operations WHERE UPPER(plate)=UPPER(?) AND is_active=1 AND id<>?",(plate,op_id)).fetchone()
    if dup: c.close(); raise HTTPException(409,'Seçilen araç zaten başka bir aktif bakım kaydında.')
    started=row['started_at'] or datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    try: base=datetime.fromisoformat(str(started).replace('Z','').replace('T',' '))
    except: base=datetime.now()
    est=(base+timedelta(hours=hours)).strftime('%Y-%m-%d %H:%M:%S') if hours>0 else None
    c.execute("""UPDATE maintenance_operations SET plate=?,reason=?,description=?,estimated_hours=?,estimated_finish_at=? WHERE id=?""",
              (plate,reason,description,hours,est,op_id))
    c.commit(); c.close()
    audit('MAINTENANCE_UPDATE',plate,f'Bakım kaydı düzeltildi. Neden: {reason}. Tahmini süre: {hours:g} saat','ARAC')
    return {'ok':True,'id':op_id,'plate':plate,'estimated_finish_at':est}

@app.post("/api/maintenance-operations/{op_id}/finish")
def maintenance_operations_finish(op_id:int,x:MaintenanceOperationIn):
    work=(x.work_done or '').strip()
    if not work: raise HTTPException(400,'Yapılan işlemler zorunludur.')
    c=db(); _ensure_maintenance_ops(c)
    row=c.execute("SELECT * FROM maintenance_operations WHERE id=? AND is_active=1",(op_id,)).fetchone()
    if not row: c.close(); raise HTTPException(404,'Aktif bakım kaydı bulunamadı.')
    user=CURRENT_AUTH_USER.get() or {}; username=user.get('username','') if isinstance(user,dict) else ''
    c.execute("UPDATE maintenance_operations SET work_done=?,finished_at=DATETIME('now','localtime'),finished_by=?,is_active=0 WHERE id=?",(work,username,op_id))
    prev=(row['previous_state'] or 'BOSTA').upper()
    if prev in ('SIRA','YUKLEMEDE','DOLU','BOS'):
        c.execute("""INSERT INTO vehicle_operations(plate,load_state,updated_at) VALUES(?,?,CURRENT_TIMESTAMP)
          ON CONFLICT(plate) DO UPDATE SET load_state=excluded.load_state,updated_at=CURRENT_TIMESTAMP""",(row['plate'],prev))
    c.commit(); c.close()
    audit('MAINTENANCE_FINISH',row['plate'],f"Bakım bitti. Önceki durum: {prev}. Yapılan işlemler: {work}",'ARAC','BAKIM',prev)
    return {'ok':True,'plate':row['plate'],'restore_state':prev}

# SAMA_ADVANCE_TRACKING_V1
class AdvanceIn(BaseModel):
    recipient_type: str = 'USTA'
    recipient_name: str = ''
    plate: str = ''
    scna: str = ''
    purpose: str = ''
    amount: float = 0
    currency: str = 'IQD'
    note: str = ''

class AdvanceSettlementIn(BaseModel):
    settlement_type: str = 'FATURA'
    amount: float = 0
    document_no: str = ''
    note: str = ''

# SAMA_ADVANCE_SETTLEMENT_EDIT_V1
class AdvanceSettlementUpdateIn(BaseModel):
    settlement_type: str = 'FATURA'
    amount: float = 0
    document_no: str = ''
    note: str = ''
    settlement_date: str = ''

class AdvanceUpdateIn(BaseModel):
    recipient_type: str = 'USTA'
    recipient_name: str = ''
    plate: str = ''
    scna: str = ''
    purpose: str = ''
    amount: float = 0
    currency: str = 'IQD'
    note: str = ''


def _ensure_advances(c):
    c.execute("""CREATE TABLE IF NOT EXISTS cash_advances(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      recipient_type TEXT NOT NULL DEFAULT 'USTA',
      recipient_name TEXT NOT NULL DEFAULT '',
      plate TEXT DEFAULT '',
      scna TEXT DEFAULT '',
      purpose TEXT DEFAULT '',
      amount REAL NOT NULL DEFAULT 0,
      settled_amount REAL NOT NULL DEFAULT 0,
      currency TEXT NOT NULL DEFAULT 'IQD',
      note TEXT DEFAULT '',
      status TEXT NOT NULL DEFAULT 'ACIK',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      created_by TEXT DEFAULT '',
      updated_at DATETIME,
      updated_by TEXT DEFAULT ''
    )""")
    c.execute("""CREATE TABLE IF NOT EXISTS cash_advance_settlements(
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      advance_id INTEGER NOT NULL,
      settlement_type TEXT NOT NULL DEFAULT 'FATURA',
      amount REAL NOT NULL DEFAULT 0,
      document_no TEXT DEFAULT '',
      note TEXT DEFAULT '',
      created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
      created_by TEXT DEFAULT ''
    )""")
    c.commit()

# SAMA_ADVANCE_DRIVER_AUTOCOMPLETE_V1
# SAMA_ADVANCE_RECIPIENT_AUTOCOMPLETE_V3
@app.get('/api/advances/recipient-search')
def advances_recipient_search(q:str='', recipient_type:str=''):
    q=(q or '').strip(); typ=(recipient_type or '').strip().upper()
    if len(q)<3: return []
    if typ not in ('SOFOR','USTA','PERSONEL','DIGER'): return []
    c=db(); _ensure_advances(c); like='%'+q+'%'
    rows=[dict(r) for r in c.execute("""
      SELECT recipient_name name, recipient_type,
             COUNT(*) record_count,
             CASE WHEN ABS(SUM(amount-settled_amount))>=0.0001 THEN 1 ELSE 0 END open_count,
             ROUND(SUM(amount-settled_amount),2) open_balance,
             MAX(id) last_id
      FROM cash_advances
      WHERE recipient_type=? AND TRIM(COALESCE(recipient_name,''))<>''
        AND UPPER(recipient_name) LIKE UPPER(?)
      GROUP BY UPPER(TRIM(recipient_name)),recipient_type
      ORDER BY CASE WHEN SUM(CASE WHEN status IN ('ACIK','KISMI') THEN 1 ELSE 0 END)>0 THEN 0 ELSE 1 END,
               MAX(id) DESC
      LIMIT 30
    """,(typ,like))]
    c.close(); return rows

@app.get('/api/advances/driver-search')
def advances_driver_search(q:str=''):
    q=(q or '').strip()
    if len(q)<3: return []
    c=db()
    like='%'+q+'%'
    rows=[]
    # fleet_drivers is the master list; plate is resolved from current fleet assignment when possible.
    try:
        rows=[dict(r) for r in c.execute("""SELECT d.name,d.phone,d.d_no,
          COALESCE((SELECT fv.plate FROM fleet_vehicles fv WHERE UPPER(TRIM(COALESCE(fv.driver_name,'')))=UPPER(TRIM(d.name)) AND COALESCE(fv.is_active,1)=1 LIMIT 1),'') plate
          FROM fleet_drivers d WHERE COALESCE(d.is_active,1)=1 AND (d.name LIKE ? OR d.phone LIKE ? OR d.d_no LIKE ?)
          ORDER BY d.name LIMIT 20""",(like,like,like))]
    except Exception:
        try:
            rows=[dict(r) for r in c.execute("SELECT name,phone,d_no,'' plate FROM drivers WHERE name LIKE ? OR phone LIKE ? OR d_no LIKE ? ORDER BY name LIMIT 20",(like,like,like))]
        except Exception:
            rows=[]
    c.close(); return rows

class AdvanceDriverCreateIn(BaseModel):
    name: str = ''
    plate: str = ''

@app.post('/api/advances/driver-create')
def advances_driver_create(x:AdvanceDriverCreateIn):
    name=(x.name or '').strip()
    if len(name)<3: raise HTTPException(400,'Şoför adı en az 3 karakter olmalı.')
    c=db()
    # Never create a duplicate exact name.
    try:
        ex=c.execute("SELECT name FROM fleet_drivers WHERE UPPER(TRIM(name))=UPPER(TRIM(?)) LIMIT 1",(name,)).fetchone()
        if ex: c.close(); return {'ok':True,'created':False,'name':ex['name']}
        cols=[r['name'] for r in c.execute('PRAGMA table_info(fleet_drivers)')]
        data={'name':name}
        if 'is_active' in cols:data['is_active']=1
        if 'created_at' in cols:data['created_at']=None
        keys=list(data.keys()); vals=[data[k] for k in keys]
        ph=','.join('?' for _ in keys)
        c.execute(f"INSERT INTO fleet_drivers({','.join(keys)}) VALUES({ph})",vals)
        c.commit()
    except Exception as e:
        c.rollback(); c.close(); raise HTTPException(400,f'Yeni şoför kaydı oluşturulamadı: {e}')
    c.close()
    return {'ok':True,'created':True,'name':name}

# SAMA_ADVANCE_PERSON_ACCOUNTS_V1
# SAMA_ADVANCE_PERSON_LEDGER_V4
# SAMA_ADVANCE_SIGNED_BALANCE_V2
# SAMA_ADVANCE_SIGNED_BALANCE_V3
class AdvancePersonSettlementIn(BaseModel):
    settlement_type: str = 'FATURA'
    amount: float = 0
    document_no: str = ''
    note: str = ''
    currency: str = 'IQD'


def _advance_person_metrics(c, typ:str, name:str, cur:str):
    typ=(typ or '').strip().upper(); name=(name or '').strip(); cur=(cur or 'IQD').strip().upper()
    row=c.execute("""SELECT COALESCE(SUM(amount),0) total_advance,
        COALESCE(SUM(settled_amount),0) total_settled,COUNT(*) record_count
      FROM cash_advances
      WHERE recipient_type=? AND UPPER(TRIM(recipient_name))=UPPER(TRIM(?)) AND currency=?""",(typ,name,cur)).fetchone()
    total=float(row['total_advance'] or 0); settled=float(row['total_settled'] or 0); bal=total-settled
    state='KAPANDI' if abs(bal)<0.0001 else ('BORCLU' if bal>0 else 'ALACAKLI')
    return {'total_advance':total,'total_settled':settled,'balance':bal,'account_status':state,'record_count':int(row['record_count'] or 0)}


def _reconcile_advance_person(c, typ:str, name:str, cur:str, username:str=''):
    """Settlement rows are the source of truth; cached settled/status fields are rebuilt safely."""
    rows=[dict(r) for r in c.execute("""SELECT id,amount FROM cash_advances
      WHERE recipient_type=? AND UPPER(TRIM(recipient_name))=UPPER(TRIM(?)) AND currency=? ORDER BY id""",
      ((typ or '').strip().upper(),(name or '').strip(),(cur or 'IQD').strip().upper()))]
    for a in rows:
        settled=float(c.execute('SELECT COALESCE(SUM(amount),0) FROM cash_advance_settlements WHERE advance_id=?',(a['id'],)).fetchone()[0] or 0)
        amount=float(a['amount'] or 0)
        st='KAPANDI' if settled>=amount-0.0001 else ('KISMI' if settled>0 else 'ACIK')
        c.execute("UPDATE cash_advances SET settled_amount=?,status=?,updated_at=COALESCE(updated_at,DATETIME('now','localtime')),updated_by=CASE WHEN ?<>'' THEN ? ELSE updated_by END WHERE id=?",
                  (settled,st,username,username,a['id']))
    return _advance_person_metrics(c,typ,name,cur)

class AdvanceOffsetIn(BaseModel):
    new_need: float = 0
    purpose: str = ''
    plate: str = ''
    scna: str = ''
    note: str = ''
    currency: str = 'IQD'

@app.get('/api/advances/person-accounts')
def advances_person_accounts(q:str='', recipient_type:str=''):
    c=db(); _ensure_advances(c)
    par=[]; where=["TRIM(COALESCE(recipient_name,''))<>''"]
    if (q or '').strip(): where.append('UPPER(recipient_name) LIKE UPPER(?)'); par.append('%'+q.strip()+'%')
    if (recipient_type or '').strip(): where.append('recipient_type=?'); par.append(recipient_type.strip().upper())
    rows=[dict(r) for r in c.execute(f"""SELECT recipient_type,recipient_name,currency,
      COUNT(*) advance_count,ROUND(SUM(amount),2) total_advance,ROUND(SUM(settled_amount),2) total_settled,
      ROUND(SUM(amount-settled_amount),2) open_balance,MAX(id) last_id
      FROM cash_advances WHERE {' AND '.join(where)}
      GROUP BY recipient_type,UPPER(TRIM(recipient_name)),currency
      ORDER BY ABS(SUM(amount-settled_amount)) DESC,recipient_name LIMIT 1000""",par)]
    for r in rows:
        b=float(r.get('open_balance') or 0); r['account_status']='KAPANDI' if abs(b)<0.0001 else ('BORCLU' if b>0 else 'ALACAKLI'); r['open_count']=1 if abs(b)>=0.0001 else 0
    c.close(); return rows

@app.get('/api/advances/person-account')
def advances_person_account(recipient_type:str,recipient_name:str,currency:str='IQD'):
    typ=(recipient_type or '').strip().upper(); name=(recipient_name or '').strip(); cur=(currency or 'IQD').strip().upper()
    c=db(); _ensure_advances(c)
    _reconcile_advance_person(c,typ,name,cur); c.commit()
    advances=[dict(r) for r in c.execute("""SELECT a.*,(a.amount-a.settled_amount) remaining FROM cash_advances a
      WHERE recipient_type=? AND UPPER(TRIM(recipient_name))=UPPER(TRIM(?)) AND currency=? ORDER BY id DESC""",(typ,name,cur))]
    movements=[]
    for a in advances:
        movements.append({'date':a.get('created_at'),'kind':'AVANS','amount':float(a.get('amount') or 0),'advance_id':a['id'],'document_no':'','note':a.get('purpose') or ''})
        for r in c.execute("SELECT * FROM cash_advance_settlements WHERE advance_id=? ORDER BY id",(a['id'],)):
            d=dict(r); movements.append({'date':d.get('created_at'),'kind':d.get('settlement_type'),'amount':-float(d.get('amount') or 0),'advance_id':a['id'],'document_no':d.get('document_no') or '','note':d.get('note') or ''})
    movements.sort(key=lambda z:str(z.get('date') or ''),reverse=True)
    total=sum(float(a.get('amount') or 0) for a in advances); settled=sum(float(a.get('settled_amount') or 0) for a in advances)
    metrics=_advance_person_metrics(c,typ,name,cur); c.close(); return {'recipient_type':typ,'recipient_name':name,'currency':cur,'total_advance':total,'total_settled':settled,'open_balance':total-settled,'account_status':metrics['account_status'],'advances':advances,'movements':movements}

@app.post('/api/advances/person-account/settlement')
def advances_person_account_settlement(recipient_type:str,recipient_name:str,x:AdvancePersonSettlementIn,request:Request):
    _require_user_perm(request,'users.manage')
    typ=(recipient_type or '').strip().upper(); name=(recipient_name or '').strip(); cur=(x.currency or 'IQD').strip().upper()
    st=(x.settlement_type or 'FATURA').strip().upper(); amt=float(x.amount or 0); doc=(x.document_no or '').strip()
    if typ not in ('SOFOR','USTA','PERSONEL','DIGER') or not name: raise HTTPException(400,'Kişi bilgisi geçersiz.')
    if st not in ('FATURA','FIS'): raise HTTPException(400,'Cari belge türü FATURA veya FİŞ olmalı.')
    if amt<=0: raise HTTPException(400,'Tutar 0 dan büyük olmalı.')
    if not doc: raise HTTPException(400,'Fatura/Fiş belge numarası zorunlu.')
    c=db(); _ensure_advances(c)
    user=_request_user(request) or {}; username=user.get('username','') if isinstance(user,dict) else ''
    # Rebuild cached totals before allocation. Existing settlement rows remain the accounting source of truth.
    _reconcile_advance_person(c,typ,name,cur,username)
    advances=[dict(r) for r in c.execute("""SELECT * FROM cash_advances WHERE recipient_type=?
      AND UPPER(TRIM(recipient_name))=UPPER(TRIM(?)) AND currency=? ORDER BY id""",(typ,name,cur))]
    left=amt; allocated=[]
    for a in advances:
        if left<=0.0001: break
        rem=max(0,float(a['amount'] or 0)-float(a['settled_amount'] or 0))
        if rem<=0: continue
        use=min(rem,left)
        c.execute("""INSERT INTO cash_advance_settlements(advance_id,settlement_type,amount,document_no,note,created_at,created_by)
          VALUES(?,?,?,?,?,DATETIME('now','localtime'),?)""",(a['id'],st,use,doc,(x.note or '').strip(),username))
        allocated.append({'advance_id':a['id'],'amount':use}); left-=use
    # Excess document value is legitimate company debt to the person. Keep it on a carrier advance
    # so signed person balance remains exact; no fake cash movement is created.
    if left>0.0001:
        if advances:
            carrier=advances[-1]['id']
        else:
            c.execute("""INSERT INTO cash_advances(recipient_type,recipient_name,purpose,amount,currency,note,status,created_at,created_by)
              VALUES(?,?,?,0,?,?,'KAPANDI',DATETIME('now','localtime'),?)""",(typ,name,'CARİ BELGE DEVİR',cur,'Avanssız cari belge taşıyıcı kaydı',username))
            carrier=c.execute('SELECT last_insert_rowid()').fetchone()[0]
        c.execute("""INSERT INTO cash_advance_settlements(advance_id,settlement_type,amount,document_no,note,created_at,created_by)
          VALUES(?,?,?,?,?,DATETIME('now','localtime'),?)""",(carrier,st,left,doc,(x.note or '').strip(),username))
        allocated.append({'advance_id':carrier,'amount':left,'excess':True}); left=0
    metrics=_reconcile_advance_person(c,typ,name,cur,username)
    c.commit(); c.close()
    audit('ADVANCE_PERSON_SETTLEMENT',name,f'{st} {amt:g} {cur} belge {doc}; cari {metrics["balance"]:g}','AVANS')
    return {'ok':True,'allocated':allocated,'total_advance':metrics['total_advance'],'total_settled':metrics['total_settled'],'balance':metrics['balance'],'account_status':metrics['account_status']}

@app.post('/api/advances/person-account/reconcile')
def advances_person_account_reconcile(recipient_type:str,recipient_name:str,currency:str='IQD',request:Request=None):
    _require_user_perm(request,'users.manage')
    typ=(recipient_type or '').strip().upper(); name=(recipient_name or '').strip(); cur=(currency or 'IQD').strip().upper()
    if not name: raise HTTPException(400,'Kişi adı zorunlu.')
    c=db(); _ensure_advances(c)
    user=_request_user(request) or {}; username=user.get('username','') if isinstance(user,dict) else ''
    m=_reconcile_advance_person(c,typ,name,cur,username); c.commit(); c.close()
    audit('ADVANCE_PERSON_RECONCILE',name,f'Cari mutabakat yapıldı. Bakiye {m["balance"]:g} {cur}','AVANS')
    return {'ok':True,**m}

@app.post('/api/advances/person-account/offset')
def advances_person_account_offset(recipient_type:str,recipient_name:str,x:AdvanceOffsetIn,request:Request):
    _require_user_perm(request,'users.manage')
    typ=(recipient_type or '').strip().upper(); name=(recipient_name or '').strip(); cur=(x.currency or 'IQD').strip().upper(); need=float(x.new_need or 0)
    if typ not in ('SOFOR','USTA','PERSONEL','DIGER') or not name: raise HTTPException(400,'Kişi bilgisi geçersiz.')
    if need<=0: raise HTTPException(400,'Yeni avans ihtiyacı 0 dan büyük olmalı.')
    if not (x.purpose or '').strip(): raise HTTPException(400,'Yeni avans nedeni zorunlu.')
    c=db(); _ensure_advances(c)
    olds=[dict(r) for r in c.execute("""SELECT * FROM cash_advances WHERE recipient_type=? AND UPPER(TRIM(recipient_name))=UPPER(TRIM(?)) AND currency=? AND status IN ('ACIK','KISMI') AND amount>settled_amount ORDER BY id""",(typ,name,cur))]
    available=sum(max(0,float(a['amount'] or 0)-float(a['settled_amount'] or 0)) for a in olds)
    person_net=sum(float(r[0] or 0)-float(r[1] or 0) for r in c.execute("SELECT amount,settled_amount FROM cash_advances WHERE recipient_type=? AND UPPER(TRIM(recipient_name))=UPPER(TRIM(?)) AND currency=?",(typ,name,cur)))
    company_debt=max(0,-person_net)
    offset=min(need,max(0,person_net))
    left=offset
    user=_request_user(request) or {}; username=user.get('username','') if isinstance(user,dict) else ''
    for a in olds:
        if left<=0: break
        rem=max(0,float(a['amount'] or 0)-float(a['settled_amount'] or 0)); use=min(rem,left)
        if use<=0: continue
        c.execute("INSERT INTO cash_advance_settlements(advance_id,settlement_type,amount,document_no,note,created_at,created_by) VALUES(?, 'MAHSUP', ?, '', ?, DATETIME('now','localtime'),?)",(a['id'],use,'Yeni avans ihtiyacına mahsup: '+(x.purpose or '').strip(),username))
        ns=float(a['settled_amount'] or 0)+use; status='KAPANDI' if ns>=float(a['amount'] or 0)-0.0001 else 'KISMI'
        c.execute("UPDATE cash_advances SET settled_amount=?,status=?,updated_at=DATETIME('now','localtime'),updated_by=? WHERE id=?",(ns,status,username,a['id']))
        left-=use
    cash_to_give=max(0,need-offset+company_debt); new_id=None
    if cash_to_give>0:
        c.execute("""INSERT INTO cash_advances(recipient_type,recipient_name,plate,scna,purpose,amount,currency,note,status,created_at,created_by)
          VALUES(?,?,?,?,?,?,?,?,'ACIK',DATETIME('now','localtime'),?)""",(typ,name,(x.plate or '').strip().upper(),(x.scna or '').strip().upper(),(x.purpose or '').strip(),cash_to_give,cur,(x.note or '').strip(),username))
        new_id=c.execute('SELECT last_insert_rowid()').fetchone()[0]
    c.commit(); c.close()
    audit('ADVANCE_PERSON_OFFSET',name,f'Yeni ihtiyaç {need:g} {cur}; mahsup {offset:g}; kasadan verilen {cash_to_give:g}','AVANS')
    return {'ok':True,'new_need':need,'offset':offset,'company_debt':company_debt,'cash_to_give':cash_to_give,'new_advance_id':new_id}

@app.get('/api/advances')
def advances_list(status:str=''):
    c=db(); _ensure_advances(c)
    sql="""SELECT a.*,(a.amount-a.settled_amount) remaining
      FROM cash_advances a WHERE 1=1"""
    par=[]
    if status: sql+=' AND a.status=?'; par.append(status.upper())
    sql+=' ORDER BY CASE WHEN a.status=\'ACIK\' THEN 0 WHEN a.status=\'KISMI\' THEN 1 ELSE 2 END,a.id DESC LIMIT 2000'
    rows=[dict(r) for r in c.execute(sql,par)]
    groups={}
    for r in c.execute("""SELECT recipient_type,UPPER(TRIM(recipient_name)) nm,currency,SUM(amount-settled_amount) balance
      FROM cash_advances GROUP BY recipient_type,UPPER(TRIM(recipient_name)),currency"""):
        b=float(r['balance'] or 0); groups[(r['recipient_type'],r['nm'],r['currency'])]=(b,'KAPANDI' if abs(b)<0.0001 else ('BORCLU' if b>0 else 'ALACAKLI'))
    for x in rows:
        b,st2=groups.get((x.get('recipient_type'),(x.get('recipient_name') or '').strip().upper(),x.get('currency')),(float(x.get('remaining') or 0),'ACIK'))
        x['record_status']=x.get('status'); x['person_balance']=b; x['account_status']=st2
    c.close(); return rows

@app.post('/api/advances')
def advances_create(x:AdvanceIn):
    typ=(x.recipient_type or 'USTA').strip().upper()
    name=(x.recipient_name or '').strip()
    amount=float(x.amount or 0); cur=(x.currency or 'IQD').strip().upper()
    if typ not in ('SOFOR','USTA','PERSONEL','DIGER'): raise HTTPException(400,'Geçersiz alıcı tipi.')
    if not name: raise HTTPException(400,'Ad Soyad / kişi adı zorunlu.')
    if amount<=0: raise HTTPException(400,'Tutar 0 dan büyük olmalı.')
    if not (x.purpose or '').strip(): raise HTTPException(400,'Avans nedeni / iş açıklaması zorunlu.')
    c=db(); _ensure_advances(c)
    user=CURRENT_AUTH_USER.get() or {}; username=user.get('username','') if isinstance(user,dict) else ''
    c.execute("""INSERT INTO cash_advances(recipient_type,recipient_name,plate,scna,purpose,amount,currency,note,status,created_at,created_by)
      VALUES(?,?,?,?,?,?,?,?,'ACIK',DATETIME('now','localtime'),?)""",
      (typ,name,(x.plate or '').strip().upper(),(x.scna or '').strip().upper(),(x.purpose or '').strip(),amount,cur,(x.note or '').strip(),username))
    aid=c.execute('SELECT last_insert_rowid()').fetchone()[0]; c.commit(); c.close()
    audit('ADVANCE_CREATE',str(aid),f'{typ} {name} {amount:g} {cur} - {(x.purpose or "").strip()}','AVANS')
    return {'ok':True,'id':aid}

@app.post('/api/advances/{advance_id}/settlement')
def advances_settlement(advance_id:int,x:AdvanceSettlementIn):
    st=(x.settlement_type or 'FATURA').strip().upper(); amt=float(x.amount or 0); doc=(x.document_no or '').strip()
    if st not in ('FATURA','FIS','NAKIT_IADE','MAHSUP'): raise HTTPException(400,'Geçersiz kapatma türü.')
    if amt<=0: raise HTTPException(400,'Tutar 0 dan büyük olmalı.')
    if st in ('FATURA','FIS') and not doc: raise HTTPException(400,'Fatura/Fiş belge numarası zorunlu.')
    c=db(); _ensure_advances(c)
    row=c.execute('SELECT * FROM cash_advances WHERE id=?',(advance_id,)).fetchone()
    if not row: c.close(); raise HTTPException(404,'Avans kaydı bulunamadı.')
    remaining=float(row['amount'] or 0)-float(row['settled_amount'] or 0)
    user=CURRENT_AUTH_USER.get() or {}; username=user.get('username','') if isinstance(user,dict) else ''
    newsettled=float(row['settled_amount'] or 0)+amt
    status='KAPANDI' if newsettled>=float(row['amount'] or 0)-0.0001 else 'KISMI'
    c.execute("""INSERT INTO cash_advance_settlements(advance_id,settlement_type,amount,document_no,note,created_at,created_by)
      VALUES(?,?,?,?,?,DATETIME('now','localtime'),?)""",(advance_id,st,amt,doc,(x.note or '').strip(),username))
    c.execute("UPDATE cash_advances SET settled_amount=?,status=?,updated_at=DATETIME('now','localtime'),updated_by=? WHERE id=?",(newsettled,status,username,advance_id))
    c.commit(); c.close()
    audit('ADVANCE_SETTLEMENT',str(advance_id),f'{st} {amt:g}','AVANS')
    return {'ok':True,'status':status,'settled_amount':newsettled,'remaining':float(row['amount'] or 0)-newsettled}

@app.post('/api/advances/{advance_id}/update')
def advances_update(advance_id:int,x:AdvanceUpdateIn):
    amount=float(x.amount or 0)
    if amount<=0: raise HTTPException(400,'Tutar 0 dan büyük olmalı.')
    if not (x.recipient_name or '').strip(): raise HTTPException(400,'Kişi adı zorunlu.')
    if not (x.purpose or '').strip(): raise HTTPException(400,'Avans nedeni zorunlu.')
    c=db(); _ensure_advances(c)
    row=c.execute('SELECT * FROM cash_advances WHERE id=?',(advance_id,)).fetchone()
    if not row: c.close(); raise HTTPException(404,'Avans kaydı bulunamadı.')
    user=CURRENT_AUTH_USER.get() or {}; username=user.get('username','') if isinstance(user,dict) else ''
    status='KAPANDI' if float(row['settled_amount'] or 0)>=amount-0.0001 else ('KISMI' if float(row['settled_amount'] or 0)>0 else 'ACIK')
    c.execute("""UPDATE cash_advances SET recipient_type=?,recipient_name=?,plate=?,scna=?,purpose=?,amount=?,currency=?,note=?,status=?,updated_at=DATETIME('now','localtime'),updated_by=? WHERE id=?""",
      ((x.recipient_type or 'USTA').strip().upper(),(x.recipient_name or '').strip(),(x.plate or '').strip().upper(),(x.scna or '').strip().upper(),(x.purpose or '').strip(),amount,(x.currency or 'IQD').strip().upper(),(x.note or '').strip(),status,username,advance_id))
    c.commit(); c.close(); audit('ADVANCE_UPDATE',str(advance_id),'Avans kaydı düzeltildi','AVANS')
    return {'ok':True,'status':status}

@app.patch('/api/advances/{advance_id}/settlements/{settlement_id}')
def advances_settlement_update(advance_id:int,settlement_id:int,x:AdvanceSettlementUpdateIn,request:Request):
    _require_user_perm(request,'users.manage')
    st=(x.settlement_type or 'FATURA').strip().upper(); amt=float(x.amount or 0); doc=(x.document_no or '').strip()
    if st not in ('FATURA','FIS','NAKIT_IADE','MAHSUP'): raise HTTPException(400,'Geçersiz kapatma türü.')
    if amt<=0: raise HTTPException(400,'Tutar 0 dan büyük olmalı.')
    if st in ('FATURA','FIS') and not doc: raise HTTPException(400,'Fatura/Fiş belge numarası zorunlu.')
    day=(x.settlement_date or '').strip()
    if day:
        try: datetime.strptime(day[:10],'%Y-%m-%d')
        except Exception: raise HTTPException(400,'Geçersiz işlem tarihi.')
    c=db(); _ensure_advances(c)
    adv=c.execute('SELECT * FROM cash_advances WHERE id=?',(advance_id,)).fetchone()
    row=c.execute('SELECT * FROM cash_advance_settlements WHERE id=? AND advance_id=?',(settlement_id,advance_id)).fetchone()
    if not adv or not row:
        c.close(); raise HTTPException(404,'Avans hareketi bulunamadı.')
    other=float(c.execute('SELECT COALESCE(SUM(amount),0) s FROM cash_advance_settlements WHERE advance_id=? AND id<>?',(advance_id,settlement_id)).fetchone()['s'] or 0)
    newsettled=other+amt
    advance_amount=float(adv['amount'] or 0)
    user=_request_user(request) or {}; username=user.get('username','') if isinstance(user,dict) else ''
    old_value=f"Tür={row['settlement_type']} | Tutar={float(row['amount'] or 0):g} | Belge={row['document_no'] or ''} | Tarih={row['created_at'] or ''} | Açıklama={row['note'] or ''}"
    if day:
        c.execute("UPDATE cash_advance_settlements SET settlement_type=?,amount=?,document_no=?,note=?,created_at=?||' 12:00:00' WHERE id=? AND advance_id=?",(st,amt,doc,(x.note or '').strip(),day[:10],settlement_id,advance_id))
    else:
        c.execute('UPDATE cash_advance_settlements SET settlement_type=?,amount=?,document_no=?,note=? WHERE id=? AND advance_id=?',(st,amt,doc,(x.note or '').strip(),settlement_id,advance_id))
    status='KAPANDI' if newsettled>=advance_amount-0.0001 else ('KISMI' if newsettled>0 else 'ACIK')
    c.execute("UPDATE cash_advances SET settled_amount=?,status=?,updated_at=DATETIME('now','localtime'),updated_by=? WHERE id=?",(newsettled,status,username,advance_id))
    c.commit(); c.close()
    new_value=f"Tür={st} | Tutar={amt:g} | Belge={doc} | Tarih={day or row['created_at'] or ''} | Açıklama={(x.note or '').strip()}"
    audit('ADVANCE_SETTLEMENT_UPDATE',str(advance_id),f'Avans hareketi düzeltildi. Hareket ID: {settlement_id}','AVANS',old_value,new_value)
    return {'ok':True,'status':status,'settled_amount':newsettled,'remaining':advance_amount-newsettled}

@app.get('/api/advances/{advance_id}/settlements')
def advances_settlements(advance_id:int):
    c=db(); _ensure_advances(c)
    rows=[dict(r) for r in c.execute('SELECT * FROM cash_advance_settlements WHERE advance_id=? ORDER BY id DESC',(advance_id,))]
    c.close(); return rows

# SAMA_CASH_CONTROL_V1
class CashExpenseIn(BaseModel):
    document_no: str=''
    amount: float=0
    currency: str='IQD'
    note: str=''
    expense_date: str=''
class CashCountIn(BaseModel):
    business_date: str=''
    currency: str='IQD'
    opening_amount: float=0
    cash_in: float=0
    actual_amount: float=0
    note: str=''

def _ensure_cash_control(c):
    c.execute("""CREATE TABLE IF NOT EXISTS cash_daily_expenses(id INTEGER PRIMARY KEY AUTOINCREMENT,document_no TEXT NOT NULL DEFAULT '',amount REAL NOT NULL DEFAULT 0,currency TEXT NOT NULL DEFAULT 'IQD',note TEXT DEFAULT '',expense_date TEXT NOT NULL DEFAULT '',created_at DATETIME DEFAULT CURRENT_TIMESTAMP,created_by TEXT DEFAULT '')""")
    c.execute("""CREATE TABLE IF NOT EXISTS cash_daily_counts(id INTEGER PRIMARY KEY AUTOINCREMENT,business_date TEXT NOT NULL,currency TEXT NOT NULL DEFAULT 'IQD',opening_amount REAL NOT NULL DEFAULT 0,cash_in REAL NOT NULL DEFAULT 0,actual_amount REAL NOT NULL DEFAULT 0,note TEXT DEFAULT '',created_at DATETIME DEFAULT CURRENT_TIMESTAMP,created_by TEXT DEFAULT '')""")
    c.commit()

@app.get('/api/cash-control')
def cash_control_summary(business_date:str='',currency:str='IQD'):
    day=(business_date or datetime.now().strftime('%Y-%m-%d')).strip(); cur=(currency or 'IQD').upper().strip()
    c=db(); _ensure_cash_control(c); _ensure_advances(c)
    last=c.execute("SELECT * FROM cash_daily_counts WHERE business_date=? AND currency=? ORDER BY id DESC LIMIT 1",(day,cur)).fetchone()
    opening=float(last['opening_amount'] or 0) if last else 0; other_cash_in=float(last['cash_in'] or 0) if last else 0; actual=float(last['actual_amount'] or 0) if last else 0
    # Sevkiyat tablosu bu modülde salt okunur. Kasa girişi, müşteriden
    # tahsil edilen/navlun/alacak değil, şoförün gerçekten teslim ettiği paradır.
    shipment_rows=[]
    if cur=='IQD':
        shipment_rows=[dict(r) for r in c.execute("""
          SELECT scna,plate,entry_cash_handed amount,
                 COALESCE(NULLIF(entry_at,''),NULLIF(delivery_time,''),NULLIF(updated_at,''),NULLIF(trip_date,''),NULLIF(created_at,'')) entry_date
          FROM trips
          WHERE entry_done=1
            AND COALESCE(is_deleted,0)=0
            AND COALESCE(entry_cash_handed,0)>0
            AND DATE(COALESCE(NULLIF(entry_at,''),NULLIF(delivery_time,''),NULLIF(updated_at,''),NULLIF(trip_date,''),NULLIF(created_at,'')))=?
          ORDER BY COALESCE(NULLIF(entry_at,''),NULLIF(delivery_time,''),NULLIF(updated_at,''),NULLIF(trip_date,''),NULLIF(created_at,'')) DESC,scna
        """,(day,))]
    shipment_cash_in=sum(float(r['amount'] or 0) for r in shipment_rows)
    # SAMA_SHIPMENT_CASH_OUT_V1
    # SAMA_CASH_OUT_DATE_FALLBACK_V2
    # Sevkiyat cikisinda sofore fiilen verilen para (trips.exit_cash) kasadan cikistir.
    # Eski kayitlarda exit_at bos olabildigi icin efektif cikis tarihi sirayla
    # exit_at -> trip_date -> created_at alanlarindan okunur. Sevkiyat tablosu salt okunurdur.
    shipment_out_rows=[]
    if cur=='IQD':
        shipment_out_rows=[dict(r) for r in c.execute("""
          SELECT t.scna,t.plate,COALESCE(d.name,'') driver_name,
                 t.exit_cash amount,
                 COALESCE(NULLIF(t.exit_at,''),NULLIF(t.trip_date,''),NULLIF(t.created_at,'')) exit_date
          FROM trips t
          LEFT JOIN drivers d ON d.id=t.driver_id
          WHERE t.exit_done=1
            AND COALESCE(t.is_deleted,0)=0
            AND COALESCE(t.exit_cash,0)>0
            AND DATE(COALESCE(NULLIF(t.exit_at,''),NULLIF(t.trip_date,''),NULLIF(t.created_at,'')))=?
          ORDER BY COALESCE(NULLIF(t.exit_at,''),NULLIF(t.trip_date,''),NULLIF(t.created_at,'')) DESC,t.scna
        """,(day,))]
    shipment_cash_out=sum(float(r['amount'] or 0) for r in shipment_out_rows)
    advances=float(c.execute("SELECT COALESCE(SUM(amount),0) FROM cash_advances WHERE currency=? AND DATE(created_at)=?",(cur,day)).fetchone()[0] or 0)
    expenses=float(c.execute("SELECT COALESCE(SUM(amount),0) FROM cash_daily_expenses WHERE currency=? AND expense_date=?",(cur,day)).fetchone()[0] or 0)
    rows=[dict(r) for r in c.execute("SELECT * FROM cash_daily_expenses WHERE currency=? AND expense_date=? ORDER BY id DESC",(cur,day))]
    expected=opening+shipment_cash_in+other_cash_in-shipment_cash_out-advances-expenses; diff=actual-expected
    c.close(); return {'business_date':day,'currency':cur,'opening_amount':opening,
      'shipment_cash_in':shipment_cash_in,'shipment_cash_entries':shipment_rows,
      'shipment_cash_out':shipment_cash_out,'shipment_cash_out_entries':shipment_out_rows,
      'other_cash_in':other_cash_in,'cash_in':other_cash_in,
      'advance_out':advances,'expense_out':expenses,'expected_amount':expected,
      'actual_amount':actual,'difference':diff,'expenses':rows}



# SAMA_CASH_DIAGNOSTIC_V1
# SAMA_CASH_DIAGNOSTIC_ROUTE_FIX_V2

# SAMA_ENTRY_CASH_DIAGNOSTIC_V1
# SAMA_ENTRY_CASH_DATE_FALLBACK_V2
@app.get('/api/entry-cash-diagnostic/{scna}')
def entry_cash_diagnostic(scna:str):
    key=_normalize_scna_value(scna)
    c=db()
    row=c.execute("""
      SELECT t.scna,t.plate,COALESCE(d.name,'') driver_name,
             t.entry_done,t.entry_cash_handed,t.entry_collection,t.entry_at,
             t.delivery_time,t.trip_date,t.created_at,t.updated_at,t.status,COALESCE(t.is_deleted,0) is_deleted
      FROM trips t
      LEFT JOIN drivers d ON d.id=t.driver_id
      WHERE UPPER(TRIM(t.scna))=UPPER(TRIM(?))
      LIMIT 1
    """,(key,)).fetchone()
    if not row:
        c.close(); raise HTTPException(404,'SCNA bulunamadı.')
    x=dict(row)
    effective_entry=x.get('entry_at') or x.get('delivery_time') or x.get('updated_at') or x.get('trip_date') or x.get('created_at') or None
    reasons=[]
    if int(x.get('is_deleted') or 0)!=0: reasons.append('Kayıt silinmiş')
    if int(x.get('entry_done') or 0)!=1: reasons.append('Giriş işlemi tamamlanmamış (entry_done != 1)')
    if float(x.get('entry_cash_handed') or 0)<=0: reasons.append('Şoförün Teslim Ettiği Para 0 veya boş')
    if not effective_entry: reasons.append('Giriş tarihi bulunamadı')
    eligible=(len(reasons)==0)
    x['effective_entry_date']=effective_entry
    x['cash_in_eligible']=eligible
    x['cash_in_block_reasons']=reasons
    x['accounting_result']='MUHASEBEYE GİRER' if eligible else 'MUHASEBEYE GİRMEZ'
    c.close(); return x


# SAMA_ENTRY_CASH_MISSING_V3
class EntryCashSetIn(BaseModel):
    amount: float = 0

@app.get('/api/entry-cash-missing')
def entry_cash_missing(business_date:str=''):
    day=(business_date or '').strip() or datetime.now().strftime('%Y-%m-%d')
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT t.scna,t.plate,COALESCE(d.name,'') driver_name,
             t.entry_cash_handed,t.entry_collection,t.entry_done,t.entry_at,t.delivery_time,t.trip_date,t.created_at,t.updated_at,t.status,
             COALESCE(NULLIF(t.entry_at,''),NULLIF(t.delivery_time,''),NULLIF(t.updated_at,''),NULLIF(t.trip_date,''),NULLIF(t.created_at,'')) effective_entry_date
      FROM trips t
      LEFT JOIN drivers d ON d.id=t.driver_id
      WHERE t.entry_done=1
        AND COALESCE(t.is_deleted,0)=0
        AND COALESCE(t.entry_cash_handed,0)<=0
        AND DATE(COALESCE(NULLIF(t.entry_at,''),NULLIF(t.delivery_time,''),NULLIF(t.updated_at,''),NULLIF(t.trip_date,''),NULLIF(t.created_at,'')))=?
      ORDER BY effective_entry_date DESC,t.scna
    """,(day,))]
    c.close(); return rows

@app.post('/api/trips/{scna}/entry-cash')
def set_entry_cash(scna:str,x:EntryCashSetIn):
    amount=float(x.amount or 0)
    if amount<=0:
        raise HTTPException(400,'Şoförün teslim ettiği para 0’dan büyük olmalı.')
    key=_normalize_scna_value(scna)
    c=db()
    row=c.execute("SELECT scna,entry_done FROM trips WHERE UPPER(TRIM(scna))=UPPER(TRIM(?)) AND COALESCE(is_deleted,0)=0 LIMIT 1",(key,)).fetchone()
    if not row:
        c.close(); raise HTTPException(404,'SCNA bulunamadı.')
    if int(row['entry_done'] or 0)!=1:
        c.close(); raise HTTPException(409,'Sevkiyat girişi tamamlanmamış.')
    c.execute("UPDATE trips SET entry_cash_handed=?,updated_at=CURRENT_TIMESTAMP WHERE scna=?",(amount,row['scna']))
    c.commit(); c.close()
    audit('ENTRY_CASH_SET',row['scna'],f'Şoförün teslim ettiği para: {amount:.2f} IQD')
    return {'ok':True,'scna':row['scna'],'entry_cash_handed':amount}

# SAMA_EXIT_CASH_MISSING_FROM_20260830_V1
# SAMA_EXIT_CASH_MISSING_EXPENSES_V2
# SAMA_EXIT_CASH_LIMIT_V3
class ExitCashSetIn(BaseModel):
    amount: float = 0

def _repair_v4_exit_cash_misimports():
    c=db()
    rows=c.execute("""
      SELECT id,scna,exit_cash,exit_allowance,
             (COALESCE(exit_official_fuel_total,0)+COALESCE(exit_commercial_fuel_total,0)+
              COALESCE(exit_baghdad_fuel_total,0)+COALESCE(exit_allowance,0)+
              COALESCE(exit_premium,0)+COALESCE(exit_other,0)+COALESCE(dock_fee,0)+
              COALESCE(port_fee,0)+COALESCE(sonar,0)) expected_cash
      FROM trips
      WHERE COALESCE(is_deleted,0)=0
        AND exit_done=1
        AND DATE(COALESCE(NULLIF(exit_at,''),NULLIF(trip_date,''),NULLIF(created_at,'')))>=DATE('2026-08-30')
        AND COALESCE(exit_cash,0)>0
        AND ABS(COALESCE(exit_cash,0)-COALESCE(exit_allowance,0))<0.01
    """).fetchall()
    fixed=[]
    for r in rows:
        expected=float(r['expected_cash'] or 0)
        current=float(r['exit_cash'] or 0)
        if expected>current+0.01:
            c.execute("UPDATE trips SET exit_cash=?,updated_at=CURRENT_TIMESTAMP WHERE id=?",(expected,r['id']))
            fixed.append((r['scna'],current,expected))
    c.commit(); c.close()
    return fixed

@app.get('/api/exit-cash-missing')
def exit_cash_missing():
    c=db()
    rows=[dict(r) for r in c.execute("""
      SELECT t.scna,t.plate,COALESCE(d.name,'') driver_name,
             t.trip_date,t.exit_at,t.created_at,t.exit_allowance,t.exit_premium,t.exit_other,
             t.exit_official_fuel_total,t.exit_commercial_fuel_total,t.exit_baghdad_fuel_total,
             t.dock_fee,t.port_fee,t.sonar,
             (COALESCE(t.exit_official_fuel_total,0)+COALESCE(t.exit_commercial_fuel_total,0)+COALESCE(t.exit_baghdad_fuel_total,0)) fuel_total,
             (COALESCE(t.exit_allowance,0)+COALESCE(t.exit_premium,0)+COALESCE(t.exit_other,0)+COALESCE(t.dock_fee,0)+COALESCE(t.port_fee,0)+COALESCE(t.sonar,0)+
              COALESCE(t.exit_official_fuel_total,0)+COALESCE(t.exit_commercial_fuel_total,0)+COALESCE(t.exit_baghdad_fuel_total,0)) visible_expense_total,
             t.exit_cash,t.exit_done,t.status
      FROM trips t
      LEFT JOIN drivers d ON d.id=t.driver_id
      WHERE t.exit_done=1
        AND COALESCE(t.is_deleted,0)=0
        AND COALESCE(t.exit_cash,0)<=0
        AND DATE(COALESCE(NULLIF(t.exit_at,''),NULLIF(t.trip_date,''),NULLIF(t.created_at,'')))>=DATE('2026-08-30')
      ORDER BY DATE(COALESCE(NULLIF(t.exit_at,''),NULLIF(t.trip_date,''),NULLIF(t.created_at,''))) DESC,t.id DESC
    """)]
    c.close(); return rows

@app.post('/api/trips/{scna}/exit-cash')
def set_trip_exit_cash(scna:str,x:ExitCashSetIn):
    amount=float(x.amount or 0)
    if amount<=0:
        raise HTTPException(400,'Çıkışta şoföre verilen toplam nakit 0 dan büyük olmalı.')
    key=_normalize_scna_value(scna)
    c=db()
    row=c.execute("""
      SELECT scna,exit_done,exit_cash,
             COALESCE(exit_official_fuel_total,0) exit_official_fuel_total,
             COALESCE(exit_commercial_fuel_total,0) exit_commercial_fuel_total,
             COALESCE(exit_baghdad_fuel_total,0) exit_baghdad_fuel_total,
             COALESCE(exit_allowance,0) exit_allowance,
             COALESCE(exit_premium,0) exit_premium,
             COALESCE(exit_other,0) exit_other,
             COALESCE(dock_fee,0) dock_fee,
             COALESCE(port_fee,0) port_fee,
             COALESCE(sonar,0) sonar
      FROM trips
      WHERE UPPER(TRIM(scna))=UPPER(TRIM(?)) AND COALESCE(is_deleted,0)=0
      LIMIT 1
    """,(key,)).fetchone()
    if not row:
        c.close(); raise HTTPException(404,'SCNA bulunamadı.')
    if int(row['exit_done'] or 0)!=1:
        c.close(); raise HTTPException(400,'Bu sevkiyatın çıkış işlemi tamamlanmamış.')
    allowed_total=sum(float(row[k] or 0) for k in (
        'exit_official_fuel_total','exit_commercial_fuel_total','exit_baghdad_fuel_total',
        'exit_allowance','exit_premium','exit_other','dock_fee','port_fee','sonar'
    ))
    if allowed_total>0 and amount>allowed_total+0.01:
        c.close()
        raise HTTPException(400,f'HATA: Şoföre verilen nakit gider toplamını aşamaz. Gider toplamı {allowed_total:,.0f} IQD, girilen {amount:,.0f} IQD.')
    c.execute("UPDATE trips SET exit_cash=?,updated_at=CURRENT_TIMESTAMP WHERE UPPER(TRIM(scna))=UPPER(TRIM(?))",(amount,key))
    c.commit(); c.close()
    audit('EXIT_CASH_SET',key,f'Çıkışta şoföre verilen toplam nakit: {amount:g} IQD','SEVKIYAT')
    return {'ok':True,'scna':key,'exit_cash':amount}

@app.get('/api/cash-out-diagnostic/{scna}')
def cash_out_diagnostic(scna:str):
    key=_normalize_scna_value(scna)
    c=db()
    row=c.execute("""
      SELECT t.*,
             COALESCE(d.name,'') driver_name,
             COALESCE(a.name,'') area_name,
             COALESCE(cu.name,'') customer_name
      FROM trips t
      LEFT JOIN drivers d ON d.id=t.driver_id
      LEFT JOIN areas a ON a.id=t.area_id
      LEFT JOIN customers cu ON cu.id=t.customer_id
      WHERE UPPER(TRIM(t.scna))=UPPER(TRIM(?))
      LIMIT 1
    """,(key,)).fetchone()
    if not row:
        c.close(); raise HTTPException(404,'SCNA bulunamadı.')
    x=dict(row)
    effective_exit=x.get('exit_at') or x.get('trip_date') or x.get('created_at') or ''
    reasons=[]
    if float(x.get('exit_cash') or 0)<=0: reasons.append('exit_cash 0 veya boş')
    if int(x.get('exit_done') or 0)!=1: reasons.append('exit_done 1 değil')
    if int(x.get('is_deleted') or 0)==1: reasons.append('kayıt silinmiş')
    if not effective_exit: reasons.append('çıkış tarihi bulunamadı')
    result={
      'scna':x.get('scna'),'plate':x.get('plate'),'driver_name':x.get('driver_name'),
      'trip_date':x.get('trip_date'),'exit_cash':x.get('exit_cash'),'exit_done':x.get('exit_done'),
      'exit_at':x.get('exit_at'),'created_at':x.get('created_at'),'updated_at':x.get('updated_at'),
      'effective_exit_date':effective_exit,'is_deleted':x.get('is_deleted',0),
      'exit_allowance':x.get('exit_allowance'),'exit_premium':x.get('exit_premium'),
      'exit_other':x.get('exit_other'),'dock_fee':x.get('dock_fee'),'port_fee':x.get('port_fee'),
      'sonar':x.get('sonar'),'entry_cash_handed':x.get('entry_cash_handed'),
      'entry_collection':x.get('entry_collection'),'entry_done':x.get('entry_done'),'entry_at':x.get('entry_at'),
      'cash_out_eligible':len(reasons)==0,'cash_out_block_reasons':reasons
    }
    c.close(); return result

@app.post('/api/cash-control/expense')
def cash_control_expense(x:CashExpenseIn):
    doc=(x.document_no or '').strip(); amt=float(x.amount or 0); cur=(x.currency or 'IQD').upper().strip(); day=(x.expense_date or '').strip() or datetime.now().strftime('%Y-%m-%d')
    if not doc: raise HTTPException(400,'Fiş/Fatura numarası zorunlu.')
    if amt<=0: raise HTTPException(400,'Tutar zorunlu ve 0 dan büyük olmalı.')
    c=db(); _ensure_cash_control(c); user=CURRENT_AUTH_USER.get() or {}; username=user.get('username','') if isinstance(user,dict) else ''
    c.execute("INSERT INTO cash_daily_expenses(document_no,amount,currency,note,expense_date,created_at,created_by) VALUES(?,?,?,?,?,DATETIME('now','localtime'),?)",(doc,amt,cur,(x.note or '').strip(),day,username)); c.commit(); c.close()
    return {'ok':True}

@app.post('/api/cash-control/count')
def cash_control_count(x:CashCountIn):
    day=(x.business_date or '').strip() or datetime.now().strftime('%Y-%m-%d'); cur=(x.currency or 'IQD').upper().strip()
    c=db(); _ensure_cash_control(c); user=CURRENT_AUTH_USER.get() or {}; username=user.get('username','') if isinstance(user,dict) else ''
    c.execute("INSERT INTO cash_daily_counts(business_date,currency,opening_amount,cash_in,actual_amount,note,created_at,created_by) VALUES(?,?,?,?,?,?,DATETIME('now','localtime'),?)",(day,cur,float(x.opening_amount or 0),float(x.cash_in or 0),float(x.actual_amount or 0),(x.note or '').strip(),username)); c.commit(); c.close()
    return {'ok':True}


HTML = r"""
<!DOCTYPE html>
<html lang="tr">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>SAMA TRACK V82</title>
<style>
*{box-sizing:border-box}
body{margin:0;font-family:Segoe UI,Arial;background:#f4f7fb;color:#1f2937}
.side{position:fixed;inset:0 auto 0 0;width:245px;background:#111827;color:#fff;padding:22px 16px}
.brand{font-size:22px;font-weight:800;margin-bottom:25px}.brand span{color:#38bdf8}
.nav button{width:100%;text-align:left;border:0;background:transparent;color:#cbd5e1;padding:12px;border-radius:9px;font-weight:700;margin:4px 0;cursor:pointer}
.nav button.active,.nav button:hover{background:#1f2937;color:#fff}
.main{margin-left:245px}
.top{background:#fff;padding:18px 26px;border-bottom:1px solid #e5e7eb;display:flex;justify-content:space-between}
.content{padding:24px}
.panel{display:none}.panel.active{display:block}
.cards{display:grid;grid-template-columns:repeat(5,1fr);gap:14px;margin-bottom:16px}
.card{background:white;padding:18px;border-radius:14px;box-shadow:0 4px 15px #0000000d}
.card b{display:block;font-size:24px;margin-top:6px}
.toolbar{display:flex;gap:10px;background:white;padding:14px;border-radius:14px;margin-bottom:14px;box-shadow:0 4px 15px #0000000d}
.toolbar input{flex:1}
input,select,textarea{border:1px solid #d1d5db;border-radius:9px;padding:9px 11px;font-size:14px}
.btn{border:0;border-radius:9px;padding:9px 13px;font-weight:700;cursor:pointer}
.primary{background:#2563eb;color:#fff}
.green{background:#059669;color:#fff}
.orange{background:#d97706;color:#fff}
.secondary{background:#e5e7eb}
.danger{background:#fee2e2;color:#991b1b}
.table{background:white;border-radius:14px;overflow:auto;box-shadow:0 4px 15px #0000000d}
table{width:100%;border-collapse:collapse;min-width:1150px}
th,td{padding:12px;border-bottom:1px solid #eef2f7;text-align:left;font-size:13px;white-space:nowrap}
th{background:#f8fafc;color:#64748b;font-size:11px}
.modal{display:none;position:fixed;inset:0;background:#0f172a8c;align-items:center;justify-content:center;z-index:20}
.box{background:white;width:min(1120px,96vw);max-height:93vh;overflow:auto;border-radius:16px;padding:20px}
.grid{display:grid;grid-template-columns:repeat(3,1fr);gap:12px}
.field{display:flex;flex-direction:column;gap:5px}
.field label{font-size:12px;font-weight:800;color:#64748b}
.wide{grid-column:1/-1}
.actions{display:flex;justify-content:flex-end;gap:8px;margin-top:18px}
.calc{background:#eff6ff;border:1px solid #bfdbfe;border-radius:10px;padding:14px;margin-top:14px}
.settle{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin-top:14px}
.settle div{background:#f8fafc;border:1px solid #e5e7eb;border-radius:10px;padding:12px}
.settle strong{display:block;font-size:18px;margin-top:4px}
.badge{padding:5px 9px;border-radius:99px;font-weight:800;font-size:11px}
.ok{background:#dcfce7;color:#166534}
.warn{background:#fef3c7;color:#92400e}
.over{background:#e0e7ff;color:#3730a3}
.short{background:#fee2e2;color:#991b1b}
.small{font-size:12px;color:#6b7280}
.mini{min-width:0!important}
.pager{display:flex;align-items:center;justify-content:flex-end;gap:8px;padding:12px 4px;color:#64748b;font-size:13px}
.status-select{min-width:120px;padding:6px 8px;font-weight:700}
.column-picker{position:relative}
.column-panel{
  display:none;position:absolute;right:0;top:42px;z-index:30;
  width:280px;max-height:420px;overflow:auto;background:#fff;
  border:1px solid #e5e7eb;border-radius:12px;padding:12px;
  box-shadow:0 12px 30px rgba(0,0,0,.15)
}
.column-panel.open{display:block}
.column-panel label{display:flex;gap:8px;align-items:center;padding:6px 2px;font-size:13px}
.column-panel input{width:auto}

.diff-box{
  border-radius:14px;
  padding:16px;
  font-size:14px;
  font-weight:800;
  margin-top:14px;
}
.diff-box strong{display:block;font-size:24px;margin-top:6px}
.diff-red{background:#fee2e2;border:1px solid #fecaca;color:#991b1b}
.diff-green{background:#dcfce7;border:1px solid #bbf7d0;color:#166534}
.diff-neutral{background:#f3f4f6;border:1px solid #e5e7eb;color:#374151}
.detail-grid{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin-bottom:14px}
.detail-box{background:#fff;border:1px solid #e5e7eb;border-radius:10px;padding:12px}
.detail-box span{display:block;color:#64748b;font-size:11px;font-weight:700}
.detail-box strong{display:block;margin-top:5px;font-size:15px}
.alert-red{background:#fee2e2;color:#991b1b;font-weight:800}
.alert-orange{background:#ffedd5;color:#9a3412;font-weight:800}
.filterbar{
  display:flex;gap:8px;align-items:center;flex-wrap:wrap;
  background:#fff;padding:10px 12px;border-radius:12px;margin-bottom:12px;
  box-shadow:0 4px 15px #0000000d
}
.filterbar input,.filterbar select{min-width:140px}
.filterbar .grow{flex:1;min-width:220px}
.section-note{font-size:12px;color:#64748b;margin-left:auto}
.compact-actions{display:flex;gap:6px;flex-wrap:wrap}
.empty-state{padding:26px;text-align:center;color:#64748b}
.quick-btn{padding:7px 10px;border-radius:8px;border:1px solid #d1d5db;background:#fff;cursor:pointer;font-weight:700}
.quick-btn:hover{background:#f8fafc}

.clickhead{cursor:pointer;user-select:none;white-space:nowrap}
.clickhead:hover{background:#e8eef7!important}
.clickhead span{font-size:11px;opacity:.65;margin-left:4px}
.head-filter-menu{position:fixed;z-index:99999;background:#fff;border:1px solid #cbd5e1;border-radius:10px;box-shadow:0 12px 35px rgba(15,23,42,.18);padding:10px;width:270px}
.head-filter-menu input{width:100%;box-sizing:border-box;border:1px solid #cbd5e1;border-radius:7px;padding:8px;margin-bottom:7px}
.head-menu-btn{display:block;width:100%;text-align:left;border:0;background:transparent;padding:8px;border-radius:6px;cursor:pointer;font-weight:600}
.head-menu-btn:hover{background:#f1f5f9}

/* V54 - tüm tablolar için başlıktan sıralama / filtre */
table thead th.v54-head{cursor:pointer;user-select:none;white-space:nowrap}
table thead th.v54-head:hover{background:#e8eef7!important}
.v54-sortmark{font-size:11px;opacity:.6;margin-left:4px}
.v54-filtered{box-shadow:inset 0 -3px 0 #64748b}
#v54TableMenu{
 position:fixed;z-index:100000;background:#fff;border:1px solid #cbd5e1;border-radius:10px;
 box-shadow:0 12px 35px rgba(15,23,42,.20);padding:10px;width:280px
}
#v54TableMenu input{width:100%;box-sizing:border-box;border:1px solid #cbd5e1;border-radius:7px;padding:8px;margin:7px 0}
.v54-menu-btn{display:block;width:100%;text-align:left;border:0;background:transparent;padding:8px;border-radius:6px;cursor:pointer;font-weight:600}
.v54-menu-btn:hover{background:#f1f5f9}

.grouped-nav{display:flex;flex-direction:column;gap:6px}
.nav-group{border-radius:8px;overflow:hidden}
.nav-group-title{
  width:100%;display:flex!important;justify-content:space-between;align-items:center;
  font-size:12px!important;font-weight:800!important;letter-spacing:.5px;
  padding:9px 10px!important;background:rgba(255,255,255,.06)!important;
  color:#cbd5e1!important;border:0!important
}
.nav-group-title:hover{background:rgba(255,255,255,.11)!important}
.nav-group-items{display:none;padding:4px 0 5px}
.nav-group.open .nav-group-items{display:block}
.nav-group.open .nav-group-title{color:#fff!important}
.nav-group-items button{padding-left:18px!important;font-size:13px!important}
.nav-group-title span{font-size:11px;opacity:.8}
.aside{overflow-y:auto}

.op-click{cursor:pointer;transition:transform .12s ease,box-shadow .12s ease}
.op-click:hover{transform:translateY(-2px);box-shadow:0 8px 20px rgba(15,23,42,.12)}

.dashboard-op-section{width:100%;display:block;clear:both;margin:12px 0 16px}
.dashboard-op-cards{width:100%;display:grid!important;grid-template-columns:repeat(5,minmax(150px,1fr));gap:12px}
.dashboard-op-detail{display:block;width:100%;box-sizing:border-box;margin-top:12px;padding:12px;background:#fff;border:1px solid #dbe3ef;border-radius:12px;clear:both;overflow-x:auto}
.dashboard-op-detail .table{width:100%;max-width:100%;overflow-x:auto}
.dashboard-op-detail table{width:100%}
@media(max-width:1200px){.dashboard-op-cards{grid-template-columns:repeat(3,minmax(150px,1fr))}}
@media(max-width:800px){.dashboard-op-cards{grid-template-columns:repeat(2,minmax(140px,1fr))}}

.autocomplete-wrap{position:relative;width:100%}
.autocomplete-list{position:absolute;z-index:100500;left:0;right:0;top:100%;background:#fff;border:1px solid #cbd5e1;border-radius:8px;box-shadow:0 10px 28px rgba(15,23,42,.18);max-height:240px;overflow:auto;margin-top:3px;display:none}
.autocomplete-item{padding:8px 10px;cursor:pointer;border-bottom:1px solid #f1f5f9;font-size:13px}
.autocomplete-item:hover,.autocomplete-item.active{background:#eef4ff}
.autocomplete-main{font-weight:700}
.autocomplete-sub{font-size:11px;color:#64748b;margin-top:2px}

.auth-overlay{position:fixed;inset:0;z-index:200000;background:rgba(15,23,42,.82);display:flex;align-items:center;justify-content:center;padding:20px}
.auth-card{width:min(430px,95vw);background:#fff;border-radius:16px;padding:22px;box-shadow:0 25px 70px rgba(0,0,0,.35)}
.auth-logo{font-weight:900;font-size:24px;margin-bottom:16px}.auth-logo span{color:#2563eb}
.permission-grid{display:grid;grid-template-columns:repeat(2,minmax(220px,1fr));gap:7px 14px;max-height:440px;overflow:auto;padding:8px}
.permission-item{display:flex;align-items:center;gap:8px;border:1px solid #e5e7eb;border-radius:8px;padding:7px}
@media(max-width:700px){.permission-grid{grid-template-columns:1fr}}
</style>
</head>
<body>

<div id="authOverlay" class="auth-overlay">
<!-- SAMA main UI language selector -->
<style id="sama-i18n-style-v2">
#samaLangBox{position:fixed;right:18px;top:12px;z-index:10050;display:flex;align-items:center;gap:4px;padding:4px;border:1px solid rgba(148,163,184,.25);border-radius:10px;background:rgba(15,23,42,.94);box-shadow:0 5px 18px rgba(0,0,0,.22);backdrop-filter:blur(7px)}
#samaLangBox button{border:0;border-radius:7px;padding:6px 9px;min-width:34px;cursor:pointer;font-weight:800;font-size:12px;line-height:1.1;background:transparent;color:#cbd5e1}
#samaLangBox button:hover{background:rgba(255,255,255,.10);color:#fff}
#samaLangBox button.active{background:#2563eb;color:#fff}
html[dir="rtl"] body{direction:rtl;text-align:right}
html[dir="rtl"] #samaLangBox{right:auto;left:18px}
@media(max-width:760px){#samaLangBox{right:8px;top:8px;transform:scale(.92);transform-origin:top right}html[dir="rtl"] #samaLangBox{right:auto;left:8px;transform-origin:top left}}
</style>
<div id="samaLangBox" aria-label="Language selector">
  <button type="button" data-lang="tr" title="Türkçe">TR</button>
  <button type="button" data-lang="en" title="English">EN</button>
  <button type="button" data-lang="ar" title="العربية">AR</button>
</div>
<script>
// SAMA_I18N_TR_EN_AR_V2
(function(){
  const DICT={
 en:{
  'Ana Sayfa':'Dashboard','Çıkış İşlemleri':'Exit Operations','Giriş İşlemleri':'Entry Operations','Operasyon':'Operations','Filo':'Fleet','Şoförler':'Drivers','Bakım':'Maintenance','Raporlar':'Reports','Ayarlar':'Settings','Muhasebe & Finans':'Accounting & Finance','MUHASEBE & FİNANS':'ACCOUNTING & FINANCE','Günlük Kasa':'Daily Cash','Avans Takip':'Advance Tracking','Kullanıcılar & Yetkiler':'Users & Permissions','Kullanıcı Adı':'Username','Şifre':'Password','Giriş':'Login','Çıkış':'Logout','Kaydet':'Save','İptal':'Cancel','Kapat':'Close','Düzenle':'Edit','Sil':'Delete','Yeni Kayıt':'New Record','Ara':'Search','Tarih':'Date','Plaka':'Plate','Şoför':'Driver','Tutar':'Amount','Açıklama':'Description','Durum':'Status','Beklemede':'Waiting','Bakımda':'Maintenance','Yolda':'On Road','Boşta':'Available','Gemide Çalışan':'Working on Ship','Boşaltıldı Dönüyor':'Unloaded / Returning','Toplam':'Total','Para Birimi':'Currency','Yetkiler':'Permissions','Kullanıcılar':'Users',
  'Sevkiyat':'Shipment','Sevkiyatlar':'Shipments','Yeni Sevkiyat':'New Shipment','Yeni Sevkiyat Ekle':'Add New Shipment','Çıkış':'Exit','Giriş':'Entry','Çıkış Tarihi':'Exit Date','Giriş Tarihi':'Entry Date','Çıkış Saati':'Exit Time','Giriş Saati':'Entry Time','Çıkış KM':'Exit KM','Giriş KM':'Entry KM','Müşteri':'Customer','Bölge':'Area','Teslimat Bölgesi':'Delivery Area','Mal Tipi':'Cargo Type','Net KG':'Net KG','Navlun':'Freight','Navlun Birim Fiyat':'Freight Unit Price','Harcırah':'Allowance','Prim':'Premium','Diğer':'Other','Not':'Note','Telefon':'Phone','D.No':'D.No','Depo Mazot':'Tank Fuel','Resmi Mazot':'Official Fuel','Ticari Mazot':'Commercial Fuel','Bağdat Mazot':'Baghdad Fuel','Litre':'Liter','Birim Fiyat':'Unit Price','Yakıt':'Fuel','Yakıt Alımı':'Fuel Purchase','Port Fee':'Port Fee','Dock Fee':'Dock Fee','SONAR':'SONAR','Tahsilat':'Collection','Kalan':'Remaining','Getirilen Para':'Cash Returned','Şoförün Getirdiği Para':'Cash Returned by Driver','Şoförün Vermesi Gereken Para':'Cash Driver Must Return','Şoföre Verilmesi Gereken':'Cash to Give Driver','Çıkış Parası':'Exit Cash','Çıkış Parası Eksik':'Missing Exit Cash',
  'Aktif':'Active','Pasif':'Inactive','Tamamlandı':'Completed','Bekliyor':'Waiting','Çıkışta':'Outbound','Girişte':'Inbound','TAMAMLANDI':'COMPLETED','ÇIKIŞTA':'OUTBOUND','GİRİŞTE':'INBOUND','GEMİDE ÇALIŞIYOR':'WORKING ON SHIP','BAKIMDA':'IN MAINTENANCE','YOLDA':'ON ROAD','BOŞALTILDI DÖNÜYOR':'UNLOADED / RETURNING','BEKLEMEDE':'WAITING','BOŞTA':'AVAILABLE',
  'Araç':'Vehicle','Araçlar':'Vehicles','Araç Ekle':'Add Vehicle','Araç Düzenle':'Edit Vehicle','Marka':'Brand','Model':'Model','Araç Tipi':'Vehicle Type','Garaj Durumu':'Garage Status','Gemi':'Vessel','Gemi Operasyonu':'Vessel Operation','Yükleme Sırası':'Loading Queue','Sıra No':'Queue No','Operasyon Notu':'Operation Note','Filo Durumu':'Fleet Status','Araç Durumu':'Vehicle Status','GPS Durumu':'GPS Status','Son Güncelleme':'Last Update',
  'Şoför Ekle':'Add Driver','Şoför Düzenle':'Edit Driver','Şoför No':'Driver No','Şoför Adı':'Driver Name','Telefon No':'Phone Number','Aktif Şoförler':'Active Drivers','Pasif Şoförler':'Inactive Drivers',
  'Bakım Kaydı':'Maintenance Record','Aracı Bakıma Al':'Send Vehicle to Maintenance','Bakım Bitti':'Maintenance Finished','Bakım Başlangıç':'Maintenance Start','Bakım Bitiş':'Maintenance End','Bakım Nedeni':'Maintenance Reason','Bakım Açıklaması':'Maintenance Description','Yapılan İşlem':'Work Performed','Tahmini Süre':'Estimated Duration','Önceki Durum':'Previous Status','DÜZENLE':'EDIT',
  'Rapor':'Report','Raporlar':'Reports','Filtrele':'Filter','Temizle':'Clear','Dışa Aktar':'Export','Excel Dışa Aktar':'Export Excel','Excel İçe Aktar':'Import Excel','İçe Aktar':'Import','Başlangıç Tarihi':'Start Date','Bitiş Tarihi':'End Date','Tarih Aralığı':'Date Range','Tümü':'All','Sonuç':'Result','Kayıt':'Record','Kayıtlar':'Records','Kayıt Sayısı':'Record Count','Toplam Kayıt':'Total Records',
  'Gün Başı':'Opening Cash','Sevkiyattan Gelen':'Cash from Shipments','Diğer Kasa Girişi':'Other Cash In','Diğer Giriş':'Other Inflow','Sevkiyat Çıkışı / Şoföre Verilen':'Shipment Outflow / Given to Driver','Avans':'Advance','Avanslar':'Advances','Doğrudan Harcama':'Direct Expense','Harcama':'Expense','Harcamalar':'Expenses','Beklenen Kasa':'Expected Cash','Beklenen':'Expected','Fiili Kasa':'Physical Cash','Fiili':'Physical','Açık':'Difference','Kasa':'Cash','Kasa Girişi':'Cash In','Kasa Çıkışı':'Cash Out','Belge No':'Document No','Fatura No':'Invoice No','Fiş No':'Receipt No','Fatura':'Invoice','Fiş':'Receipt','Nakit İade':'Cash Return','Mahsup':'Offset','Belge Türü':'Document Type','Hesaplaşma':'Settlement','Hesaplaşmalar':'Settlements','Kalan Avans':'Remaining Advance','Avans Tutarı':'Advance Amount','Avans Ver':'Give Advance','Avans Düzenle':'Edit Advance','Kişi':'Person','Kişi Tipi':'Person Type','Usta':'Technician','Personel':'Personnel','Diğer':'Other',
  'Kullanıcı':'User','Rol':'Role','Yetki':'Permission','Yetkiler':'Permissions','Kullanıcı Ekle':'Add User','Kullanıcı Düzenle':'Edit User','Şifre Değiştir':'Change Password','Yeni Şifre':'New Password','Mevcut Şifre':'Current Password','Çıkış Yap':'Logout','Oturum':'Session','Admin':'Admin','Operatör':'Operator','İzleyici':'Viewer',
  'İşlem Geçmişi':'Audit Log','İşlem':'Action','Oluşturulma Tarihi':'Created Date','Güncelleme Tarihi':'Updated Date','Detay':'Detail','Silindi':'Deleted','Güncellendi':'Updated','Eklendi':'Added','Başarılı':'Successful','Hata':'Error','Uyarı':'Warning','Onay':'Confirm','Evet':'Yes','Hayır':'No','Seçiniz':'Select','Seç':'Select','Yükleniyor...':'Loading...','Aranıyor...':'Searching...','Kayıt bulunamadı':'No records found','Veri bulunamadı':'No data found','Zorunlu alan':'Required field','Lütfen bekleyin':'Please wait','Yenile':'Refresh','Geri':'Back','İleri':'Next','Bugün':'Today','Dün':'Yesterday','Bu Ay':'This Month','Bu Yıl':'This Year'
 },
 ar:{
  'Ana Sayfa':'الرئيسية','Çıkış İşlemleri':'عمليات الخروج','Giriş İşlemleri':'عمليات الدخول','Operasyon':'العمليات','Filo':'الأسطول','Şoförler':'السائقون','Bakım':'الصيانة','Raporlar':'التقارير','Ayarlar':'الإعدادات','Muhasebe & Finans':'المحاسبة والمالية','MUHASEBE & FİNANS':'المحاسبة والمالية','Günlük Kasa':'الصندوق اليومي','Avans Takip':'متابعة السلف','Kullanıcılar & Yetkiler':'المستخدمون والصلاحيات','Kullanıcı Adı':'اسم المستخدم','Şifre':'كلمة المرور','Giriş':'تسجيل الدخول','Çıkış':'تسجيل الخروج','Kaydet':'حفظ','İptal':'إلغاء','Kapat':'إغلاق','Düzenle':'تعديل','Sil':'حذف','Yeni Kayıt':'سجل جديد','Ara':'بحث','Tarih':'التاريخ','Plaka':'رقم المركبة','Şoför':'السائق','Tutar':'المبلغ','Açıklama':'الوصف','Durum':'الحالة','Beklemede':'في الانتظار','Bakımda':'في الصيانة','Yolda':'على الطريق','Boşta':'متاح','Gemide Çalışan':'يعمل على السفينة','Boşaltıldı Dönüyor':'تم التفريغ / عائد','Toplam':'المجموع','Para Birimi':'العملة','Yetkiler':'الصلاحيات','Kullanıcılar':'المستخدمون',
  'Sevkiyat':'الشحنة','Sevkiyatlar':'الشحنات','Yeni Sevkiyat':'شحنة جديدة','Yeni Sevkiyat Ekle':'إضافة شحنة جديدة','Çıkış':'خروج','Giriş':'دخول','Çıkış Tarihi':'تاريخ الخروج','Giriş Tarihi':'تاريخ الدخول','Çıkış Saati':'وقت الخروج','Giriş Saati':'وقت الدخول','Çıkış KM':'عداد الخروج','Giriş KM':'عداد الدخول','Müşteri':'العميل','Bölge':'المنطقة','Teslimat Bölgesi':'منطقة التسليم','Mal Tipi':'نوع الحمولة','Net KG':'الوزن الصافي','Navlun':'أجرة النقل','Navlun Birim Fiyat':'سعر وحدة النقل','Harcırah':'مخصصات','Prim':'مكافأة','Diğer':'أخرى','Not':'ملاحظة','Telefon':'الهاتف','D.No':'رقم السائق','Depo Mazot':'وقود الخزان','Resmi Mazot':'وقود رسمي','Ticari Mazot':'وقود تجاري','Bağdat Mazot':'وقود بغداد','Litre':'لتر','Birim Fiyat':'سعر الوحدة','Yakıt':'الوقود','Yakıt Alımı':'شراء الوقود','Tahsilat':'التحصيل','Kalan':'المتبقي','Getirilen Para':'النقد المعاد','Şoförün Getirdiği Para':'النقد الذي أعاده السائق','Şoförün Vermesi Gereken Para':'المبلغ المطلوب من السائق','Şoföre Verilmesi Gereken':'المبلغ المطلوب إعطاؤه للسائق','Çıkış Parası':'نقد الخروج','Çıkış Parası Eksik':'نقد الخروج مفقود',
  'Aktif':'نشط','Pasif':'غير نشط','Tamamlandı':'مكتمل','Bekliyor':'قيد الانتظار','Çıkışta':'في الخروج','Girişte':'في الدخول','TAMAMLANDI':'مكتمل','ÇIKIŞTA':'في الخروج','GİRİŞTE':'في الدخول','GEMİDE ÇALIŞIYOR':'يعمل على السفينة','BAKIMDA':'في الصيانة','YOLDA':'على الطريق','BOŞALTILDI DÖNÜYOR':'تم التفريغ / عائد','BEKLEMEDE':'في الانتظار','BOŞTA':'متاح',
  'Araç':'المركبة','Araçlar':'المركبات','Araç Ekle':'إضافة مركبة','Araç Düzenle':'تعديل المركبة','Marka':'الماركة','Model':'الموديل','Araç Tipi':'نوع المركبة','Garaj Durumu':'حالة المرآب','Gemi':'السفينة','Gemi Operasyonu':'عملية السفينة','Yükleme Sırası':'ترتيب التحميل','Sıra No':'رقم الدور','Operasyon Notu':'ملاحظة العملية','Filo Durumu':'حالة الأسطول','Araç Durumu':'حالة المركبة','GPS Durumu':'حالة GPS','Son Güncelleme':'آخر تحديث',
  'Şoför Ekle':'إضافة سائق','Şoför Düzenle':'تعديل السائق','Şoför No':'رقم السائق','Şoför Adı':'اسم السائق','Telefon No':'رقم الهاتف','Aktif Şoförler':'السائقون النشطون','Pasif Şoförler':'السائقون غير النشطين',
  'Bakım Kaydı':'سجل الصيانة','Aracı Bakıma Al':'إرسال المركبة للصيانة','Bakım Bitti':'انتهت الصيانة','Bakım Başlangıç':'بدء الصيانة','Bakım Bitiş':'نهاية الصيانة','Bakım Nedeni':'سبب الصيانة','Bakım Açıklaması':'وصف الصيانة','Yapılan İşlem':'العمل المنفذ','Tahmini Süre':'المدة التقديرية','Önceki Durum':'الحالة السابقة','DÜZENLE':'تعديل',
  'Rapor':'تقرير','Filtrele':'تصفية','Temizle':'مسح','Dışa Aktar':'تصدير','Excel Dışa Aktar':'تصدير Excel','Excel İçe Aktar':'استيراد Excel','İçe Aktar':'استيراد','Başlangıç Tarihi':'تاريخ البداية','Bitiş Tarihi':'تاريخ النهاية','Tarih Aralığı':'نطاق التاريخ','Tümü':'الكل','Sonuç':'النتيجة','Kayıt':'سجل','Kayıtlar':'السجلات','Kayıt Sayısı':'عدد السجلات','Toplam Kayıt':'إجمالي السجلات',
  'Gün Başı':'رصيد أول اليوم','Sevkiyattan Gelen':'الوارد من الشحنات','Diğer Kasa Girişi':'إدخال نقدي آخر','Diğer Giriş':'إدخال آخر','Sevkiyat Çıkışı / Şoföre Verilen':'خروج الشحنة / المدفوع للسائق','Avans':'سلفة','Avanslar':'السلف','Doğrudan Harcama':'مصروف مباشر','Harcama':'مصروف','Harcamalar':'المصروفات','Beklenen Kasa':'الرصيد المتوقع','Beklenen':'متوقع','Fiili Kasa':'النقد الفعلي','Fiili':'فعلي','Açık':'الفرق','Kasa':'الصندوق','Kasa Girişi':'دخول نقدي','Kasa Çıkışı':'خروج نقدي','Belge No':'رقم المستند','Fatura No':'رقم الفاتورة','Fiş No':'رقم الإيصال','Fatura':'فاتورة','Fiş':'إيصال','Nakit İade':'إرجاع نقدي','Mahsup':'تسوية','Belge Türü':'نوع المستند','Hesaplaşma':'تسوية','Hesaplaşmalar':'التسويات','Kalan Avans':'السلفة المتبقية','Avans Tutarı':'مبلغ السلفة','Avans Ver':'إعطاء سلفة','Avans Düzenle':'تعديل السلفة','Kişi':'الشخص','Kişi Tipi':'نوع الشخص','Usta':'فني','Personel':'موظف',
  'Kullanıcı':'المستخدم','Rol':'الدور','Yetki':'الصلاحية','Kullanıcı Ekle':'إضافة مستخدم','Kullanıcı Düzenle':'تعديل المستخدم','Şifre Değiştir':'تغيير كلمة المرور','Yeni Şifre':'كلمة المرور الجديدة','Mevcut Şifre':'كلمة المرور الحالية','Çıkış Yap':'تسجيل الخروج','Oturum':'الجلسة','Admin':'مدير','Operatör':'مشغل','İzleyici':'مشاهد',
  'İşlem Geçmişi':'سجل العمليات','İşlem':'العملية','Oluşturulma Tarihi':'تاريخ الإنشاء','Güncelleme Tarihi':'تاريخ التحديث','Detay':'التفاصيل','Silindi':'تم الحذف','Güncellendi':'تم التحديث','Eklendi':'تمت الإضافة','Başarılı':'نجاح','Hata':'خطأ','Uyarı':'تحذير','Onay':'تأكيد','Evet':'نعم','Hayır':'لا','Seçiniz':'اختر','Seç':'اختر','Yükleniyor...':'جارٍ التحميل...','Aranıyor...':'جارٍ البحث...','Kayıt bulunamadı':'لم يتم العثور على سجلات','Veri bulunamadı':'لا توجد بيانات','Zorunlu alan':'حقل مطلوب','Lütfen bekleyin':'يرجى الانتظار','Yenile':'تحديث','Geri':'رجوع','İleri':'التالي','Bugün':'اليوم','Dün':'أمس','Bu Ay':'هذا الشهر','Bu Yıl':'هذا العام'
 }
};
  // SAMA_I18N_FULL_TRANSLATION_V4
  // SAMA_I18N_COMPREHENSIVE_V5
  const EXTRA={
    en:{
      'Ana Sayfa Gör':'View Dashboard','Sevkiyat Gör':'View Shipments','Sevkiyat / Çıkış / Giriş Düzenle':'Edit Shipment / Exit / Entry','Sevkiyat Sil':'Delete Shipment','Excel / OneDrive İçe Aktar':'Import Excel / OneDrive','Operasyon / Gemi / Yükleme Sırası Gör':'View Operations / Vessel / Loading Queue','Operasyon / Gemi / Yükleme Sırası Düzenle':'Edit Operations / Vessel / Loading Queue','Filo Gör':'View Fleet','Filo Düzenle':'Edit Fleet','Şoförler Gör':'View Drivers','Şoförler Düzenle':'Edit Drivers','Bakım Gör':'View Maintenance','Bakım Düzenle':'Edit Maintenance','Raporlar Gör':'View Reports','Rapor / Excel Dışa Aktar':'Export Report / Excel','İşlem Geçmişi Gör':'View Audit Log','Günlük Kasa Harcama Düzelt':'Edit Daily Cash Expense','Kullanıcılar & Yetkiler Yönet':'Manage Users & Permissions',
      'İlk Admin Hesabını Oluştur':'Create First Admin Account','Ad Soyad':'Full Name','Admin Oluştur':'Create Admin','Giriş Yap':'Sign In','Oturum gerekli.':'Session required.','Şifrenizi değiştirmeden diğer işlemleri kullanamazsınız.':'You must change your password before using other functions.','Bu işlem için yetkiniz yok':'You do not have permission for this action',
      'YÜKLEME & SEVKİYAT':'LOADING & SHIPMENT','FİLO & OPERASYON':'FLEET & OPERATIONS','RAPORLAMA':'REPORTING','SİSTEM':'SYSTEM','YÖNETİM':'MANAGEMENT','OPERASYON MERKEZİ':'OPERATIONS CENTER','FİLO YÖNETİMİ':'FLEET MANAGEMENT','ŞOFÖR YÖNETİMİ':'DRIVER MANAGEMENT','BAKIM YÖNETİMİ':'MAINTENANCE MANAGEMENT',
      'Yeni Çıkış':'New Exit','Yeni Giriş':'New Entry','Çıkış Kaydı':'Exit Record','Giriş Kaydı':'Entry Record','Çıkışı Kaydet':'Save Exit','Girişi Kaydet':'Save Entry','Sevkiyat Bilgileri':'Shipment Information','Çıkış Bilgileri':'Exit Information','Giriş Bilgileri':'Entry Information','Araç ve Şoför':'Vehicle and Driver','Teslimat Bilgileri':'Delivery Information','Yakıt ve Masraflar':'Fuel and Expenses','Mali Bilgiler':'Financial Information','Ekstra Masraflar':'Extra Expenses','Dönüş Bilgileri':'Return Information','Dönüş Harcamaları':'Return Expenses','Müşteri Tahsilatı':'Customer Collection','Şoförün Getirdiği Nakit':'Cash Returned by Driver','Depo Başlangıç Mazot':'Starting Tank Fuel','Depo Bitiş Mazot':'Ending Tank Fuel','Toplam Yakıt':'Total Fuel','Toplam Masraf':'Total Expense','Toplam Tutar':'Total Amount','Toplam Araç':'Total Vehicles','Toplam Şoför':'Total Drivers','Toplam Sevkiyat':'Total Shipments','Toplam Avans':'Total Advances','Toplam Harcama':'Total Expenses','Toplam Kalan':'Total Remaining',
      'GEMİDE ÇALIŞAN':'WORKING ON SHIP','GEMİDE':'ON SHIP','BOŞALTILDI':'UNLOADED','DÖNÜYOR':'RETURNING','BEKLEME':'WAITING','GARAj':'GARAGE','GARJDA':'IN GARAGE','BAKIM':'MAINTENANCE','YOLDA':'ON ROAD','BOŞTA':'AVAILABLE','AKTİF':'ACTIVE','PASİF':'INACTIVE','TAMAMLANDI':'COMPLETED','BEKLİYOR':'WAITING','İPTAL':'CANCELLED',
      'Durum Özeti':'Status Summary','Filo Özeti':'Fleet Summary','Operasyon Özeti':'Operations Summary','Günlük Özet':'Daily Summary','Günlük Yönetici Özeti':'Daily Manager Summary','Aktif Sevkiyatlar':'Active Shipments','Son Sevkiyatlar':'Recent Shipments','Son İşlemler':'Recent Actions','Bekleyen İşler':'Pending Tasks','Uyarılar':'Alerts','Anomaliler':'Anomalies','Performans':'Performance','Yakıt Performansı':'Fuel Performance','Rota Performansı':'Route Performance',
      'Plaka Ara':'Search Plate','Şoför Ara':'Search Driver','SCNA Ara':'Search SCNA','Müşteri Ara':'Search Customer','Bölge Ara':'Search Area','Kayıt Ara':'Search Records','Arama':'Search','Filtre':'Filter','Filtreleri Temizle':'Clear Filters','Sonuçları Göster':'Show Results','Listeyi Yenile':'Refresh List','Listeye Dön':'Back to List','Detayları Gör':'View Details','Detay Göster':'Show Details','Görüntüle':'View','Düzelt':'Correct','Güncelle':'Update','Ekle':'Add','Kaldır':'Remove','Pasife Al':'Deactivate','Aktife Al':'Activate','Geri Al':'Undo','Onayla':'Confirm','Tamamla':'Complete','Başlat':'Start','Bitir':'Finish',
      'Kayıt Tarihi':'Record Date','Oluşturan':'Created By','Güncelleyen':'Updated By','Başlatan':'Started By','Bitiren':'Finished By','Çıkışı Yapan':'Exit Recorded By','Girişi Yapan':'Entry Recorded By','Son İşlem':'Last Action','Son Durum':'Latest Status','Önceki Durum':'Previous Status','Yeni Durum':'New Status','Başlangıç':'Start','Bitiş':'End','Başlangıç Saati':'Start Time','Bitiş Saati':'End Time','Tahmini Çıkış':'Estimated Exit','Tahmini Bitiş':'Estimated Finish','Süre':'Duration','Saat':'Hour','Gün':'Day',
      'Mazot':'Diesel','Yakıt Türü':'Fuel Type','Yakıt Tutarı':'Fuel Amount','Yakıt Litresi':'Fuel Liters','Yakıt Birim Fiyatı':'Fuel Unit Price','Toplam Yakıt Tutarı':'Total Fuel Amount','Resmi':'Official','Ticari':'Commercial','Bağdat':'Baghdad','Depo':'Tank','Başlangıç LT':'Starting Liters','Bitiş LT':'Ending Liters','Alınan LT':'Purchased Liters','Toplam LT':'Total Liters',
      'Şoföre Verilen':'Given to Driver','Şoförden Gelen':'Returned by Driver','Şoför Bakiyesi':'Driver Balance','Müşteri Alacağı':'Customer Receivable','Tahsil Edilen':'Collected','Tahsil Edilecek':'To Be Collected','Kalan Alacak':'Remaining Receivable','Ödenen':'Paid','Ödenecek':'To Be Paid','Bakiye':'Balance','Fark':'Difference','Eksik':'Shortage','Fazla':'Surplus','Eksik Tutar':'Short Amount','Fazla Tutar':'Excess Amount','Nakit':'Cash','Nakit Giriş':'Cash In','Nakit Çıkış':'Cash Out','Fiili Sayım':'Physical Count','Kasa Sayımı':'Cash Count','Kasa Bakiyesi':'Cash Balance','Gün Sonu':'End of Day','Gün Sonu Kasa':'End-of-Day Cash','Açılış Bakiyesi':'Opening Balance','Kapanış Bakiyesi':'Closing Balance',
      'Avans Takibi':'Advance Tracking','Yeni Avans':'New Advance','Avans Ekle':'Add Advance','Avans Alan':'Advance Recipient','Avans Tarihi':'Advance Date','Avans Açıklaması':'Advance Description','Belgelendirilen Harcama':'Documented Expense','Nakit İade / Mahsup':'Cash Return / Offset','Kalan Bakiye':'Remaining Balance','Hesap Kapat':'Close Account','Hesaplaşma Ekle':'Add Settlement','Toplu Fiş/Fatura Girişi':'Bulk Receipt/Invoice Entry','Belge Tarihi':'Document Date','Belge Açıklaması':'Document Description','Makbuz':'Receipt','FİŞ':'RECEIPT','FATURA':'INVOICE','NAKİT İADE':'CASH RETURN','MAHSUP':'OFFSET','USTA':'TECHNICIAN','ŞOFÖR':'DRIVER','PERSONEL':'PERSONNEL','DİĞER':'OTHER',
      'Bakım Geçmişi':'Maintenance History','Bakımda Olan Araçlar':'Vehicles in Maintenance','Bakım Girişi':'Maintenance Entry','Bakım Çıkışı':'Maintenance Exit','Bakım Süresi':'Maintenance Duration','Tahmini Bakım Süresi':'Estimated Maintenance Duration','Bakım Notu':'Maintenance Note','Yapılacak İş':'Work to Do','Yapılan İş':'Work Performed','Bakımı Bitir':'Finish Maintenance','Bakım Kaydını Düzenle':'Edit Maintenance Record',
      'Gemi Adı':'Vessel Name','Yükleme Durumu':'Loading Status','Yük Durumu':'Load Status','Sıra':'Queue','Yükleme Sırasında':'In Loading Queue','Yüklemede':'Loading','Yüklendi':'Loaded','Boş Araç':'Empty Vehicle','Dolu Araç':'Loaded Vehicle','Operasyon Açıklaması':'Operation Description','Operasyon Durumu':'Operation Status',
      'Rota':'Route','Rotalar':'Routes','Rota Standardı':'Route Standard','Beklenen KM':'Expected KM','Beklenen Saat':'Expected Hours','Beklenen Tüketim':'Expected Consumption','Tolerans':'Tolerance','Tüketim':'Consumption','Ortalama':'Average','Ortalama Tüketim':'Average Consumption','Sapma':'Deviation','Risk':'Risk','Yüksek Risk':'High Risk','Normal':'Normal',
      'Veri Kalitesi':'Data Quality','Eksik Veri':'Missing Data','Hatalı Veri':'Invalid Data','Mükerrer':'Duplicate','Mükerrer Kayıt':'Duplicate Record','Kontrol Et':'Check','Kontrol Sonucu':'Check Result','Uygun':'Valid','Uygun Değil':'Invalid','Zorunlu':'Required','Seçim Yapın':'Make a selection','Lütfen bir seçim yapın':'Please make a selection','Lütfen zorunlu alanları doldurun':'Please fill in required fields','İşlem başarılı':'Operation successful','Kayıt başarıyla oluşturuldu':'Record created successfully','Kayıt başarıyla güncellendi':'Record updated successfully','Kayıt silindi':'Record deleted','Bir hata oluştu':'An error occurred','Sunucu hatası':'Server error','Bağlantı hatası':'Connection error','Yetkiniz yok':'Permission denied','Emin misiniz?':'Are you sure?','Bu işlem geri alınamaz':'This action cannot be undone'
    },
    ar:{
      'Ana Sayfa Gör':'عرض الرئيسية','Sevkiyat Gör':'عرض الشحنات','Sevkiyat / Çıkış / Giriş Düzenle':'تعديل الشحنة / الخروج / الدخول','Sevkiyat Sil':'حذف الشحنة','Excel / OneDrive İçe Aktar':'استيراد Excel / OneDrive','Filo Gör':'عرض الأسطول','Filo Düzenle':'تعديل الأسطول','Şoförler Gör':'عرض السائقين','Şoförler Düzenle':'تعديل السائقين','Bakım Gör':'عرض الصيانة','Bakım Düzenle':'تعديل الصيانة','Raporlar Gör':'عرض التقارير','Rapor / Excel Dışa Aktar':'تصدير التقرير / Excel','Kullanıcılar & Yetkiler Yönet':'إدارة المستخدمين والصلاحيات',
      'İlk Admin Hesabını Oluştur':'إنشاء حساب المدير الأول','Ad Soyad':'الاسم الكامل','Admin Oluştur':'إنشاء مدير','Giriş Yap':'تسجيل الدخول','Oturum gerekli.':'الجلسة مطلوبة.','Bu işlem için yetkiniz yok':'ليست لديك صلاحية لهذا الإجراء',
      'YÜKLEME & SEVKİYAT':'التحميل والشحن','FİLO & OPERASYON':'الأسطول والعمليات','RAPORLAMA':'التقارير','SİSTEM':'النظام','YÖNETİM':'الإدارة','OPERASYON MERKEZİ':'مركز العمليات','FİLO YÖNETİMİ':'إدارة الأسطول','ŞOFÖR YÖNETİMİ':'إدارة السائقين','BAKIM YÖNETİMİ':'إدارة الصيانة',
      'Yeni Çıkış':'خروج جديد','Yeni Giriş':'دخول جديد','Çıkış Kaydı':'سجل الخروج','Giriş Kaydı':'سجل الدخول','Çıkışı Kaydet':'حفظ الخروج','Girişi Kaydet':'حفظ الدخول','Sevkiyat Bilgileri':'معلومات الشحنة','Çıkış Bilgileri':'معلومات الخروج','Giriş Bilgileri':'معلومات الدخول','Araç ve Şoför':'المركبة والسائق','Teslimat Bilgileri':'معلومات التسليم','Yakıt ve Masraflar':'الوقود والمصاريف','Mali Bilgiler':'المعلومات المالية','Ekstra Masraflar':'مصاريف إضافية','Dönüş Bilgileri':'معلومات العودة','Dönüş Harcamaları':'مصاريف العودة','Müşteri Tahsilatı':'تحصيل العميل','Şoförün Getirdiği Nakit':'النقد الذي أعاده السائق','Toplam Yakıt':'إجمالي الوقود','Toplam Masraf':'إجمالي المصاريف','Toplam Tutar':'إجمالي المبلغ','Toplam Araç':'إجمالي المركبات','Toplam Şoför':'إجمالي السائقين','Toplam Sevkiyat':'إجمالي الشحنات','Toplam Avans':'إجمالي السلف','Toplam Harcama':'إجمالي المصاريف','Toplam Kalan':'إجمالي المتبقي',
      'GEMİDE ÇALIŞAN':'يعمل على السفينة','GEMİDE':'على السفينة','BOŞALTILDI':'تم التفريغ','DÖNÜYOR':'عائد','BEKLEME':'انتظار','BAKIM':'صيانة','YOLDA':'على الطريق','BOŞTA':'متاح','AKTİF':'نشط','PASİF':'غير نشط','TAMAMLANDI':'مكتمل','BEKLİYOR':'قيد الانتظار','İPTAL':'ملغي',
      'Durum Özeti':'ملخص الحالة','Filo Özeti':'ملخص الأسطول','Operasyon Özeti':'ملخص العمليات','Günlük Özet':'الملخص اليومي','Aktif Sevkiyatlar':'الشحنات النشطة','Son Sevkiyatlar':'أحدث الشحنات','Son İşlemler':'آخر العمليات','Bekleyen İşler':'الأعمال المعلقة','Uyarılar':'التنبيهات','Anomaliler':'الحالات غير الطبيعية','Performans':'الأداء','Yakıt Performansı':'أداء الوقود','Rota Performansı':'أداء المسار',
      'Plaka Ara':'بحث برقم المركبة','Şoför Ara':'بحث عن سائق','SCNA Ara':'بحث SCNA','Müşteri Ara':'بحث عن عميل','Bölge Ara':'بحث عن منطقة','Kayıt Ara':'بحث في السجلات','Arama':'بحث','Filtre':'تصفية','Filtreleri Temizle':'مسح عوامل التصفية','Sonuçları Göster':'عرض النتائج','Listeyi Yenile':'تحديث القائمة','Listeye Dön':'العودة للقائمة','Detayları Gör':'عرض التفاصيل','Görüntüle':'عرض','Düzelt':'تصحيح','Güncelle':'تحديث','Ekle':'إضافة','Kaldır':'إزالة','Pasife Al':'تعطيل','Aktife Al':'تفعيل','Geri Al':'تراجع','Onayla':'تأكيد','Tamamla':'إكمال','Başlat':'بدء','Bitir':'إنهاء',
      'Kayıt Tarihi':'تاريخ السجل','Oluşturan':'أنشأ بواسطة','Güncelleyen':'حدّث بواسطة','Başlatan':'بدأ بواسطة','Bitiren':'أنهى بواسطة','Son İşlem':'آخر إجراء','Son Durum':'آخر حالة','Önceki Durum':'الحالة السابقة','Yeni Durum':'الحالة الجديدة','Başlangıç':'البداية','Bitiş':'النهاية','Süre':'المدة','Saat':'ساعة','Gün':'يوم',
      'Mazot':'ديزل','Yakıt Türü':'نوع الوقود','Yakıt Tutarı':'مبلغ الوقود','Yakıt Litresi':'لترات الوقود','Yakıt Birim Fiyatı':'سعر وحدة الوقود','Toplam Yakıt Tutarı':'إجمالي مبلغ الوقود','Resmi':'رسمي','Ticari':'تجاري','Bağdat':'بغداد','Depo':'الخزان','Başlangıç LT':'لترات البداية','Bitiş LT':'لترات النهاية','Alınan LT':'اللترات المشتراة','Toplam LT':'إجمالي اللترات',
      'Şoföre Verilen':'المعطى للسائق','Şoförden Gelen':'المعاد من السائق','Şoför Bakiyesi':'رصيد السائق','Müşteri Alacağı':'مستحقات العميل','Tahsil Edilen':'المحصّل','Tahsil Edilecek':'المطلوب تحصيله','Kalan Alacak':'المستحق المتبقي','Ödenen':'المدفوع','Ödenecek':'المطلوب دفعه','Bakiye':'الرصيد','Fark':'الفرق','Eksik':'نقص','Fazla':'زيادة','Nakit':'نقد','Nakit Giriş':'دخول نقدي','Nakit Çıkış':'خروج نقدي','Fiili Sayım':'الجرد الفعلي','Kasa Sayımı':'جرد الصندوق','Kasa Bakiyesi':'رصيد الصندوق','Gün Sonu':'نهاية اليوم','Açılış Bakiyesi':'رصيد الافتتاح','Kapanış Bakiyesi':'رصيد الإغلاق',
      'Avans Takibi':'متابعة السلف','Yeni Avans':'سلفة جديدة','Avans Ekle':'إضافة سلفة','Avans Alan':'مستلم السلفة','Avans Tarihi':'تاريخ السلفة','Avans Açıklaması':'وصف السلفة','Belgelendirilen Harcama':'المصروف الموثق','Nakit İade / Mahsup':'إعادة نقد / تسوية','Kalan Bakiye':'الرصيد المتبقي','Hesap Kapat':'إغلاق الحساب','Hesaplaşma Ekle':'إضافة تسوية','Toplu Fiş/Fatura Girişi':'إدخال جماعي للإيصالات/الفواتير','Belge Tarihi':'تاريخ المستند','Belge Açıklaması':'وصف المستند','FİŞ':'إيصال','FATURA':'فاتورة','NAKİT İADE':'إعادة نقد','MAHSUP':'تسوية','USTA':'فني','ŞOFÖR':'سائق','PERSONEL':'موظف','DİĞER':'أخرى',
      'Bakım Geçmişi':'سجل الصيانة','Bakımda Olan Araçlar':'المركبات في الصيانة','Bakım Girişi':'دخول الصيانة','Bakım Çıkışı':'خروج الصيانة','Bakım Süresi':'مدة الصيانة','Tahmini Bakım Süresi':'مدة الصيانة التقديرية','Bakım Notu':'ملاحظة الصيانة','Yapılacak İş':'العمل المطلوب','Yapılan İş':'العمل المنفذ','Bakımı Bitir':'إنهاء الصيانة','Bakım Kaydını Düzenle':'تعديل سجل الصيانة',
      'Gemi Adı':'اسم السفينة','Yükleme Durumu':'حالة التحميل','Yük Durumu':'حالة الحمولة','Sıra':'الدور','Yükleme Sırasında':'في طابور التحميل','Yüklemede':'قيد التحميل','Yüklendi':'تم التحميل','Boş Araç':'مركبة فارغة','Dolu Araç':'مركبة محملة','Operasyon Açıklaması':'وصف العملية','Operasyon Durumu':'حالة العملية',
      'Rota':'المسار','Rotalar':'المسارات','Rota Standardı':'معيار المسار','Beklenen KM':'المسافة المتوقعة','Beklenen Saat':'الساعات المتوقعة','Beklenen Tüketim':'الاستهلاك المتوقع','Tolerans':'السماحية','Tüketim':'الاستهلاك','Ortalama':'المتوسط','Ortalama Tüketim':'متوسط الاستهلاك','Sapma':'الانحراف','Risk':'المخاطر','Yüksek Risk':'مخاطر عالية','Normal':'طبيعي',
      'Veri Kalitesi':'جودة البيانات','Eksik Veri':'بيانات ناقصة','Hatalı Veri':'بيانات خاطئة','Mükerrer':'مكرر','Mükerrer Kayıt':'سجل مكرر','Kontrol Et':'تحقق','Kontrol Sonucu':'نتيجة التحقق','Uygun':'صحيح','Uygun Değil':'غير صحيح','Zorunlu':'إلزامي','Seçim Yapın':'اختر','Lütfen bir seçim yapın':'يرجى الاختيار','Lütfen zorunlu alanları doldurun':'يرجى ملء الحقول المطلوبة','İşlem başarılı':'تمت العملية بنجاح','Kayıt başarıyla oluşturuldu':'تم إنشاء السجل بنجاح','Kayıt başarıyla güncellendi':'تم تحديث السجل بنجاح','Kayıt silindi':'تم حذف السجل','Bir hata oluştu':'حدث خطأ','Sunucu hatası':'خطأ في الخادم','Bağlantı hatası':'خطأ في الاتصال','Yetkiniz yok':'لا توجد صلاحية','Emin misiniz?':'هل أنت متأكد؟','Bu işlem geri alınamaz':'لا يمكن التراجع عن هذا الإجراء'
    }
  };
  Object.assign(DICT.en,EXTRA.en); Object.assign(DICT.ar,EXTRA.ar);
  // SAMA_I18N_EXHAUSTIVE_V7B
  const UI7B={en:{
'Akıllı Anormallik Merkezi':'Smart Anomaly Center','Araç Yönetimi':'Vehicle Management','Araç Getir':'Load Vehicle','Araç Kartı':'Vehicle Card','Araç Performansı':'Vehicle Performance','Aylık Özet':'Monthly Summary','Bakım Takibi':'Maintenance Tracking','Bakıma Yakın':'Maintenance Due Soon','Bakımı Geçmiş':'Maintenance Overdue','Bağlantıyı Test Et':'Test Connection','Canlı Operasyon':'Live Operations','DB Kaydını Aç':'Open DB Record','DURUM / DEĞİŞTİR':'STATUS / CHANGE','Durum Değiştir':'Change Status','Düzeltilecek Alan':'Field to Correct','Düzeltmeyi Kaydet':'Save Correction','Efektif Çıkış Tarihi':'Effective Exit Date','Eksik Veri Merkezi':'Missing Data Center','EKSİK SAYISI':'MISSING COUNT','EXCEL GÜNCELLEMESİNE AÇIK':'OPEN TO EXCEL UPDATE','Excel Korumasını Aç':'Enable Excel Protection','Excel Korumasını Kaldır':'Remove Excel Protection','Excel ↔ Veritabanı Karşılaştırma':'Excel ↔ Database Comparison','FATURA / FİŞ / İADE':'INVOICE / RECEIPT / RETURN','FAZLA PARA':'EXCESS CASH','Fiilen Saydığım Kasa':'Physical Cash Count','Finans / Kârlılık':'Finance / Profitability','Gecikmiş Sefer':'Delayed Trip','Gemi Atanmamış':'Vessel Not Assigned','GERÇEK KM':'ACTUAL KM','GERÇEK LT/100':'ACTUAL L/100','Giriş İşlemini Geri Al':'Undo Entry Operation','Giriş Tarihi Var / KM Yok':'Entry Date Exists / KM Missing','Giriş Tarihlerini Excel’den Onar':'Repair Entry Dates from Excel','GÖRÜNEN GİDER TOPLAMI':'VISIBLE EXPENSE TOTAL','Gün Sonu / Yönetici Özeti':'End of Day / Manager Summary','Güncel Bakiye / Borç':'Current Balance / Debt','Günlük Doğrudan Harcama':'Daily Direct Expense','Günlük Kasa Kontrolü':'Daily Cash Control','HATA NEDENİ':'ERROR REASON','Hatalı SCNA Listesini Excel İndir':'Download Invalid SCNA List as Excel','İş / Avans Nedeni':'Job / Advance Reason','İşlem Geçmişi / Denetim':'Audit Log','İşlem Sayısı':'Action Count','İşlem Türü':'Action Type','Kalan Açık Bakiye':'Remaining Open Balance','KALAN GÜN':'DAYS LEFT','Kasa Farkı':'Cash Difference','Kasa Teslim':'Cash Handover','KM ANORMALLİKLERİ':'KM ANOMALIES','KM Uyarısı':'KM Warning','Kontrol Et ve Kaydet':'Check and Save','Kâr Marjı':'Profit Margin','MAL CİNSİ':'CARGO TYPE','MAZOT TAKIP':'FUEL TRACKING','MAZOT TOPLAM':'TOTAL FUEL','Muhasebeye Düşmeyen Dönüşler':'Returns Missing from Accounting','MUHASEBEYE GİRER':'INCLUDED IN ACCOUNTING','MUHASEBEYE GİRMEZ':'NOT INCLUDED IN ACCOUNTING','MÜŞTERİDEN KALAN':'CUSTOMER BALANCE','MÜŞTERİDEN TAHSİL EDİLEN':'COLLECTED FROM CUSTOMER','Operasyon Uyarıları':'Operation Alerts','ORT. GİDER':'AVG. EXPENSE','ORT. SEFER SÜRESİ':'AVG. TRIP DURATION','ORT. SÜRE (SAAT)':'AVG. DURATION (HOURS)','Para Eksiği':'Cash Shortage','Para Farkı':'Cash Difference','PDF / Yazdır':'PDF / Print','Performans Raporları':'Performance Reports','PLAKA SEÇİLMEDİ':'NO PLATE SELECTED','Profilim / Şifre Değiştir':'My Profile / Change Password','Rota Standartları':'Route Standards','ROTA YAKIT UYARISI':'ROUTE FUEL WARNING','SCNA Düzenleme Merkezi':'SCNA Editing Center','SCNA Kasa Tanı':'SCNA Cash Diagnostics','SCNA Tanı':'SCNA Diagnostics','Sefer Kârı':'Trip Profit','Sefer Sayısı':'Trip Count','SEFERİ KAYDET':'SAVE TRIP','Sevkiyat / Yakıt Tüketim / Para Hesabı':'Shipment / Fuel Consumption / Cash Calculation','SEVKİYAT SONUCU MÜŞTERİDEN ALINACAK KALAN PARA':'REMAINING AMOUNT RECEIVABLE FROM CUSTOMER AFTER SHIPMENT','Sıra Bekliyor':'Waiting in Queue','Şoförün Teslim Ettiği Para':'Cash Handed Over by Driver','Tüm Araçlar':'All Vehicles','Tüm Durumlar':'All Statuses','Tüm Kişiler':'All People','Tüm Şoförler':'All Drivers','Verilen Avans':'Advance Given','YAKITI KAYDET':'SAVE FUEL','Yakıt Ekle':'Add Fuel','YENİ SEVKİYAT':'NEW SHIPMENT','ÇIKIŞ TARİHİ':'EXIT DATE','GİRİŞ TARİHİ':'ENTRY DATE','TİP':'TYPE','NAVLUN':'FREIGHT','GİDER':'EXPENSE','NET KÂR':'NET PROFIT','KÂR MARJI':'PROFIT MARGIN','TOPLAM KM':'TOTAL KM','BİRİM':'UNIT','BİRİM FİYAT':'UNIT PRICE','EKSİK PARA':'CASH SHORTAGE','GECİKMİŞ':'DELAYED','AKTİF ARAÇ':'ACTIVE VEHICLE','AKTİF ŞOFÖR':'ACTIVE DRIVER','GEMİ ATANMADI':'NO VESSEL ASSIGNED','KİŞİ':'PERSON','KULLANICI':'USER','MÜKERRER':'DUPLICATE'
},ar:{
'Akıllı Anormallik Merkezi':'مركز الشذوذ الذكي','Araç Yönetimi':'إدارة المركبات','Aylık Özet':'الملخص الشهري','Bakım Takibi':'متابعة الصيانة','Bakıma Yakın':'الصيانة قريبة','Bakımı Geçmiş':'الصيانة متأخرة','Bağlantıyı Test Et':'اختبار الاتصال','Canlı Operasyon':'العمليات المباشرة','DURUM / DEĞİŞTİR':'الحالة / تغيير','Durum Değiştir':'تغيير الحالة','Düzeltilecek Alan':'الحقل المراد تصحيحه','Düzeltmeyi Kaydet':'حفظ التصحيح','Eksik Veri Merkezi':'مركز البيانات الناقصة','EXCEL GÜNCELLEMESİNE AÇIK':'مفتوح لتحديث Excel','Excel Korumasını Aç':'تفعيل حماية Excel','Excel Korumasını Kaldır':'إزالة حماية Excel','FATURA / FİŞ / İADE':'فاتورة / إيصال / إرجاع','FAZLA PARA':'نقد زائد','Finans / Kârlılık':'المالية / الربحية','Gecikmiş Sefer':'رحلة متأخرة','Gemi Atanmamış':'لم يتم تعيين سفينة','GERÇEK KM':'الكيلومترات الفعلية','Giriş İşlemini Geri Al':'التراجع عن عملية الدخول','Giriş Tarihi Var / KM Yok':'تاريخ الدخول موجود / العداد مفقود','GÖRÜNEN GİDER TOPLAMI':'إجمالي المصروف الظاهر','Gün Sonu / Yönetici Özeti':'نهاية اليوم / ملخص المدير','Günlük Doğrudan Harcama':'مصروف نقدي مباشر يومي','Günlük Kasa Kontrolü':'مراقبة الصندوق اليومي','HATA NEDENİ':'سبب الخطأ','İş / Avans Nedeni':'سبب العمل / السلفة','İşlem Geçmişi / Denetim':'سجل التدقيق','İşlem Sayısı':'عدد العمليات','İşlem Türü':'نوع العملية','Kalan Açık Bakiye':'الرصيد المفتوح المتبقي','Kasa Farkı':'فرق الصندوق','Kasa Teslim':'تسليم النقد','KM ANORMALLİKLERİ':'شذوذات الكيلومترات','KM Uyarısı':'تحذير الكيلومترات','Kontrol Et ve Kaydet':'تحقق واحفظ','Kâr Marjı':'هامش الربح','MAL CİNSİ':'نوع الحمولة','MAZOT TAKIP':'متابعة الوقود','MAZOT TOPLAM':'إجمالي الوقود','Muhasebeye Düşmeyen Dönüşler':'العودات غير المسجلة في المحاسبة','MUHASEBEYE GİRER':'يدخل المحاسبة','MUHASEBEYE GİRMEZ':'لا يدخل المحاسبة','MÜŞTERİDEN KALAN':'المتبقي على العميل','MÜŞTERİDEN TAHSİL EDİLEN':'المحصل من العميل','Operasyon Uyarıları':'تنبيهات العمليات','ORT. GİDER':'متوسط المصروف','ORT. SEFER SÜRESİ':'متوسط مدة الرحلة','Para Eksiği':'نقص نقدي','Para Farkı':'فرق نقدي','PDF / Yazdır':'PDF / طباعة','Performans Raporları':'تقارير الأداء','PLAKA SEÇİLMEDİ':'لم يتم اختيار لوحة','Profilim / Şifre Değiştir':'ملفي / تغيير كلمة المرور','Rota Standartları':'معايير المسار','ROTA YAKIT UYARISI':'تحذير وقود المسار','SCNA Düzenleme Merkezi':'مركز تعديل SCNA','SCNA Kasa Tanı':'تشخيص صندوق SCNA','SCNA Tanı':'تشخيص SCNA','Sefer Kârı':'ربح الرحلة','Sefer Sayısı':'عدد الرحلات','SEFERİ KAYDET':'حفظ الرحلة','Sevkiyat / Yakıt Tüketim / Para Hesabı':'الشحنة / استهلاك الوقود / حساب النقد','Sıra Bekliyor':'بانتظار الدور','Şoförün Teslim Ettiği Para':'النقد الذي سلّمه السائق','Tüm Araçlar':'كل المركبات','Tüm Durumlar':'كل الحالات','Tüm Kişiler':'كل الأشخاص','Tüm Şoförler':'كل السائقين','Verilen Avans':'السلفة المعطاة','YAKITI KAYDET':'حفظ الوقود','Yakıt Ekle':'إضافة وقود','YENİ SEVKİYAT':'شحنة جديدة','ÇIKIŞ TARİHİ':'تاريخ الخروج','GİRİŞ TARİHİ':'تاريخ الدخول','TİP':'النوع','NAVLUN':'أجرة النقل','GİDER':'المصروف','NET KÂR':'صافي الربح','KÂR MARJI':'هامش الربح','TOPLAM KM':'إجمالي الكيلومترات','BİRİM':'الوحدة','BİRİM FİYAT':'سعر الوحدة','EKSİK PARA':'نقص نقدي','GECİKMİŞ':'متأخر','AKTİF ARAÇ':'مركبة نشطة','AKTİF ŞOFÖR':'سائق نشط','GEMİ ATANMADI':'لم تعين سفينة','KİŞİ':'الشخص','KULLANICI':'المستخدم','MÜKERRER':'مكرر'
}};
  for(const l of ['en','ar']) DICT[l]=Object.assign(DICT[l]||{},UI7B[l]||{});
  const WORD7B={en:{'Kaydet':'Save','Kayıt':'Record','Araç':'Vehicle','Şoför':'Driver','Giriş':'Entry','Çıkış':'Exit','Tarih':'Date','Saat':'Time','Durum':'Status','Bakım':'Maintenance','Mazot':'Fuel','Yakıt':'Fuel','Harcama':'Expense','Gider':'Expense','Toplam':'Total','Eksik':'Missing','Fazla':'Excess','Beklenen':'Expected','Gerçek':'Actual','Fiili':'Physical','Kasa':'Cash','Avans':'Advance','Müşteri':'Customer','Bölge':'Area','Gemi':'Vessel','Operasyon':'Operation','Rapor':'Report','Kullanıcı':'User','Yetki':'Permission','Açıklama':'Description','Neden':'Reason','Tutar':'Amount','Bakiye':'Balance','Kalan':'Remaining','Başlangıç':'Start','Bitiş':'End','Günlük':'Daily','Aylık':'Monthly','Aktif':'Active','Pasif':'Inactive','Tamamlandı':'Completed','Bekliyor':'Waiting','Düzelt':'Correct','Düzenle':'Edit','Sil':'Delete','Ekle':'Add','Ara':'Search','Seç':'Select','Tüm':'All','Para':'Cash','Hata':'Error','Uyarı':'Warning','Veri':'Data','Fark':'Difference','Sefer':'Trip','Sevkiyat':'Shipment','Yükleme':'Loading','Sıra':'Queue','Gün':'Day','Kişi':'Person','Belge':'Document','Fatura':'Invoice','Fiş':'Receipt','İade':'Return','Koruma':'Protection','Karşılaştırma':'Comparison','Performans':'Performance','Rota':'Route','Özet':'Summary','Yönetim':'Management','Kontrol':'Control','İşlem':'Action','Yeni':'New','Eski':'Old'},ar:{'Kaydet':'حفظ','Kayıt':'سجل','Araç':'المركبة','Şoför':'السائق','Giriş':'دخول','Çıkış':'خروج','Tarih':'التاريخ','Saat':'الوقت','Durum':'الحالة','Bakım':'الصيانة','Mazot':'الوقود','Yakıt':'الوقود','Harcama':'المصروف','Gider':'المصروف','Toplam':'الإجمالي','Eksik':'ناقص','Fazla':'زائد','Beklenen':'المتوقع','Gerçek':'الفعلي','Fiili':'الفعلي','Kasa':'الصندوق','Avans':'السلفة','Müşteri':'العميل','Bölge':'المنطقة','Gemi':'السفينة','Operasyon':'العملية','Rapor':'التقرير','Kullanıcı':'المستخدم','Yetki':'الصلاحية','Açıklama':'الوصف','Neden':'السبب','Tutar':'المبلغ','Bakiye':'الرصيد','Kalan':'المتبقي','Başlangıç':'البداية','Bitiş':'النهاية','Günlük':'يومي','Aylık':'شهري','Aktif':'نشط','Pasif':'غير نشط','Tamamlandı':'مكتمل','Bekliyor':'بانتظار','Düzelt':'تصحيح','Düzenle':'تعديل','Sil':'حذف','Ekle':'إضافة','Ara':'بحث','Seç':'اختيار','Tüm':'كل','Para':'النقد','Hata':'خطأ','Uyarı':'تحذير','Veri':'بيانات','Fark':'فرق','Sefer':'رحلة','Sevkiyat':'شحنة','Yükleme':'تحميل','Sıra':'دور','Gün':'يوم','Kişi':'شخص','Belge':'مستند','Fatura':'فاتورة','Fiş':'إيصال','İade':'إرجاع','Koruma':'حماية','Karşılaştırma':'مقارنة','Performans':'أداء','Rota':'مسار','Özet':'ملخص','Yönetim':'إدارة','Kontrol':'تحقق','İşlem':'إجراء','Yeni':'جديد','Eski':'قديم'}};
  const originals=new WeakMap();
  const attrOriginals=new WeakMap();
  function lang(){const x=localStorage.getItem('sama_lang');return ['tr','en','ar'].includes(x)?x:'tr'}
  function translateString(value,l,lexical=false){
    const raw=value==null?'':String(value), key=raw.trim();
    if(l==='tr'||!key)return raw;
    const dict=DICT[l]||{};
    if(dict[key]) return raw.replace(key,dict[key]);
    let out=raw;
    for(const k of Object.keys(dict).filter(k=>k.length>=5).sort((a,b)=>b.length-a.length)) if(out.includes(k)) out=out.split(k).join(dict[k]);
    if(lexical && out===raw){
      for(const [k,v] of Object.entries(WORD7B[l]||{}).sort((a,b)=>b[0].length-a[0].length)){
        const re=new RegExp('(^|[^\p{L}])'+k.replace(/[.*+?^${}()|[\]\\]/g,'\\$&')+'(?=$|[^\p{L}])','gu');
        out=out.replace(re,(m,p1)=>p1+v);
      }
    }
    return out;
  }
  function translateNode(node,l){
    if(!node||!node.parentElement||node.parentElement.closest('#samaLangBox,script,style,noscript'))return;
    if(!originals.has(node))originals.set(node,node.nodeValue);
    const tag=(node.parentElement.tagName||'').toUpperCase();
    const lexical=['BUTTON','LABEL','H1','H2','H3','H4','H5','H6','TH','OPTION','LEGEND','SMALL','STRONG','B','A','SPAN'].includes(tag)||(tag==='DIV'&&String(originals.get(node)||'').trim().length<=80);
    node.nodeValue=translateString(originals.get(node),l,lexical);
  }
  function translateAttrs(el,l){
    if(!el||el.closest('#samaLangBox'))return;
    let store=attrOriginals.get(el);if(!store){store={};attrOriginals.set(el,store)}
    for(const a of ['placeholder','title','aria-label','data-label']){
      if(el.hasAttribute(a)){if(!(a in store))store[a]=el.getAttribute(a);el.setAttribute(a,translateString(store[a],l))}
    }
    if(el.tagName==='INPUT'&&['button','submit','reset'].includes((el.type||'').toLowerCase())&&el.hasAttribute('value')){
      if(!('value' in store))store.value=el.getAttribute('value');el.setAttribute('value',translateString(store.value,l));
    }
  }
  function applyLanguage(l){
    l=['tr','en','ar'].includes(l)?l:'tr';
    localStorage.setItem('sama_lang',l);
    document.documentElement.lang=l;document.documentElement.dir=l==='ar'?'rtl':'ltr';
    document.querySelectorAll('#samaLangBox button').forEach(b=>b.classList.toggle('active',b.dataset.lang===l));
    const walker=document.createTreeWalker(document.body,NodeFilter.SHOW_TEXT);let n;
    while((n=walker.nextNode()))translateNode(n,l);
    document.querySelectorAll('input,textarea,button,select,option,[title],[aria-label],[data-label]').forEach(el=>translateAttrs(el,l));
  }
  // SAMA_I18N_SELECTOR_MAIN_PAGE_V3
  // The selector was rendered inside authOverlay, so hiding the login overlay also hid the selector.
  // Move only the selector to body; it then stays visible on the authenticated main UI.
  const samaLangBox=document.getElementById('samaLangBox');
  if(samaLangBox && samaLangBox.parentElement!==document.body){document.body.appendChild(samaLangBox);}
  const _samaAlert=window.alert.bind(window), _samaConfirm=window.confirm.bind(window), _samaPrompt=window.prompt.bind(window);
  window.alert=(m)=>_samaAlert(translateString(m,lang()));
  window.confirm=(m)=>_samaConfirm(translateString(m,lang()));
  window.prompt=(m,d)=>_samaPrompt(translateString(m,lang()),d);
  window.setSamaLanguage=applyLanguage;
  document.querySelectorAll('#samaLangBox button').forEach(b=>b.addEventListener('click',()=>applyLanguage(b.dataset.lang)));
  let timer=null;
  new MutationObserver(()=>{clearTimeout(timer);timer=setTimeout(()=>applyLanguage(lang()),40)}).observe(document.body,{childList:true,subtree:true});
  applyLanguage(lang());
})();
</script>

  <div class="auth-card">
    <div class="auth-logo">SAMA <span>TRACK</span></div>
    <div id="authSetupBox" style="display:none">
      <h3>İlk Admin Hesabını Oluştur</h3>
      <div class="field"><label>Ad Soyad</label><input id="setupFullName"></div>
      <div class="field"><label>Kullanıcı Adı</label><input id="setupUsername"></div>
      <div class="field"><label>Şifre</label><input id="setupPassword" type="password"></div>
      <button class="btn primary" style="width:100%" onclick="createFirstAdmin()">Admin Oluştur</button>
    </div>
    <div id="authLoginBox">
      <h3>Giriş</h3>
      <div class="field"><label>Kullanıcı Adı</label><input id="loginUsername" onkeydown="if(event.key==='Enter')loginPassword.focus()"></div>
      <div class="field"><label>Şifre</label><input id="loginPassword" type="password" onkeydown="if(event.key==='Enter')authLogin()"></div>
      <button class="btn primary" style="width:100%" onclick="authLogin()">Giriş Yap</button>
    </div>
    <div id="authMsg" class="small" style="margin-top:10px"></div>
  </div>
</div>
<aside class="side">
<div class="brand">SAMA <span>TRACK</span></div>
<div class="nav grouped-nav">

<div class="nav-group open">
  <button class="nav-group-title" onclick="toggleNavGroup(this)">GENEL <span>▾</span></button>
  <div class="nav-group-items">
    <button class="active" onclick="show('dash',this)">Ana Sayfa</button>
  </div>
</div>

<div class="nav-group">
  <button class="nav-group-title" onclick="toggleNavGroup(this)">SEVKİYAT <span>▸</span></button>
  <div class="nav-group-items">
    <button onclick="show('trips',this)">Sevkiyatlar</button>
    <button onclick="show('exit',this)">Çıkış İşlemi</button>
    <button onclick="show('entry',this)">Giriş İşlemi</button>
    <button onclick="show('detail',this)">SCNA Detay</button>
    <button onclick="show('editcenter',this)">SCNA Düzenleme</button>
    <button onclick="show('quality',this)">Eksik Veri</button>
    <button onclick="show('compare',this)">Excel ↔ DB</button>
    <button onclick="show('bulkfix',this)">Toplu Düzeltme</button>
  </div>
</div>

<div class="nav-group">
  <button class="nav-group-title" onclick="toggleNavGroup(this)">OPERASYON <span>▸</span></button>
  <div class="nav-group-items">
    <button onclick="show('opcenter',this)">Operasyon Kontrol</button>
    <button onclick="show('liveops',this)">Canlı Operasyon</button>
    <button onclick="show('loadqueue',this)">Yükleme Sırası</button>
    <button onclick="show('vesselops',this)">Gemi Operasyonu</button>
    <button onclick="show('maintops',this)">Bakımda</button>
    <button onclick="show('routes',this)">Rota Standartları</button>
  </div>
</div>

<div class="nav-group">
  <button class="nav-group-title" onclick="toggleNavGroup(this)">FİLO <span>▸</span></button>
  <div class="nav-group-items">
    <button onclick="show('fleetmanage',this)">Filo Yönetimi</button>
    <button onclick="show('vehiclecard',this)">Araç Kartı</button>
    <button onclick="show('maintenance',this)">Bakım Takibi</button>
  </div>
</div>

<div class="nav-group">
  <button class="nav-group-title" onclick="toggleNavGroup(this)">ŞOFÖR <span>▸</span></button>
  <div class="nav-group-items">
    <button onclick="show('drivermanage',this)">Şoför Yönetimi</button>
    <button onclick="show('drivercard',this)">Şoför Kartı</button>
  </div>
</div>

<div class="nav-group">
  <button class="nav-group-title" onclick="toggleNavGroup(this)">RAPORLAR <span>▸</span></button>
  <div class="nav-group-items">
    <button onclick="show('anomaly',this)">Anormallik Merkezi</button>
    <button onclick="show('daily',this)">Gün Sonu Özeti</button>
    <button onclick="show('performance',this)">Performans</button>
    <button onclick="show('finance',this)">Finans / Kârlılık</button>
    <button onclick="show('alerts',this)">Uyarılar</button>
    <button onclick="show('audit',this)">İşlem Geçmişi</button>
  </div>
</div>


<!-- SAMA_FINANCE_NAV_V1 -->
<div class="nav-group">
  <button class="nav-group-title" onclick="toggleNavGroup(this)">MUHASEBE & FİNANS <span>▸</span></button>
  <div class="nav-group-items">
    <button onclick="show('advances',this)">Avans Takip</button>
    <button onclick="show('cashcontrol',this)">Günlük Kasa</button>
  </div>
</div>

<div class="nav-group" id="adminNavGroup" style="display:none">
  <button class="nav-group-title" onclick="toggleNavGroup(this)">YÖNETİM <span>▸</span></button>
  <div class="nav-group-items">
    <button onclick="show('useradmin',this)">Kullanıcılar & Yetkiler</button>
    <button onclick="show('deleted',this)">Silinenler</button>
  </div>
</div>
</div>
</aside>

<main class="main">
<div class="top">
<h2 style="margin:0">Sevkiyat / Yakıt Tüketim / Para Hesabı</h2>
<div style="display:flex;align-items:center;gap:10px">
  <span>V82 ADMIN FIX</span>
  <button class="btn secondary" id="authUserBtn" onclick="openMyProfile()">Kullanıcı</button>
  <button class="btn secondary" onclick="authLogout()">Çıkış</button>
</div>
</div>

<div class="content">

<section id="dash" class="panel active">
<div class="cards">
<div class="card">Toplam Sevkiyat<b id="dTotal">0</b></div>
<div class="dashboard-op-section">
  <div class="cards dashboard-op-cards">
    <div class="card op-click" onclick="openDashboardFleet('GEMIDE')">Gemide Çalışan<b id="dOpVessel">0</b></div>
    <div class="card op-click" onclick="openDashboardFleet('BAKIM')">Bakımda<b id="dOpMaintenance">0</b></div>
    <div class="card op-click" onclick="openDashboardFleet('YOLDA')">Yolda<b id="dOpRoad">0</b></div>
    <div class="card op-click" onclick="openDashboardFleet('DONUYOR')">Boşaltıldı Dönüyor<b id="dOpReturning">0</b></div>
    <div class="card op-click" onclick="openDashboardFleet('BEKLEMEDE')">Beklemede<b id="dOpWaiting">0</b></div>
    <div class="card op-click" onclick="openDashboardFleet('BOSTA')">Boşta<b id="dOpIdle">0</b></div>
  </div>
  <div id="dashFleetDetail" class="dashboard-op-detail" style="display:none"></div>
</div>

<div class="card">Toplam Navlun<b id="dFreight">0</b></div>
<div class="card">Toplam Mazot LT<b id="dFuelLiters">0</b></div>
<div class="card">SEVKİYAT SONUCU MÜŞTERİDEN ALINACAK KALAN PARA<b id="dDue">0</b></div>
<div class="card">Şoför Hesap Farkı<b id="dDiff">0</b></div>
</div>

<div class="table">
<table>
<thead>
<tr>
<th>SCNA</th>
<th>PLAKA</th>
<th>GERÇEK KM</th>
<th>TÜKETİLEN LT</th>
<th>LT/100 KM</th>
<th>KM/LT</th>
<th>TOPLAM NAVLUN</th>
<th>MÜŞTERİDEN TAHSİL EDİLEN</th>
<th>MÜŞTERİDEN KALAN</th>
<th>ŞOFÖR HESAP FARKI</th>
<th>DURUM</th>
</tr>
</thead>
<tbody id="dashRows"></tbody>
</table>
</div>
</section>

<section id="trips" class="panel">
<div class="filterbar">
<input id="q" class="grow" placeholder="SCNA / plaka / şoför ara..." oninput="tripPage=1;loadTrips()">
<select id="statusFilter" onchange="tripPage=1;loadTrips()">
  <option value="Tümü">Tüm Durumlar</option>
  <option value="Bekliyor">Bekliyor</option>
  <option value="Yolda">Yolda</option>
  <option value="Tamamlandı">Tamamlandı</option>
</select>
<select id="tripCustomerFilter" onchange="tripPage=1;loadTrips()"><option value="">Tüm Müşteriler</option></select>
<select id="tripAreaFilter" onchange="tripPage=1;loadTrips()"><option value="">Tüm Bölgeler</option></select>
<select id="tripCargoFilter" onchange="tripPage=1;loadTrips()"><option value="">Tüm Mallar</option></select>
<button class="btn secondary" onclick="clearTripFilters()">Filtreyi Temizle</button>
<input id="excelFile" type="file" accept=".xlsx,.xlsm" style="display:none" onchange="importExcel(this)">
<button class="btn green" onclick="excelImportMode='sync';document.getElementById('excelFile').click()">Excel Kontrol / Senkronize Et</button>
<button class="btn primary" onclick="toggleDirectOneDrive()">OneDrive Linkinden Oku</button>
<div id="directOneDrivePanel" class="column-panel" style="width:min(780px,95vw);padding:14px">
  <b>OneDrive Paylaşım Linkinden Canlı Okuma</b>
  <div class="small" style="margin:6px 0 10px">
    Microsoft Client ID gerekmez. Program yalnızca paylaşılan Excel'i indirmeyi dener; kaynak dosyaya yazmaz.
  </div>
  <div class="field wide">
    <label>OneDrive Excel Linki</label>
    <input id="directOneDriveUrl" value="https://1drv.ms/x/c/81409aab255dc057/IQBiVbpNV1kOQZEV8Vhp0CrgAe4nAh2ez_k4aJbuxv6mfD8?e=FHexAF">
  </div>
  <div class="compact-actions" style="margin-top:10px">
    <button class="btn secondary" onclick="testDirectOneDrive()">Bağlantıyı Test Et</button>
    <button class="btn green" onclick="syncDirectOneDrive()">Excel'i Oku ve Önizle</button>
  </div>
  <div id="directOneDriveInfo" class="calc" style="display:none;margin-top:10px"></div>
</div>

<button class="btn orange" onclick="document.getElementById('entryDateRepairFile').click()">Giriş Tarihlerini Excel'den Onar</button>
<input id="entryDateRepairFile" type="file" accept=".xlsx,.xlsm" style="display:none" onchange="repairEntryDatesFromExcel(this)">

<button class="btn secondary" onclick="exportAllExcel()">Tümünü Excel'e Aktar</button>
<button class="btn secondary" onclick="exportFilteredExcel()">Filtreyi Excel'e Aktar</button>
<button class="btn orange" onclick="document.getElementById('remainRepairFile').click()">REMAIN Toplu Onar</button>
<button class="btn orange" onclick="document.getElementById('freightRepairFile').click()">FREIGHT Toplu Onar</button>
<input id="freightRepairFile" type="file" accept=".xlsx,.xlsm" style="display:none" onchange="repairFreightFromExcel(this)">

<input id="remainRepairFile" type="file" accept=".xlsx,.xlsm" style="display:none" onchange="repairRemainFromExcel(this)">


<div class="column-picker">
  <button class="btn secondary" onclick="toggleColumnPanel(event)">Gösterilecek Bilgiler</button>
  <div id="columnPanel" class="column-panel">
    <label><input type="checkbox" data-col="scna" checked onchange="saveColumnPrefs()">SCNA</label>
    <label><input type="checkbox" data-col="plate" checked onchange="saveColumnPrefs()">Plaka</label>
    <label><input type="checkbox" data-col="driver" checked onchange="saveColumnPrefs()">Şoför</label>
    <label><input type="checkbox" data-col="customer" checked onchange="saveColumnPrefs()">Müşteri</label>
    <label><input type="checkbox" data-col="area" checked onchange="saveColumnPrefs()">Bölge</label>
    <label><input type="checkbox" data-col="vessel" checked onchange="saveColumnPrefs()">Gemi</label>
    <label><input type="checkbox" data-col="cargo" checked onchange="saveColumnPrefs()">Mal Cinsi</label>
    <label><input type="checkbox" data-col="cargo_type" checked onchange="saveColumnPrefs()">Yük Tipi</label>
    <label><input type="checkbox" data-col="kg" checked onchange="saveColumnPrefs()">Net KG</label>
    <label><input type="checkbox" data-col="freight" checked onchange="saveColumnPrefs()">Navlun</label>
    <label><input type="checkbox" data-col="exit_at" checked onchange="saveColumnPrefs()">Çıkış Tarihi</label>
    <label><input type="checkbox" data-col="entry_at" checked onchange="saveColumnPrefs()">Giriş Tarihi</label>
    <label><input type="checkbox" data-col="status" checked onchange="saveColumnPrefs()">Durum</label>
  </div>
</div>

<button class="btn primary" onclick="newTrip()">+ Yeni Sevkiyat</button>
</div>
<div id="excelResult" class="calc" style="display:none;margin-bottom:14px"></div>
<div class="table">
<table>
<thead>
<tr>
<th data-col="scna" class="clickhead" onclick="tripHeadClick(event,\'scna\')">SCNA <span id="sort_scna">↕</span></th>
<th data-col="plate" class="clickhead" onclick="tripHeadClick(event,\'plate\')">PLAKA <span id="sort_plate">↕</span></th>
<th data-col="driver" class="clickhead" onclick="tripHeadClick(event,\'driver\')">ŞOFÖR <span id="sort_driver">↕</span></th>
<th data-col="customer">MÜŞTERİ</th>
<th data-col="area" class="clickhead" onclick="tripHeadClick(event,\'area\')">BÖLGE <span id="sort_area">↕</span></th><th data-col="vessel">GEMİ</th>
<th data-col="cargo">MAL</th>
<th data-col="cargo_type" class="clickhead" onclick="tripHeadClick(event,\'cargo_type\')">TİP <span id="sort_cargo_type">↕</span></th>
<th data-col="kg">KG</th>
<th data-col="freight">NAVLUN</th>
<th data-col="exit_at" class="clickhead" onclick="tripHeadClick(event,\'exit_at\')">ÇIKIŞ TARİHİ <span id="sort_exit_at">↕</span></th>
<th data-col="entry_at" class="clickhead" onclick="tripHeadClick(event,\'entry_at\')">GİRİŞ TARİHİ <span id="sort_entry_at">↕</span></th>
<th data-col="status">DURUM / DEĞİŞTİR</th>
</tr>
</thead>
<tbody id="tripRows"></tbody>
</table>
</div>
<div id="tripHeadMenu" class="head-filter-menu" style="display:none">
  <div id="tripHeadMenuTitle" style="font-weight:800;margin-bottom:8px"></div>
  <button class="head-menu-btn" onclick="applyHeadSort('asc')">▲ Artan Sırala</button>
  <button class="head-menu-btn" onclick="applyHeadSort('desc')">▼ Azalan Sırala</button>
  <div style="border-top:1px solid #e5e7eb;margin:8px 0"></div>
  <input id="tripHeadSearch" placeholder="Bu sütunda ara..." oninput="applyHeadTextFilter()">
  <button class="head-menu-btn" onclick="clearHeadFilter()">Bu Filtreyi Temizle</button>
  <button class="head-menu-btn" onclick="clearAllHeadFilters()">Tüm Başlık Filtrelerini Temizle</button>
</div>
<div class="pager"><button class="btn secondary" onclick="prevTripPage()">← Önceki</button><span id="tripPageInfo">1 / 1</span><button class="btn secondary" onclick="nextTripPage()">Sonraki →</button></div>
</section>

<section id="exit" class="panel">
<div class="filterbar">
<b>Çıkış İşlemi</b>
<input id="exitQ" class="grow" placeholder="SCNA / plaka ara" oninput="loadExit()">
<span id="exitCount" style="font-weight:800;white-space:nowrap;padding:0 8px">0 İŞLEM</span>
<input id="exitFrom" type="date" onchange="applyPanelDateFilter('exit')">
<input id="exitTo" type="date" onchange="applyPanelDateFilter('exit')">
<select id="exitStatusFilter" onchange="loadExit()">
  <option value="Bekliyor">Bekliyor</option>
  <option value="Yolda">Yolda</option>
  <option value="Tümü">Tümü</option>
  <option value="Tamamlandı">Tamamlandı</option>
</select>
<button class="btn secondary" onclick="exitQ.value='';exitStatusFilter.value='Bekliyor';loadExit()">Temizle</button>
<span class="section-note">Çıkış bekleyen kayıtlar varsayılan gelir.</span>
</div>
<div id="exitCashMissingBox" class="calc" style="display:none;border-color:#f59e0b;background:#fffbeb;margin-bottom:14px"></div>
<div class="table">
<table>
<thead>
<tr><th>SCNA</th><th>PLAKA</th><th>NAVLUN</th><th>ÇIKIŞ TARİHİ</th><th>DURUM</th><th>İŞLEM</th></tr>
</thead>
<tbody id="exitRows"></tbody>
</table>
</div>
</section>

<section id="entry" class="panel">
<div class="filterbar">
<b>Giriş İşlemi</b>
<input id="entryQ" class="grow" placeholder="SCNA / plaka ara" oninput="loadEntry()">
<span id="entryCount" style="font-weight:800;white-space:nowrap;padding:0 8px">0 İŞLEM</span>
<input id="entryFrom" type="date" onchange="applyPanelDateFilter('entry')">
<input id="entryTo" type="date" onchange="applyPanelDateFilter('entry')">
<select id="entryStatusFilter" onchange="loadEntry()">
  <option value="Yolda">Yolda</option>
  <option value="Tamamlandı">Tamamlandı</option>
  <option value="Tümü">Tümü</option>
</select>
<button class="btn secondary" onclick="entryQ.value='';entryStatusFilter.value='Yolda';loadEntry()">Temizle</button>
<span class="section-note">Yoldaki kayıtlar varsayılan gelir.</span>
</div>
<div class="table">
<table>
<thead>
<tr><th>SCNA</th><th>PLAKA</th><th>ŞOFÖR</th><th>ÇIKIŞ KM</th><th>ÇIKIŞ TARİHİ</th><th>GİRİŞ TARİHİ</th><th>NAVLUN</th><th>İŞLEM</th></tr>
</thead>
<tbody id="entryRows"></tbody>
</table>
</div>
</section>



<section id="editcenter" class="panel">
<div class="filterbar">
  <b>SCNA Düzenleme Merkezi</b>
  <input id="editScna" class="grow" placeholder="SCNA yaz..." onkeydown="if(event.key==='Enter')loadEditCenter()">
  <button class="btn primary" onclick="loadEditCenter()">Kaydı Getir</button>
</div>
<div id="editCenterContent">
  <div class="empty-state">Düzenlemek için SCNA gir.</div>
</div>
</section>


<section id="quality" class="panel">
<div class="filterbar">
  <b>Eksik Veri Merkezi</b>
  <input id="qualityQ" class="grow" placeholder="SCNA / plaka / şoför ara" oninput="renderQuality()">
<input id="qualityFrom" type="date" onchange="renderQuality()">
<input id="qualityTo" type="date" onchange="renderQuality()">
  <select id="qualityType" onchange="renderQuality()">
    <option value="">Tüm Sorunlar</option>
    <option value="entry_date_no_km">Giriş Tarihi Var / KM Yok</option>
    <option value="exit_km_no_date">Çıkış KM Var / Tarih Yok</option>
    <option value="collection_missing">COLLECTION Eksik</option>
    <option value="driver_missing">Şoför Eksik</option>
    <option value="customer_missing">Müşteri Eksik</option>
    <option value="area_missing">Bölge Eksik</option>
    <option value="cargo_missing">Mal Eksik</option>
  </select>
  <button class="btn secondary" onclick="loadQuality()">Yenile</button>
</div>
<div class="cards">
  <div class="card">Toplam Sorunlu<b id="qualityTotal">0</b></div>
  <div class="card">Giriş Tarihi/KM<b id="qualityEntryKm">0</b></div>
  <div class="card">Collection Eksik<b id="qualityCollection">0</b></div>
  <div class="card">Şoför Eksik<b id="qualityDriver">0</b></div>
</div>
<div class="table"><table><thead><tr>
<th>SCNA</th><th>PLAKA</th><th>ŞOFÖR</th><th>MÜŞTERİ</th><th>BÖLGE</th>
<th>DURUM</th><th>ÇIKIŞ</th><th>GİRİŞ</th><th>SORUN</th><th>İŞLEM</th>
</tr></thead><tbody id="qualityRows"></tbody></table></div>
</section>

<section id="compare" class="panel">
<div class="filterbar">
  <b>Excel ↔ Veritabanı Karşılaştırma</b>
  <input id="compareExcelFile" type="file" accept=".xlsx,.xlsm" onchange="compareExcelDb(this)">
  <select id="compareFilter" onchange="renderCompare()">
    <option value="">Tüm Farklar</option>
    <option value="FARKLI">Aynı SCNA / Değer Farklı</option>
    <option value="EXCEL_VAR_DB_YOK">Excel'de Var / DB'de Yok</option>
    <option value="DB_VAR_EXCEL_YOK">DB'de Var / Excel'de Yok</option>
  </select>
</div>
<div class="cards">
  <div class="card">Değeri Farklı<b id="cmpDifferent">0</b></div>
  <div class="card">Excel'de Var / DB Yok<b id="cmpExcelOnly">0</b></div>
  <div class="card">DB'de Var / Excel Yok<b id="cmpDbOnly">0</b></div>
</div>
<div id="compareContent"></div>
</section>

<section id="bulkfix" class="panel">
<div class="filterbar"><b>Toplu Düzeltme</b><span class="section-note">Her satıra bir SCNA yaz.</span></div>
<div class="grid">
  <div class="field wide"><label>SCNA Listesi</label>
    <textarea id="bulkScnas" rows="12" placeholder="90812&#10;KENMOONYS40"></textarea>
  </div>
  <div class="field"><label>Düzeltilecek Alan</label>
    <select id="bulkField" onchange="refreshBulkValue()">
      <option value="status">Durum</option><option value="customer_id">Müşteri</option>
      <option value="area_id">Bölge</option><option value="cargo_category_id">Mal Cinsi</option>
      <option value="driver_id">Şoför</option><option value="cargo_type">Yük Tipi</option>
      <option value="trip_date">Sevkiyat Tarihi</option>
    </select>
  </div>
  <div class="field"><label>Yeni Değer</label><div id="bulkValueBox"></div></div>
</div>
<div class="compact-actions" style="margin-top:12px">
  <button class="btn orange" onclick="runBulkFix()">Toplu Düzeltmeyi Uygula</button>
</div>
<div id="bulkResult" class="calc" style="display:none;margin-top:12px"></div>
</section>

<section id="detail" class="panel">
<div class="toolbar">
  <input id="detailScna" placeholder="SCNA yaz..." onkeydown="if(event.key==='Enter')loadScnaDetail()">
  <button class="btn primary" onclick="loadScnaDetail()">SCNA Getir</button>
</div>
<div id="detailContent"></div>
</section>


<section id="vehiclecard" class="panel">
<div class="filterbar">
  <b>Araç Kartı</b>
  <input id="vehicleCardPlate" class="grow" placeholder="Plaka yaz..." onkeydown="if(event.key==='Enter')loadVehicleCard()">
  <button class="btn primary" onclick="loadVehicleCard()">Araç Getir</button>
</div>
<div id="vehicleCardContent"></div>
</section>

<section id="drivercard" class="panel">
<div class="filterbar">
  <b>Şoför Kartı</b>
  <input id="driverCardName" class="grow" placeholder="Şoför adı yaz..." onkeydown="if(event.key==='Enter')loadDriverCard()">
  <button class="btn primary" onclick="loadDriverCard()">Şoför Getir</button>
</div>
<div id="driverCardContent"></div>
</section>


<section id="daily" class="panel">
<div class="filterbar">
  <b>Gün Sonu / Yönetici Özeti</b>
  <input id="dailyDate" type="date">
  <button class="btn primary" onclick="loadDailySummary()">Raporu Getir</button>
  <button class="quick-btn" onclick="dailyToday()">Bugün</button>
  <button class="btn secondary" onclick="exportDailyExcel()">Excel</button>
  <button class="btn secondary" onclick="printDailyReport()">PDF / Yazdır</button>
</div>
<div id="dailyReport">
<div class="cards">
  <div class="card">Çıkan Araç<b id="dailyExited">0</b></div>
  <div class="card">Giren Araç<b id="dailyEntered">0</b></div>
  <div class="card">Şu An Yolda<b id="dailyOnroad">0</b></div>
  <div class="card">Gecikmiş<b id="dailyDelayed">0</b></div>
  <div class="card">Günlük Ton<b id="dailyTon">0</b></div>
</div>
<div class="cards">
  <div class="card">Navlun<b id="dailyFreight">0</b></div>
  <div class="card">Tahsilat<b id="dailyCollected">0</b></div>
  <div class="card">Gider<b id="dailyCost">0</b></div>
  <div class="card">Net Kâr<b id="dailyProfit">0</b></div>
  <div class="card">Şoför Para Farkı<b id="dailyDriverDiff">0</b></div>
</div>
<div id="dailyTable"></div>
</div>
</section>

<section id="anomaly" class="panel">
<div class="filterbar">
  <b>Akıllı Anormallik Merkezi</b>
  <select id="anomalyType" onchange="renderAnomalies()">
    <option value="all">Tümü</option>
    <option value="fuel">Yakıt</option>
    <option value="km">KM</option>
    <option value="cash">Para Farkı</option>
    <option value="driver">Tekrarlayan Şoför Para Eksiği</option>
  </select>
  <button class="btn secondary" onclick="loadAnomalies()">Yenile</button>
</div>
<div class="cards">
  <div class="card">Yakıt Uyarısı<b id="anFuelCount">0</b></div>
  <div class="card">KM Uyarısı<b id="anKmCount">0</b></div>
  <div class="card">Para Eksiği<b id="anCashCount">0</b></div>
  <div class="card">Tekrarlayan Şoför<b id="anDriverCount">0</b></div>
</div>
<div id="anomalyContent"></div>
</section>

<section id="performance" class="panel">
<div class="filterbar">
  <b>Performans Raporları</b>
  <input id="perfQ" class="grow" placeholder="Plaka / şoför ara" oninput="renderPerformanceFilter()">
  <button class="btn secondary" onclick="perfQ.value='';renderPerformanceFilter()">Temizle</button>
  <button class="btn secondary" onclick="loadVehiclePerformance()">Araç Performansı</button>
  <button class="btn secondary" onclick="loadDriverPerformance()">Şoför Performansı</button>
</div>
<div id="performanceContent"></div>
</section>



<section id="finance" class="panel">
<div class="filterbar">
  <b>Finans / Kârlılık</b>
  <input id="finFrom" type="date">
  <input id="finTo" type="date">
  <button class="btn primary" onclick="loadFinance()">Uygula</button>
  <button class="quick-btn" onclick="setFinancePreset('today')">Bugün</button>
  <button class="quick-btn" onclick="setFinancePreset('month')">Bu Ay</button>
  <button class="quick-btn" onclick="setFinancePreset('year')">Bu Yıl</button>
  <button class="btn secondary" onclick="clearFinanceDates()">Tüm Zamanlar</button>
</div>

<div class="cards">
  <div class="card">Toplam Navlun<b id="finFreight">0</b></div>
  <div class="card">Toplam Gider<b id="finCost">0</b></div>
  <div class="card">Net Kâr<b id="finProfit">0</b></div>
  <div class="card">Kâr Marjı<b id="finMargin">0%</b></div>
  <div class="card">Toplam Ton<b id="finTon">0</b></div>
</div>

<div class="cards">
  <div class="card">Yakıt Gideri<b id="finFuel">0</b></div>
  <div class="card">Diğer Giderler<b id="finOther">0</b></div>
  <div class="card">Toplam KM<b id="finKm">0</b></div>
  <div class="card">Kâr / KM<b id="finProfitKm">0</b></div>
  <div class="card">Kâr / Ton<b id="finProfitTon">0</b></div>
</div>

<div class="cards">
  <div class="card">Sefer Sayısı<b id="finTrips">0</b></div>
  <div class="card">Tahsil Edilen<b id="finCollected">0</b></div>
  <div class="card">Kalan Alacak<b id="finReceivable">0</b></div>
</div>

<div class="toolbar">
  <button class="btn secondary" onclick="loadFinanceTable('customer')">Müşteriye Göre</button>
  <button class="btn secondary" onclick="loadFinanceTable('area')">Bölgeye Göre</button>
  <button class="btn secondary" onclick="loadFinanceTable('cargo')">Mal Cinsine Göre</button>
  <button class="btn secondary" onclick="loadFinanceMonthly()">Aylık Özet</button>
</div>

<div id="financeTable"></div>
</section>


<section id="opcenter" class="panel">
<div class="filterbar">
  <b>Operasyon Kontrol Merkezi</b>
  <input id="opQ" class="grow" placeholder="SCNA / plaka / şoför / bölge ara" oninput="renderOperationCenter()">
  <select id="opState" onchange="renderOperationCenter()">
    <option value="">Tüm Aktifler</option>
    <option value="Yolda">Yolda</option>
    <option value="Bekliyor">Bekliyor</option>
    <option value="Gecikmiş">Sadece Gecikmiş</option>
  </select>
  <button class="btn secondary" onclick="loadOperationCenter()">Yenile</button>
</div>

<div class="cards">
  <div class="card">Toplam Kayıt<b id="opTotal">0</b></div>
  <div class="card">Yolda<b id="opOnRoad">0</b></div>
  <div class="card">Bekliyor<b id="opWaiting">0</b></div>
  <div class="card">Gecikmiş<b id="opDelayed">0</b></div>
  <div class="card">Bugün Çıkış<b id="opTodayExit">0</b></div>
  <div class="card">Bugün Giriş<b id="opTodayEntry">0</b></div>
</div>

<div class="table"><table><thead><tr>
<th>SCNA</th><th>PLAKA</th><th>ŞOFÖR</th><th>BÖLGE</th><th>DURUM</th>
<th>ÇIKIŞ</th><th>YOLDA SAAT</th><th>ROTA ORT. SAAT</th><th>ROTA ORT. KM</th>
<th>ROTA ORT. LT/100</th><th>GEÇMİŞ SEFER</th><th>UYARI</th>
</tr></thead><tbody id="opRows"></tbody></table></div>

<div class="filterbar" style="margin-top:16px">
  <b>Rota Öğrenme</b>
  <input id="routeLearnQ" class="grow" placeholder="Bölge ara" oninput="renderRouteLearning()">
</div>
<div class="table"><table><thead><tr>
<th>BÖLGE</th><th>SEFER</th><th>ORT. KM</th><th>ORT. SÜRE (SAAT)</th>
<th>ORT. LT/100</th><th>ORT. GİDER</th>
</tr></thead><tbody id="routeLearnRows"></tbody></table></div>

<div class="filterbar" style="margin-top:16px">
  <b>Araç / Şoför Performans Karşılaştırması</b>
  <button class="btn secondary" onclick="renderOpPerformance('vehicle')">Araçlar</button>
  <button class="btn secondary" onclick="renderOpPerformance('driver')">Şoförler</button>
</div>
<div id="opPerformanceContent"></div>
</section>



<section id="fleetmanage" class="panel">
<div class="filterbar">
  <b>Araç Yönetimi</b>
  <input id="fleetManageQ" class="grow" placeholder="Plaka / marka / model ara" oninput="renderFleetManage()">
  <select id="fleetManageState" onchange="renderFleetManage()">
    <option value="">Tüm Araçlar</option><option value="1">Aktif</option><option value="0">Pasif</option>
  </select>
  <button class="btn secondary" onclick="fleetManageQ.value='';fleetManageState.value='';renderFleetManage()">Temizle</button>
  <button class="btn primary" onclick="openFleetVehicle()">+ Araç Ekle</button>
</div>
<div class="cards">
  <div class="card">Toplam Araç<b id="fmTotal">0</b></div>
  <div class="card">Aktif Araç<b id="fmActive">0</b></div>
  <div class="card">Pasif / Çıkmış<b id="fmPassive">0</b></div>
</div>
<div class="table"><table><thead><tr>
<th>PLAKA</th><th>MARKA</th><th>MODEL</th><th>ARAÇ TİPİ</th><th>SON ŞOFÖR</th><th>GARAJ DURUMU</th><th>AKTİF</th><th>NOT</th><th>İŞLEM</th>
</tr></thead><tbody id="fleetManageRows"></tbody></table></div>
</section>

<section id="drivermanage" class="panel">
<div style="margin-bottom:14px"><a href="/driver-bulk" target="_blank" class="btn orange" style="display:inline-block;text-decoration:none">📋 ŞOFÖR LİSTESİNİ TOPLU GÜNCELLE</a></div>
<div class="filterbar">
  <b>Şoför Yönetimi</b>
  <input id="driverManageQ" class="grow" placeholder="Şoför / telefon / D.No ara" oninput="renderDriverManage()">
  <select id="driverManageState" onchange="renderDriverManage()">
    <option value="">Tüm Şoförler</option><option value="1">Aktif</option><option value="0">Pasif</option>
  </select>
  <button class="btn secondary" onclick="driverManageQ.value='';driverManageState.value='';renderDriverManage()">Temizle</button>
  <button class="btn primary" onclick="openFleetDriver()">+ Şoför Ekle</button>
</div>
<div class="cards">
  <div class="card">Toplam Şoför<b id="dmTotal">0</b></div>
  <div class="card">Aktif Şoför<b id="dmActive">0</b></div>
  <div class="card">Pasif / Ayrılmış<b id="dmPassive">0</b></div>
</div>
<div class="table"><table><thead><tr>
<th>ŞOFÖR</th><th>TELEFON</th><th>D.NO</th><th>SON PLAKA</th><th>AKTİF</th><th>NOT</th><th>İŞLEM</th>
</tr></thead><tbody id="driverManageRows"></tbody></table></div>
</section>

<section id="loadqueue" class="panel">
<div class="filterbar"><b>Yükleme Sırası / Araç Durumu</b>
<input id="fleetQ" class="grow" placeholder="Plaka / şoför / SCNA / bölge" oninput="renderFleetOps()">
<select id="fleetState" onchange="renderFleetOps()"><option value="">Tüm Araçlar</option><option value="BOS">Boş</option><option value="DOLU">Dolu</option><option value="SIRA">Sıra Bekliyor</option><option value="YUKLEMEDE">Yüklemede</option></select>
<button class="btn secondary" onclick="fleetQ.value='';fleetState.value='';renderFleetOps()">Temizle</button>
<button class="btn secondary" onclick="loadFleetOps()">Yenile</button>
<button class="btn primary" onclick="resequenceFleet()">Sırayı 1'den Düzenle</button></div>
<div class="cards dashboard-op-cards"><div class="card op-click" onclick="setFleetCardFilter('')">Toplam Araç<b id="fleetTotal">0</b></div><div class="card op-click" onclick="setFleetCardFilter('BOS')">Boş<b id="fleetEmpty">0</b></div><div class="card op-click" onclick="setFleetCardFilter('DOLU')">Dolu<b id="fleetLoaded">0</b></div><div class="card op-click" onclick="setFleetCardFilter('SIRA')">Sıra Bekliyor<b id="fleetWaiting">0</b></div><div class="card op-click" onclick="setFleetCardFilter('YUKLEMEDE')">Yüklemede<b id="fleetLoading">0</b></div></div>
<div class="table"><table><thead><tr><th>SIRA</th><th>PLAKA</th><th>DURUM</th><th>ŞOFÖR</th><th>SCNA</th><th>BÖLGE</th><th>GEMİ</th><th>NOT</th><th>İŞLEM</th></tr></thead><tbody id="fleetRows"></tbody></table></div>
</section>

<section id="vesselops" class="panel">
<div class="filterbar">
  <b>Gemi Operasyonu</b>
  <input id="vesselQ" class="grow" placeholder="Gemi / plaka / şoför ara" oninput="renderVesselOps()">
  <select id="vesselState" onchange="renderVesselOps()">
    <option value="">Tüm Araçlar</option>
    <option value="CALISIYOR">Gemide Çalışıyor</option>
    <option value="CIKTI">Operasyondan Çıktı</option>
  </select>
  <button class="btn secondary" onclick="vesselQ.value='';vesselState.value='';renderVesselOps()">Temizle</button>
  <button class="btn secondary" onclick="loadFleetOps()">Yenile</button>
</div>

<div class="cards">
  <div class="card">Gemide Çalışan<b id="vesselWorkingCount">0</b></div>
  <div class="card">Operasyondan Çıkan<b id="vesselOutCount">0</b></div>
  <div class="card">Gemi Atanmamış<b id="vesselUnassignedCount">0</b></div>
</div>

<div id="vesselSummary" class="cards"></div>

<div class="table"><table><thead><tr>
<th>GEMİ</th><th>PLAKA</th><th>ŞOFÖR</th><th>OPERASYON DURUMU</th><th>NOT</th><th>İŞLEM</th>
</tr></thead><tbody id="vesselRows"></tbody></table></div>
</section>


<section id="maintops" class="panel">
<div class="filterbar">
  <b>🔧 Bakım Operasyonu</b>
  <input id="maintOpsQ" class="grow" placeholder="Plaka / neden / açıklama / yapılan işlem ara" oninput="renderMaintenanceOps()">
  <button class="btn primary" onclick="openMaintenanceStart()">+ ARACI BAKIMA AL</button>
  <button class="btn secondary" onclick="loadMaintenanceOps()">Yenile</button>
</div>
<div class="cards dashboard-op-cards"><div class="card">Şu An Bakımda<b id="maintOpsActive">0</b></div><div class="card">Toplam Bakım Kaydı<b id="maintOpsTotal">0</b></div></div>
<div class="table"><table><thead><tr><th>PLAKA</th><th>DURUM</th><th>BAKIM GİRİŞ</th><th>TAHMİNİ SÜRE</th><th>TAHMİNİ ÇIKIŞ</th><th>ÖNCEKİ DURUM</th><th>BAKIM NEDENİ</th><th>AÇIKLAMA</th><th>YAPILAN İŞLEMLER</th><th>BAKIMA ALAN</th><th>BAKIM ÇIKIŞ</th><th>KAPATAN</th><th>İŞLEM</th></tr></thead><tbody id="maintOpsRows"></tbody></table></div>
</section>

<section id="liveops" class="panel">
<div class="filterbar">
<b>Canlı Operasyon</b>
<input id="liveQ" class="grow" placeholder="SCNA / plaka / şoför / bölge" oninput="renderLiveOps()">
<select id="liveState" onchange="renderLiveOps()">
  <option value="">Tümü</option>
  <option value="Yolda">Yolda</option>
  <option value="Bekliyor">Bekliyor</option>
  <option value="Gecikmiş">Gecikmiş</option>
</select>
<button class="btn secondary" onclick="liveQ.value='';liveState.value='';renderLiveOps()">Temizle</button>
<button class="btn secondary" onclick="loadLiveOps()">Yenile</button>
</div>
<div class="cards dashboard-op-cards">
<div class="card op-click" onclick="setLiveCardFilter('')">Toplam<b id="liveTotal">0</b></div>
<div class="card op-click" onclick="setLiveCardFilter('Yolda')">Yolda<b id="liveOnRoad">0</b></div>
<div class="card op-click" onclick="setLiveCardFilter('Bekliyor')">Bekliyor<b id="liveWaiting">0</b></div>
<div class="card op-click" onclick="setLiveCardFilter('Gecikmiş')">Gecikmiş<b id="liveDelayed">0</b></div>
<div class="card op-click" onclick="setLiveCardFilter('BAKIM')">Bakıma Yakın<b id="liveMaint">0</b></div>
</div>
<div class="table"><table><thead><tr>
<th>SCNA</th><th>PLAKA</th><th>ŞOFÖR</th><th>BÖLGE</th><th>DURUM</th>
<th>ÇIKIŞ</th><th>YOLDA SAAT</th><th>BEKLENEN SAAT</th><th>UYARI</th>
</tr></thead><tbody id="liveRows"></tbody></table></div>
</section>

<section id="routes" class="panel">
<div class="filterbar">
<b>Rota Standartları</b>
<input id="routeQ" class="grow" placeholder="Bölge ara" oninput="renderRouteStandards()">
<button class="btn secondary" onclick="routeQ.value='';renderRouteStandards()">Temizle</button>
<button class="btn secondary" onclick="loadRouteStandards()">Yenile</button>
</div>
<div class="table"><table><thead><tr>
<th>BÖLGE</th><th>BEKLENEN SAAT</th><th>BEKLENEN KM</th><th>HEDEF LT/100</th><th>TOLERANS %</th><th>İŞLEM</th>
</tr></thead><tbody id="routeRows"></tbody></table></div>
</section>

<section id="maintenance" class="panel">
<div class="filterbar">
<b>Bakım Takibi</b>
<input id="maintQ" class="grow" placeholder="Plaka ara" oninput="renderMaintenance()">
<select id="maintState" onchange="renderMaintenance()">
  <option value="">Tümü</option>
  <option value="due">Bakımı Geçmiş</option>
  <option value="soon">2.000 KM İçinde</option>
  <option value="ok">Normal</option>
</select>
<button class="btn secondary" onclick="maintQ.value='';maintState.value='';renderMaintenance()">Temizle</button>
<button class="btn primary" onclick="openMaintenance()">+ Bakım Kaydı</button>
<button class="btn secondary" onclick="loadMaintenance()">Yenile</button>
</div>
<div class="table"><table><thead><tr>
<th>PLAKA</th><th>GÜNCEL KM</th><th>SON BAKIM KM</th><th>SONRAKİ BAKIM KM</th><th>KALAN KM</th>
<th>YAĞ</th><th>LASTİK</th><th>FREN</th><th>MOTOR</th><th>İŞLEM</th>
</tr></thead><tbody id="maintRows"></tbody></table></div>
</section>

<section id="alerts" class="panel">
<div class="toolbar">
  <b>Operasyon Uyarıları</b>
  <button class="btn secondary" onclick="loadAlerts()">Yenile</button>
</div>
<div class="cards">
  <div class="card">Gecikmiş Sefer<b id="alertDelayCount">0</b></div>
  <div class="card">Yakıt Anormalliği<b id="alertFuelCount">0</b></div>
</div>
<div id="alertsContent"></div>
</section>


<section id="useradmin" class="panel">
<div class="filterbar">
  <b>Kullanıcılar & Yetkiler</b>
  <button class="btn primary" onclick="openAuthUserCreate()">+ Kullanıcı Ekle</button>
  <button class="btn secondary" onclick="loadAuthUsers()">Yenile</button>
</div>
<div class="cards">
  <div class="card">Toplam Kullanıcı<b id="authUserTotal">0</b></div>
  <div class="card">Aktif<b id="authUserActive">0</b></div>
</div>
<div class="table"><table><thead><tr>
<th>KULLANICI</th><th>AD SOYAD</th><th>ROL</th><th>DURUM</th><th>ŞİFRE DEĞİŞİMİ</th><th>İŞLEM</th>
</tr></thead><tbody id="authUserRows"></tbody></table></div>
</section>


<section id="advances" class="panel">
<div class="filterbar"><b>Avans Takip</b><input id="advQ" class="grow" data-no-autocomplete="1" autocomplete="off" placeholder="Kişi / plaka / iş / SCNA ara" oninput="renderAdvances()"><select id="advType" onchange="renderAdvances()"><option value="">Tüm Kişiler</option><option value="USTA">Usta</option><option value="SOFOR">Şoför</option><option value="PERSONEL">Personel</option><option value="DIGER">Diğer</option></select><select id="advStatus" onchange="renderAdvances()"><option value="">Tüm Durumlar</option><option value="BORCLU">Borçlu</option><option value="ALACAKLI">Alacaklı</option><option value="KAPANDI">Kapalı</option></select><button class="btn primary" onclick="openAdvanceNew()">+ Yeni Avans</button><button class="btn secondary" onclick="openAdvancePersonAccounts()">KİŞİ / CARİ HESAPLARI</button><button class="btn secondary" onclick="loadAdvances()">Yenile</button></div>
<div class="cards"><div class="card">Açık Kayıt<b id="advOpen">0</b></div><div class="card">Toplam Verilen<b id="advTotal">0</b></div><div class="card">Belgelenen / İade<b id="advSettled">0</b></div><div class="card">Kalan Açık Bakiye<b id="advRemaining">0</b></div></div>
<div id="advPersonSummary" style="display:none;margin:12px 0;padding:12px;border:1px solid var(--border);border-radius:12px"></div>
<div class="table"><table><thead><tr><th>TARİH</th><th>TİP</th><th>KİŞİ</th><th>PLAKA</th><th>SCNA</th><th>İŞ / NEDEN</th><th>VERİLEN</th><th>BELGELENEN/İADE</th><th>KALAN</th><th>PARA</th><th>NOT</th><th>DURUM</th><th>İŞLEM</th></tr></thead><tbody id="advRows"></tbody></table></div>
</section>


<section id="cashcontrol" class="panel">
<div class="filterbar"><b>Günlük Kasa Kontrolü</b><input id="cashDiagScna" placeholder="SCNA Tanı" style="max-width:150px"><button class="btn secondary" onclick="openCashDiagnostic()">SCNA Kasa Tanı</button><input id="cashDay" type="date" onchange="loadCashControl()"><select id="cashCur" onchange="loadCashControl()"><option>IQD</option><option>USD</option><option>TRY</option></select><button class="btn primary" onclick="openCashCount()">Kasa Tutarı Gir / Hesapla</button><button class="btn secondary" onclick="cashAddExpenseRow()">+ Harcama Satırı</button><button class="btn secondary" onclick="for(let i=0;i<10;i++)cashAddExpenseRow()">+ 10 Satır</button></div>
<div class="cards"><div class="card">Gün Başı Kasa<b id="cashOpening">0</b></div><div class="card" role="button" tabindex="0" style="cursor:pointer" onclick="openShipmentCashDetails()" onkeydown="if(event.key==='Enter'||event.key===' '){event.preventDefault();openShipmentCashDetails()}">Sevkiyattan Gelen ↗<b id="cashShipmentIn">0</b></div><div class="card">Diğer Kasa Girişi<b id="cashIn">0</b></div><div class="card" role="button" tabindex="0" style="cursor:pointer" onclick="openShipmentCashOutDetails()">Sevkiyat Çıkışı / Şoföre Verilen ↗<b id="cashShipmentOut">0</b></div><div class="card">Verilen Avans<b id="cashAdvance">0</b></div><div class="card">Günlük Doğrudan Harcama<b id="cashExpense">0</b></div><div class="card">Beklenen Kasa<b id="cashExpected">0</b></div><div class="card">Fiili Kasa<b id="cashActual">0</b></div><div class="card">Kasa Farkı<b id="cashDiff">0</b></div></div>
<div class="section-note">Beklenen Kasa = Gün Başı Kasa + Sevkiyattan Gelen + Diğer Kasa Girişi − Sevkiyat Çıkışında Şoföre Verilen − Avanslar − Günlük Doğrudan Harcamalar. Sevkiyat giriş/çıkış kartlarına tıklayarak SCNA detaylarını görebilirsiniz.</div>
<div class="section-note">Doğrudan kasadan ödenen fiş/faturaları girin. Avans karşılığı getirilen belgeleri buraya tekrar girmeyin; avans zaten kasadan çıkış sayılır.</div>
<div class="table"><table><thead><tr><th>Tarih (boşsa bugün)</th><th>Fiş/Fatura No *</th><th>Tutar *</th><th>Açıklama</th><th></th></tr></thead><tbody id="cashExpenseEntry"></tbody></table></div><div style="margin:10px 0"><button class="btn green" onclick="saveCashExpenses()">Harcama Satırlarını Kaydet</button></div>
<div class="table"><table><thead><tr><th>Tarih</th><th>Fiş/Fatura No</th><th>Tutar</th><th>Açıklama</th><th>Giren</th></tr></thead><tbody id="cashExpenseHistory"></tbody></table></div>
</section>

<section id="deleted" class="panel">
<div class="filterbar"><b>Silinen Sevkiyatlar</b><span class="section-note">7 gün içinde geri alınabilir.</span><button class="btn secondary" onclick="loadDeletedTrips()">Yenile</button></div>
<div class="table"><table><thead><tr><th>SCNA</th><th>PLAKA</th><th>SİLİNME TARİHİ</th><th>SİLEN</th><th>KALAN GÜN</th><th>İŞLEM</th></tr></thead><tbody id="deletedRows"></tbody></table></div>
</section>

<section id="audit" class="panel">
<div class="filterbar">
  <b>İşlem Geçmişi / Denetim</b>
  <input id="auditQ" class="grow" placeholder="İşlem / detay / eski-yeni değer ara" oninput="loadAudit()">
  <select id="auditUser" onchange="loadAudit()"><option value="">Tüm Kullanıcılar</option></select>
  <select id="auditAction" onchange="loadAudit()"><option value="">Tüm İşlemler</option></select>
  <input id="auditScna" placeholder="SCNA" oninput="loadAudit()">
  <input id="auditFrom" type="date" onchange="loadAudit()">
  <input id="auditTo" type="date" onchange="loadAudit()">
  <button class="btn secondary" onclick="clearAuditFilters()">Temizle</button>
  <button class="btn secondary" onclick="loadAudit()">Yenile</button>
</div>
<div class="table"><table><thead><tr>
<th>TARİH/SAAT</th><th>KULLANICI</th><th>İŞLEM</th><th>KAYIT</th><th>DETAY</th><th>ESKİ DEĞER</th><th>YENİ DEĞER</th>
</tr></thead><tbody id="auditRows"></tbody></table></div>
</section>

</div>
</main>


<div id="v54TableMenu" style="display:none">
  <div id="v54MenuTitle" style="font-weight:800"></div>
  <button class="v54-menu-btn" onclick="v54ApplySort('asc')">▲ Artan Sırala</button>
  <button class="v54-menu-btn" onclick="v54ApplySort('desc')">▼ Azalan Sırala</button>
  <div style="border-top:1px solid #e5e7eb;margin:7px 0"></div>
  <input id="v54FilterInput" placeholder="Bu sütunda ara..." oninput="v54ApplyFilter()">
  <button class="v54-menu-btn" onclick="v54ClearColumn()">Bu Sütun Filtresini Temizle</button>
  <button class="v54-menu-btn" onclick="v54ClearTable()">Bu Tablodaki Tüm Filtreleri Temizle</button>
</div>
<div class="modal" id="modal">
<div class="box">
<h3 id="mTitle"></h3>
<div id="mBody"></div>
<div class="actions">
<button class="btn secondary" onclick="closeM()">Kapat</button>
<button class="btn primary" id="mSave">Kaydet</button>
</div>
</div>
</div>




<script>
let AUTH_USER=null;
let AUTH_CATALOG=null;
let AUTH_USERS=[];
let AUTH_DRIVER_OPTIONS=[];

function hasPerm(p){return !!AUTH_USER && (AUTH_USER.role==='ADMIN'||(AUTH_USER.permissions||[]).includes(p));}

async function authFetch(url,opt={}){
  const r=await fetch(url,opt);
  let d={};
  try{d=await r.json();}catch(e){}
  if(!r.ok)throw new Error(d.detail||'Hata');
  return d;
}

async function authBoot(){
  try{
    const s=await authFetch('/api/auth/status');
    if(s.setup_required){
      authSetupBox.style.display='block';authLoginBox.style.display='none';authOverlay.style.display='flex';return;
    }
    if(!s.authenticated){
      authSetupBox.style.display='none';authLoginBox.style.display='block';authOverlay.style.display='flex';return;
    }
    AUTH_USER=s.user;
    authOverlay.style.display='none';
    applyAuthUI();
    if(AUTH_USER.must_change_password){
      setTimeout(()=>openMyProfile(true),100);
    }else{
      await init();
    }
  }catch(e){authMsg.innerText=e.message;}
}

async function createFirstAdmin(){
  authMsg.innerText='';
  try{
    await authFetch('/api/auth/setup',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({
      username:setupUsername.value,full_name:setupFullName.value,password:setupPassword.value
    })});
    authSetupBox.style.display='none';authLoginBox.style.display='block';
    authMsg.innerText='Admin oluşturuldu. Şimdi giriş yap.';
  }catch(e){authMsg.innerText=e.message;}
}

async function authLogin(){
  authMsg.innerText='';
  try{
    const d=await authFetch('/api/auth/login',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({
      username:loginUsername.value,password:loginPassword.value
    })});
    AUTH_USER=d.user;authOverlay.style.display='none';applyAuthUI();
    if(AUTH_USER.must_change_password){
      setTimeout(()=>openMyProfile(true),100);
    }else{
      await init();
    }
  }catch(e){authMsg.innerText=e.message;}
}
async function authLogout(){
  try{await authFetch('/api/auth/logout',{method:'POST'});}catch(e){}
  location.reload();
}

function applyAuthUI(){
  if(!AUTH_USER)return;
  authUserBtn.innerText=(AUTH_USER.full_name||AUTH_USER.username)+' | '+AUTH_USER.role;
  const pagePerm={
    dash:'dashboard.view',
    trips:'shipment.view',exit:'shipment.view',entry:'shipment.view',detail:'shipment.view',
    editcenter:'shipment.edit',quality:'shipment.view',compare:'shipment.view',bulkfix:'shipment.edit',
    opcenter:'operation.view',liveops:'operation.view',loadqueue:'operation.view',vesselops:'operation.view',routes:'operation.view',
    fleetmanage:'fleet.view',vehiclecard:'fleet.view',maintenance:'maintenance.view',
    drivermanage:'driver.view',drivercard:'driver.view',
    anomaly:'report.view',daily:'report.view',performance:'report.view',finance:'report.view',alerts:'report.view',
    audit:'audit.view',useradmin:'users.manage',deleted:'shipment.delete',advances:'users.manage'
  };
  document.querySelectorAll('.nav button[onclick*="show("]').forEach(btn=>{
    const m=(btn.getAttribute('onclick')||'').match(/show\('([^']+)'/);
    if(!m)return;
    const perm=pagePerm[m[1]];
    if(perm&&!hasPerm(perm))btn.style.display='none';
  });
  const ag=document.getElementById('adminNavGroup');
  if(ag)ag.style.display=hasPerm('users.manage')?'':'none';

  // Mutation buttons
  document.querySelectorAll('button').forEach(b=>{
    const t=(b.innerText||'').toLocaleUpperCase('tr-TR');
    if((t.includes('YENİ SEVKİYAT')||t.includes('+ YENİ SEVKİYAT'))&&!hasPerm('shipment.create'))b.style.display='none';
    if((t.includes('EXCEL')||t.includes('ONEDRIVE'))&&t.includes('AKTAR')&&!hasPerm('excel.import'))b.style.display='none';
  });
}

function openMyProfile(force=false){
  openM('Profilim / Şifre Değiştir',`
    <div class="calc"><b>${AUTH_USER.full_name||AUTH_USER.username}</b><br>
    Kullanıcı: ${AUTH_USER.username} | Rol: ${AUTH_USER.role}</div>
    ${force?'<div class="calc" style="border-color:#f59e0b"><b>İlk girişte şifrenizi değiştirmeniz gerekiyor.</b></div>':''}
    <div class="grid">
      <div class="field"><label>Mevcut Şifre</label><input id="myOldPass" type="password"></div>
      <div class="field"><label>Yeni Şifre</label><input id="myNewPass" type="password"></div>
      <div class="field"><label>Yeni Şifre Tekrar</label><input id="myNewPass2" type="password"></div>
    </div>`,
    async()=>{
      if(myNewPass.value!==myNewPass2.value){alert('Yeni şifreler aynı değil.');return;}
      try{
        await authFetch('/api/auth/change-password',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({
          current_password:myOldPass.value,new_password:myNewPass.value
        })});
        AUTH_USER.must_change_password=0;closeM();alert('Şifre değiştirildi.');await init();
      }catch(e){alert(e.message);}
    }
  );
}

async function loadAuthUsers(){
  if(!hasPerm('users.manage'))return;
  AUTH_USERS=await authFetch('/api/auth/users');
  if(!AUTH_CATALOG)AUTH_CATALOG=await authFetch('/api/auth/permission-catalog');
  if(!AUTH_DRIVER_OPTIONS.length)AUTH_DRIVER_OPTIONS=await authFetch('/api/auth/driver-options');
  authUserTotal.innerText=AUTH_USERS.length;
  authUserActive.innerText=AUTH_USERS.filter(x=>x.is_active).length;
  authUserRows.innerHTML=AUTH_USERS.map(x=>`<tr>
    <td><b>${x.username}</b></td><td>${x.full_name||''}</td><td>${x.role}${x.driver_id?' • ŞOFÖR':''}</td>
    <td>${x.is_active?'AKTİF':'PASİF'}</td>
    <td>${x.must_change_password?'ZORUNLU':'-'}</td>
    <td>
      <button class="btn secondary" onclick="openAuthUserEdit(${x.id})">Düzenle</button>
      <button class="btn secondary" onclick="openAuthPermissions(${x.id})">Yetkiler</button>
      <button class="btn orange" onclick="openAuthReset(${x.id})">Şifre Sıfırla</button>
    </td></tr>`).join('');
  if(typeof v54InitTables==='function')v54InitTables();
}

function openAuthUserCreate(){
  openM('Yeni Kullanıcı',`<div class="grid">
    <div class="field"><label>Ad Soyad</label><input id="auFull"></div>
    <div class="field"><label>Kullanıcı Adı</label><input id="auUser"></div>
    <div class="field"><label>Başlangıç Şifresi</label><input id="auPass" type="password"></div>
    <div class="field"><label>Rol</label><select id="auRole"><option>OPERATOR</option><option>VIEWER</option><option>DRIVER</option><option>ADMIN</option></select></div>
    <div class="field"><label>Şoför Hesabı (opsiyonel)</label><select id="auDriver"><option value="">Şoför değil</option>${AUTH_DRIVER_OPTIONS.map(d=>`<option value="${d.id}">${d.d_no?d.d_no+' - ':''}${d.name}</option>`).join('')}</select></div>
    <div class="field"><label>İlk girişte şifre değiştir</label><select id="auForce"><option value="1">EVET</option><option value="0">HAYIR</option></select></div>
  </div>`,async()=>{
    try{
      await authFetch('/api/auth/users',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({
        username:auUser.value,full_name:auFull.value,password:auPass.value,role:auRole.value,
        is_active:1,must_change_password:Number(auForce.value),driver_id:auDriver.value?Number(auDriver.value):null
      })});
      closeM();await loadAuthUsers();
    }catch(e){alert(e.message);}
  });
}

function openAuthUserEdit(id){
  const x=AUTH_USERS.find(z=>z.id===id);if(!x)return;
  openM('Kullanıcı Düzenle - '+x.username,`<div class="grid">
    <div class="field"><label>Ad Soyad</label><input id="aeFull" value="${String(x.full_name||'').replace(/"/g,'&quot;')}"></div>
    <div class="field"><label>Rol</label><select id="aeRole"><option ${x.role==='ADMIN'?'selected':''}>ADMIN</option><option ${x.role==='OPERATOR'?'selected':''}>OPERATOR</option><option ${x.role==='VIEWER'?'selected':''}>VIEWER</option><option ${x.role==='DRIVER'?'selected':''}>DRIVER</option></select></div>
    <div class="field"><label>Durum</label><select id="aeActive"><option value="1" ${x.is_active?'selected':''}>AKTİF</option><option value="0" ${!x.is_active?'selected':''}>PASİF</option></select></div>
    <div class="field"><label>Şoför Hesabı</label><select id="aeDriver"><option value="">Şoför değil</option>${AUTH_DRIVER_OPTIONS.map(d=>`<option value="${d.id}" ${Number(x.driver_id)===Number(d.id)?'selected':''}>${d.d_no?d.d_no+' - ':''}${d.name}</option>`).join('')}</select></div>
    <div class="field"><label>Şifre değişimi zorunlu</label><select id="aeForce"><option value="1" ${x.must_change_password?'selected':''}>EVET</option><option value="0" ${!x.must_change_password?'selected':''}>HAYIR</option></select></div>
  </div>`,async()=>{
    try{
      await authFetch('/api/auth/users/'+id,{method:'PATCH',headers:{'Content-Type':'application/json'},body:JSON.stringify({
        full_name:aeFull.value,role:aeRole.value,is_active:Number(aeActive.value),must_change_password:Number(aeForce.value),
        driver_id:aeDriver.value?Number(aeDriver.value):null
      })});
      closeM();await loadAuthUsers();
    }catch(e){alert(e.message);}
  });
}

async function openAuthPermissions(id){
  if(!AUTH_CATALOG)AUTH_CATALOG=await authFetch('/api/auth/permission-catalog');
  const x=AUTH_USERS.find(z=>z.id===id);if(!x)return;
  const set=new Set(x.permissions||[]);
  openM('Yetkiler - '+x.username,`
    <div class="calc">Rol: <b>${x.role}</b>. Aşağıdaki kutular kullanıcıya özel son yetkilerdir.</div>
    <div class="permission-grid">
      ${AUTH_CATALOG.permissions.map(p=>`<label class="permission-item">
        <input type="checkbox" class="permCheck" value="${p.key}" ${set.has(p.key)?'checked':''}> ${p.label}
      </label>`).join('')}
    </div>`,async()=>{
      const permissions={};
      document.querySelectorAll('.permCheck').forEach(e=>permissions[e.value]=e.checked);
      try{
        await authFetch('/api/auth/users/'+id+'/permissions',{method:'PUT',headers:{'Content-Type':'application/json'},body:JSON.stringify({permissions})});
        closeM();await loadAuthUsers();
      }catch(e){alert(e.message);}
    }
  );
}

function openAuthReset(id){
  const x=AUTH_USERS.find(z=>z.id===id);if(!x)return;
  openM('Şifre Sıfırla - '+x.username,`<div class="grid">
    <div class="field"><label>Yeni Geçici Şifre</label><input id="arPass" type="password"></div>
    <div class="field"><label>İlk girişte değiştirsin</label><select id="arForce"><option value="1">EVET</option><option value="0">HAYIR</option></select></div>
  </div>`,async()=>{
    try{
      await authFetch('/api/auth/users/'+id+'/reset-password',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({
        new_password:arPass.value,must_change_password:Number(arForce.value)
      })});
      closeM();alert('Şifre sıfırlandı. Kullanıcının açık oturumları kapatıldı.');await loadAuthUsers();
    }catch(e){alert(e.message);}
  });
}


let samaAutocompleteSeq=0;

async function samaAutocompleteItems(type,q){
  q=String(q||'').trim();
  if(q.length<3)return [];

  try{
    if(type==='plate'){
      const rows=await api('/api/autocomplete/plates?q='+encodeURIComponent(q));
      return (rows||[]).map(x=>({
        value:String(x.plate||'').trim().toUpperCase(),
        sub:[x.brand,x.model,x.vehicle_type].filter(Boolean).join(' • ')
      })).filter(x=>x.value);
    }

    if(type==='customer'){
    const Q=q.toLocaleUpperCase('tr-TR');

    // Asıl kaynak: veritabanındaki müşteri kayıtları ve geçmiş sevkiyatlar.
    try{
      const rows=await api('/api/autocomplete/customers?q='+encodeURIComponent(q));
      const server=(rows||[]).map(x=>({
        value:String(x.name||x.customer||x.customer_name||'').trim(),
        sub:''
      })).filter(x=>x.value);
      if(server.length)return server;
    }catch(e){
      console.warn('Müşteri autocomplete API:',e);
    }

    // Sunucu sonucu yoksa tarayıcıdaki lookups listesini de tara.
    const seen=new Set(),out=[];
    const arr=((window.L&&window.L.customers)||[]);
    arr.forEach(x=>{
      const v=String(x.name||x.customer||x.customer_name||x||'').trim();
      const k=v.toLocaleUpperCase('tr-TR');
      if(v&&!seen.has(k)&&k.includes(Q)){
        seen.add(k);out.push({value:v,sub:''});
      }
    });
    return out.slice(0,50);
  }

  if(type==='driver'){
      const Q=q.toLocaleUpperCase('tr-TR');

      // Bu endpoint Şoför Yönetimi ekranında zaten kullanılan ve çalışan ana listedir.
      const fleetResp=await api('/api/fleet/drivers');
      const fleetRows=Array.isArray(fleetResp) ? fleetResp : (fleetResp?.rows||[]);
      const fleet=fleetRows.map(x=>({
        value:String(x.name||x.driver_name||x.full_name||'').trim(),
        sub:[x.phone,x.d_no,x.driver_no,x.last_plate].filter(Boolean).join(' • ')
      })).filter(x=>
        x.value &&
        Number(x.is_active)!==0 &&
        (x.value+' '+(x.sub||'')).toLocaleUpperCase('tr-TR').includes(Q)
      ).slice(0,50);

      if(fleet.length)return fleet;

      // Eski/ithal kayıtlar için ikinci kaynak.
      const rows=await api('/api/autocomplete/drivers?q='+encodeURIComponent(q));
      const server=(rows||[]).map(x=>({
        value:String(x.name||x.driver_name||x.full_name||'').trim(),
        sub:[x.phone,x.d_no,x.driver_no].filter(Boolean).join(' • ')
      })).filter(x=>x.value);
      if(server.length)return server;
    }
  }catch(e){
    console.warn('Autocomplete API fallback:',e);
  }

  // Sunucu erişilemezse tarayıcıdaki mevcut listelere geri dön.
  if(type==='plate'){
    const merged=[
      ...(window.fleetManageData||[]),
      ...((window.L&&window.L.vehicles)||[]),
      ...(window.fleetData||[])
    ];
    const seen=new Set(),out=[];
    const Q=q.toLocaleUpperCase('tr-TR');
    merged.forEach(x=>{
      const plate=String(x.plate||x.plate_no||'').trim().toUpperCase();
      if(!plate||seen.has(plate)||!plate.includes(Q))return;
      seen.add(plate);
      out.push({value:plate,sub:[x.driver_name,x.driver,x.brand,x.model].filter(Boolean).join(' • ')});
    });
    return out.slice(0,40);
  }

  if(type==='driver'){
    const Q=q.toLocaleUpperCase('tr-TR'),seen=new Set(),out=[];
    [...(window.driverManageData||[]),...((window.L&&window.L.vehicles)||[]),...(window.fleetData||[])].forEach(x=>{
      const v=String(x.name||x.driver_name||x.driver||'').trim();
      const k=v.toLocaleUpperCase('tr-TR');
      if(!v||seen.has(k)||!k.includes(Q))return;
      seen.add(k);
      out.push({value:v,sub:[x.phone,x.d_no,x.plate].filter(Boolean).join(' • ')});
    });
    return out.slice(0,40);
  }

  return [];
}

function samaAttachAutocomplete(input,type){
  if(!input || input.dataset.samaAutocompleteReady==='1') return;
  let wrap=input.parentElement;
  if(!wrap.classList.contains('autocomplete-wrap')){
    const w=document.createElement('div');w.className='autocomplete-wrap';
    input.parentNode.insertBefore(w,input);w.appendChild(input);wrap=w;
  }
  const list=document.createElement('div');list.className='autocomplete-list';wrap.appendChild(list);
  input.dataset.samaAutocompleteReady='1';input.setAttribute('autocomplete','off');
  let active=-1;

  let renderSeq=0;
  const render=async()=>{
    const q=String(input.value||'').trim();
    const seq=++renderSeq;
    active=-1;
    if(q.length<3){
      list.innerHTML='<div class="autocomplete-item"><div class="autocomplete-sub">Aramak için en az 3 karakter yazın</div></div>';
      list.style.display='block';
      return;
    }

    list.innerHTML='<div class="autocomplete-item"><div class="autocomplete-sub">Aranıyor...</div></div>';
    list.style.display='block';

    const a=await samaAutocompleteItems(type,q);
    if(seq!==renderSeq)return;

    list.innerHTML=a.length?a.slice(0,25).map((x,i)=>`<div class="autocomplete-item" data-value="${String(x.value).replace(/"/g,'&quot;')}"><div class="autocomplete-main">${x.value}</div>${x.sub?`<div class="autocomplete-sub">${x.sub}</div>`:''}</div>`).join(''):'<div class="autocomplete-item"><div class="autocomplete-sub">Eşleşme bulunamadı</div></div>';
    list.style.display='block';
    list.querySelectorAll('[data-value]').forEach(el=>el.onmousedown=e=>{e.preventDefault();input.value=el.dataset.value;
        input.dispatchEvent(new Event('input',{bubbles:true}));
        input.dispatchEvent(new Event('change',{bubbles:true}));
        if(type==='plate' && input.id==='nPlate') setTimeout(()=>newTripPlateDriverSync(),0);
        list.style.display='none';});
  };
  input.addEventListener('focus',()=>{
    if(type==='driver' && input.id==='nDriver' && input.value){
      setTimeout(()=>input.select(),0);
    }
    render();
  });
  input.addEventListener('input',render);
  input.addEventListener('keydown',e=>{
    const items=[...list.querySelectorAll('[data-value]')];
    if(!items.length)return;
    if(e.key==='ArrowDown'){e.preventDefault();active=Math.min(active+1,items.length-1);}
    else if(e.key==='ArrowUp'){e.preventDefault();active=Math.max(active-1,0);}
    else if(e.key==='Enter'&&active>=0){e.preventDefault();input.value=items[active].dataset.value;
      input.dispatchEvent(new Event('input',{bubbles:true}));
      input.dispatchEvent(new Event('change',{bubbles:true}));
      if(type==='plate' && input.id==='nPlate') setTimeout(()=>newTripPlateDriverSync(),0);
      list.style.display='none';return;}
    else if(e.key==='Escape'){list.style.display='none';return;} else return;
    items.forEach((el,i)=>el.classList.toggle('active',i===active));
  });
  input.addEventListener('blur',()=>setTimeout(()=>list.style.display='none',150));
}

function samaInitAllAutocompletes(root=document){
  root.querySelectorAll('input:not([readonly]):not([type="file"]):not([type="hidden"])').forEach(el=>{
    if(el.dataset.noAutocomplete==='1')return;
    const id=(el.id||'').toLowerCase(), ph=(el.placeholder||'').toLocaleLowerCase('tr-TR');
    const label=(el.closest('.field')?.querySelector('label')?.innerText||'').toLocaleLowerCase('tr-TR');
    const isSearch=ph.includes('ara')||id.endsWith('q')||id.includes('filter');
    // SAMA_SEARCH_AUTOCOMPLETE_FIX_V1
    if(!isSearch&&(id.includes('plate')||ph.includes('plaka')||label.includes('plaka'))) samaAttachAutocomplete(el,'plate');
    else if(!isSearch&&(id.includes('driver')||ph.includes('şoför')||ph.includes('sofor')||label.includes('şoför')||label.includes('sofor'))) samaAttachAutocomplete(el,'driver');
    else if(!isSearch&&(id.includes('customer')||ph.includes('müşteri')||ph.includes('musteri')||label.includes('müşteri')||label.includes('musteri'))) samaAttachAutocomplete(el,'customer');
  });
}
document.addEventListener('DOMContentLoaded',()=>{samaInitAllAutocompletes();new MutationObserver(()=>samaInitAllAutocompletes()).observe(document.body,{childList:true,subtree:true});});

let samaServerWasDown=false;

async function samaHeartbeat(){
  try{
    const r=await fetch('/api/health?ts='+Date.now(),{cache:'no-store'});
    if(r.ok){
      if(samaServerWasDown){
        location.reload();
        return;
      }
      samaServerWasDown=false;
    }else{
      samaServerWasDown=true;
    }
  }catch(e){
    samaServerWasDown=true;
  }
}

setInterval(samaHeartbeat,2000);
setTimeout(samaHeartbeat,1000);

let v54Ctx=null;
const v54States=new WeakMap();

function v54Text(td){ return (td?.innerText||td?.textContent||'').trim(); }
function v54Num(s){
  let x=(s||'').trim().replace(/\s/g,'');
  if(!x) return null;
  // TR/European formatted number: 1.234.567,89
  if(/^[-+]?\d{1,3}(\.\d{3})*(,\d+)?$/.test(x)) x=x.replace(/\./g,'').replace(',','.');
  else if(/^[-+]?\d+(,\d+)?$/.test(x)) x=x.replace(',','.');
  else return null;
  const n=Number(x); return Number.isFinite(n)?n:null;
}
function v54Date(s){
  let m=(s||'').match(/^(\d{1,2})[.\/-](\d{1,2})[.\/-](\d{4})(?:\s+(\d{1,2}):(\d{2}))?/);
  if(!m) return null;
  return new Date(+m[3],+m[2]-1,+m[1],+(m[4]||0),+(m[5]||0)).getTime();
}
function v54Cmp(a,b){
  const da=v54Date(a),db=v54Date(b);
  if(da!==null&&db!==null) return da-db;
  const na=v54Num(a),nb=v54Num(b);
  if(na!==null&&nb!==null) return na-nb;
  return String(a).localeCompare(String(b),'tr',{numeric:true,sensitivity:'base'});
}
function v54State(table){
  if(!v54States.has(table)) v54States.set(table,{filters:{},sortCol:null,sortDir:null});
  return v54States.get(table);
}
function v54InitTables(root=document){
  root.querySelectorAll('table').forEach((table,ti)=>{
    if(table.dataset.v54Ready==='1') return;
    table.dataset.v54Ready='1';
    const heads=[...(table.tHead?.rows?.[0]?.cells||[])];
    heads.forEach((th,i)=>{
      // Trips has its own V53 richer handler. Avoid double menu.
      if(th.classList.contains('clickhead')) return;
      const title=v54Text(th);
      if(!title || ['İŞLEM','DÜZENLE'].includes(title.toLocaleUpperCase('tr-TR'))) return;
      th.classList.add('v54-head');
      th.dataset.v54Col=i;
      th.dataset.v54Title=title;
      if(!th.querySelector('.v54-sortmark')){
        const sp=document.createElement('span');sp.className='v54-sortmark';sp.textContent=' ↕';th.appendChild(sp);
      }
      th.addEventListener('click',e=>v54Open(e,table,i,th));
    });
  });
}
function v54Open(ev,table,col,th){
  v54Ctx={table,col,th};
  const st=v54State(table);
  v54MenuTitle.textContent=th.dataset.v54Title||v54Text(th);
  v54FilterInput.value=st.filters[col]||'';
  const r=th.getBoundingClientRect();
  v54TableMenu.style.left=Math.min(r.left,window.innerWidth-300)+'px';
  v54TableMenu.style.top=Math.min(r.bottom+4,window.innerHeight-260)+'px';
  v54TableMenu.style.display='block';
  setTimeout(()=>v54FilterInput.focus(),30);
}
function v54ApplySort(dir){
  if(!v54Ctx)return;
  const st=v54State(v54Ctx.table);st.sortCol=v54Ctx.col;st.sortDir=dir;
  v54TableMenu.style.display='none';v54Render(v54Ctx.table);
}
function v54ApplyFilter(){
  if(!v54Ctx)return;
  const st=v54State(v54Ctx.table),v=v54FilterInput.value.trim();
  if(v)st.filters[v54Ctx.col]=v;else delete st.filters[v54Ctx.col];
  v54Render(v54Ctx.table);
}
function v54ClearColumn(){
  if(!v54Ctx)return;
  const st=v54State(v54Ctx.table);delete st.filters[v54Ctx.col];
  v54FilterInput.value='';v54TableMenu.style.display='none';v54Render(v54Ctx.table);
}
function v54ClearTable(){
  if(!v54Ctx)return;
  v54States.set(v54Ctx.table,{filters:{},sortCol:null,sortDir:null});
  v54TableMenu.style.display='none';v54Render(v54Ctx.table);
}
function v54Render(table){
  const body=table.tBodies?.[0];if(!body)return;
  const st=v54State(table);
  let rows=[...body.rows];

  // Filtering
  rows.forEach(row=>{
    let ok=true;
    for(const [c,q] of Object.entries(st.filters)){
      const txt=v54Text(row.cells[+c]).toLocaleUpperCase('tr-TR');
      if(!txt.includes(String(q).toLocaleUpperCase('tr-TR'))){ok=false;break;}
    }
    row.style.display=ok?'':'none';
  });

  // Sorting keeps hidden rows too, so clearing filter restores correct order.
  if(st.sortCol!==null){
    const c=st.sortCol,m=st.sortDir==='desc'?-1:1;
    rows.sort((a,b)=>v54Cmp(v54Text(a.cells[c]),v54Text(b.cells[c]))*m);
    rows.forEach(r=>body.appendChild(r));
  }

  [...(table.tHead?.rows?.[0]?.cells||[])].forEach((th,i)=>{
    if(th.classList.contains('clickhead'))return;
    const mark=th.querySelector('.v54-sortmark');
    if(mark)mark.textContent=st.sortCol===i?(st.sortDir==='asc'?' ▲':' ▼'):' ↕';
    th.classList.toggle('v54-filtered',Object.prototype.hasOwnProperty.call(st.filters,i));
  });
}
document.addEventListener('click',e=>{
  if(v54TableMenu.style.display!=='none'&&!v54TableMenu.contains(e.target)&&!e.target.closest('.v54-head'))v54TableMenu.style.display='none';
});
document.addEventListener('DOMContentLoaded',()=>v54InitTables());
const v54Observer=new MutationObserver(()=>v54InitTables());
document.addEventListener('DOMContentLoaded',()=>v54Observer.observe(document.body,{childList:true,subtree:true}));

let L={vehicles:[],areas:[],cargo_categories:[],customers:[]},cache=[],cx=null;
let excelImportMode='sync';
let lastImportFailedRecords=[];
let lastExcelFile=null;
let lastExcelPreview=null;

const api=async(u,o={})=>{
  let r=await fetch(u,o);
  if(!r.ok){
    let e=await r.json();
    throw new Error(e.detail||'Hata');
  }
  return r.json();
};

const money=v=>Number(v||0).toLocaleString('tr-TR',{maximumFractionDigits:2});

function parseMoney(value){
  if(value===null || value===undefined) return 0;
  let s=String(value).trim().replace(/\s/g,'').replace(/[₺$€£₺]/g,'');
  if(!s) return 0;

  const hasDot=s.includes('.');
  const hasComma=s.includes(',');

  if(hasDot && hasComma){
    // Son görülen ayıracı ondalık kabul et, diğerini binlik kabul et.
    if(s.lastIndexOf(',') > s.lastIndexOf('.')){
      s=s.replace(/\./g,'').replace(',','.');
    }else{
      s=s.replace(/,/g,'');
    }
  }else if(hasDot){
    const parts=s.split('.');
    // 469.770 gibi 3 haneli son grup varsa binlik ayırıcı kabul et.
    if(parts.length>1 && parts.slice(1).every(p=>p.length===3)){
      s=parts.join('');
    }
  }else if(hasComma){
    const parts=s.split(',');
    // 469,770 gibi 3 haneli son grup varsa binlik ayırıcı kabul et.
    if(parts.length>1 && parts.slice(1).every(p=>p.length===3)){
      s=parts.join('');
    }else{
      s=s.replace(',','.');
    }
  }

  s=s.replace(/[^\d.-]/g,'');
  const n=Number(s);
  return Number.isFinite(n)?n:0;
}

function parseQty(value){
  if(value===null || value===undefined) return 0;
  let s=String(value).trim().replace(/\s/g,'');
  if(!s) return 0;
  // Miktarlarda virgülü ondalık kabul et.
  if(s.includes(',') && !s.includes('.')) s=s.replace(',','.');
  const n=Number(s);
  return Number.isFinite(n)?n:0;
}


function fmtDateTime(v){
  if(!v) return '';
  let s=String(v).replace('T',' ');
  // YYYY-MM-DD HH:mm:ss -> DD.MM.YYYY HH:mm
  const m=s.match(/^(\d{4})-(\d{2})-(\d{2})[ T](\d{2}):(\d{2})/);
  if(m) return `${m[3]}.${m[2]}.${m[1]} ${m[4]}:${m[5]}`;
  return s;
}

function getColumnPrefs(){
  const def=['scna','plate','driver','customer','area','cargo','cargo_type','kg','freight','exit_at','entry_at','status'];
  try{
    const v=JSON.parse(localStorage.getItem('tripColumnPrefs')||'null');
    return Array.isArray(v)&&v.length?v:def;
  }catch(e){ return def; }
}

function applyColumnPrefs(){
  const selected=getColumnPrefs();
  document.querySelectorAll('#trips [data-col]').forEach(el=>{
    el.style.display=selected.includes(el.dataset.col)?'':'none';
  });
  document.querySelectorAll('#columnPanel input[data-col]').forEach(cb=>{
    cb.checked=selected.includes(cb.dataset.col);
  });
}

function saveColumnPrefs(){
  const selected=[...document.querySelectorAll('#columnPanel input[data-col]:checked')].map(x=>x.dataset.col);
  if(!selected.length) return;
  localStorage.setItem('tripColumnPrefs',JSON.stringify(selected));
  applyColumnPrefs();
}

function toggleColumnPanel(e){
  if(e) e.stopPropagation();
  columnPanel.classList.toggle('open');
}

document.addEventListener('click',e=>{
  const p=document.querySelector('.column-picker');
  if(p && !p.contains(e.target)) columnPanel.classList.remove('open');
});

function sBadge(s){
  if(s==='HESAP TAMAM') return '<span class="badge ok">HESAP TAMAM</span>';
  if(s==='EKSİK PARA') return '<span class="badge short">EKSİK PARA</span>';
  if(s==='FAZLA PARA') return '<span class="badge over">FAZLA PARA</span>';
  return '<span class="badge warn">AÇIK</span>';
}


function toggleNavGroup(btn){
  const group=btn.closest('.nav-group');
  if(!group) return;
  const willOpen=!group.classList.contains('open');

  document.querySelectorAll('.nav-group').forEach(g=>{
    g.classList.remove('open');
    const s=g.querySelector('.nav-group-title span');
    if(s) s.textContent='▸';
  });

  if(willOpen){
    group.classList.add('open');
    const s=group.querySelector('.nav-group-title span');
    if(s) s.textContent='▾';
  }
}

function openNavGroupForButton(btn){
  if(!btn) return;
  const group=btn.closest('.nav-group');
  if(!group) return;
  document.querySelectorAll('.nav-group').forEach(g=>{
    g.classList.remove('open');
    const s=g.querySelector('.nav-group-title span');
    if(s) s.textContent='▸';
  });
  group.classList.add('open');
  const s=group.querySelector('.nav-group-title span');
  if(s) s.textContent='▾';
}

function show(id,b){
  document.querySelectorAll('.panel').forEach(x=>x.classList.remove('active'));
  document.getElementById(id).classList.add('active');
  document.querySelectorAll('.nav button').forEach(x=>x.classList.remove('active'));
  b.classList.add('active');
  openNavGroupForButton(b);

  if(id==='dash') loadDash();
  if(id==='trips') loadTrips();
  if(id==='exit') loadExit();
  if(id==='entry') loadEntry();
  if(id==='detail') document.getElementById('detailScna').focus();
  if(id==='editcenter') document.getElementById('editScna').focus();
  if(id==='quality') loadQuality();
  if(id==='bulkfix') refreshBulkValue();
  if(id==='vehiclecard') document.getElementById('vehicleCardPlate').focus();
  if(id==='drivercard') document.getElementById('driverCardName').focus();
  if(id==='anomaly') loadAnomalies();
  if(id==='daily') dailyToday();
  if(id==='performance') loadVehiclePerformance();
  if(id==='finance') loadFinance();
  if(id==='fleetmanage') loadFleetManage();
  if(id==='drivermanage') loadDriverManage();
  if(id==='opcenter') loadOperationCenter();
  if(id==='loadqueue') loadFleetOps();
  if(id==='vesselops') loadFleetOps();
  if(id==='maintops') loadMaintenanceOps();
  if(id==='liveops') loadLiveOps();
  if(id==='routes') loadRouteStandards();
  if(id==='maintenance') loadMaintenance();
  if(id==='alerts') loadAlerts();
  if(id==='useradmin') loadAuthUsers();
  if(id==='advances') loadAdvances();
  if(id==='cashcontrol') loadCashControl();
  if(id==='audit') loadAudit();
}

function openM(t,b,s){
  mTitle.innerText=t;
  mBody.innerHTML=b;
  mSave.onclick=s;
  modal.style.display='flex';
}

function closeM(){
  modal.style.display='none';
}

let tripPage=1;
const tripPageSize=50;
let tripTotal=0;

async function ft(q='',status='',limit=50,offset=0,customer='',area='',cargo=''){
  let url='/api/trips?q='+encodeURIComponent(q||'')+
          '&status='+encodeURIComponent(status||'')+
          '&customer='+encodeURIComponent(customer||'')+
          '&area='+encodeURIComponent(area||'')+
          '&cargo='+encodeURIComponent(cargo||'')+
          '&limit='+limit+'&offset='+offset;
  cache=await api(url);
  return cache;
}

async function init(){
  L=await api('/api/lookups');

  const addOpts=(id,items,labelKey='name')=>{
    const el=document.getElementById(id);
    if(!el) return;
    items.forEach(x=>{
      const o=document.createElement('option');
      o.value=x[labelKey]||'';
      o.textContent=x[labelKey]||'';
      el.appendChild(o);
    });
  };

  addOpts('tripCustomerFilter',L.customers||[]);
  addOpts('tripAreaFilter',L.areas||[]);
  addOpts('tripCargoFilter',L.cargo_categories||[]);

  loadDash();
}


let integratedFleetRows=[];

async function loadIntegratedFleetStatus(){
  const d=await api('/api/integrated-fleet-status');
  integratedFleetRows=d.rows||[];
  const c=d.counts||{};
  if(document.getElementById('dOpVessel')){
    dOpVessel.innerText=c.GEMIDE||0;
    dOpMaintenance.innerText=c.BAKIM||0;
    dOpRoad.innerText=c.YOLDA||0;
    if(document.getElementById('dOpReturning')) dOpReturning.innerText=c.DONUYOR||0;
    dOpWaiting.innerText=c.BEKLEMEDE||0;
    dOpIdle.innerText=c.BOSTA||0;
  }
  try{const mo=await api('/api/maintenance-operations?active_only=1');if(document.getElementById('dOpMaintenance'))dOpMaintenance.innerText=mo.length;}catch(e){}
  return d;
}
function integratedStateLabel(s){
  return {GEMIDE:'GEMİDE ÇALIŞIYOR',BAKIM:'BAKIMDA',YOLDA:'YOLDA',DONUYOR:'BOŞALTILDI DÖNÜYOR',BEKLEMEDE:'BEKLEMEDE',BOSTA:'BOŞTA'}[s]||s;
}
async function openDashboardFleet(state){
  if(state==='BAKIM'){return openDashboardMaintenance();}
  if(!integratedFleetRows.length)await loadIntegratedFleetStatus();
  const rows=integratedFleetRows.filter(x=>x.operation_state===state);
  const box=document.getElementById('dashFleetDetail');
  box.style.display='block';
  box.innerHTML=`<div style="display:flex;justify-content:space-between;align-items:center;margin-bottom:10px;gap:10px">
    <b>${integratedStateLabel(state)} — ${rows.length} ARAÇ</b>
    <button class="btn secondary" onclick="document.getElementById('dashFleetDetail').style.display='none'">Kapat</button>
  </div>
  <div class="table"><table><thead><tr>
    <th>PLAKA</th><th>ŞOFÖR</th><th>SCNA</th><th>BÖLGE</th><th>GEMİ</th><th>DURUM</th><th>ÇIKIŞ</th><th>İŞLEM</th>
  </tr></thead><tbody>
  ${rows.map(x=>`<tr>
    <td><b>${x.plate||''}</b></td><td>${x.driver||''}</td><td>${x.scna||''}</td>
    <td>${x.area||''}</td><td>${x.vessel||''}</td><td>${integratedStateLabel(x.operation_state)}</td>
    <td>${fmtDateTime(x.exit_at)}</td>
    <td>${state==='YOLDA'?`<button class="btn orange" onclick="markGpsReturning('${String(x.plate||'').replace(/'/g,"\\'")}')">BOŞALTILDI DÖNÜYOR</button>`:state==='DONUYOR'?`<button class="btn secondary" onclick="undoGpsReturning('${String(x.plate||'').replace(/'/g,"\\'")}')">↩ YOLDA'YA GERİ AL</button>`:''}</td>
  </tr>`).join('')}
  </tbody></table></div>`;
  if(typeof v54InitTables==='function')v54InitTables();
}

async function markGpsReturning(plate){
  if(!confirm(`${plate} aracı BOŞALTILDI DÖNÜYOR olarak işaretlensin mi?`)) return;
  try{
    await api('/api/fleet/gps-state',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({plate,state:'DONUYOR'})});
    await loadIntegratedFleetStatus();
    openDashboardFleet('YOLDA');
  }catch(e){alert(e.message);}
}

async function undoGpsReturning(plate){
  if(!confirm(`${plate} aracı tekrar YOLDA durumuna alınsın mı?`)) return;
  try{
    await api('/api/fleet/gps-state',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({plate,state:''})});
    await loadIntegratedFleetStatus();
    openDashboardFleet('DONUYOR');
  }catch(e){alert(e.message);}
}


let maintenanceOpsRows=[];
async function loadMaintenanceOps(){maintenanceOpsRows=await api('/api/maintenance-operations');const active=maintenanceOpsRows.filter(x=>Number(x.is_active)===1);const a=document.getElementById('maintOpsActive'),t=document.getElementById('maintOpsTotal');if(a)a.innerText=active.length;if(t)t.innerText=maintenanceOpsRows.length;const d=document.getElementById('dOpMaintenance');if(d)d.innerText=active.length;renderMaintenanceOps();}
function renderMaintenanceOps(){const body=document.getElementById('maintOpsRows');if(!body)return;const q=(document.getElementById('maintOpsQ')?.value||'').trim().toLocaleUpperCase('tr-TR');const rows=maintenanceOpsRows.filter(x=>!q||[x.plate,x.reason,x.description,x.work_done,x.previous_state].join(' ').toLocaleUpperCase('tr-TR').includes(q));body.innerHTML=rows.map(x=>`<tr><td><b>${x.plate||''}</b></td><td>${Number(x.is_active)===1?'<span class="badge warn">BAKIMDA</span>':'<span class="badge ok">TAMAMLANDI</span>'}</td><td>${fmtDateTime(x.started_at)}</td><td>${Number(x.estimated_hours||0)?Number(x.estimated_hours).toLocaleString('tr-TR')+' saat':'-'}</td><td>${fmtDateTime(x.estimated_finish_at)||'-'}</td><td>${x.previous_state||''}</td><td>${x.reason||''}</td><td>${x.description||''}</td><td>${x.work_done||''}</td><td>${x.started_by||''}</td><td>${fmtDateTime(x.finished_at)||'-'}</td><td>${x.finished_by||''}</td><td><div class="compact-actions"><button class="btn secondary" onclick="editMaintenanceOp(${x.id})">DÜZENLE</button>${Number(x.is_active)===1?`<button class="btn green" onclick="finishMaintenanceOp(${x.id},'${String(x.plate||'').replace(/'/g,"\\'")}')">BAKIM BİTTİ</button>`:''}</div></td></tr>`).join('');if(typeof v54InitTables==='function')v54InitTables();}
async function openMaintenanceStart(){openM('Aracı Bakıma Al',`<div class="grid"><div class="field"><label>Plaka</label><div class="autocomplete-wrap"><input id="moPlate" autocomplete="off" placeholder="En az 3 karakter yazın..." oninput="maintenancePlateSearch(this)"><div id="moPlateList" class="autocomplete-list"></div></div></div><div class="field"><label>Tahmini Bakım Süresi (Saat)</label><input id="moHours" type="number" min="0" step="0.5" placeholder="Örn: 6"></div><div class="field wide"><label>Bakım Nedeni</label><input id="moReason" placeholder="Örn: Yağ kaçağı / fren / lastik / periyodik bakım"></div><div class="field wide"><label>Açıklama</label><textarea id="moDescription" rows="4" placeholder="Arıza veya bakım hakkında detay..."></textarea></div></div>`,async()=>{try{await api('/api/maintenance-operations/start',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({plate:moPlate.value,reason:moReason.value,description:moDescription.value,estimated_hours:Number(moHours.value||0)})});closeM();await loadMaintenanceOps();await loadIntegratedFleetStatus();alert('Araç bakıma alındı.');}catch(e){alert(e.message);}});}
async function maintenancePlateSearch(input){
  const list=document.getElementById('moPlateList');if(!list)return;
  const q=(input.value||'').trim().toUpperCase();
  if(q.length<3){list.style.display='none';list.innerHTML='';return;}
  try{
    const rows=await api('/api/autocomplete/plates?q='+encodeURIComponent(q));
    const active=new Set((maintenanceOpsRows||[]).filter(x=>Number(x.is_active)===1).map(x=>String(x.plate||'').toUpperCase()));
    const usable=(rows||[]).filter(x=>!active.has(String(x.plate||'').toUpperCase()));
    list.innerHTML=usable.slice(0,20).map(x=>`<div class="autocomplete-item" onclick="selectMaintenancePlate('${String(x.plate||'').replace(/'/g,"\\'")}')"><div class="autocomplete-main">${x.plate||''}</div><div class="autocomplete-sub">${[x.brand,x.model,x.vehicle_type].filter(Boolean).join(' • ')}</div></div>`).join('')||'<div class="autocomplete-item">Uygun plaka bulunamadı.</div>';
    list.style.display='block';
  }catch(e){list.style.display='none';}
}
function selectMaintenancePlate(plate){const i=document.getElementById('moPlate'),l=document.getElementById('moPlateList');if(i)i.value=plate;if(l)l.style.display='none';}
async function editMaintenanceOp(id){
  const x=(maintenanceOpsRows||[]).find(r=>Number(r.id)===Number(id));if(!x)return alert('Bakım kaydı bulunamadı.');
  openM(`${x.plate} — Bakım Kaydını Düzelt`,`<div class="grid"><div class="field"><label>Plaka</label><div class="autocomplete-wrap"><input id="moEditPlate" value="${x.plate||''}" autocomplete="off" oninput="maintenanceEditPlateSearch(this)"><div id="moEditPlateList" class="autocomplete-list"></div></div></div><div class="field"><label>Tahmini Bakım Süresi (Saat)</label><input id="moEditHours" type="number" min="0" step="0.5" value="${Number(x.estimated_hours||0)}"></div><div class="field wide"><label>Bakım Nedeni</label><input id="moEditReason" value="${String(x.reason||'').replace(/\"/g,'&quot;')}"></div><div class="field wide"><label>Açıklama</label><textarea id="moEditDescription" rows="4">${x.description||''}</textarea></div></div>`,async()=>{try{await api('/api/maintenance-operations/'+id+'/update',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({plate:moEditPlate.value,reason:moEditReason.value,description:moEditDescription.value,estimated_hours:Number(moEditHours.value||0)})});closeM();await loadMaintenanceOps();await loadIntegratedFleetStatus();alert('Bakım kaydı düzeltildi.');}catch(e){alert(e.message);}});
}
async function maintenanceEditPlateSearch(input){
  const list=document.getElementById('moEditPlateList');if(!list)return;const q=(input.value||'').trim().toUpperCase();if(q.length<3){list.style.display='none';return;}
  try{const rows=await api('/api/autocomplete/plates?q='+encodeURIComponent(q));list.innerHTML=(rows||[]).slice(0,20).map(x=>`<div class="autocomplete-item" onclick="document.getElementById('moEditPlate').value='${String(x.plate||'').replace(/'/g,"\\'")}';document.getElementById('moEditPlateList').style.display='none'"><div class="autocomplete-main">${x.plate||''}</div><div class="autocomplete-sub">${[x.brand,x.model,x.vehicle_type].filter(Boolean).join(' • ')}</div></div>`).join('');list.style.display='block';}catch(e){list.style.display='none';}
}
async function finishMaintenanceOp(id,plate){openM(`${plate} — Bakım Bitir`,`<div class="field"><label>Yapılan İşlemler</label><textarea id="moWork" rows="7" placeholder="Değişen parçalar, onarım, kontrol, yağ, filtre vb..."></textarea></div>`,async()=>{try{await api('/api/maintenance-operations/'+id+'/finish',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({plate:plate,work_done:moWork.value})});closeM();await loadMaintenanceOps();await loadIntegratedFleetStatus();await openDashboardMaintenance();alert('Bakım tamamlandı. Araç önceki operasyon akışına döndü.');}catch(e){alert(e.message);}});}
async function openDashboardMaintenance(){const rows=await api('/api/maintenance-operations?active_only=1');const box=document.getElementById('dashFleetDetail');box.style.display='block';box.innerHTML=`<div style="display:flex;justify-content:space-between;align-items:center;margin-bottom:10px;gap:10px"><b>BAKIMDA — ${rows.length} ARAÇ</b><button class="btn secondary" onclick="document.getElementById('dashFleetDetail').style.display='none'">Kapat</button></div><div class="table"><table><thead><tr><th>PLAKA</th><th>BAKIM GİRİŞ</th><th>TAHMİNİ SÜRE</th><th>TAHMİNİ ÇIKIŞ</th><th>BAKIM NEDENİ</th><th>AÇIKLAMA</th><th>YAPILAN İŞLEMLER</th><th>BAKIMA ALAN</th><th>İŞLEM</th></tr></thead><tbody>${rows.map(x=>`<tr><td><b>${x.plate||''}</b></td><td>${fmtDateTime(x.started_at)}</td><td>${Number(x.estimated_hours||0)?Number(x.estimated_hours).toLocaleString('tr-TR')+' saat':'-'}</td><td>${fmtDateTime(x.estimated_finish_at)||'-'}</td><td>${x.reason||''}</td><td>${x.description||''}</td><td>${x.work_done||''}</td><td>${x.started_by||''}</td><td><button class="btn green" onclick="finishMaintenanceOp(${x.id},'${String(x.plate||'').replace(/'/g,"\\'")}')">BAKIM BİTTİ</button></td></tr>`).join('')}</tbody></table></div>`;if(typeof v54InitTables==='function')v54InitTables();}

function setFleetCardFilter(state){
  const el=document.getElementById('fleetState');
  if(el) el.value=state||'';
  renderFleetOps();
}
function setLiveCardFilter(state){
  const el=document.getElementById('liveState');
  if(!el)return;
  if(state==='BAKIM'){
    el.value='';
    renderLiveOps();
    const rows=[...document.querySelectorAll('#liveRows tr')];
    rows.forEach(r=>{r.style.display=(r.innerText||'').toLocaleUpperCase('tr-TR').includes('BAKIM')?'':'none';});
    return;
  }
  el.value=state||'';
  renderLiveOps();
}

async function loadDash(){
  await loadIntegratedFleetStatus();
  const s=await api('/api/stats');
  const d=await ft('', '', 25, 0);

  dTotal.innerText=s.total||0;
  dFreight.innerText=money(s.freight_total||0);
  dFuelLiters.innerText=money(s.fuel_consumed||0);
  dDue.innerText=money(s.due_total||0);
  dDiff.innerText=money(s.driver_diff||0);

  let totalDiff=Number(s.driver_diff||0);
  let diffCard=dDiff.closest('.card');
  diffCard.classList.remove('diff-red','diff-green');
  if(totalDiff<0) diffCard.classList.add('diff-red');
  else if(totalDiff>0) diffCard.classList.add('diff-green');

  dashRows.innerHTML='';
  d.forEach(x=>{
    let diffClass=Number(x.driver_cash_diff||0)<0?'style="background:#fee2e2;color:#991b1b;font-weight:800"':
                  Number(x.driver_cash_diff||0)>0?'style="background:#dcfce7;color:#166534;font-weight:800"':'';
    dashRows.innerHTML+=`<tr>
      <td><b>${x.scna}</b></td><td>${x.plate}</td><td>${money(x.actual_km)}</td>
      <td>${money(x.fuel_consumed_liters)}</td><td>${money(x.liters_per_100km)}</td><td>${money(x.km_per_liter)}</td>
      <td>${money(x.freight_total)}</td><td>${money(x.expected_cash_handover)}</td>
      <td style="font-weight:800">${money(x.amount_due)}</td><td ${diffClass}>${money(x.driver_cash_diff)}</td>
      <td>${sBadge(x.driver_cash_status)}</td></tr>`;
  });
}


let tripHeadField='';
let tripSortField='';
let tripSortDir='';
let tripColumnFilters={};
let tripOriginalRows=[];

function tripHeadClick(ev,field){
  tripHeadField=field;
  const menu=document.getElementById('tripHeadMenu');
  tripHeadMenuTitle.textContent={scna:'SCNA',plate:'PLAKA',driver:'ŞOFÖR',area:'BÖLGE',cargo_type:'TİP',exit_at:'ÇIKIŞ TARİHİ',entry_at:'GİRİŞ TARİHİ'}[field]||field;
  tripHeadSearch.value=tripColumnFilters[field]||'';
  const r=ev.currentTarget.getBoundingClientRect();
  menu.style.left=Math.min(r.left,window.innerWidth-290)+'px';
  menu.style.top=(r.bottom+4)+'px';
  menu.style.display='block';
}
function applyHeadSort(dir){tripSortField=tripHeadField;tripSortDir=dir;updateTripSortIcons();tripHeadMenu.style.display='none';renderCurrentTripRows();}
function applyHeadTextFilter(){const v=tripHeadSearch.value.trim();if(v)tripColumnFilters[tripHeadField]=v;else delete tripColumnFilters[tripHeadField];renderCurrentTripRows();}
function clearHeadFilter(){delete tripColumnFilters[tripHeadField];tripHeadSearch.value='';tripHeadMenu.style.display='none';renderCurrentTripRows();}
function clearAllHeadFilters(){tripColumnFilters={};tripSortField='';tripSortDir='';tripHeadMenu.style.display='none';updateTripSortIcons();renderCurrentTripRows();}
function updateTripSortIcons(){
 ['scna','plate','driver','area','cargo_type','exit_at','entry_at'].forEach(f=>{const e=document.getElementById('sort_'+f);if(e)e.textContent=tripSortField===f?(tripSortDir==='asc'?'▲':'▼'):'↕';});
}
function tripVal(x,f){
 if(f==='entry_at')return x.entry_display_at||x.entry_at||x.delivery_time||'';
 if(f==='cargo_type')return x.cargo_type_name||x.cargo_type||'';
 if(f==='driver')return x.driver_name||x.driver||'';
 if(f==='area')return x.area_name||x.area||'';
 if(f==='customer')return x.customer_name||x.customer||'';
 if(f==='cargo')return x.cargo_name||x.cargo||'';
 if(f==='vessel')return x.vessel_name||x.vessel||'';
 return x[f]??'';
}
function currentTripRows(){
 let a=[...tripOriginalRows];
 for(const [f,q] of Object.entries(tripColumnFilters)){
   const n=String(q).toLocaleUpperCase('tr-TR');
   a=a.filter(x=>String(tripVal(x,f)).toLocaleUpperCase('tr-TR').includes(n));
 }
 if(tripSortField){
   const f=tripSortField,m=tripSortDir==='desc'?-1:1;
   a.sort((x,y)=>{
     let A=tripVal(x,f),B=tripVal(y,f);
     if(f==='exit_at'||f==='entry_at'){A=A?new Date(A).getTime():0;B=B?new Date(B).getTime():0;return(A-B)*m;}
     return String(A).localeCompare(String(B),'tr',{numeric:true,sensitivity:'base'})*m;
   });
 }
 return a;
}
function renderCurrentTripRows(){
 const d=currentTripRows();
 tripRows.innerHTML='';
 d.forEach(x=>{
    tripRows.innerHTML+=`<tr>
      <td data-col="scna"><b>${x.scna}</b></td>
      <td data-col="plate">${x.plate}</td>
      <td data-col="driver">${x.driver_name||x.driver||""}</td>
      <td data-col="customer">${x.customer_name||x.customer||""}</td>
      <td data-col="area">${x.area_name||x.area||""}</td>
      <td data-col="vessel">${x.vessel_name||''}</td>
      <td data-col="cargo">${x.cargo_name||x.cargo||""}</td>
      <td data-col="cargo_type">${x.cargo_type_name||x.cargo_type||''}</td>
      <td data-col="kg">${money(x.kg)}</td>
      <td data-col="freight">${money(x.freight_total||x.freight)}</td>
      <td data-col="exit_at">${fmtDateTime(x.exit_at)}</td>
      <td data-col="entry_at">${fmtDateTime(x.entry_display_at||x.entry_at||x.delivery_time)}</td>
      <td data-col="status">
        <div style="display:flex;gap:8px;align-items:center">
          <select onchange="changeTripStatus('${x.scna}',this.value)">
            <option value="Bekliyor" ${(x.status_display||x.status)==='Bekliyor'?'selected':''}>Bekliyor</option>
            <option value="Yolda" ${(x.status_display||x.status)==='Yolda'?'selected':''}>Yolda</option>
            <option value="Tamamlandı" ${(x.status_display||x.status)==='Tamamlandı'?'selected':''}>Tamamlandı</option>
          </select>
          <button class="btn secondary" onclick="openEditFrom('${x.scna}')">Düzenle</button>
          ${AUTH_USER&&AUTH_USER.role==='ADMIN'?`<button class="btn" style="background:#dc2626;color:white" onclick="softDeleteTrip('${x.scna}')">Sil</button>`:''}
        </div>
      </td></tr>`;
 });
 applyColumnPrefs();
}
document.addEventListener('click',e=>{
 const m=document.getElementById('tripHeadMenu');
 if(m&&m.style.display!=='none'&&!m.contains(e.target)&&!e.target.closest('.clickhead'))m.style.display='none';
});

async function loadTrips(){
  const query=q.value||'';
  const status=statusFilter.value||'Tümü';
  const customer=tripCustomerFilter.value||'';
  const area=tripAreaFilter.value||'';
  const cargo=tripCargoFilter.value||'';
  const cnt=await api('/api/trips-count?q='+encodeURIComponent(query)+
    '&status='+encodeURIComponent(status)+
    '&customer='+encodeURIComponent(customer)+
    '&area='+encodeURIComponent(area)+
    '&cargo='+encodeURIComponent(cargo));
  tripTotal=cnt.count||0;
  const maxPage=Math.max(1,Math.ceil(tripTotal/tripPageSize));
  if(tripPage>maxPage) tripPage=maxPage;
  const d=await ft(query,status,tripPageSize,(tripPage-1)*tripPageSize,customer,area,cargo);
  tripOriginalRows=d||[];
  renderCurrentTripRows();
  tripPageInfo.innerText=tripPage+' / '+maxPage+'  ('+tripTotal+' kayıt)';
  applyColumnPrefs();
}


function openEditFrom(scna){
  switchPanel('editcenter');
  editScna.value=scna;
  loadEditCenter();
}

function prevTripPage(){ if(tripPage>1){tripPage--;loadTrips();} }
function nextTripPage(){ const m=Math.max(1,Math.ceil(tripTotal/tripPageSize)); if(tripPage<m){tripPage++;loadTrips();} }

function clearTripFilters(){
  tripColumnFilters={};tripSortField='';tripSortDir='';updateTripSortIcons();
  q.value='';
  statusFilter.value='Tümü';
  tripCustomerFilter.value='';
  tripAreaFilter.value='';
  tripCargoFilter.value='';
  tripPage=1;
  loadTrips();
}

async function changeTripStatus(scna,status){
  try{
    await api('/api/trips/'+encodeURIComponent(scna)+'/status',{
      method:'PATCH',headers:{'Content-Type':'application/json'},body:JSON.stringify({status})
    });
    await loadTrips();
  }catch(e){ alert(e.message); }
}



async function loadExitCashMissing(){
  const box=document.getElementById('exitCashMissingBox');
  if(!box)return;
  try{
    const rows=await api('/api/exit-cash-missing');
    if(!rows.length){box.style.display='none';box.innerHTML='';return;}
    box.style.display='block';
    box.innerHTML=`<div style="display:flex;justify-content:space-between;align-items:center;gap:10px;margin-bottom:10px">
      <div><b style="color:#92400e">⚠ ÇIKIŞ PARASI EKSİK</b><br><span class="small">30.08.2026 ve sonrası çıkışı tamamlanmış fakat şoföre verilen toplam nakit girilmemiş ${rows.length} kayıt.</span></div>
      <span class="badge warn">${rows.length} KAYIT</span>
    </div>
    <div class="table"><table class="mini"><thead><tr><th>SCNA</th><th>PLAKA</th><th>ŞOFÖR</th><th>ÇIKIŞ TARİHİ</th><th>RESMİ MAZOT</th><th>TİCARİ MAZOT</th><th>BAĞDAT MAZOT</th><th>MAZOT TOPLAM</th><th>HARCIRAH</th><th>PRİM</th><th>OTHER</th><th>DOCK</th><th>PORT</th><th>SONAR</th><th>GÖRÜNEN GİDER TOPLAMI</th><th>ŞOFÖRE VERİLEN TOPLAM NAKİT</th><th>İŞLEM</th></tr></thead><tbody>
    ${rows.map(x=>`<tr class="alert-orange"><td><b>${x.scna}</b></td><td>${x.plate||''}</td><td>${x.driver_name||''}</td><td>${fmtDateTime(x.exit_at||x.trip_date||x.created_at)}</td><td>${money(x.exit_official_fuel_total||0)}</td><td>${money(x.exit_commercial_fuel_total||0)}</td><td>${money(x.exit_baghdad_fuel_total||0)}</td><td><b>${money(x.fuel_total||0)}</b></td><td>${money(x.exit_allowance||0)}</td><td>${money(x.exit_premium||0)}</td><td>${money(x.exit_other||0)}</td><td>${money(x.dock_fee||0)}</td><td>${money(x.port_fee||0)}</td><td>${money(x.sonar||0)}</td><td><b>${money(x.visible_expense_total||0)}</b></td><td><input id="missingCash_${String(x.scna).replace(/[^A-Za-z0-9_]/g,'_')}" type="text" inputmode="decimal" placeholder="Toplam IQD" style="width:150px"></td><td><button class="btn orange" onclick="saveMissingExitCash('${String(x.scna).replace(/'/g,"\\'")}',Number(x.visible_expense_total||0))">Kaydet</button></td></tr>`).join('')}
    </tbody></table></div>`;
  }catch(e){
    box.style.display='block';box.innerHTML='<b style="color:#991b1b">Çıkış parası eksik kayıtlar yüklenemedi: '+e.message+'</b>';
  }
}

async function saveMissingExitCash(scna,allowedTotal=0){
  const id='missingCash_'+String(scna).replace(/[^A-Za-z0-9_]/g,'_');
  const el=document.getElementById(id);
  const amount=Number(String(el?.value||'').replace(/\./g,'').replace(',','.'));
  if(!amount||amount<=0){alert('Şoföre verilen toplam nakdi girin.');return;}
  allowedTotal=Number(allowedTotal||0);
  if(allowedTotal>0 && amount>allowedTotal+0.01){
    alert('HATA: Fazla para girdiniz.\n\nGider toplamı: '+money(allowedTotal)+' IQD\nGirilen: '+money(amount)+' IQD\nFazla: '+money(amount-allowedTotal)+' IQD\n\nKayıt yapılmadı.');
    el.focus();
    return;
  }
  if(!confirm(scna+' için şoföre verilen toplam nakit '+money(amount)+' IQD olarak kaydedilsin mi?'))return;
  try{
    await api('/api/trips/'+encodeURIComponent(scna)+'/exit-cash',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({amount})});
    await loadExitCashMissing();
    if(typeof loadCashControl==='function'){
      try{await loadCashControl();}catch(e){}
    }
  }catch(e){alert(e.message);}
}

async function loadExit(){
  await loadExitCashMissing();
  const status=exitStatusFilter.value||'Bekliyor';
  const d=await ft(exitQ.value||'',status,100,0);
  exitRows.innerHTML='';
  d.forEach(x=>{
    exitRows.innerHTML+=`<tr>
      <td><b>${x.scna}</b></td>
      <td>${x.plate}</td>
      <td>${money(x.freight_total)}</td>
      <td>${fmtDateTime(x.exit_at)}</td>
      <td>${x.status}</td>
      <td><button class="btn orange" onclick="openExit('${x.scna}')">${x.exit_done?'Çıkışı Düzenle':'Çıkış Yap'}</button> <button class="btn secondary" onclick="printTrip('${x.scna}','ÇIKIŞ')">🖨 PRINT</button></td>
    </tr>`;
  });
  if(document.getElementById('exitCount')) exitCount.innerText=d.length+' İŞLEM';
  applyPanelDateFilter('exit');
}

async function loadEntry(){
  const status=entryStatusFilter.value||'Yolda';
  const q=entryQ.value||'';
  const [d,countData]=await Promise.all([
    ft(q,status,100,0),
    api('/api/trips-count?q='+encodeURIComponent(q)+'&status='+encodeURIComponent(status))
  ]);
  const total=Number((countData&&countData.count)??0);
  if(document.getElementById('entryCount')) entryCount.innerText=total+' İŞLEM';
  entryRows.innerHTML='';
  d.forEach(x=>{
    entryRows.innerHTML+=`<tr>
      <td><b>${x.scna}</b></td>
      <td>${x.plate}</td>
      <td>${x.driver_name||''}</td>
      <td>${money(x.exit_km)}</td>
      <td>${fmtDateTime(x.exit_at)}</td>
      <td>${fmtDateTime(x.entry_display_at||x.entry_at||x.delivery_time)}</td>
      <td>${money(x.freight_total)}</td>
      <td><button class="btn green" onclick="openEntry('${x.scna}')">${x.entry_done?'Girişi Düzenle':'Giriş Yap'}</button> <button class="btn secondary" onclick="printTrip('${x.scna}','GİRİŞ')">🖨 PRINT</button></td>
    </tr>`;
  });
  applyPanelDateFilter('entry');
}




let qualityCache={rows:[],counts:{},total:0};
let compareCache={rows:[],counts:{},total:0};

async function loadQuality(){
  qualityCache=await api('/api/data-quality');
  qualityTotal.innerText=qualityCache.total||0;
  qualityEntryKm.innerText=qualityCache.counts.entry_date_no_km||0;
  qualityCollection.innerText=qualityCache.counts.collection_missing||0;
  qualityDriver.innerText=qualityCache.counts.driver_missing||0;
  renderQuality();
}
function renderQuality(){
  const qv=(document.getElementById('qualityQ')?.value||'').toLowerCase();
  const type=document.getElementById('qualityType')?.value||'';
  const rows=(qualityCache.rows||[]).filter(x=>{
    const hay=[x.scna,x.plate,x.driver_name,x.customer_name,x.area_name].join(' ').toLowerCase();
    if(qv&&!hay.includes(qv)) return false;
    if(type&&!x[type]) return false;
    const df=document.getElementById('qualityFrom')?.value||''; const dt=document.getElementById('qualityTo')?.value||'';
    const raw=x.entry_display_at||x.entry_at||x.delivery_time||x.exit_at||x.trip_date||''; const ds=String(raw).slice(0,10);
    if(df&&ds&&ds<df)return false; if(dt&&ds&&ds>dt)return false;
    return true;
  });
  qualityRows.innerHTML='';
  if(!rows.length){
    qualityRows.innerHTML='<tr><td colspan="10" class="empty-state">Bu filtrede sorunlu kayıt yok.</td></tr>'; return;
  }
  rows.forEach(x=>{
    qualityRows.innerHTML+=`<tr class="alert-orange">
      <td><b>${x.scna}</b></td><td>${x.plate||''}</td><td>${x.driver_name||''}</td>
      <td>${x.customer_name||''}</td><td>${x.area_name||''}</td><td>${x.status||''}</td>
      <td>${fmtDateTime(x.exit_at)}</td><td>${fmtDateTime(x.entry_display_at||x.entry_at||x.delivery_time)}</td>
      <td>${(x.problems||[]).join('<br>')}</td>
      <td><button class="quick-btn" onclick="openEditFrom('${x.scna}')">Düzelt</button></td></tr>`;
  });
}

async function compareExcelDb(input){
  const file=input.files&&input.files[0]; if(!file) return;
  compareContent.innerHTML='<div class="calc"><b>Karşılaştırılıyor...</b></div>';
  try{
    const raw=await file.arrayBuffer();
    const r=await fetch('/api/compare-excel-db',{method:'POST',headers:{'Content-Type':'application/octet-stream'},body:raw});
    const data=await r.json(); if(!r.ok) throw new Error(data.detail||'Karşılaştırma hatası');
    compareCache=data;
    cmpDifferent.innerText=data.counts.different||0;
    cmpExcelOnly.innerText=data.counts.excel_only||0;
    cmpDbOnly.innerText=data.counts.db_only||0;
    renderCompare();
  }catch(e){ compareContent.innerHTML='<div class="calc"><b style="color:#991b1b">'+e.message+'</b></div>'; }
}
function renderCompare(){
  const type=document.getElementById('compareFilter')?.value||'';
  const rows=(compareCache.rows||[]).filter(x=>!type||x.result===type);
  let h=`<div class="table"><table><thead><tr><th>EXCEL SATIR</th><th>SCNA</th><th>SONUÇ</th><th>FARKLAR</th><th>İŞLEM</th></tr></thead><tbody>`;
  if(!rows.length) h+='<tr><td colspan="5" class="empty-state">Fark bulunmadı.</td></tr>';
  rows.forEach(x=>{
    h+=`<tr class="${x.result==='FARKLI'?'alert-orange':'alert-red'}">
      <td>${x.row||''}</td><td><b>${x.scna||''}</b></td><td>${x.result||''}</td>
      <td>${(x.differences||[]).join('<br>')}</td>
      <td>${x.result!=='DB_VAR_EXCEL_YOK'?`<button class="quick-btn" onclick="openEditFrom('${x.scna}')">DB Kaydını Aç</button>`:''}</td></tr>`;
  });
  h+='</tbody></table></div>'; compareContent.innerHTML=h;
}

function refreshBulkValue(){
  const f=bulkField.value; let h='';
  if(f==='status') h='<select id="bulkValue"><option>Bekliyor</option><option>Yolda</option><option>Tamamlandı</option></select>';
  else if(f==='customer_id') h='<select id="bulkValue"><option value="">Seçin</option>'+(L.customers||[]).map(x=>`<option value="${x.id}">${x.name}</option>`).join('')+'</select>';
  else if(f==='area_id') h='<select id="bulkValue"><option value="">Seçin</option>'+(L.areas||[]).map(x=>`<option value="${x.id}">${x.name}</option>`).join('')+'</select>';
  else if(f==='cargo_category_id') h='<select id="bulkValue"><option value="">Seçin</option>'+(L.cargo_categories||[]).map(x=>`<option value="${x.id}">${x.name}</option>`).join('')+'</select>';
  else if(f==='driver_id'){
    const seen=new Set(), arr=(L.vehicles||[]).filter(x=>{const id=x.current_driver_id||x.driver_id;if(!id||seen.has(id))return false;seen.add(id);return true;});
    h='<select id="bulkValue"><option value="">Seçin</option>'+arr.map(x=>`<option value="${x.current_driver_id||x.driver_id}">${x.driver_name||''}</option>`).join('')+'</select>';
  } else if(f==='cargo_type') h='<select id="bulkValue"><option>BULK</option><option>BAG</option></select>';
  else if(f==='trip_date') h='<input id="bulkValue" type="date">';
  bulkValueBox.innerHTML=h;
}
async function runBulkFix(){
  const scnas=(bulkScnas.value||'').split(/\r?\n|,|;/).map(x=>x.trim().toUpperCase()).filter(Boolean);
  if(!scnas.length){alert('SCNA listesi boş.');return;}
  if(!confirm(scansText(scans=scnas)))return;
  const value=document.getElementById('bulkValue')?.value||'';
  const r=await api('/api/bulk-fix',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({scnas,field:bulkField.value,value})});
  bulkResult.style.display='block';
  bulkResult.innerHTML='<b style="color:#166534">'+r.updated+' kayıt güncellendi.</b>'+(r.failed?.length?'<br>'+r.failed.map(x=>x.scna+': '+x.reason).join('<br>'):'');
}
function scansText(scans){return scans.length+' kayıt toplu olarak değiştirilecek. Emin misin?';}


let editCx=null;

async function loadEditCenter(){
  const scna=(editScna.value||'').trim().toUpperCase();
  if(!scna) return;

  try{
    editCx=await api('/api/trips/'+encodeURIComponent(scna)+'/edit-data');

    const compactLegacy=String(editCx.scna||'').replace(/[^A-Z0-9]/gi,'');
    const legacyBad=(!editCx.scna) || (/^\d+$/.test(String(editCx.scna)) ? String(editCx.scna).length<5 : compactLegacy.length<5);
    if(legacyBad){
      setTimeout(()=>{
        if(window.editValidation){
          editValidation.style.display='block';
          editValidation.innerHTML='<b style="color:#9a3412">Eski kayıt uyarısı:</b> Bu SCNA eksik/hatalı formatta kaydedilmiş. Kayıt açıldı; yeni kayıtlar artık bu şekilde kaydedilemez.';
        }
      },50);
    }


    const drivers=(L.vehicles||[])
      .filter((v,i,a)=>v.driver_id && a.findIndex(z=>z.driver_id===v.driver_id)===i)
      .map(v=>`<option value="${v.driver_id}">${v.driver_name||''}</option>`).join('');

    const areas=(L.areas||[]).map(a=>`<option value="${a.id}">${a.name}</option>`).join('');
    const cargos=(L.cargo_categories||[]).map(a=>`<option value="${a.id}">${a.name}</option>`).join('');
    const customers=(L.customers||[]).map(a=>`<option value="${a.id}">${a.name}</option>`).join('');

    editCenterContent.innerHTML=`
      <div class="calc">
        <b>${scna}</b> kaydı düzenleniyor.
        <span class="small">Kaydetmeden önce program mantık kontrolleri yapacak.</span>
        <div style="margin-top:8px">
          ${Number(editCx.manual_lock||0)===1
            ? `<span class="badge" style="background:#dcfce7;color:#166534">MANUEL KORUMALI</span>
               <span class="small">Excel güncellemesi bu SCNA'nın üzerine yazamaz.</span>`
            : `<span class="badge">EXCEL GÜNCELLEMESİNE AÇIK</span>
               <span class="small">İlk manuel kayıtta otomatik korunacak.</span>`}
        </div>
      </div>

      <div class="grid" style="margin-top:14px">
        <div class="field"><label>Plaka</label><input id="ePlate" value="${editCx.plate||''}"></div>
        <div class="field"><label>Şoför</label><select id="eDriver"><option value="">Seçin</option>${drivers}</select></div>
        <div class="field"><label>Müşteri</label><select id="eCustomer"><option value="">Seçin</option>${customers}</select></div>
        <div class="field"><label>Bölge</label><select id="eArea"><option value="">Seçin</option>${areas}</select></div>

        <div class="field"><label>Mal Cinsi</label><select id="eCargo"><option value="">Seçin</option>${cargos}</select></div>
        <div class="field"><label>Yük Tipi</label><select id="eCargoType"><option>BULK</option><option>BAG</option></select></div>
        <div class="field"><label>Sevkiyat Tarihi</label><input id="eDate" type="date" value="${editCx.trip_date||''}"></div>
        <div class="field"><label>Durum</label><select id="eStatus"><option>Bekliyor</option><option>Yolda</option><option>Tamamlandı</option></select></div>

        <div class="field"><label>Net KG</label><input id="eKg" type="number" value="${editCx.net_kg||0}"></div>
        <div class="field"><label>Navlun Birim Fiyat</label><input id="eRate" type="text" inputmode="decimal" value="${editCx.freight_rate||0}"></div>
        <div class="field"><label>Navlun Bazı</label><select id="eBasis"><option>KG</option><option>TON</option></select></div>
        <div class="field"><label>Çıkış KM</label><input id="eExitKm" type="number" value="${editCx.exit_km||0}"></div>

        <div class="field"><label>Başlangıç Depo Mazot</label><input id="eTankStart" type="number" value="${editCx.tank_start_liters||0}"></div>
        <div class="field"><label>Prim</label><input id="ePremium" type="text" inputmode="decimal" value="${editCx.exit_premium||0}"></div>
        <div class="field"><label>Dock Fee</label><input id="eDock" type="text" inputmode="decimal" value="${editCx.dock_fee||0}"></div>
        <div class="field"><label>Port Fee</label><input id="ePort" type="text" inputmode="decimal" value="${editCx.port_fee||0}"></div>

        <div class="field"><label>SONAR</label><input id="eSonar" type="text" inputmode="decimal" value="${editCx.sonar||0}"></div>
        <div class="field"><label>Harcırah</label><input id="eAllowance" type="text" inputmode="decimal" value="${editCx.exit_allowance||0}"></div>
        <div class="field"><label>Çıkış Diğer</label><input id="eOther" type="text" inputmode="decimal" value="${editCx.exit_other||0}"></div>
        <div class="field"><label>Giriş KM</label><input id="eEntryKm" type="number" value="${editCx.entry_km||0}"></div>

        <div class="field"><label>Dönüş Depoda Kalan LT</label><input id="eTankEnd" type="number" value="${editCx.tank_end_liters||0}"></div>
        <div class="field"><label>Ekstra Gider 1</label><input id="eExtra1" type="text" inputmode="decimal" value="${editCx.entry_extra_expense_1||0}"></div>
        <div class="field"><label>Ekstra Gider 2</label><input id="eExtra2" type="text" inputmode="decimal" value="${editCx.entry_extra_expense_2||0}"></div>
        <div class="field"><label>Ekstra Gider 3</label><input id="eExtra3" type="text" inputmode="decimal" value="${editCx.entry_extra_expense_3||0}"></div>

        <div class="field"><label>Müşteriden Alınan Para</label><input id="eCollection" type="text" inputmode="decimal" value="${editCx.entry_collection||0}"></div>
        <div class="field"><label>Şoförün Verdiği Para</label><input id="eHand" type="text" inputmode="decimal" value="${editCx.entry_cash_handed||0}"></div>
        <div class="field wide"><label>Not</label><textarea id="eNote">${editCx.entry_note||''}</textarea></div>
      </div>

      <div id="editValidation" class="calc" style="display:none;margin-top:12px"></div>

      <div class="compact-actions" style="margin-top:14px">
        <button class="btn green" onclick="validateAndSaveEdit()">Kontrol Et ve Kaydet</button>
        ${Number(editCx.manual_lock||0)===1
          ? `<button class="btn secondary" onclick="toggleManualLock(false)">Excel Korumasını Kaldır</button>`
          : `<button class="btn secondary" onclick="toggleManualLock(true)">Excel Korumasını Aç</button>`}
        <button class="btn orange" onclick="undoEntryEdit()">Giriş İşlemini Geri Al</button>
        <button class="btn secondary" onclick="undoExitEdit()">Çıkış İşlemini Geri Al</button>
      </div>
    `;

    eDriver.value=editCx.driver_id||'';
    eCustomer.value=editCx.customer_id||'';
    eArea.value=editCx.area_id||'';
    eCargo.value=editCx.cargo_category_id||'';
    eCargoType.value=editCx.cargo_type||'BULK';
    eBasis.value=editCx.freight_basis||'KG';
    eStatus.value=editCx.status||'Bekliyor';

  }catch(e){
    editCenterContent.innerHTML='<div class="calc"><b style="color:#991b1b">'+e.message+'</b></div>';
  }
}

function editPayload(){
  return {
    plate:ePlate.value,
    driver_id:eDriver.value?Number(eDriver.value):null,
    area_id:eArea.value?Number(eArea.value):null,
    cargo_category_id:eCargo.value?Number(eCargo.value):null,
    cargo_type:eCargoType.value,
    customer_id:eCustomer.value?Number(eCustomer.value):null,
    trip_date:eDate.value,
    net_kg:Number(eKg.value||0),
    freight_rate:parseMoney(eRate.value),
    freight_basis:eBasis.value,
    exit_km:Number(eExitKm.value||0),
    tank_start_liters:Number(eTankStart.value||0),
    exit_premium:parseMoney(ePremium.value),
    dock_fee:parseMoney(eDock.value),
    port_fee:parseMoney(ePort.value),
    sonar:parseMoney(eSonar.value),
    exit_allowance:parseMoney(eAllowance.value),
    exit_other:parseMoney(eOther.value),
    entry_km:Number(eEntryKm.value||0),
    tank_end_liters:Number(eTankEnd.value||0),
    entry_extra_expense_1:parseMoney(eExtra1.value),
    entry_extra_expense_2:parseMoney(eExtra2.value),
    entry_extra_expense_3:parseMoney(eExtra3.value),
    entry_collection:parseMoney(eCollection.value),
    entry_cash_handed:parseMoney(eHand.value),
    status:eStatus.value,
    note:eNote.value
  };
}


async function toggleManualLock(lockIt){
  const scna=(editScna.value||'').trim().toUpperCase();
  if(!scna) return;
  if(!lockIt){
    if(!confirm(scna+' için Excel koruması kaldırılacak. Sonraki Excel Güncelle işleminde bu kayıt değişebilir. Emin misin?')) return;
  }
  try{
    await api('/api/trips/'+encodeURIComponent(scna)+(lockIt?'/manual-lock':'/manual-unlock'),{method:'POST'});
    await loadEditCenter();
  }catch(e){
    alert(e.message);
  }
}

async function validateAndSaveEdit(){
  const scna=(editScna.value||'').trim().toUpperCase();
  if(!scna) return;
  const payload=editPayload();

  try{
    const v=await api('/api/trips/'+encodeURIComponent(scna)+'/validate',{
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify(payload)
    });

    editValidation.style.display='block';

    if(v.errors && v.errors.length){
      editValidation.innerHTML='<b style="color:#991b1b">Kayıt engellendi:</b><br>'+v.errors.join('<br>');
      return;
    }

    if(v.warnings && v.warnings.length){
      editValidation.innerHTML='<b style="color:#9a3412">Uyarılar:</b><br>'+v.warnings.join('<br>')+
        '<br><br><button class="btn orange" onclick="saveEditConfirmed()">Uyarıya Rağmen Kaydet</button>';
      return;
    }

    await saveEditConfirmed();
  }catch(e){
    alert(e.message);
  }
}

async function saveEditConfirmed(){
  const scna=(editScna.value||'').trim().toUpperCase();
  const payload=editPayload();
  try{
    const r=await api('/api/trips/'+encodeURIComponent(scna),{
      method:'PUT',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify(payload)
    });
    editValidation.style.display='block';
    editValidation.innerHTML='<b style="color:#166534">Kayıt güncellendi ve Excel güncellemelerine karşı korumaya alındı.</b>';
    await loadEditCenter();
  }catch(e){
    alert(e.message);
  }
}

async function undoEntryEdit(){
  const scna=(editScna.value||'').trim().toUpperCase();
  if(!confirm(scna+' giriş işlemi geri alınacak. Emin misin?')) return;
  await api('/api/trips/'+encodeURIComponent(scna)+'/undo-entry',{method:'POST'});
  await loadEditCenter();
}

async function undoExitEdit(){
  const scna=(editScna.value||'').trim().toUpperCase();
  if(!confirm(scna+' çıkış + giriş + yolda mazot kayıtları geri alınacak. Emin misin?')) return;
  await api('/api/trips/'+encodeURIComponent(scna)+'/undo-exit',{method:'POST'});
  await loadEditCenter();
}

async function loadScnaDetail(){
  const scna=(detailScna.value||'').trim().toUpperCase();
  if(!scna) return;
  try{
    const d=await api('/api/scna-detail/'+encodeURIComponent(scna));
    const x=d.trip;

    let html=`
      <div class="detail-grid">
        <div class="detail-box"><span>SCNA</span><strong>${x.scna}</strong></div>
        <div class="detail-box"><span>Plaka</span><strong>${x.plate||''}</strong></div>
        <div class="detail-box"><span>Şoför</span><strong>${x.driver_name||''}</strong></div>
        <div class="detail-box"><span>Durum</span><strong>${x.status||''}</strong></div>

        <div class="detail-box"><span>Müşteri</span><strong>${x.customer_name||''}</strong></div>
        <div class="detail-box"><span>Bölge</span><strong>${x.area_name||''}</strong></div>
        <div class="detail-box"><span>Mal / Tip</span><strong>${x.cargo_name||''} ${x.cargo_type_name||''}</strong></div>
        <div class="detail-box"><span>Net KG</span><strong>${money(x.net_kg)}</strong></div>

        <div class="detail-box"><span>Toplam Navlun</span><strong>${money(x.freight_total)}</strong></div>
        <div class="detail-box"><span>Müşteriden Alınan</span><strong>${money(x.entry_collection)}</strong></div>
        <div class="detail-box"><span>Müşteriden Kalan</span><strong>${money(x.amount_due)}</strong></div>
        <div class="detail-box"><span>Şoförün Vermesi Gereken</span><strong>${money(x.expected_cash_handover)}</strong></div>

        <div class="detail-box"><span>Çıkış</span><strong>${fmtDateTime(x.exit_at)}</strong></div>
        <div class="detail-box"><span>Giriş</span><strong>${fmtDateTime(x.entry_at)}</strong></div>
        <div class="detail-box"><span>Gerçek KM</span><strong>${money(x.actual_km)}</strong></div>
        <div class="detail-box"><span>LT / 100 KM</span><strong>${money(x.liters_per_100km)}</strong></div>

        <div class="detail-box"><span>Toplam Gider</span><strong>${money(x.total_expense)}</strong></div>
        <div class="detail-box"><span>Sefer Kârı</span><strong>${money(x.trip_profit)}</strong></div>
        <div class="detail-box"><span>Şoför Para Farkı</span><strong>${money(x.driver_cash_diff)}</strong></div>
        <div class="detail-box"><span>Sefer Süresi (Saat)</span><strong>${money(x.trip_duration_hours)}</strong></div>
      </div>`;

    html+=`<div class="table"><table><thead><tr>
      <th>YOLDA MAZOT LT</th><th>TOPLAM</th><th>BİRİM</th><th>NOT</th><th>TARİH</th>
    </tr></thead><tbody>`;

    if(d.fuels.length){
      d.fuels.forEach(f=>{
        html+=`<tr><td>${money(f.liters)}</td><td>${money(f.total)}</td>
          <td>${money(Number(f.total||0)/Number(f.liters||1))}</td>
          <td>${f.note||''}</td><td>${f.created_at||''}</td></tr>`;
      });
    }else{
      html+='<tr><td colspan="5">Yolda mazot kaydı yok.</td></tr>';
    }
    html+='</tbody></table></div>';

    html+=`<div class="table" style="margin-top:14px"><table><thead><tr>
      <th>TARİH</th><th>İŞLEM</th><th>DETAY</th>
    </tr></thead><tbody>`;

    if(d.logs.length){
      d.logs.forEach(l=>{
        html+=`<tr><td>${l.created_at||''}</td><td>${l.action||''}</td><td>${l.detail||''}</td></tr>`;
      });
    }else{
      html+='<tr><td colspan="3">İşlem geçmişi yok.</td></tr>';
    }
    html+='</tbody></table></div>';

    detailContent.innerHTML=html;
  }catch(e){
    detailContent.innerHTML='<div class="calc"><b style="color:#991b1b">'+e.message+'</b></div>';
  }
}


function switchPanel(id){
  document.querySelectorAll('.panel').forEach(x=>x.classList.remove('active'));
  const p=document.getElementById(id);
  if(p) p.classList.add('active');
  document.querySelectorAll('.nav button').forEach(x=>x.classList.remove('active'));
}

function openVehicleCardFrom(plate){
  switchPanel('vehiclecard');
  vehicleCardPlate.value=plate;
  loadVehicleCard();
}

function openDriverCardFrom(name){
  switchPanel('drivercard');
  driverCardName.value=name;
  loadDriverCard();
}

async function loadVehicleCard(){
  const plate=(vehicleCardPlate.value||'').trim().toUpperCase();
  if(!plate) return;

  try{
    const d=await api('/api/vehicle-card/'+encodeURIComponent(plate));
    const v=d.vehicle,s=d.summary,m=d.maintenance;

    let html=`
      <div class="detail-grid">
        <div class="detail-box"><span>Plaka</span><strong>${v.plate||''}</strong></div>
        <div class="detail-box"><span>Son / Aktif Sefer Şoförü</span><strong>${v.driver_name||''}</strong></div>
        <div class="detail-box"><span>Telefon</span><strong>${v.phone||''}</strong></div>
        <div class="detail-box"><span>D.No</span><strong>${v.d_no||''}</strong></div>

        <div class="detail-box"><span>Toplam Sefer</span><strong>${s.trip_count||0}</strong></div>
        <div class="detail-box"><span>Tamamlanan</span><strong>${s.completed_count||0}</strong></div>
        <div class="detail-box"><span>Toplam Ton</span><strong>${money(s.total_ton)}</strong></div>
        <div class="detail-box"><span>Toplam KM</span><strong>${money(s.total_km)}</strong></div>

        <div class="detail-box"><span>Ort. LT/100</span><strong>${money(s.avg_l100)}</strong></div>
        <div class="detail-box"><span>Toplam Navlun</span><strong>${money(s.total_freight)}</strong></div>
        <div class="detail-box"><span>Toplam Kâr</span><strong>${money(s.total_profit)}</strong></div>
        <div class="detail-box"><span>Sonraki Bakım</span><strong>${m?money(m.next_service_km):''}</strong></div>
      </div>`;

    if(m){
      html+=`<div class="calc"><b>Bakım:</b>
        Son bakım ${money(m.last_service_km)} KM |
        Sonraki bakım ${money(m.next_service_km)} KM |
        Yağ ${money(m.oil_change_km)} KM
        <br><span class="small">${m.note||''}</span>
      </div>`;
    }

    html+=`<div class="table" style="margin-top:14px"><table><thead><tr>
      <th>SCNA</th><th>TARİH</th><th>ŞOFÖR</th><th>BÖLGE</th><th>KG</th>
      <th>KM</th><th>LT/100</th><th>NAVLUN</th><th>KÂR</th><th>DURUM</th>
    </tr></thead><tbody>`;

    d.trips.forEach(x=>{
      html+=`<tr>
        <td><b>${x.scna}</b></td>
        <td>${x.trip_date||''}</td>
        <td>${x.driver_name||''}</td>
        <td>${x.area_name||''}</td>
        <td>${money(x.net_kg)}</td>
        <td>${money(x.actual_km)}</td>
        <td>${money(x.liters_per_100km)}</td>
        <td>${money(x.freight_total)}</td>
        <td>${money(x.trip_profit)}</td>
        <td>${x.status||''}</td>
      </tr>`;
    });

    html+='</tbody></table></div>';
    vehicleCardContent.innerHTML=html;
  }catch(e){
    vehicleCardContent.innerHTML='<div class="calc"><b style="color:#991b1b">'+e.message+'</b></div>';
  }
}

async function loadDriverCard(){
  const name=(driverCardName.value||'').trim();
  if(!name) return;

  try{
    const d=await api('/api/driver-card/'+encodeURIComponent(name));
    const x=d.driver,s=d.summary;

    let html=`
      <div class="detail-grid">
        <div class="detail-box"><span>Şoför</span><strong>${x.name||''}</strong></div>
        <div class="detail-box"><span>Telefon</span><strong>${x.phone||''}</strong></div>
        <div class="detail-box"><span>D.No</span><strong>${x.d_no||''}</strong></div>
        <div class="detail-box"><span>Kullandığı Araçlar</span><strong>${(d.plates||[]).join(', ')}</strong></div>

        <div class="detail-box"><span>Toplam Sefer</span><strong>${s.trip_count||0}</strong></div>
        <div class="detail-box"><span>Tamamlanan</span><strong>${s.completed_count||0}</strong></div>
        <div class="detail-box"><span>Aktif Sefer</span><strong>${s.active_trips||0}</strong></div>
        <div class="detail-box"><span>Toplam Ton</span><strong>${money(s.total_ton)}</strong></div>

        <div class="detail-box"><span>Toplam KM</span><strong>${money(s.total_km)}</strong></div>
        <div class="detail-box"><span>Ort. LT/100</span><strong>${money(s.avg_l100)}</strong></div>
        <div class="detail-box"><span>Toplam Para Farkı</span><strong>${money(s.cash_diff)}</strong></div>
      </div>`;

    html+=`<div class="table"><table><thead><tr>
      <th>SCNA</th><th>PLAKA</th><th>TARİH</th><th>BÖLGE</th><th>KG</th>
      <th>KM</th><th>LT/100</th><th>ŞOFÖR FARK</th><th>DURUM</th>
    </tr></thead><tbody>`;

    d.trips.forEach(t=>{
      const cls=Number(t.driver_cash_diff||0)<0?'class="alert-red"':'';
      html+=`<tr ${cls}>
        <td><b>${t.scna}</b></td>
        <td>${t.plate||''}</td>
        <td>${t.trip_date||''}</td>
        <td>${t.area_name||''}</td>
        <td>${money(t.net_kg)}</td>
        <td>${money(t.actual_km)}</td>
        <td>${money(t.liters_per_100km)}</td>
        <td>${money(t.driver_cash_diff)}</td>
        <td>${t.status||''}</td>
      </tr>`;
    });

    html+='</tbody></table></div>';
    driverCardContent.innerHTML=html;
  }catch(e){
    driverCardContent.innerHTML='<div class="calc"><b style="color:#991b1b">'+e.message+'</b></div>';
  }
}

let anomalyCache={fuel:[],km:[],cash:[],repeated_driver_cash:[]};


let dailyCache=null;

function localISODate(d=new Date()){
  const y=d.getFullYear(),m=String(d.getMonth()+1).padStart(2,'0'),da=String(d.getDate()).padStart(2,'0');
  return `${y}-${m}-${da}`;
}
function dailyToday(){
  dailyDate.value=localISODate();
  loadDailySummary();
}
async function loadDailySummary(){
  const day=dailyDate.value||localISODate();
  dailyCache=await api('/api/daily-manager-summary?day='+encodeURIComponent(day));
  const d=dailyCache;
  dailyExited.innerText=d.exited||0;
  dailyEntered.innerText=d.entered||0;
  dailyOnroad.innerText=d.onroad||0;
  dailyDelayed.innerText=d.delayed||0;
  dailyTon.innerText=money(d.ton);
  dailyFreight.innerText=money(d.freight);
  dailyCollected.innerText=money(d.collected);
  dailyCost.innerText=money(d.cost);
  dailyProfit.innerText=money(d.profit);
  dailyDriverDiff.innerText=money(d.driver_diff);

  let h=`<div class="table"><table><thead><tr>
    <th>SCNA</th><th>PLAKA</th><th>ŞOFÖR</th><th>BÖLGE</th><th>DURUM</th>
    <th>ÇIKIŞ</th><th>GİRİŞ</th><th>KG</th><th>NAVLUN</th>
  </tr></thead><tbody>`;
  if(d.rows.length){
    d.rows.forEach(x=>{
      h+=`<tr><td><b>${x.scna}</b></td><td>${x.plate||''}</td><td>${x.driver_name||''}</td>
      <td>${x.area_name||''}</td><td>${x.status||''}</td><td>${fmtDateTime(x.exit_at)}</td>
      <td>${fmtDateTime(x.entry_display_at||x.entry_at||x.delivery_time)}</td><td>${money(x.net_kg)}</td><td>${money(x.freight_total)}</td></tr>`;
    });
  } else h+='<tr><td colspan="9" class="empty-state">Bu tarihte giriş/çıkış hareketi yok.</td></tr>';
  h+='</tbody></table></div>';
  dailyTable.innerHTML=h;
}

function exportDailyExcel(){
  if(!dailyCache){ alert('Önce raporu getir.'); return; }
  const d=dailyCache;
  const rows=[
    ['GÜN SONU YÖNETİCİ ÖZETİ',d.day],
    ['Çıkan Araç',d.exited],['Giren Araç',d.entered],['Şu An Yolda',d.onroad],
    ['Gecikmiş',d.delayed],['Toplam Ton',d.ton],['Navlun',d.freight],
    ['Tahsilat',d.collected],['Gider',d.cost],['Net Kâr',d.profit],
    ['Şoför Para Farkı',d.driver_diff],[],
    ['SCNA','PLAKA','ŞOFÖR','BÖLGE','DURUM','ÇIKIŞ','GİRİŞ','KG','NAVLUN']
  ];
  d.rows.forEach(x=>rows.push([
    x.scna,x.plate,x.driver_name,x.area_name,x.status,x.exit_at,x.entry_at,x.net_kg,x.freight_total
  ]));
  const ws=XLSX.utils.aoa_to_sheet(rows);
  const wb=XLSX.utils.book_new();
  XLSX.utils.book_append_sheet(wb,ws,'GUN SONU');
  XLSX.writeFile(wb,'GUN_SONU_'+d.day+'.xlsx');
}

function printDailyReport(){
  if(!dailyCache){ alert('Önce raporu getir.'); return; }
  const w=window.open('','_blank');
  w.document.write(`<html><head><title>Gün Sonu ${dailyCache.day}</title>
  <style>body{font-family:Arial;padding:25px}h2{margin-bottom:5px}.cards{display:flex;flex-wrap:wrap;gap:10px;margin:15px 0}
  .card{border:1px solid #ddd;padding:10px;min-width:150px}.card b{display:block;font-size:18px;margin-top:5px}
  table{width:100%;border-collapse:collapse;font-size:11px}th,td{border:1px solid #bbb;padding:6px;text-align:left}
  th{background:#eee}</style></head><body>
  <h2>Gün Sonu / Yönetici Özeti</h2><div>${dailyCache.day}</div>${dailyReport.innerHTML}</body></html>`);
  w.document.close();
  setTimeout(()=>w.print(),300);
}

async function loadAnomalies(){
  anomalyCache=await api('/api/anomalies');
  anFuelCount.innerText=anomalyCache.fuel.length;
  anKmCount.innerText=anomalyCache.km.length;
  anCashCount.innerText=anomalyCache.cash_total_count ?? anomalyCache.cash.length;
  anDriverCount.innerText=anomalyCache.repeated_driver_cash.length;
  renderAnomalies();
}

function renderAnomalies(){
  const type=anomalyType.value||'all';
  let html='';

  const renderTable=(title,rows,kind)=>{
    html+=`<div class="table" style="margin-top:14px"><table><thead><tr>
      <th colspan="7">${title}</th></tr><tr>
      <th>SCNA</th><th>PLAKA</th><th>ŞOFÖR</th><th>BÖLGE</th>
      <th>GERÇEK</th><th>REFERANS</th><th>SAPMA</th>
    </tr></thead><tbody>`;

    if(rows.length){
      rows.forEach(x=>{
        html+=`<tr class="${kind==='cash'?'alert-red':'alert-orange'}">
          <td><b>${x.scna||''}</b></td>
          <td>${x.plate||''}</td>
          <td>${x.driver_name||''}</td>
          <td>${x.area_name||''}</td>
          <td>${money(x.value)}</td>
          <td>${money(x.baseline)}</td>
          <td>${kind==='cash'?(x.source||''):('%'+money(x.deviation))}</td>
        </tr>`;
      });
    }else{
      html+='<tr><td colspan="7">Kayıt yok.</td></tr>';
    }
    html+='</tbody></table></div>';
  };

  if(type==='all' || type==='fuel') renderTable('YAKIT ANORMALLİKLERİ',anomalyCache.fuel,'fuel');
  if(type==='all' || type==='km') renderTable('KM ANORMALLİKLERİ',anomalyCache.km,'km');
  if(type==='all' || type==='cash'){
    renderTable('ŞOFÖR PARA EKSİKLERİ',anomalyCache.cash,'cash');
    if((anomalyCache.cash_total_count||0)>(anomalyCache.cash||[]).length){
      html+=`<div class="calc"><b>Toplam ${(anomalyCache.cash_total_count||0)} para eksiği kaydı var.</b>
      Performans için tabloda en yüksek 500 kayıt gösteriliyor.</div>`;
    }
  }

  if(type==='all' || type==='driver'){
    html+=`<div class="table" style="margin-top:14px"><table><thead><tr>
      <th>TEKRARLAYAN ŞOFÖR PARA EKSİĞİ</th><th>SON KONTROL</th><th>EKSİK SAYISI</th>
    </tr></thead><tbody>`;
    if(anomalyCache.repeated_driver_cash.length){
      anomalyCache.repeated_driver_cash.forEach(x=>{
        html+=`<tr class="alert-red">
          <td><b>${x.driver_name||''}</b></td>
          <td>${x.checked_count||0} sefer</td>
          <td>${x.short_count||0}</td>
        </tr>`;
      });
    }else{
      html+='<tr><td colspan="3">Tekrarlayan şoför para eksiği bulunmadı.</td></tr>';
    }
    html+='</tbody></table></div>';
  }

  anomalyContent.innerHTML=html;
}

async function loadVehiclePerformance(){
  const d=await api('/api/performance/vehicles');
  let html=`<div class="table"><table><thead><tr>
    <th>PLAKA</th><th>SEFER</th><th>TAMAMLANAN</th><th>TON</th><th>TOPLAM KM</th>
    <th>ORT. LT/100</th><th>TOPLAM NAVLUN</th><th>TOPLAM KÂR</th>
  </tr></thead><tbody>`;
  d.forEach(x=>{
    html+=`<tr>
      <td><b class="perf-key">${x.plate||''}</b></td>
      <td>${x.trip_count||0}</td>
      <td>${x.completed_count||0}</td>
      <td>${money(x.total_ton)}</td>
      <td>${money(x.total_km)}</td>
      <td>${money(x.avg_l100)}</td>
      <td>${money(x.total_freight)}</td>
      <td>${money(x.total_profit)}</td>
    </tr>`;
  });
  html+='</tbody></table></div>';
  performanceContent.innerHTML=html;
  renderPerformanceFilter();
}


function renderPerformanceFilter(){
  const qv=(document.getElementById('perfQ')?.value||'').toLowerCase();
  document.querySelectorAll('#performanceContent tbody tr').forEach(tr=>{
    const key=(tr.querySelector('.perf-key')?.textContent||'').toLowerCase();
    tr.style.display=!qv || key.includes(qv)?'':'none';
  });
}

async function loadDriverPerformance(){
  const d=await api('/api/performance/drivers');
  let html=`<div class="table"><table><thead><tr>
    <th>ŞOFÖR</th><th>SEFER</th><th>TAMAMLANAN</th><th>TON</th><th>TOPLAM KM</th>
    <th>ORT. LT/100</th><th>TOPLAM NAVLUN</th><th>PARA FARKI</th>
  </tr></thead><tbody>`;
  d.forEach(x=>{
    let cls=Number(x.cash_diff||0)<0?'alert-red':Number(x.cash_diff||0)>0?'style="background:#dcfce7;color:#166534;font-weight:800"':'';
    html+=`<tr>
      <td><b class="perf-key">${x.driver_name||''}</b></td>
      <td>${x.trip_count||0}</td>
      <td>${x.completed_count||0}</td>
      <td>${money(x.total_ton)}</td>
      <td>${money(x.total_km)}</td>
      <td>${money(x.avg_l100)}</td>
      <td>${money(x.total_freight)}</td>
      <td ${cls}>${money(x.cash_diff)}</td>
    </tr>`;
  });
  html+='</tbody></table></div>';
  performanceContent.innerHTML=html;
  renderPerformanceFilter();
}



function financeParams(){
  const f=finFrom.value||'';
  const t=finTo.value||'';
  return 'date_from='+encodeURIComponent(f)+'&date_to='+encodeURIComponent(t);
}


function setFinancePreset(type){
  const now=new Date();
  const iso=d=>d.toISOString().slice(0,10);
  let from=new Date(now),to=new Date(now);
  if(type==='month') from=new Date(now.getFullYear(),now.getMonth(),1);
  if(type==='year') from=new Date(now.getFullYear(),0,1);
  finFrom.value=iso(from);
  finTo.value=iso(to);
  loadFinance();
}

function clearFinanceDates(){
  finFrom.value='';
  finTo.value='';
  loadFinance();
}

async function loadFinance(){
  const s=await api('/api/finance-summary?'+financeParams());

  finFreight.innerText=money(s.total_freight);
  finCost.innerText=money(s.total_cost);
  finProfit.innerText=money(s.total_profit);
  finMargin.innerText='%'+money(s.margin_percent);
  finTon.innerText=money(s.total_ton);

  finFuel.innerText=money(s.fuel_cost);
  finOther.innerText=money(s.other_cost);
  finKm.innerText=money(s.total_km);
  finProfitKm.innerText=money(s.profit_per_km);
  finProfitTon.innerText=money(s.profit_per_ton);

  finTrips.innerText=s.trip_count||0;
  finCollected.innerText=money(s.collected);
  finReceivable.innerText=money(s.receivable);

  const profitCard=finProfit.closest('.card');
  profitCard.classList.remove('diff-red','diff-green');
  if(Number(s.total_profit||0)<0) profitCard.classList.add('diff-red');
  else if(Number(s.total_profit||0)>0) profitCard.classList.add('diff-green');

  await loadFinanceTable('customer');
}

async function loadFinanceTable(type){
  let url='/api/finance-by-customer';
  let first='MÜŞTERİ';

  if(type==='area'){
    url='/api/finance-by-area';
    first='BÖLGE';
  }else if(type==='cargo'){
    url='/api/finance-by-cargo';
    first='MAL CİNSİ';
  }

  const d=await api(url+'?'+financeParams());

  let html=`<div class="table"><table><thead><tr>
    <th>${first}</th><th>SEFER</th><th>TON</th>`;
  if(type==='area') html+='<th>TOPLAM KM</th>';
  html+='<th>NAVLUN</th><th>GİDER</th><th>NET KÂR</th><th>KÂR MARJI</th></tr></thead><tbody>';

  d.forEach(x=>{
    const margin=Number(x.freight||0)>0?(Number(x.profit||0)/Number(x.freight))*100:0;
    const cls=Number(x.profit||0)<0?'class="alert-red"':'';
    html+=`<tr ${cls}>
      <td><b>${x.name||''}</b></td>
      <td>${x.trip_count||0}</td>
      <td>${money(x.total_ton)}</td>`;
    if(type==='area') html+=`<td>${money(x.total_km)}</td>`;
    html+=`
      <td>${money(x.freight)}</td>
      <td>${money(x.cost)}</td>
      <td>${money(x.profit)}</td>
      <td>%${money(margin)}</td>
    </tr>`;
  });

  html+='</tbody></table></div>';
  financeTable.innerHTML=html;
}

async function loadFinanceMonthly(){
  const d=await api('/api/finance-monthly');
  let html=`<div class="table"><table><thead><tr>
    <th>AY</th><th>SEFER</th><th>TON</th><th>NAVLUN</th><th>GİDER</th><th>NET KÂR</th><th>MARJ</th>
  </tr></thead><tbody>`;

  d.forEach(x=>{
    const margin=Number(x.freight||0)>0?(Number(x.profit||0)/Number(x.freight))*100:0;
    html+=`<tr ${Number(x.profit||0)<0?'class="alert-red"':''}>
      <td><b>${x.month||''}</b></td>
      <td>${x.trip_count||0}</td>
      <td>${money(x.total_ton)}</td>
      <td>${money(x.freight)}</td>
      <td>${money(x.cost)}</td>
      <td>${money(x.profit)}</td>
      <td>%${money(margin)}</td>
    </tr>`;
  });

  html+='</tbody></table></div>';
  financeTable.innerHTML=html;
}


let opCache={summary:{},rows:[]};
let routeLearnCache=[];
let opPerfCache={vehicles:[],drivers:[]};

async function loadOperationCenter(){
  const [op,routes,perf]=await Promise.all([
    api('/api/operation-center'),
    api('/api/route-learning'),
    api('/api/operation-performance')
  ]);
  opCache=op;
  routeLearnCache=routes;
  opPerfCache=perf;

  opTotal.innerText=op.summary.total||0;
  opOnRoad.innerText=op.summary.onroad||0;
  opWaiting.innerText=op.summary.waiting||0;
  opDelayed.innerText=op.summary.delayed||0;
  opTodayExit.innerText=op.summary.today_exit||0;
  opTodayEntry.innerText=op.summary.today_entry||0;

  renderOperationCenter();
  renderRouteLearning();
  renderOpPerformance('vehicle');
}

function renderOperationCenter(){
  const qv=(document.getElementById('opQ')?.value||'').toLowerCase();
  const state=document.getElementById('opState')?.value||'';

  const rows=(opCache.rows||[]).filter(x=>{
    const hay=[x.scna,x.plate,x.driver_name,x.area_name].join(' ').toLowerCase();
    if(qv&&!hay.includes(qv)) return false;
    if(state==='Gecikmiş'&&!x.delayed) return false;
    if(state&&state!=='Gecikmiş'&&x.status!==state) return false;
    return true;
  });

  opRows.innerHTML='';
  if(!rows.length){
    opRows.innerHTML='<tr><td colspan="12" class="empty-state">Bu filtrede aktif operasyon yok.</td></tr>';
    return;
  }

  rows.forEach(x=>{
    const refHours=Number(x.route_avg_hours||x.expected_hours||0);
    let warning='';
    if(x.delayed){
      warning='<span class="badge short">GECİKMİŞ</span>';
    }else if(x.status==='Yolda'){
      warning='<span class="badge warn">YOLDA</span>';
    }

    opRows.innerHTML+=`<tr class="${x.delayed?'alert-red':''}">
      <td><b>${x.scna}</b></td>
      <td><button class="quick-btn" onclick="openVehicleCardFrom('${x.plate}')">${x.plate||''}</button></td>
      <td>${x.driver_name?`<button class="quick-btn" onclick="openDriverCardFrom('${String(x.driver_name).replace(/'/g,"\\'")}')">${x.driver_name}</button>`:''}</td>
      <td>${x.area_name||''}</td>
      <td>${x.status||''}</td>
      <td>${fmtDateTime(x.exit_at)}</td>
      <td>${money(x.hours_on_road)}</td>
      <td>${money(refHours)}</td>
      <td>${money(x.route_avg_km||x.expected_km||0)}</td>
      <td>${money(x.route_avg_l100||x.expected_l100||0)}</td>
      <td>${x.completed_trips||0}</td>
      <td>${warning}</td>
    </tr>`;
  });
}

function renderRouteLearning(){
  const qv=(document.getElementById('routeLearnQ')?.value||'').toLowerCase();
  const rows=routeLearnCache.filter(x=>!qv||String(x.area_name||'').toLowerCase().includes(qv));
  routeLearnRows.innerHTML='';
  if(!rows.length){
    routeLearnRows.innerHTML='<tr><td colspan="6" class="empty-state">Rota geçmişi bulunamadı.</td></tr>';
    return;
  }
  rows.forEach(x=>{
    routeLearnRows.innerHTML+=`<tr>
      <td><b>${x.area_name||''}</b></td>
      <td>${x.trip_count||0}</td>
      <td>${money(x.avg_km)}</td>
      <td>${money(x.avg_hours)}</td>
      <td>${money(x.avg_l100)}</td>
      <td>${money(x.avg_cost)}</td>
    </tr>`;
  });
}

function renderOpPerformance(type){
  const rows=type==='driver'?(opPerfCache.drivers||[]):(opPerfCache.vehicles||[]);
  const first=type==='driver'?'ŞOFÖR':'PLAKA';
  let h=`<div class="table"><table><thead><tr>
    <th>${first}</th><th>SEFER</th><th>TOPLAM KM</th><th>ORT. LT/100</th><th>ORT. SEFER SÜRESİ</th>
  </tr></thead><tbody>`;
  if(!rows.length){
    h+='<tr><td colspan="5" class="empty-state">Performans verisi yok.</td></tr>';
  }else{
    rows.forEach(x=>{
      h+=`<tr>
        <td><b>${x.name||''}</b></td>
        <td>${x.trip_count||0}</td>
        <td>${money(x.total_km)}</td>
        <td>${money(x.avg_l100)}</td>
        <td>${money(x.avg_hours)}</td>
      </tr>`;
    });
  }
  h+='</tbody></table></div>';
  opPerformanceContent.innerHTML=h;
}


let liveCache={summary:{},rows:[]};



let fleetManageData=[];
let driverManageData=[];

async function loadFleetManage(){
  const d=await api('/api/fleet/vehicles');fleetManageData=d.rows||[];
  fmTotal.innerText=d.total||0;fmActive.innerText=d.active||0;fmPassive.innerText=d.passive||0;
  renderFleetManage();
}
function renderFleetManage(){
  if(!document.getElementById('fleetManageRows'))return;
  const q=(fleetManageQ.value||'').toLocaleUpperCase('tr-TR'),s=fleetManageState.value;
  const a=fleetManageData.filter(x=>(s===''||String(x.is_active)===s)&&(!q||[x.plate,x.brand,x.model,x.vehicle_type,x.current_driver,x.note].some(v=>String(v||'').toLocaleUpperCase('tr-TR').includes(q))));
  fleetManageRows.innerHTML=a.map(x=>`<tr><td><b>${x.plate}</b></td><td>${x.brand||''}</td><td>${x.model||''}</td><td>${x.vehicle_type||''}</td><td>${x.current_driver||''}</td><td>${x.garage_state||''}</td><td>${x.is_active?'AKTİF':'PASİF'}</td><td>${x.note||''}</td><td><button class="btn secondary" onclick="openFleetVehicle('${x.plate}')">Düzenle</button> <button class="btn ${x.is_active?'orange':'green'}" onclick="toggleFleetVehicle('${x.plate}',${x.is_active?0:1})">${x.is_active?'Pasife Al':'Aktif Et'}</button></td></tr>`).join('');
  if(typeof v54InitTables==='function')v54InitTables();
}
function openFleetVehicle(plate=''){
  const x=fleetManageData.find(z=>z.plate===plate)||{plate:'',brand:'',model:'',vehicle_type:'DAMPER',is_active:1,garage_state:'GARAGE',note:''};
  openM(plate?'Araç Düzenle':'Yeni Araç',`<div class="grid">
  <div class="field"><label>Plaka</label><input id="fvPlate" value="${x.plate||''}" ${plate?'readonly':''}></div>
  <div class="field"><label>Marka</label><input id="fvBrand" value="${x.brand||''}"></div>
  <div class="field"><label>Model</label><input id="fvModel" value="${x.model||''}"></div>
  <div class="field"><label>Araç Tipi</label><select id="fvType"><option ${x.vehicle_type==='DAMPER'?'selected':''}>DAMPER</option><option ${x.vehicle_type==='TANKER'?'selected':''}>TANKER</option><option ${x.vehicle_type==='TRAILER'?'selected':''}>TRAILER</option><option ${x.vehicle_type==='OTHER'?'selected':''}>OTHER</option></select></div>
  <div class="field"><label>Garaj Durumu</label><select id="fvGarage"><option value="GARAGE" ${x.garage_state==='GARAGE'?'selected':''}>GARAGE</option><option value="DISARIDA" ${x.garage_state==='DISARIDA'?'selected':''}>DIŞARIDA</option><option value="BAKIM" ${x.garage_state==='BAKIM'?'selected':''}>BAKIM</option></select></div>
  <div class="field"><label>Aktif</label><select id="fvActive"><option value="1" ${x.is_active?'selected':''}>AKTİF</option><option value="0" ${!x.is_active?'selected':''}>PASİF</option></select></div>
  <div class="field wide"><label>Not</label><input id="fvNote" value="${String(x.note||'').replace(/"/g,'&quot;')}"></div></div>`,
  async()=>{await api('/api/fleet/vehicles',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({plate:fvPlate.value,brand:fvBrand.value,model:fvModel.value,vehicle_type:fvType.value,is_active:Number(fvActive.value),garage_state:fvGarage.value,note:fvNote.value})});closeM();await loadFleetManage();await loadFleetOps();});
}
async function toggleFleetVehicle(plate,active){
  await api('/api/fleet/vehicles/'+encodeURIComponent(plate)+'/active',{method:'PATCH',headers:{'Content-Type':'application/json'},body:JSON.stringify({is_active:active})});
  await loadFleetManage();await loadFleetOps();
}

async function loadDriverManage(){
  const d=await api('/api/fleet/drivers');driverManageData=d.rows||[];
  dmTotal.innerText=d.total||0;dmActive.innerText=d.active||0;dmPassive.innerText=d.passive||0;
  renderDriverManage();
}
function renderDriverManage(){
  if(!document.getElementById('driverManageRows'))return;
  const q=(driverManageQ.value||'').toLocaleUpperCase('tr-TR'),s=driverManageState.value;
  const a=driverManageData.filter(x=>(s===''||String(x.is_active)===s)&&(!q||[x.name,x.phone,x.d_no,x.last_plate,x.note].some(v=>String(v||'').toLocaleUpperCase('tr-TR').includes(q))));
  driverManageRows.innerHTML=a.map(x=>`<tr><td><b>${x.name}</b></td><td>${x.phone||''}</td><td>${x.d_no||''}</td><td>${x.last_plate||''}</td><td>${x.is_active?'AKTİF':'PASİF'}</td><td>${x.note||''}</td><td><button class="btn secondary" onclick="openFleetDriver(${x.id})">Düzenle</button> <button class="btn ${x.is_active?'orange':'green'}" onclick="toggleFleetDriver(${x.id},${x.is_active?0:1})">${x.is_active?'Pasife Al':'Aktif Et'}</button></td></tr>`).join('');
  if(typeof v54InitTables==='function')v54InitTables();
}
function openFleetDriver(id=0){
  const x=driverManageData.find(z=>z.id===id)||{name:'',phone:'',d_no:'',is_active:1,note:''};
  openM(id?'Şoför Düzenle':'Yeni Şoför',`<div class="grid">
  <div class="field"><label>Ad Soyad</label><input id="fdName" value="${String(x.name||'').replace(/"/g,'&quot;')}"></div>
  <div class="field"><label>Telefon</label><input id="fdPhone" value="${x.phone||''}"></div>
  <div class="field"><label>D.No</label><input id="fdNo" value="${x.d_no||''}"></div>
  <div class="field"><label>Aktif</label><select id="fdActive"><option value="1" ${x.is_active?'selected':''}>AKTİF</option><option value="0" ${!x.is_active?'selected':''}>PASİF</option></select></div>
  <div class="field wide"><label>Not</label><input id="fdNote" value="${String(x.note||'').replace(/"/g,'&quot;')}"></div></div>`,
  async()=>{await api('/api/fleet/drivers',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({name:fdName.value,phone:fdPhone.value,d_no:fdNo.value,is_active:Number(fdActive.value),note:fdNote.value})});closeM();await loadDriverManage();});
}
async function toggleFleetDriver(id,active){
  await api('/api/fleet/drivers/'+id+'/active',{method:'PATCH',headers:{'Content-Type':'application/json'},body:JSON.stringify({is_active:active})});
  await loadDriverManage();
}

let fleetData=[];
async function loadFleetOps(){
 const d=await api('/api/fleet-operations');fleetData=d.rows||[];const c=d.counts||{};
 if(document.getElementById('fleetTotal')){fleetTotal.innerText=c.TOPLAM||0;fleetEmpty.innerText=c.BOS||0;fleetLoaded.innerText=c.DOLU||0;fleetWaiting.innerText=c.SIRA||0;fleetLoading.innerText=c.YUKLEMEDE||0;}
 renderFleetOps();renderVesselOps();
}
function fleetLabel(s){return {BOS:'BOŞ',DOLU:'DOLU',SIRA:'SIRA BEKLİYOR',YUKLEMEDE:'YÜKLEMEDE'}[s]||s;}
function fleetFiltered(q,state){q=(q||'').trim().toLocaleUpperCase('tr-TR');return fleetData.filter(x=>!state||x.load_state===state).filter(x=>!q||[x.plate,x.driver,x.scna,x.area,x.vessel,x.operation_note].some(v=>String(v||'').toLocaleUpperCase('tr-TR').includes(q)));}
function renderFleetOps(){
 if(!document.getElementById('fleetRows'))return;const a=fleetFiltered(fleetQ.value,fleetState.value);
 fleetRows.innerHTML=a.map(x=>`<tr><td><b>${x.load_state==='SIRA'?(x.queue_no||'-'):'-'}</b></td><td><b>${x.plate}</b></td><td>${fleetLabel(x.load_state)}</td><td>${x.driver||''}</td><td>${x.scna||''}</td><td>${x.area||''}</td><td>${x.vessel||''}</td><td>${x.operation_note||''}</td><td><button class="btn secondary" onclick="editFleetOp('${x.plate}')">Durum Değiştir</button></td></tr>`).join('');
 if(typeof v54InitTables==='function')v54InitTables();
}
function vesselOpState(x){
  const vessel=String(x.vessel||'').trim();
  if(!vessel)return 'ATANMADI';
  if(x.vessel_state==='CIKTI')return 'CIKTI';
  return 'CALISIYOR';
}
function renderVesselOps(){
  if(!document.getElementById('vesselRows'))return;

  const q=(vesselQ.value||'').trim().toLocaleUpperCase('tr-TR');
  const state=vesselState.value;
  let a=fleetData.filter(x=>{
    const s=vesselOpState(x);
    if(state && s!==state)return false;
    if(!q)return true;
    return [x.vessel,x.plate,x.driver,x.operation_note].some(v=>
      String(v||'').toLocaleUpperCase('tr-TR').includes(q)
    );
  });

  const working=fleetData.filter(x=>vesselOpState(x)==='CALISIYOR' && String(x.vessel||'').trim()).length;
  const out=fleetData.filter(x=>vesselOpState(x)==='CIKTI').length;
  const unassigned=fleetData.filter(x=>!String(x.vessel||'').trim()).length;
  if(document.getElementById('vesselWorkingCount'))vesselWorkingCount.innerText=working;
  if(document.getElementById('vesselOutCount'))vesselOutCount.innerText=out;
  if(document.getElementById('vesselUnassignedCount'))vesselUnassignedCount.innerText=unassigned;

  const groups={};
  fleetData.filter(x=>vesselOpState(x)==='CALISIYOR' && String(x.vessel||'').trim())
    .forEach(x=>{const g=x.vessel.trim();groups[g]=(groups[g]||0)+1;});

  vesselSummary.innerHTML=Object.entries(groups)
    .sort((a,b)=>b[1]-a[1])
    .map(([g,n])=>`<div class="card">${g}<b>${n} araç</b></div>`).join('');

  vesselRows.innerHTML=a.map(x=>{
    const s=vesselOpState(x);
    return `<tr>
      <td><b>${x.vessel||'GEMİ ATANMADI'}</b></td>
      <td><b>${x.plate}</b></td>
      <td>${x.driver||''}</td>
      <td>${s==='CALISIYOR'?'GEMİDE ÇALIŞIYOR':(s==='CIKTI'?'OPERASYONDAN ÇIKTI':'GEMİ ATANMADI')}</td>
      <td>${x.operation_note||''}</td>
      <td><button class="btn secondary" onclick="editVesselAssignment('${x.plate}')">Düzenle</button></td>
    </tr>`;
  }).join('');

  if(typeof v54InitTables==='function')v54InitTables();
}

function editVesselAssignment(plate){
  const x=fleetData.find(z=>z.plate===plate)||{plate,vessel:'',vessel_state:'CALISIYOR',operation_note:''};
  const state=vesselOpState(x);

  openM('Gemi Operasyonu - '+plate,`<div class="grid">
    <div class="field"><label>Plaka</label><input id="voPlate" value="${plate}" readonly></div>
    <div class="field"><label>Gemi</label><input id="voVessel" value="${String(x.vessel||'').replace(/"/g,'&quot;')}" placeholder="Gemi adı"></div>
    <div class="field"><label>Durum</label><select id="voState">
      <option value="CALISIYOR" ${state==='CALISIYOR'?'selected':''}>GEMİDE ÇALIŞIYOR</option>
      <option value="CIKTI" ${state==='CIKTI'?'selected':''}>OPERASYONDAN ÇIKTI</option>
    </select></div>
    <div class="field wide"><label>Not</label><input id="voNote" value="${String(x.operation_note||'').replace(/"/g,'&quot;')}"></div>
  </div>`,
  async()=>{
    await api('/api/vessel-operations',{
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify({
        plate:voPlate.value,
        load_state:voState.value,
        queue_no:0,
        vessel:voVessel.value,
        operation_note:voNote.value
      })
    });
    closeM();
    await loadFleetOps();
    await loadIntegratedFleetStatus();
    try{await loadTrips();}catch(e){}
  });
}

function editFleetOp(plate){
 const x=fleetData.find(z=>z.plate===plate)||{plate,load_state:'BOS',queue_no:0,vessel:'',operation_note:''};
 openM('Araç Operasyon Durumu - '+plate,`<div class="grid">
 <div class="field"><label>Plaka</label><input id="foPlate" value="${plate}" readonly></div>
 <div class="field"><label>Durum</label><select id="foState" onchange="foQueue.disabled=this.value!=='SIRA'"><option value="BOS" ${x.load_state==='BOS'?'selected':''}>BOŞ</option><option value="DOLU" ${x.load_state==='DOLU'?'selected':''}>DOLU</option><option value="SIRA" ${x.load_state==='SIRA'?'selected':''}>SIRA BEKLİYOR</option><option value="YUKLEMEDE" ${x.load_state==='YUKLEMEDE'?'selected':''}>YÜKLEMEDE</option></select></div>
 <div class="field"><label>Sıra No</label><input id="foQueue" type="number" min="0" value="${x.queue_no||0}" ${x.load_state==='SIRA'?'':'disabled'}></div>
 <div class="field"><label>Gemi</label><input id="foVessel" value="${String(x.vessel||'').replace(/"/g,'&quot;')}"></div>
 <div class="field wide"><label>Operasyon Notu</label><input id="foNote" value="${String(x.operation_note||'').replace(/"/g,'&quot;')}"></div></div>`,
 async()=>{await api('/api/fleet-operations',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({plate:foPlate.value,load_state:foState.value,queue_no:Number(foQueue.value||0),vessel:foVessel.value,operation_note:foNote.value})});closeM();await loadFleetOps();await loadIntegratedFleetStatus();try{await loadTrips();}catch(e){}});
}
async function resequenceFleet(){await api('/api/fleet-operations/resequence',{method:'POST'});await loadFleetOps();}

async function loadLiveOps(){
  const d=await api('/api/live-operations');
  liveCache=d;
  liveTotal.innerText=d.summary.total||0;
  liveOnRoad.innerText=d.summary.onroad||0;
  liveWaiting.innerText=d.summary.waiting||0;
  liveDelayed.innerText=d.summary.delayed||0;
  liveMaint.innerText=d.summary.maintenance_due||0;
  renderLiveOps();
}

function renderLiveOps(){
  const qv=(document.getElementById('liveQ')?.value||'').toLowerCase();
  const state=(document.getElementById('liveState')?.value||'');
  const rows=(liveCache.rows||[]).filter(x=>{
    const hay=[x.scna,x.plate,x.driver_name,x.area_name].join(' ').toLowerCase();
    if(qv && !hay.includes(qv)) return false;
    if(state==='Gecikmiş' && !x.delayed) return false;
    if(state && state!=='Gecikmiş' && x.status!==state) return false;
    return true;
  });

  liveRows.innerHTML='';
  if(!rows.length){
    liveRows.innerHTML='<tr><td colspan="9" class="empty-state">Bu filtrede kayıt yok.</td></tr>';
    return;
  }
  rows.forEach(x=>{
    const badge=x.delayed?'<span class="badge short">GECİKMİŞ</span>':
      (x.status==='Yolda'?'<span class="badge warn">YOLDA</span>':'');
    liveRows.innerHTML+=`<tr class="${x.delayed?'alert-red':''}">
      <td><b>${x.scna}</b></td><td>${x.plate||''}</td><td>${x.driver_name||''}</td>
      <td>${x.area_name||''}</td><td>${x.status||''}</td><td>${fmtDateTime(x.exit_at)}</td>
      <td>${money(x.hours_on_road)}</td><td>${money(x.expected_hours||0)}</td><td>${badge}</td></tr>`;
  });
}

let routeCache=[];

async function loadRouteStandards(){
  const d=await api('/api/route-standards');
  routeCache=d;
  renderRouteStandards();
}

function renderRouteStandards(){
  const qv=(document.getElementById('routeQ')?.value||'').toLowerCase();
  const d=routeCache.filter(x=>!qv || String(x.area_name||'').toLowerCase().includes(qv));
  routeRows.innerHTML='';
  if(!d.length){
    routeRows.innerHTML='<tr><td colspan="6" class="empty-state">Bu filtrede rota yok.</td></tr>';
    return;
  }
  d.forEach(x=>{
    routeRows.innerHTML+=`<tr>
      <td><b>${x.area_name||''}</b></td>
      <td><input id="rh_${x.area_id}" type="number" value="${x.expected_hours||0}" style="width:95px"></td>
      <td><input id="rk_${x.area_id}" type="number" value="${x.expected_km||0}" style="width:105px"></td>
      <td><input id="rl_${x.area_id}" type="number" value="${x.expected_l100||0}" style="width:95px"></td>
      <td><input id="rt_${x.area_id}" type="number" value="${x.tolerance_percent||0}" style="width:90px"></td>
      <td><button class="btn green" onclick="saveRoute(${x.area_id})">Kaydet</button></td></tr>`;
  });
}

async function saveRoute(areaId){
  try{
    await api('/api/route-standards',{
      method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({
        area_id:areaId,
        expected_hours:Number(document.getElementById('rh_'+areaId).value||0),
        expected_km:Number(document.getElementById('rk_'+areaId).value||0),
        expected_l100:Number(document.getElementById('rl_'+areaId).value||0),
        tolerance_percent:Number(document.getElementById('rt_'+areaId).value||0)
      })
    });
    loadRouteStandards();
  }catch(e){alert(e.message)}
}

let maintenanceCache=[];

async function loadMaintenance(){
  const d=await api('/api/maintenance');
  maintenanceCache=d;
  renderMaintenance();
}

function renderMaintenance(){
  const qv=(document.getElementById('maintQ')?.value||'').toLowerCase();
  const state=(document.getElementById('maintState')?.value||'');
  const d=maintenanceCache.filter(x=>{
    if(qv && !String(x.plate||'').toLowerCase().includes(qv)) return false;
    const km=Number(x.km_left||0);
    if(state==='due' && !(x.next_service_km>0 && km<=0)) return false;
    if(state==='soon' && !(x.next_service_km>0 && km>0 && km<=2000)) return false;
    if(state==='ok' && !(x.next_service_km<=0 || km>2000)) return false;
    return true;
  });
  maintRows.innerHTML='';
  if(!d.length){
    maintRows.innerHTML='<tr><td colspan="10" class="empty-state">Bu filtrede bakım kaydı yok.</td></tr>';
    return;
  }
  d.forEach(x=>{
    const km=Number(x.km_left||0);
    const cls=km<=0?'alert-red':(km<=2000?'alert-orange':'');
    maintRows.innerHTML+=`<tr class="${cls}">
      <td><b>${x.plate}</b></td><td>${money(x.current_km)}</td><td>${money(x.last_service_km)}</td>
      <td>${money(x.next_service_km)}</td><td>${x.next_service_km>0?money(x.km_left):''}</td>
      <td>${money(x.oil_change_km)}</td><td>${x.tire_note||''}</td><td>${x.brake_note||''}</td>
      <td>${x.engine_note||''}</td>
      <td><button class="btn secondary" onclick='openMaintenance(${JSON.stringify(x)})'>Düzenle</button></td></tr>`;
  });
}

function openMaintenance(x=null){
  const plates=(L.vehicles||[]).map(v=>`<option value="${v.plate}">${v.plate}</option>`).join('');
  openM(x?'Bakım Kaydı Düzenle':'Yeni Bakım Kaydı',`
    <div class="grid">
      <div class="field"><label>Plaka</label><select id="mPlate"><option value="">Seçin</option>${plates}</select></div>
      <div class="field"><label>Son Bakım KM</label><input id="mLast" type="number" value="${x?x.last_service_km||0:0}"></div>
      <div class="field"><label>Sonraki Bakım KM</label><input id="mNext" type="number" value="${x?x.next_service_km||0:0}"></div>
      <div class="field"><label>Yağ Değişim KM</label><input id="mOil" type="number" value="${x?x.oil_change_km||0:0}"></div>
      <div class="field"><label>Lastik</label><input id="mTire" value="${x?x.tire_note||'':''}"></div>
      <div class="field"><label>Fren</label><input id="mBrake" value="${x?x.brake_note||'':''}"></div>
      <div class="field"><label>Motor</label><input id="mEngine" value="${x?x.engine_note||'':''}"></div>
      <div class="field wide"><label>Not</label><textarea id="mNote">${x?x.note||'':''}</textarea></div>
    </div>
  `,async()=>{
    await api('/api/maintenance',{
      method:'POST',headers:{'Content-Type':'application/json'},
      body:JSON.stringify({
        plate:mPlate.value,last_service_km:Number(mLast.value||0),next_service_km:Number(mNext.value||0),
        oil_change_km:Number(mOil.value||0),tire_note:mTire.value,brake_note:mBrake.value,
        engine_note:mEngine.value,note:mNote.value
      })
    });
    closeM(); loadMaintenance();
  });
  if(x) mPlate.value=x.plate;
}

async function loadAlerts(){
  const d=await api('/api/alerts');
  alertDelayCount.innerText=d.delayed.length;
  alertFuelCount.innerText=d.fuel_alerts.length;

  let html=`<div class="calc">
    <b>Varsayılan kurallar:</b>
    Yolda 48 saati aşan sefer gecikmiş sayılır.
    Yakıt tüketimi aracın kendi geçmiş ortalamasının %20 üstüne çıkarsa uyarı oluşur.
  </div>`;

  html+=`<div class="table" style="margin-top:14px"><table><thead><tr>
    <th>GECİKMİŞ SCNA</th><th>PLAKA</th><th>ŞOFÖR</th><th>BÖLGE</th><th>ÇIKIŞ</th><th>YOLDA SAAT</th>
  </tr></thead><tbody>`;
  if(d.delayed.length){
    d.delayed.forEach(x=>{
      html+=`<tr class="alert-red">
        <td><b>${x.scna}</b></td><td>${x.plate}</td><td>${x.driver_name||''}</td>
        <td>${x.area_name||''}</td><td>${fmtDateTime(x.exit_at)}</td><td>${money(x.hours_on_road)}</td>
      </tr>`;
    });
  }else{
    html+='<tr><td colspan="6">48 saati aşan yolda araç yok.</td></tr>';
  }
  html+='</tbody></table></div>';

  html+=`<div class="table" style="margin-top:14px"><table><thead><tr>
    <th>YAKIT UYARISI SCNA</th><th>PLAKA</th><th>ŞOFÖR</th><th>SEFER LT/100</th><th>ARAÇ ORT.</th><th>FAZLA %</th>
  </tr></thead><tbody>`;
  if(d.fuel_alerts.length){
    d.fuel_alerts.forEach(x=>{
      html+=`<tr class="alert-orange">
        <td><b>${x.scna}</b></td><td>${x.plate}</td><td>${x.driver_name||''}</td>
        <td>${money(x.trip_l100)}</td><td>${money(x.avg_l100)}</td><td>%${money(x.percent_over)}</td>
      </tr>`;
    });
  }else{
    html+='<tr><td colspan="6">Yakıt anormalliği bulunmadı.</td></tr>';
  }
  html+='</tbody></table></div>';

  try{
    const rf=await api('/api/route-fuel-alerts');
    html+=`<div class="table" style="margin-top:14px"><table><thead><tr>
      <th>ROTA YAKIT UYARISI</th><th>PLAKA</th><th>ŞOFÖR</th><th>BÖLGE</th>
      <th>GERÇEK LT/100</th><th>HEDEF</th><th>TOLERANS %</th>
    </tr></thead><tbody>`;
    if(rf.length){
      rf.forEach(x=>{
        html+=`<tr class="alert-orange"><td><b>${x.scna}</b></td><td>${x.plate}</td>
          <td>${x.driver_name||''}</td><td>${x.area_name||''}</td><td>${money(x.actual_l100)}</td>
          <td>${money(x.expected_l100)}</td><td>%${money(x.tolerance_percent)}</td></tr>`;
      });
    }else html+='<tr><td colspan="7">Rota standardına göre yakıt uyarısı yok.</td></tr>';
    html+='</tbody></table></div>';
  }catch(e){}

  alertsContent.innerHTML=html;
}

let auditMetaLoaded=false;
async function loadAudit(){
  if(!hasPerm('audit.view'))return;
  const url='/api/audit?q='+encodeURIComponent(auditQ.value||'')+
    '&username='+encodeURIComponent(auditUser.value||'')+
    '&action='+encodeURIComponent(auditAction.value||'')+
    '&scna='+encodeURIComponent(auditScna.value||'')+
    '&date_from='+encodeURIComponent(auditFrom.value||'')+
    '&date_to='+encodeURIComponent(auditTo.value||'')+
    '&limit=1000';
  const d=await api(url);
  if(!auditMetaLoaded){
    auditUser.innerHTML='<option value="">Tüm Kullanıcılar</option>'+
      (d.users||[]).map(x=>`<option value="${x}">${x}</option>`).join('');
    auditAction.innerHTML='<option value="">Tüm İşlemler</option>'+
      (d.actions||[]).map(x=>`<option value="${x}">${x}</option>`).join('');
    auditMetaLoaded=true;
  }
  auditRows.innerHTML=(d.rows||[]).map(x=>`<tr>
    <td>${fmtDateTime(x.created_at)}</td><td><b>${x.username||'-'}</b></td>
    <td>${x.action||''}</td><td>${x.scna||''}</td><td>${x.detail||''}</td>
    <td>${x.old_value||''}</td><td>${x.new_value||''}</td>
  </tr>`).join('');
  if(typeof v54InitTables==='function')v54InitTables();
}
function clearAuditFilters(){
  auditQ.value='';auditUser.value='';auditAction.value='';auditScna.value='';
  auditFrom.value='';auditTo.value='';loadAudit();
}



function exportAllExcel(){
  window.location.href='/api/export-excel?mode=all';
}

function exportFilteredExcel(){
  const query=(document.getElementById('q')?.value||'');
  const status=(document.getElementById('statusFilter')?.value||'Tümü');
  window.location.href=
    '/api/export-excel?mode=filtered&q='+encodeURIComponent(query)+
    '&status='+encodeURIComponent(status);
}




async function repairFreightFromExcel(input){
  const file=input.files&&input.files[0];
  if(!file) return;

  if(!confirm('Excel AU/FREIGHT değerleri mevcut SCNA kayıtlarının navlun birim fiyatını düzeltecek. Devam edilsin mi?')){
    input.value=''; return;
  }

  const result=document.getElementById('excelResult');
  result.style.display='block';
  result.innerHTML='<b>FREIGHT değerleri Excel’den düzeltiliyor...</b>';

  try{
    const raw=await file.arrayBuffer();
    const r=await fetch('/api/repair-freight-from-excel',{
      method:'POST',
      headers:{'Content-Type':'application/octet-stream'},
      body:raw
    });
    const d=await r.json();
    if(!r.ok) throw new Error(d.detail||'FREIGHT onarım hatası');

    let html='<b style="color:#166534">FREIGHT toplu onarım tamamlandı.</b><br>'+
      'Güncellenen: <b>'+d.updated+'</b>'+
      ' | Zaten doğru: <b>'+d.unchanged+'</b>'+
      ' | FREIGHT bulunamayan: <b>'+d.missing+'</b>'+
      ' | DB’de SCNA bulunamadı: <b>'+d.db_not_found+'</b>';

    if(d.samples&&d.samples.length){
      html+='<br><br><b>Örnek düzeltmeler:</b><br>'+
        d.samples.map(x=>x.scna+' | '+money(x.old_rate)+' → '+money(x.new_rate)+' | AMOUNT '+money(x.amount)).join('<br>');
    }

    result.innerHTML=html;
    try{ await loadAnomalies(); }catch(e){}
  }catch(e){
    result.innerHTML='<b style="color:#991b1b">'+e.message+'</b>';
  }finally{
    input.value='';
  }
}

async function repairRemainFromExcel(input){
  const file=input.files&&input.files[0];
  if(!file) return;

  if(!confirm('Excel REMAIN değerleri aynı SCNA kayıtlarına toplu olarak yazılacak. Devam edilsin mi?')){
    input.value='';
    return;
  }

  const result=document.getElementById('excelResult');
  result.style.display='block';
  result.innerHTML='<b>REMAIN değerleri Excel’den yeniden okunuyor...</b>';

  try{
    const raw=await file.arrayBuffer();
    const r=await fetch('/api/repair-remain-from-excel',{
      method:'POST',
      headers:{'Content-Type':'application/octet-stream'},
      body:raw
    });
    const d=await r.json();
    if(!r.ok) throw new Error(d.detail||'REMAIN onarım hatası');

    let html='<b style="color:#166534">REMAIN toplu onarım tamamlandı.</b><br>'+
      'Sayfa: '+d.sheet+
      ' | Başlık satırı: '+d.header_row+
      ' | REMAIN kolonu: '+d.remain_column_index+
      ' | REMAIN okunan: <b>'+d.parsed+'</b>'+
      ' | DB güncellenen: <b>'+d.updated+'</b>'+
      ' | Excel’de REMAIN boş: <b>'+d.blank+'</b>'+
      ' | DB’de SCNA bulunamadı: <b>'+d.db_not_found+'</b>';

    if(!d.remain_header_found){
      html+='<br><b style="color:#9a3412">Uyarı:</b> REMAIN başlığı bulunamadı; eski format kolon fallback’i kullanıldı.';
    }

    if(d.samples&&d.samples.length){
      html+='<br><br><b>Örnek güncellenenler:</b><br>'+
        d.samples.map(x=>'Satır '+x.row+' | '+x.scna+' → '+money(x.remain)).join('<br>');
    }

    result.innerHTML=html;

    try{
      await loadAnomalies();
    }catch(e){}
  }catch(e){
    result.innerHTML='<b style="color:#991b1b">'+e.message+'</b>';
  }finally{
    input.value='';
  }
}

async function downloadImportErrors(){
  if(!lastImportFailedRecords.length){
    alert('Hatalı kayıt yok.');
    return;
  }

  const r=await fetch('/api/export-import-errors',{
    method:'POST',
    headers:{'Content-Type':'application/json'},
    body:JSON.stringify({failed_records:lastImportFailedRecords})
  });

  if(!r.ok){
    alert('Hata Excel dosyası oluşturulamadı.');
    return;
  }

  const blob=await r.blob();
  const url=URL.createObjectURL(blob);
  const a=document.createElement('a');
  a.href=url;
  a.download='EXCEL_AKTARIM_HATALARI.xlsx';
  document.body.appendChild(a);
  a.click();
  a.remove();
  URL.revokeObjectURL(url);
}



function toggleDirectOneDrive(){
  const p=document.getElementById('directOneDrivePanel');
  p.classList.toggle('open');
  const saved=localStorage.getItem('sama_direct_onedrive_url');
  if(saved) directOneDriveUrl.value=saved;
}

async function testDirectOneDrive(){
  const url=(directOneDriveUrl.value||'').trim();
  if(!url){alert('OneDrive linki gir.');return;}
  localStorage.setItem('sama_direct_onedrive_url',url);
  directOneDriveInfo.style.display='block';
  directOneDriveInfo.innerHTML='<b>OneDrive bağlantısı test ediliyor...</b>';

  try{
    const r=await fetch('/api/public-onedrive/preview',{
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify({share_url:url})
    });
    const d=await r.json();
    if(!r.ok){
      const detail=d.detail||d;
      let msg=typeof detail==='string'?detail:(detail.message||JSON.stringify(detail));
      throw new Error(msg);
    }
    directOneDriveInfo.innerHTML=
      '<b style="color:#166534">BAĞLANTI BAŞARILI.</b><br>'+
      'Excel boyutu: '+money(d.size)+' byte'+
      ' | Sayfa: <b>'+d.sheet+'</b>'+
      ' | Yöntem: '+d.method+
      '<br><span class="small">Kaynak OneDrive dosyasına yazma yapılmadı.</span>';
  }catch(e){
    directOneDriveInfo.innerHTML=
      '<b style="color:#991b1b">Bağlantı kurulamadı:</b> '+e.message+
      '<br><span class="small">Dosya paylaşımını "Bağlantıya sahip herkes görüntüleyebilir" yapıp tekrar dene.</span>';
  }
}

async function syncDirectOneDrive(){
  const url=(directOneDriveUrl.value||'').trim();
  if(!url){alert('OneDrive linki gir.');return;}
  localStorage.setItem('sama_direct_onedrive_url',url);

  excelResult.style.display='block';
  excelResult.innerHTML='<b>OneDrive Web Excel okunuyor...</b>';

  try{
    const fr=await fetch('/api/public-onedrive/file',{
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify({share_url:url})
    });

    if(!fr.ok){
      const d=await fr.json();
      const detail=d.detail||d;
      throw new Error(typeof detail==='string'?detail:(detail.message||JSON.stringify(detail)));
    }

    const blob=await fr.blob();
    const file=new File([blob],'WEB_EXCEL.xlsx',{
      type:'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    });
    lastExcelFile=file;

    const raw=await file.arrayBuffer();
    const r=await fetch(
      '/api/import-excel?filename=WEB_EXCEL.xlsx&update_existing=1&preview_only=1&skip_warnings=0',
      {
        method:'POST',
        headers:{'Content-Type':'application/octet-stream'},
        body:raw
      }
    );
    const data=await r.json();
    if(!r.ok) throw new Error(data.detail||'Web Excel önizleme hatası');

    lastExcelPreview=data;
    lastImportFailedRecords=data.failed_records||[];
    renderExcelPreview(data);

    excelResult.insertAdjacentHTML(
      'afterbegin',
      '<div class="calc" style="margin-bottom:8px"><b>Kaynak: ONEDRIVE PAYLAŞIM LİNKİ</b> '+
      '<span class="small">Excel yalnızca okundu. Kaynak dosya değiştirilmedi.</span></div>'
    );
  }catch(e){
    excelResult.innerHTML='<b style="color:#991b1b">OneDrive okuma hatası:</b> '+e.message;
  }
}

async function repairEntryDatesFromExcel(input){
  const file=input.files&&input.files[0];
  if(!file) return;

  if(!confirm(
    'Bu işlem Excel’den yalnızca SCNA, GİRİŞ TARİHİ ve eksikse GİRİŞ KM okuyacak. '+
    'Programda giriş tarihi boş olan kayıtları onaracak. Diğer alanlara ve kaynak Excel’e dokunmayacak. Devam edilsin mi?'
  )){
    input.value='';
    return;
  }

  const result=document.getElementById('excelResult');
  result.style.display='block';
  result.innerHTML='<b>Eksik giriş tarihleri kontrol ediliyor ve onarılıyor...</b>';

  try{
    const raw=await file.arrayBuffer();
    const r=await fetch('/api/repair-entry-dates-from-excel',{
      method:'POST',
      headers:{'Content-Type':'application/octet-stream'},
      body:raw
    });
    const d=await r.json();

    if(!r.ok) throw new Error(d.detail||'Giriş tarihi onarım hatası');

    let html='<b style="color:#166534">Giriş tarihi onarımı tamamlandı.</b><br>'+
      'Excel’de giriş tarihi olan: <b>'+d.excel_dates+'</b>'+
      ' | DB’de zaten doğru: <b>'+d.already_ok+'</b>'+
      ' | Onarılan: <b>'+d.repaired+'</b>'+
      ' | DB’de SCNA bulunamadı: <b>'+d.db_missing+'</b>'+
      ' | Son kontrolde hâlâ eksik: <b style="color:'+
        (d.verify_missing_count?'#991b1b':'#166534')+'">'+d.verify_missing_count+'</b>';

    if(d.repaired_rows&&d.repaired_rows.length){
      html+='<br><br><b>Onarılan kayıtlar:</b>'+
        '<div style="max-height:300px;overflow:auto;margin-top:6px;border:1px solid #bbf7d0;border-radius:8px;padding:8px">';
      html+=d.repaired_rows.map(x=>
        '<div class="small"><b>'+x.scna+'</b> → '+fmtDateTime(x.entry_date)+
        (x.entry_km?' | Giriş KM '+money(x.entry_km):' | Giriş KM Excel’de boş')+
        (x.was_manual_lock?' | <b>Manuel korumalıydı, yalnızca eksik giriş bilgisi dolduruldu</b>':'')+
        '</div>'
      ).join('');
      html+='</div>';
    }

    if(d.verify_missing&&d.verify_missing.length){
      html+='<br><br><b style="color:#991b1b">Hâlâ giriş tarihi eksik kalan SCNA:</b>'+
        '<div style="max-height:220px;overflow:auto;margin-top:6px;border:1px solid #fecaca;border-radius:8px;padding:8px">';
      html+=d.verify_missing.map(x=>
        '<div class="small">Satır '+x.row+' | <b>'+x.scna+'</b> | Excel '+x.excel_date+'</div>'
      ).join('');
      html+='</div>';
    }

    result.innerHTML=html;

    L=await api('/api/lookups');
    tripPage=1;
    await loadTrips();
    await loadDash();

  }catch(e){
    result.innerHTML='<b style="color:#991b1b">'+e.message+'</b>';
  }finally{
    input.value='';
  }
}

async function importExcel(input){
  const file=input.files && input.files[0];
  if(!file) return;

  lastExcelFile=file;
  const result=document.getElementById('excelResult');
  result.style.display='block';
  result.innerHTML='<b>Excel kontrol ediliyor...</b><br><span class="small">Yeni SCNA eklenecek, mevcut SCNA güncellenecek, manuel korumalı kayıtlar değiştirilmeyecek. Henüz kayıt yapılmıyor.</span>';

  try{
    const raw=await file.arrayBuffer();
    const r=await fetch(
      '/api/import-excel?filename='+encodeURIComponent(file.name)+
      '&update_existing=1'+
      '&preview_only=1',
      {
        method:'POST',
        headers:{'Content-Type':'application/octet-stream'},
        body:raw
      }
    );

    const data=await r.json();
    if(!r.ok) throw new Error(data.detail || 'Excel önizleme hatası');

    lastExcelPreview=data;
    lastImportFailedRecords=data.failed_records||[];
    renderExcelPreview(data);

  }catch(e){
    result.innerHTML='<b style="color:#991b1b">Excel kontrolü başarısız:</b> '+e.message;
  }finally{
    input.value='';
  }
}

function renderExcelPreview(data){
  const result=document.getElementById('excelResult');
  const c=data.preview_counts||{};

  let html=`
    <div style="display:flex;gap:12px;flex-wrap:wrap;align-items:center">
      <b>Excel Kontrol Merkezi</b>
      <span class="badge">Okunan: ${data.processed||0}</span>
      <span class="badge" style="background:#dcfce7;color:#166534">Temiz: ${c.clean||0}</span>
      <span class="badge" style="background:#ffedd5;color:#9a3412">Uyarı: ${c.warning||0}</span>
      <span class="badge" style="background:#fee2e2;color:#991b1b">Hatalı: ${c.error||0}</span>
      <span class="badge">Mükerrer: ${c.duplicate||0}</span>
      <span class="badge">Güncellenecek: ${c.update||0}</span>
      <span class="badge">REMAIN okunan: ${data.remain_updated||0}</span>
      <span class="badge">REMAIN boş/yok: ${data.remain_missing||0}</span>
    </div>

    <div class="filterbar" style="margin-top:10px">
      <select id="previewStatusFilter" onchange="renderPreviewRows()">
        <option value="all">Tüm Önizleme</option>
        <option value="problem">Sadece Sorunlular</option>
        <option value="HATALI">Sadece Hatalı</option>
        <option value="UYARI">Sadece Uyarı</option>
                <option value="MANUEL KORUMALI">Sadece Manuel Korumalı</option>
      </select>
      <span class="small"><b>Uyarılar kayıt aktarımını engellemez.</b> Yalnızca gerçek hatalı satırlar durdurulur.</span>
      <button class="btn green" onclick="confirmExcelImport()">Onayla ve Aktar</button>
      <button class="btn secondary" onclick="cancelExcelPreview()">İptal</button>
      ${(data.failed_records||[]).length?'<button class="btn orange" onclick="downloadImportErrors()">Hatalı SCNA Listesini Excel İndir</button>':''}
    </div>

    <div id="previewRowsArea"></div>
  `;

  result.innerHTML=html;
  renderPreviewRows();
}

function renderPreviewRows(){
  if(!lastExcelPreview) return;
  const mode=document.getElementById('previewStatusFilter')?.value||'all';

  let rows=[];
  if(mode==='problem'){
    rows=lastExcelPreview.problem_records||[];
  }else if(mode==='all'){
    rows=lastExcelPreview.preview_records||[];
  }else{
    const base=(mode==='HATALI'||mode==='UYARI'||mode==='MÜKERRER'||mode==='MANUEL KORUMALI')
      ? (lastExcelPreview.problem_records||[])
      : (lastExcelPreview.preview_records||[]);
    rows=base.filter(x=>x.status===mode);
  }

  let h=`<div class="table"><table><thead><tr>
    <th>SATIR</th><th>SCNA</th><th>PLAKA</th><th>ŞOFÖR</th><th>MÜŞTERİ</th>
    <th>BÖLGE</th><th>MAL</th><th>KG</th><th>PRICE (K)</th><th>FREIGHT (AU)</th><th>AMOUNT</th>
    <th>COLLECTION</th><th>REMAIN</th><th>GİRİŞ TARİHİ</th><th>ÇIKIŞ KM</th><th>GİRİŞ KM</th>
    <th>SONUÇ</th><th>UYARI / HATA</th>
  </tr></thead><tbody>`;

  if(!rows.length){
    h+='<tr><td colspan="18" class="empty-state">Bu filtrede kayıt yok.</td></tr>';
  }else{
    rows.forEach(x=>{
      let cls=x.status==='HATALI'?'alert-red':(x.status==='UYARI'?'alert-orange':'');
      h+=`<tr class="${cls}">
        <td>${x.row||''}</td>
        <td><b>${x.scna||''}</b></td>
        <td>${x.plate||''}</td>
        <td>${x.driver||''}</td>
        <td>${x.customer||''}</td>
        <td>${x.area||''}</td>
        <td>${x.cargo||''}</td>
        <td>${money(x.kg)}</td>
        <td>${money(x.price)}</td>
        <td>${money(x.freight_rate)}</td>
        <td>${money(x.amount)}</td>
        <td>${money(x.collection)}</td>
        <td>${x.excel_remain===null||x.excel_remain===undefined?'':money(x.excel_remain)}</td>
        <td>${fmtDateTime(x.delivery_time)}</td>
        <td>${money(x.exit_km)}</td>
        <td>${money(x.entry_km)}</td>
        <td><b>${x.status||''}</b></td>
        <td>${(x.messages||[]).join('<br>')}</td>
      </tr>`;
    });
  }
  h+='</tbody></table></div>';
  previewRowsArea.innerHTML=h;
}

async function confirmExcelImport(){
  if(!lastExcelFile){
    alert('Excel dosyası bulunamadı. Tekrar seç.');
    return;
  }

  const skipWarnings=0;
  const result=document.getElementById('excelResult');
  result.innerHTML='<b>Onaylanan kayıtlar aktarılıyor...</b>';

  try{
    const raw=await lastExcelFile.arrayBuffer();
    const r=await fetch(
      '/api/import-excel?filename='+encodeURIComponent(lastExcelFile.name)+
      '&update_existing=1'+
      '&preview_only=0&skip_warnings='+skipWarnings,
      {
        method:'POST',
        headers:{'Content-Type':'application/octet-stream'},
        body:raw
      }
    );
    const data=await r.json();
    if(!r.ok) throw new Error(data.detail || 'Excel aktarım hatası');

    lastImportFailedRecords=data.failed_records||[];

    let html=
      '<b style="color:#166534">Excel aktarımı tamamlandı.</b><br>'+
      'Sayfa: '+data.sheet+
      ' | Okunan: <b>'+data.processed+'</b>'+
      ' | Eklenen: <b>'+data.added+'</b>'+
      ' | Güncellenen: <b>'+(data.updated||0)+'</b>'+
      ' | Eski mükerrer atlama: <b>'+data.skipped+'</b>'+
      ' | Manuel korunan: <b>'+(data.manual_protected_skipped||0)+'</b>'+
      ' | Uyarılı atlanan: <b>'+(data.warning_skipped||0)+'</b>'+
      ' | REMAIN okunan: <b>'+(data.remain_updated||0)+'</b>'+
      ' | REMAIN boş/yok: <b>'+(data.remain_missing||0)+'</b>'+
      ' | Excel dönüş tarihi: <b>'+(data.delivery_date_source_count||0)+'</b>'+
      ' | Başarıyla yazılan dönüş tarihi: <b>'+(data.delivery_date_written_count||0)+'</b>'+
      ' | Yazılamayan dönüş tarihi: <b style="color:'+
        ((data.delivery_date_not_written_count||0)>0?'#991b1b':'#166534')+'">'+
        (data.delivery_date_not_written_count||0)+'</b>'+
      ' | Geçersiz SCNA atlanan: <b>'+(data.invalid_scna_skipped||0)+'</b>'+
      ' | Excel içi mükerrer: <b>'+(data.excel_duplicate_count||0)+'</b>'+
      ' | Kaydedilemeyen: <b style="color:#991b1b">'+data.error_count+'</b>';

    if(data.failed_records && data.failed_records.length){
      html+='<br><br><b style="color:#991b1b">Kaydedilemeyen SCNA kayıtları:</b>';
      html+='<div style="max-height:260px;overflow:auto;margin-top:6px;border:1px solid #fecaca;border-radius:8px;padding:8px;background:#fff7f7">';
      html+=data.failed_records.map(x=>
        '<div class="small"><b>Satır '+x.row+'</b> | <b>SCNA: '+(x.scna||'BOŞ')+'</b> | '+x.reason+'</div>'
      ).join('');
      html+='</div><div style="margin-top:8px"><button class="btn orange" onclick="downloadImportErrors()">Hatalı SCNA Listesini Excel İndir</button></div>';
    }

    // Aynı Excel ile giriş tarihlerini ikinci kez çapraz kontrol et.
    try{
      const verifyRaw=await lastExcelFile.arrayBuffer();
      const vr=await fetch('/api/check-entry-dates-from-excel',{
        method:'POST',
        headers:{'Content-Type':'application/octet-stream'},
        body:verifyRaw
      });
      const vd=await vr.json();
      if(vr.ok){
        html+='<br><br><b>Giriş Tarihi Son Kontrol:</b> '+
          'Excel tarihli: <b>'+vd.excel_with_entry_date+'</b>'+
          ' | DB tarihli: <b>'+vd.db_with_entry_date+'</b>'+
          ' | DB’de tarihi eksik: <b style="color:'+(vd.missing_date_count?'#991b1b':'#166534')+'">'+vd.missing_date_count+'</b>';

        if(vd.missing_dates&&vd.missing_dates.length){
          html+='<div style="max-height:220px;overflow:auto;margin-top:6px;border:1px solid #fecaca;border-radius:8px;padding:8px">';
          html+=vd.missing_dates.map(x=>
            '<div class="small">Satır '+x.row+' | <b>'+x.scna+'</b> | Excel: '+x.excel_date+
            ' | DB: BOŞ | '+(x.manual_lock?'MANUEL KORUMALI':'')+'</div>'
          ).join('');
          html+='</div>';
        }
        result.innerHTML=html;
      }
    }catch(e){}

    result.innerHTML=html;
    L=await api('/api/lookups');
    tripPage=1;
    await loadTrips();
    await loadDash();

  }catch(e){
    result.innerHTML='<b style="color:#991b1b">Excel aktarımı başarısız:</b> '+e.message;
  }
}

function cancelExcelPreview(){
  lastExcelFile=null;
  lastExcelPreview=null;
  lastImportFailedRecords=[];
  const result=document.getElementById('excelResult');
  result.innerHTML='<span class="small">Excel aktarımı iptal edildi.</span>';
}

let newTripLastPlateLookup='';
async function newTripPlateDriverSync(){
  const p=document.getElementById('nPlate');
  const d=document.getElementById('nDriver');
  if(!p||!d)return;

  const plate=String(p.value||'').trim().toUpperCase();
  if(plate.length<5)return;
  if(plate===newTripLastPlateLookup)return;

  // Önce tarayıcıdaki mevcut listeden hızlıca dene.
  const local=(L.vehicles||[]).find(x=>
    String(x.plate||'').trim().toUpperCase()===plate
  );
  if(local && local.driver_name && !String(d.value||'').trim()){
    d.value=local.driver_name;
  }

  // Asıl kaynak: veritabanındaki bu plakaya ait SON SEVKİYAT.
  try{
    const x=await api('/api/autocomplete/plate-last-driver?plate='+encodeURIComponent(plate));
    if(String(p.value||'').trim().toUpperCase()!==plate)return;

    if(x && x.found && x.driver_name){
      // Kullanıcı şoförü elle değiştirmediyse otomatik doldur.
      // Plaka değiştirilince yeni aracın son şoförü tekrar getirilebilir.
      const current=String(d.value||'').trim();
      const localName=local&&local.driver_name?String(local.driver_name).trim():'';
      if(!current || current===localName || newTripLastPlateLookup!==plate){
        d.value=x.driver_name;
        d.dispatchEvent(new Event('change',{bubbles:true}));
      }
      d.title='Son sefer: '+(x.scna||'')+' | '+x.driver_name;
    }
    newTripLastPlateLookup=plate;
  }catch(e){
    console.warn('Son şoför bilgisi alınamadı:',e);
  }
}

async function newTrip(){
  let ao=L.areas.map(a=>`<option value="${a.id}">${a.name}</option>`).join('');
  let co=L.cargo_categories.map(c=>`<option value="${c.id}">${c.name}</option>`).join('');
  let cu=L.customers.map(c=>`<option value="${c.id}">${c.name}</option>`).join('');

  openM('Yeni Sevkiyat',`
  <div class="grid">
    <div class="field"><label>SCNA</label><input id="nScna"></div>
    <div class="field"><label>Plaka</label><input id="nPlate" placeholder="En az 3 karakter yazın" oninput="newTripPlateDriverSync()"></div>
    <div class="field"><label>Şoför</label><input id="nDriver" placeholder="En az 3 karakter yazın" onfocus="samaAttachAutocomplete(this,'driver')"></div>
    <div class="field"><label>Bölge</label><select id="nArea"><option value="">Seçin</option>${ao}</select></div>

    <div class="field"><label>Taşınan Mal Cinsi</label><select id="nCargo"><option value="">Seçin</option>${co}</select></div>
    <div class="field"><label>Yük Tipi</label><select id="nCargoType"><option value="BULK">BULK</option><option value="BAG">BAG</option></select></div>
    <div class="field"><label>Müşteri</label><input id="nCustomer" placeholder="En az 3 karakter yazın" onfocus="samaAttachAutocomplete(this,'customer')"></div>
    <div class="field"><label>Tarih</label><input id="nDate" type="date"></div>
    <div class="field"><label>Net KG</label><input id="nKg" type="number" oninput="calcFreight()"></div>

    <div class="field"><label>Navlun Birimi</label><select id="nBasis" onchange="calcFreight()"><option value="TON">TON</option><option value="KG">KG</option></select></div>
    <div class="field"><label>Navlun Birim Fiyat</label><input id="nRate" type="text" inputmode="decimal" oninput="calcFreight()"></div>

    <div class="field"><label>Çıkış KM</label><input id="nExitKm" type="number"></div>
    <div class="field"><label>Depodaki Mevcut Mazot (LT)</label><input id="nTankStart" type="number"></div>
    <div class="field"><label>Prim</label><input id="nPremium" type="text" inputmode="decimal"></div>
    <div class="field"><label>Dock Fee</label><input id="nDockFee" type="text" inputmode="decimal"></div>
    <div class="field"><label>Port Fee</label><input id="nPortFee" type="text" inputmode="decimal"></div>
    <div class="field"><label>SONAR</label><input id="nSonar" type="text" inputmode="decimal"></div>
  </div>

  <div class="calc">
    Toplam Navlun: <b id="nTotal">0</b>
  </div>
  `, async()=>{
    let v=L.vehicles.find(x=>String(x.plate||'').toUpperCase()===String(nPlate.value||'').trim().toUpperCase());
    let driverName=String(nDriver.value||'').trim();
    let customerName=String(nCustomer.value||'').trim();

    const scnaVal=String(nScna.value||'').trim().toUpperCase();
    const compactScna=scnaVal.replace(/[^A-Z0-9]/g,'');
    if(!scnaVal || (/^\d+$/.test(scnaVal) ? scnaVal.length<5 : compactScna.length<5)){
      alert('SCNA eksik görünüyor. En az 5 hane/karakter gir.');
      nScna.focus();
      return;
    }

    try{
      let selectedDriverId=null;
      let selectedCustomerId=null;

      if(driverName){
        const dr=await api('/api/resolve/driver',{
          method:'POST',headers:{'Content-Type':'application/json'},
          body:JSON.stringify({name:driverName})
        });
        selectedDriverId=dr.id||null;
      }else if(v){
        selectedDriverId=v.current_driver_id||v.driver_id||null;
      }

      if(customerName){
        const cu=await api('/api/resolve/customer',{
          method:'POST',headers:{'Content-Type':'application/json'},
          body:JSON.stringify({name:customerName})
        });
        selectedCustomerId=cu.id||null;
      }

      await api('/api/trips',{
        method:'POST',
        headers:{'Content-Type':'application/json'},
        body:JSON.stringify({
          scna:nScna.value.trim(),
          plate:nPlate.value,
          driver_id:selectedDriverId,
          area_id:nArea.value?Number(nArea.value):null,
          cargo_category_id:nCargo.value?Number(nCargo.value):null,
          cargo_type:nCargoType.value,
          customer_id:selectedCustomerId,
          trip_date:nDate.value,
          net_kg:parseQty(nKg.value),
          freight_rate:parseMoney(nRate.value),
          freight_basis:nBasis.value,
          exit_km:parseQty(nExitKm.value),
          tank_start_liters:parseQty(nTankStart.value),
          exit_premium:parseMoney(nPremium.value),
          dock_fee:parseMoney(nDockFee.value),
          port_fee:parseMoney(nPortFee.value),
          sonar:parseMoney(nSonar.value)
        })
      });

      closeM();
      loadTrips();

    }catch(e){
      alert(e.message);
    }
  });
}

function calcFreight(){
  let kg=parseQty(nKg.value);
  let rate=parseMoney(nRate.value);
  nTotal.innerText=money(nBasis.value==='KG'?kg*rate:(kg/1000)*rate);
}

async function openExit(scna){
  cx=await api('/api/trips/'+scna);

  openM('Çıkış / Yakıt - '+scna,`
  <div class="grid">
    <div class="field"><label>Plaka</label><input value="${cx.plate}" disabled></div>
    <div class="field"><label>Toplam Navlun</label><input value="${money(cx.freight_total)}" disabled></div>
    <div class="field"><label>Çıkış Tarihi / Saati</label><input value="${fmtDateTime(cx.exit_at)||'Çıkış kaydedilince otomatik oluşur'}" disabled></div>
    <div class="field"><label>Çıkış KM</label><input value="${money(cx.exit_km)}" disabled></div>

    <div class="field"><label>Şoföre Verilen Avans</label><input id="xCash" type="text" inputmode="decimal" value="${cx.exit_cash||0}" oninput="calcExitFuel()"></div>
    <div class="field"><label>Çıkışta Depodaki Mevcut Mazot (LT)</label><input value="${money(cx.tank_start_liters)}" disabled></div>

    <div class="field"><label>Resmi Mazot (LT)</label><input id="xOffLt" type="number" value="${cx.exit_official_fuel_liters||0}" oninput="calcExitFuel()"></div>
    <div class="field"><label>Resmi Mazot Toplam Tutar</label><input id="xOffTotal" type="text" inputmode="decimal" value="${cx.exit_official_fuel_total||0}" oninput="calcExitFuel()"></div>

    <div class="field"><label>Ticari Mazot (LT)</label><input id="xComLt" type="number" value="${cx.exit_commercial_fuel_liters||0}" oninput="calcExitFuel()"></div>
    <div class="field"><label>Ticari Mazot Toplam Tutar</label><input id="xComTotal" type="text" inputmode="decimal" value="${cx.exit_commercial_fuel_total||0}" oninput="calcExitFuel()"></div>

    <div class="field"><label>Bağdat Mazotu (LT)</label><input id="xBagLt" type="number" value="${cx.exit_baghdad_fuel_liters||0}" oninput="calcExitFuel()"></div>
    <div class="field"><label>Bağdat Mazotu Toplam Tutar</label><input id="xBagTotal" type="text" inputmode="decimal" value="${cx.exit_baghdad_fuel_total||0}" oninput="calcExitFuel()"></div>

    <div class="field"><label>Harcırah</label><input id="xAllow" type="text" inputmode="decimal" value="${cx.exit_allowance||0}" oninput="calcExitFuel()"></div>
    <div class="field"><label>Prim</label><input value="${money(cx.exit_premium)}" disabled></div>
    <div class="field"><label>Dock Fee</label><input value="${money(cx.dock_fee)}" disabled></div>
    <div class="field"><label>Port Fee</label><input value="${money(cx.port_fee)}" disabled></div>
    <div class="field"><label>SONAR</label><input value="${money(cx.sonar)}" disabled></div>
    <div class="field"><label>OTHER / Ekstra Para</label><input id="xOther" type="text" inputmode="decimal" value="${cx.exit_other||0}" oninput="calcExitFuel()"></div>
    <div class="field wide"><label>OTHER Açıklaması</label><input id="xOtherNote" value="${cx.exit_other_note||''}" placeholder="Neden verildiğini yazın"></div>

    <div class="field wide"><label>Çıkış Notu</label><textarea id="xNote">${cx.exit_note||''}</textarea></div>
  </div>

  <div class="settle">
    <div>Çıkışta Depo<strong id="xTankStartShow">0 LT</strong></div>
    <div>Çıkışta Alınan Toplam<strong id="xBoughtLt">0 LT</strong></div>
    <div>Yola Çıkarken Toplam Yakıt<strong id="xAvailableStart">0 LT</strong></div>
    <div>Çıkış Mazot Maliyeti<strong id="xFuelCost">0</strong></div>
    <div style="border:2px solid #2563eb;background:#eff6ff">Şoföre Verilmesi Gereken<strong id="xRequiredCash">0 IQD</strong><span class="small">Mazot + Harcırah + Prim + Dock + Port + SONAR + OTHER</span></div>
    <div id="xCashDiffBox" style="border:2px solid #cbd5e1"><span id="xCashDiffLabel">Girilen Para Kontrolü</span><strong id="xCashDiff">0 IQD</strong></div>
  </div>
  # SAMA_EXIT_CASH_SUMMARY_V4
  `, async()=>{
    try{
      if(parseMoney(xOther.value)>0 && !(document.getElementById('xOtherNote')?.value||'').trim()){alert('OTHER tutarı için açıklama zorunlu.');return;}
      const requiredCash=parseMoney(xOffTotal.value)+parseMoney(xComTotal.value)+parseMoney(xBagTotal.value)+parseMoney(xAllow.value)+Number(cx.exit_premium||0)+Number(cx.dock_fee||0)+Number(cx.port_fee||0)+Number(cx.sonar||0)+parseMoney(xOther.value);
      const enteredCash=parseMoney(xCash.value);
      if(enteredCash>requiredCash+0.01){
        alert('HATA: Fazla para girdiniz.\n\nŞoföre verilmesi gereken: '+money(requiredCash)+' IQD\nGirilen: '+money(enteredCash)+' IQD\nFazla: '+money(enteredCash-requiredCash)+' IQD\n\nKayıt yapılmadı.');
        xCash.focus(); return;
      }
      await api('/api/trips/'+scna+'/exit',{
        method:'PATCH',
        headers:{'Content-Type':'application/json'},
        body:JSON.stringify({
          exit_cash:parseMoney(xCash.value),

          exit_official_fuel_liters:parseQty(xOffLt.value),
          exit_official_fuel_total:parseMoney(xOffTotal.value),

          exit_commercial_fuel_liters:parseQty(xComLt.value),
          exit_commercial_fuel_total:parseMoney(xComTotal.value),

          exit_baghdad_fuel_liters:parseQty(xBagLt.value),
          exit_baghdad_fuel_total:parseMoney(xBagTotal.value),

          exit_allowance:parseMoney(xAllow.value),
          exit_other:parseMoney(xOther.value),
          exit_other_note:(document.getElementById('xOtherNote')?.value||'').trim(),
          exit_note:xNote.value
        })
      });

      closeM();
      loadExit();

    }catch(e){
      alert(e.message);
    }
  });

  calcExitFuel();
}

function calcExitFuel(){
  let tank=Number(cx.tank_start_liters||0);

  let offLt=parseQty(xOffLt.value);
  let comLt=parseQty(xComLt.value);
  let bagLt=parseQty(xBagLt.value);

  let bought=offLt+comLt+bagLt;

  let fuelCost=
    parseMoney(xOffTotal.value)+
    parseMoney(xComTotal.value)+
    parseMoney(xBagTotal.value);

  xTankStartShow.innerText=money(tank)+' LT';
  xBoughtLt.innerText=money(bought)+' LT';
  xAvailableStart.innerText=money(tank+bought)+' LT';
  xFuelCost.innerText=money(fuelCost);

  // Şoföre çıkışta fiilen verilmesi gereken toplam nakit.
  // Resmi/ticari/Bağdat mazot tutarları bu iş akışında şoföre verilen nakdin parçasıdır.
  const requiredCash=
    fuelCost+
    Number(cx.exit_premium||0)+
    Number(cx.dock_fee||0)+
    Number(cx.port_fee||0)+
    Number(cx.sonar||0)+
    parseMoney(xAllow.value)+
    parseMoney(xOther.value);
  const enteredCash=parseMoney(xCash.value);
  const cashDiff=enteredCash-requiredCash;
  if(document.getElementById('xRequiredCash')) xRequiredCash.innerText=money(requiredCash)+' IQD';
  if(document.getElementById('xCashDiff')) xCashDiff.innerText=money(Math.abs(cashDiff))+' IQD';
  if(document.getElementById('xCashDiffLabel')){
    if(cashDiff>0.01){xCashDiffLabel.innerText='FAZLA PARA';xCashDiffBox.style.borderColor='#ef4444';xCashDiffBox.style.background='#fef2f2';xCashDiff.style.color='#b91c1c';}
    else if(cashDiff < -0.01){xCashDiffLabel.innerText='EKSİK PARA';xCashDiffBox.style.borderColor='#f59e0b';xCashDiffBox.style.background='#fffbeb';xCashDiff.style.color='#b45309';}
    else{xCashDiffLabel.innerText='PARA TAMAM';xCashDiffBox.style.borderColor='#22c55e';xCashDiffBox.style.background='#f0fdf4';xCashDiff.style.color='#166534';}
  }
}


async function openEntry(scna){
  cx=await api('/api/trips/'+scna);
  let fuels=await api('/api/trips/'+scna+'/fuel');

  openM('Giriş / Yakıt Tüketim / Hesap - '+scna,`
  <div class="grid">
    <div class="field"><label>Plaka</label><input value="${cx.plate}" disabled></div>
    <div class="field"><label>Toplam Navlun</label><input value="${money(cx.freight_total)}" disabled></div>
    <div class="field"><label>Çıkış Tarihi / Saati</label><input value="${fmtDateTime(cx.exit_at)}" disabled></div>
    <div class="field"><label>Giriş Tarihi / Saati</label><input value="${fmtDateTime(cx.entry_at)||'Giriş kaydedilince otomatik oluşur'}" disabled></div>
    <div class="field"><label>Çıkış KM</label><input id="gExit" value="${cx.exit_km}" disabled></div>

    <div class="field"><label>Giriş KM</label><input id="gKm" type="number" value="${cx.entry_km||0}" oninput="calcEntry()"></div>
    <div class="field"><label>Dönüşte Depoda Kalan Mazot (LT)</label><input id="gTankEnd" type="number" value="${cx.tank_end_liters||0}" oninput="calcEntry()"></div>

    <div class="field"><label>Ekstra Gider 1</label><input id="gExtra1" type="text" inputmode="decimal" value="${cx.entry_extra_expense_1||0}" oninput="calcEntry()"></div>
    <div class="field"><label>Ekstra Gider 2</label><input id="gExtra2" type="text" inputmode="decimal" value="${cx.entry_extra_expense_2||0}" oninput="calcEntry()"></div>
    <div class="field"><label>Ekstra Gider 3</label><input id="gExtra3" type="text" inputmode="decimal" value="${cx.entry_extra_expense_3||0}" oninput="calcEntry()"></div>
    <div class="field"><label>Müşteriden Alınan Para</label><input id="gCollect" type="text" inputmode="decimal" value="${cx.entry_collection||0}" oninput="calcEntry()"></div>
    <div class="field"><label>Şoförün Teslim Ettiği Para</label><input id="gHand" type="text" inputmode="decimal" value="${cx.entry_cash_handed||0}" oninput="calcEntry()"></div>

    <div class="field wide"><label>Giriş Notu</label><textarea id="gNote">${cx.entry_note||''}</textarea></div>
  </div>

  <h4>Yolda Alınan Ek Mazot</h4>

  <div class="grid">
    <div class="field"><label>Alınan Mazot (LT)</label><input id="fLit" type="number" oninput="calcRoadFuelPreview()"></div>
    <div class="field"><label>Ödenen Toplam Tutar</label><input id="fTotal" type="text" inputmode="decimal" oninput="calcRoadFuelPreview()"></div>
    <div class="field"><label>Programın Hesapladığı Birim Fiyat</label><input id="fUnitPreview" disabled></div>
    <div class="field wide"><label>Açıklama</label><input id="fNote"></div>
  </div>

  <div style="margin-top:10px">
    <button class="btn green" onclick="addFuel('${scna}')">+ Mazot Alımı Ekle</button>
  </div>

  <div id="fuelList" style="margin-top:12px">${fuelHtml(fuels)}</div>

  <div class="settle">
    <div>Gerçek KM<strong id="gActualKm">0</strong></div>
    <div>Toplam Kullanılabilir Yakıt<strong id="gFuelAvailable">0 LT</strong></div>
    <div>Tüketilen Yakıt<strong id="gConsumed">0 LT</strong></div>
    <div>Dönüşte Kalan<strong id="gTankEndShow">0 LT</strong></div>
  </div>

  <div class="settle">
    <div>100 KM Tüketim<strong id="gL100">0 LT</strong></div>
    <div>KM / LT<strong id="gKmPerLt">0</strong></div>
    <div>Yolda Alınan Mazot<strong id="gRoadFuel">0 LT</strong></div>
  </div>

  <div class="settle">
    <div>Toplam Gider<strong id="gTotalExpense">0</strong></div>
    <div>Sefer Kârı<strong id="gProfit">0</strong></div>
    <div>Müşteriden Kalan Alacak<strong id="gDue">0</strong></div>
    <div>Dönüş Harcamaları<strong id="gReturnSpend">0</strong></div>
  </div>

  <div id="driverExpectedBox" class="diff-box diff-green">
    ŞOFÖRÜN VERMESİ GEREKEN PARA
    <strong id="gExpected">0</strong>
    <span class="small">Toplam Navlun − (Yolda Mazot + Ekstra Giderler)</span>
    <div class="small" id="gFormulaTrace" style="margin-top:8px"></div>
  </div>

  <div class="settle">
    <div>Şoförün Verdiği Para<strong id="gHandShow">0</strong></div>
    <div>Şoför Para Farkı<strong id="gDiffSmall">0</strong></div>
  </div>

  <div id="driverDiffBox" class="diff-box diff-neutral">
    ŞOFÖR PARA FARKI
    <strong id="gDiff">0</strong>
  </div>

  <div class="calc">
    Şoför Hesabı: <span id="gStatus"></span>
  </div>
  `, async()=>{
    try{
      await api('/api/trips/'+scna+'/entry',{
        method:'PATCH',
        headers:{'Content-Type':'application/json'},
        body:JSON.stringify({
          entry_km:parseQty(gKm.value),
          tank_end_liters:parseQty(gTankEnd.value),

          entry_extra_expense_1:parseMoney(gExtra1.value),
          entry_extra_expense_2:parseMoney(gExtra2.value),
          entry_extra_expense_3:parseMoney(gExtra3.value),

          entry_collection:parseMoney(gCollect.value),
          entry_cash_handed:parseMoney(gHand.value),
          entry_note:gNote.value
        })
      });

      closeM();
      loadEntry();
      loadDash();

    }catch(e){
      alert(e.message);
    }
  });

  calcEntry();
}

function calcRoadFuelPreview(){
  let lt=parseQty(fLit.value);
  let total=parseMoney(fTotal.value);
  fUnitPreview.value=lt>0?money(total/lt):'0';
}

function fuelHtml(f){
  if(!f.length){
    return '<span class="small">Henüz yolda alınmış mazot kaydı yok.</span>';
  }

  return `
  <table class="mini">
    <thead>
      <tr><th>LT</th><th>TOPLAM</th><th>BİRİM FİYAT</th><th>NOT</th><th></th></tr>
    </thead>
    <tbody>
      ${f.map(x=>`
        <tr>
          <td>${money(x.liters)}</td>
          <td>${money(x.total)}</td>
          <td>${money(Number(x.total||0)/Number(x.liters||1))}</td>
          <td>${x.note||''}</td>
          <td>
            <button class="btn danger" onclick="delFuel(${x.id},'${x.scna}')">Sil</button>
          </td>
        </tr>
      `).join('')}
    </tbody>
  </table>`;
}

async function addFuel(scna){
  try{
    await api('/api/trips/'+scna+'/fuel',{
      method:'POST',
      headers:{'Content-Type':'application/json'},
      body:JSON.stringify({
        liters:parseQty(fLit.value),
        total:parseMoney(fTotal.value),
        note:fNote.value
      })
    });

    let f=await api('/api/trips/'+scna+'/fuel');
    fuelList.innerHTML=fuelHtml(f);

    fLit.value='';
    fTotal.value='';
    fUnitPreview.value='';
    fNote.value='';

    cx=await api('/api/trips/'+scna);
    calcEntry();

  }catch(e){
    alert(e.message);
  }
}

async function delFuel(id,scna){
  await api('/api/fuel/'+id,{method:'DELETE'});

  let f=await api('/api/trips/'+scna+'/fuel');
  fuelList.innerHTML=fuelHtml(f);

  cx=await api('/api/trips/'+scna);
  calcEntry();
}

function calcEntry(){
  let actualKm=Math.max(0,parseQty(gKm.value)-Number(gExit.value||0));

  let roadFuelLt=Number(cx.road_fuel_liters||0);
  let roadFuelTotal=Number(cx.road_fuel_total||0);

  let totalAvailable=
    Number(cx.tank_start_liters||0)+
    Number(cx.exit_official_fuel_liters||0)+
    Number(cx.exit_commercial_fuel_liters||0)+
    Number(cx.exit_baghdad_fuel_liters||0)+
    roadFuelLt;

  let endTank=parseQty(gTankEnd.value);
  let consumed=Math.max(0,totalAvailable-endTank);

  let l100=actualKm>0?(consumed/actualKm)*100:0;
  let kmPerLt=consumed>0?actualKm/consumed:0;

  let extra=
    parseMoney(gExtra1.value)+
    parseMoney(gExtra2.value)+
    parseMoney(gExtra3.value);

  let exitFuelCost=
    Number(cx.exit_official_fuel_total||0)+
    Number(cx.exit_commercial_fuel_total||0)+
    Number(cx.exit_baghdad_fuel_total||0);

  let totalFuelCost=exitFuelCost+roadFuelTotal;

  let totalExpense=
    totalFuelCost+
    Number(cx.exit_allowance||0)+
    Number(cx.exit_premium||0)+
    Number(cx.dock_fee||0)+
    Number(cx.exit_other||0)+
    extra;

  // Üç ayrı para:
  // 1) Toplam Navlun
  // 2) Müşteriden fiilen alınan para
  // 3) Şoförün şirkete vermesi gereken para
  let customerMoney=parseMoney(gCollect.value);
  let hand=parseMoney(gHand.value);

  // Dönüş harcamaları:
  // Yolda alınan mazot + Ekstra Gider 1/2/3.
  let returnSpend=roadFuelTotal+extra;

  // Şoförün vermesi gereken:
  // TOPLAM NAVLUN - DÖNÜŞ HARCAMALARI.
  // Müşteriden tahsil edilen para bu hesaba dahil değildir.
  let freightTotal=Number(cx.freight_total||0);
  let expected=freightTotal-returnSpend;
  let diff=hand-expected;

  // Eski raporlarla uyumluluk için müşteriden kalan hesaplanmaya devam eder,
  // ancak şoför hesabında kullanılmaz.
  let due=freightTotal-customerMoney;

  // Sefer kârı tüm gerçek giderleri içerir.
  let profit=Number(cx.freight_total||0)-totalExpense;

  gActualKm.innerText=money(actualKm);
  gFuelAvailable.innerText=money(totalAvailable)+' LT';
  gConsumed.innerText=money(consumed)+' LT';
  gTankEndShow.innerText=money(endTank)+' LT';

  gL100.innerText=money(l100)+' LT';
  gKmPerLt.innerText=money(kmPerLt);
  gRoadFuel.innerText=money(roadFuelLt)+' LT';

  gTotalExpense.innerText=money(totalExpense);
  gProfit.innerText=money(profit);
  gDue.innerText=money(due);
  gReturnSpend.innerText=money(returnSpend);
  gExpected.innerText=money(expected);
  gFormulaTrace.innerText=money(freightTotal)+' − '+money(returnSpend)+' = '+money(expected);
  gHandShow.innerText=money(hand);
  gDiffSmall.innerText=money(diff);
  gDiff.innerText=money(diff);

  driverExpectedBox.classList.remove('diff-red','diff-green','diff-neutral');
  if(expected>0){
    driverExpectedBox.classList.add('diff-green');
  }else if(expected<0){
    driverExpectedBox.classList.add('diff-red');
  }else{
    driverExpectedBox.classList.add('diff-neutral');
  }


  driverDiffBox.classList.remove('diff-red','diff-green','diff-neutral');

  if(diff<0){
    driverDiffBox.classList.add('diff-red');
    gStatus.innerHTML='<span class="badge short">EKSİK PARA</span>';
  }else if(diff>0){
    driverDiffBox.classList.add('diff-green');
    gStatus.innerHTML='<span class="badge over">FAZLA PARA</span>';
  }else{
    driverDiffBox.classList.add('diff-neutral');
    gStatus.innerHTML='<span class="badge ok">HESAP TAMAM</span>';
  }
}

function filterAuditRows(){
  const qv=(document.getElementById('auditQ')?.value||'').toLowerCase();
  document.querySelectorAll('#auditRows tr').forEach(tr=>{
    const txt=(tr.textContent||'').toLowerCase();
    tr.style.display=!qv || txt.includes(qv)?'':'none';
  });
}


authBoot();
setTimeout(installUniversalListTools,1200);

function parseRowDate(text){
  const m=String(text||'').match(/(\d{2})[.\/]([0-1]?\d)[.\/](\d{4})/);
  if(m)return `${m[3]}-${String(m[2]).padStart(2,'0')}-${String(m[1]).padStart(2,'0')}`;
  const y=String(text||'').match(/(\d{4})-(\d{2})-(\d{2})/); return y?`${y[1]}-${y[2]}-${y[3]}`:'';
}
function applyPanelDateFilter(panelId){
  const panel=document.getElementById(panelId); if(!panel)return;
  const from=panel.querySelector('input[id$="From"][type="date"]')?.value||'';
  const to=panel.querySelector('input[id$="To"][type="date"]')?.value||'';
  panel.querySelectorAll('tbody tr').forEach(tr=>{
    const dates=[...tr.querySelectorAll('td')].map(td=>parseRowDate(td.innerText)).filter(Boolean);
    const ok=!dates.length||dates.some(d=>(!from||d>=from)&&(!to||d<=to)); tr.style.display=ok?'':'none';
  });
  updatePanelCounters(panel);
}
function updatePanelCounters(panel){
  if(!panel)return; const rows=[...panel.querySelectorAll('tbody tr')].filter(r=>r.style.display!=='none'&&!r.querySelector('.empty-state'));
  const count=rows.length; const id=panel.id;
  const known=document.getElementById(id+'Count'); if(known)known.innerText=count+' İŞLEM';
  let badge=panel.querySelector('.auto-list-count');
  const q=panel.querySelector('.filterbar input[type="text"],.filterbar input:not([type])');
  if(q&&!badge){badge=document.createElement('span');badge.className='auto-list-count';badge.style.cssText='font-weight:800;white-space:nowrap;padding:0 8px';q.insertAdjacentElement('afterend',badge);}
  if(badge)badge.innerText=count+' KAYIT';
}
function installUniversalListTools(){
  document.querySelectorAll('.panel').forEach(panel=>{
    const bar=panel.querySelector('.filterbar'); const table=panel.querySelector('table'); if(!bar||!table)return;
    if(!bar.querySelector('input[type="date"]') && !['daily','audit'].includes(panel.id)){
      const f=document.createElement('input');f.type='date';f.id='autoFrom_'+panel.id;f.title='Başlangıç Tarihi';f.onchange=()=>applyPanelDateFilter(panel.id);
      const t=document.createElement('input');t.type='date';t.id='autoTo_'+panel.id;t.title='Bitiş Tarihi';t.onchange=()=>applyPanelDateFilter(panel.id);
      bar.appendChild(f);bar.appendChild(t);
    }
    updatePanelCounters(panel);
    const tb=panel.querySelector('tbody');if(tb)new MutationObserver(()=>applyPanelDateFilter(panel.id)).observe(tb,{childList:true,subtree:true});
  });
}
async function softDeleteTrip(scna){
  if(!AUTH_USER||AUTH_USER.role!=='ADMIN')return;
  if(!confirm(scna+' sevkiyatı silinsin mi? 7 gün içinde geri alınabilir.'))return;
  await api('/api/trips/'+encodeURIComponent(scna),{method:'DELETE'}); await loadTrips();
}

let cashSummary=null;
function cashToday(){return new Date().toISOString().slice(0,10)}
async function loadCashControl(){
 if(!cashDay.value)cashDay.value=cashToday(); cashSummary=await api('/api/cash-control?business_date='+encodeURIComponent(cashDay.value)+'&currency='+encodeURIComponent(cashCur.value));
 const c=cashSummary.currency; cashOpening.innerText=money(cashSummary.opening_amount)+' '+c; cashShipmentIn.innerText=money(cashSummary.shipment_cash_in)+' '+c; cashIn.innerText=money(cashSummary.other_cash_in)+' '+c; cashShipmentOut.innerText=money(cashSummary.shipment_cash_out||0)+' '+c; cashAdvance.innerText=money(cashSummary.advance_out)+' '+c; cashExpense.innerText=money(cashSummary.expense_out)+' '+c; cashExpected.innerText=money(cashSummary.expected_amount)+' '+c; cashActual.innerText=money(cashSummary.actual_amount)+' '+c;
 const d=Number(cashSummary.difference||0); cashDiff.innerText=(d<0?'AÇIK ':'FAZLA ')+money(Math.abs(d))+' '+c; cashDiff.style.color=d<0?'#ef4444':(d>0?'#22c55e':'inherit');
 const canEditCashExpense=!!AUTH_USER && (AUTH_USER.role==='ADMIN' || (AUTH_USER.permissions||[]).includes('cash.expense.edit'));
 cashExpenseHistory.innerHTML=(cashSummary.expenses||[]).map(x=>`<tr><td>${x.expense_date||''}</td><td>${x.document_no||''}</td><td>${money(x.amount)} ${x.currency||''}</td><td>${x.note||''}</td><td>${x.created_by||''}</td><td>${canEditCashExpense?`<button class="btn" onclick='openCashExpenseEdit(${JSON.stringify(x)})'>Düzelt</button>`:''}</td></tr>`).join('');
}

function openCashExpenseEdit(x){
 const can=!!AUTH_USER && (AUTH_USER.role==='ADMIN' || (AUTH_USER.permissions||[]).includes('cash.expense.edit'));
 if(!can)return alert('Bu işlem için yetkiniz yok.');
 openM('Günlük Kasa Harcaması Düzelt',`<div class="grid">
   <label>Fatura/Fiş No*<input id="ceDoc" value="${String(x.document_no||'').replace(/&/g,'&amp;').replace(/"/g,'&quot;')}"></label>
   <label>Tutar*<input id="ceAmount" inputmode="decimal" value="${Number(x.amount||0)}"></label>
   <label>Para Birimi<select id="ceCur"><option>IQD</option><option>USD</option><option>TRY</option></select></label>
   <label>Tarih<input id="ceDate" type="date" value="${String(x.expense_date||'').slice(0,10)}"></label>
   <label style="grid-column:1/-1">Açıklama<input id="ceNote" value="${String(x.note||'').replace(/&/g,'&amp;').replace(/"/g,'&quot;')}"></label>
 </div><div style="margin-top:12px"><button class="btn green" onclick="saveCashExpenseEdit(${Number(x.id)})">Düzeltmeyi Kaydet</button></div>`,null);
 setTimeout(()=>{const e=document.getElementById('ceCur');if(e)e.value=(x.currency||'IQD')},0);
}
async function saveCashExpenseEdit(id){
 const body={document_no:(ceDoc.value||'').trim(),amount:parseMoney(ceAmount.value||0),currency:ceCur.value,expense_date:ceDate.value,note:(ceNote.value||'').trim()};
 if(!body.document_no)return alert('Fatura/Fiş No gerekli.');
 if(!(body.amount>0))return alert('Tutar 0’dan büyük olmalı.');
 if(!confirm('Bu harcama satırı düzeltilecek. Devam edilsin mi?'))return;
 try{await api('/api/cash-control/expense/'+id,{method:'PATCH',headers:{'Content-Type':'application/json'},body:JSON.stringify(body)});closeM();await loadCashControl();}
 catch(e){alert(e.message)}
}

async function openEntryCashMissing(){
 const day=(document.getElementById('cashDate')?.value||new Date().toISOString().slice(0,10));
 try{
   const rows=await api('/api/entry-cash-missing?business_date='+encodeURIComponent(day));
   const body=(rows||[]).length ? (rows||[]).map(x=>`<tr>
     <td><b>${x.scna||''}</b></td><td>${x.plate||''}</td><td>${x.driver_name||''}</td>
     <td>${fmtDateTime(x.effective_entry_date)}</td><td>${money(x.entry_collection||0)} IQD</td>
     <td><input id="ecm_${x.scna}" type="text" inputmode="decimal" placeholder="Teslim edilen para" style="min-width:150px"></td>
     <td><button class="btn green" onclick="saveMissingEntryCash('${x.scna}')">Kaydet</button></td>
   </tr>`).join('') : '<tr><td colspan="7">Bu tarihte muhasebeye düşmeyen tamamlanmış dönüş yok.</td></tr>';
   openM('Muhasebeye Düşmeyen Dönüşler — '+day,`<div class="small" style="margin-bottom:10px">Girişi tamamlanmış ama “Şoförün Teslim Ettiği Para” 0/boş olan kayıtlar.</div><div class="table"><table><thead><tr><th>SCNA</th><th>Plaka</th><th>Şoför</th><th>Giriş Tarihi</th><th>Müşteriden Alınan</th><th>Şoförün Teslim Ettiği</th><th>İşlem</th></tr></thead><tbody>${body}</tbody></table></div>`,null);
 }catch(e){alert(e.message)}
}

async function saveMissingEntryCash(scna){
 const el=document.getElementById('ecm_'+scna);
 const amount=parseMoney(el?.value||0);
 if(!(amount>0))return alert('Teslim edilen para 0’dan büyük olmalı.');
 if(!confirm(scna+' için '+money(amount)+' IQD teslim edilen para olarak kaydedilsin mi?'))return;
 try{
   await api('/api/trips/'+encodeURIComponent(scna)+'/entry-cash',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({amount})});
   alert('Kaydedildi. Günlük Kasa otomatik güncellenecek.');
   closeM();
   if(typeof loadCashControl==='function')await loadCashControl();
 }catch(e){alert(e.message)}
}

function openEntryCashDiagnostic(){
 const scna=(document.getElementById('cashDiagScna')?.value||'').trim();
 if(!scna)return alert('SCNA girin.');
 api('/api/entry-cash-diagnostic/'+encodeURIComponent(scna)).then(x=>{
   const reasons=(x.cash_in_block_reasons||[]).length?(x.cash_in_block_reasons||[]).join('<br>'):'Yok';
   openM('Giriş / Muhasebe Tanı — '+(x.scna||scna),`<div class="table"><table><tbody>
   <tr><th>SCNA</th><td><b>${x.scna||''}</b></td></tr><tr><th>Plaka</th><td>${x.plate||''}</td></tr><tr><th>Şoför</th><td>${x.driver_name||''}</td></tr>
   <tr><th>Durum</th><td>${x.status||''}</td></tr><tr><th>entry_done</th><td>${x.entry_done}</td></tr>
   <tr><th>Şoförün Teslim Ettiği Para</th><td><b>${money(x.entry_cash_handed||0)} IQD</b></td></tr><tr><th>Müşteriden Alınan</th><td>${money(x.entry_collection||0)} IQD</td></tr>
   <tr><th>entry_at</th><td>${fmtDateTime(x.entry_at)}</td></tr><tr><th>updated_at</th><td>${fmtDateTime(x.updated_at)}</td></tr><tr><th>Efektif Muhasebe Tarihi</th><td><b>${fmtDateTime(x.effective_entry_date)}</b></td></tr><tr><th>delivery_time</th><td>${fmtDateTime(x.delivery_time)}</td></tr><tr><th>trip_date</th><td>${x.trip_date||''}</td></tr><tr><th>created_at</th><td>${fmtDateTime(x.created_at)}</td></tr>
   <tr><th>Muhasebeye Uygun</th><td><b style="color:${x.cash_in_eligible?'#166534':'#b91c1c'}">${x.cash_in_eligible?'EVET':'HAYIR'}</b></td></tr>
   <tr><th>Sonuç</th><td><b>${x.accounting_result||''}</b></td></tr><tr><th>Engel Nedeni</th><td>${reasons}</td></tr>
   </tbody></table></div>`,null);
 }).catch(e=>alert(e.message));
}

function openCashDiagnostic(){
 const scna=(document.getElementById('cashDiagScna')?.value||'').trim();
 if(!scna)return alert('SCNA girin.');
 api('/api/cash-out-diagnostic/'+encodeURIComponent(scna)).then(x=>{
   const reasons=(x.cash_out_block_reasons||[]).length?(x.cash_out_block_reasons||[]).join(', '):'Yok';
   openM('SCNA Kasa Tanı — '+(x.scna||scna),`<div class="table"><table><tbody>
   <tr><th>SCNA</th><td>${x.scna||''}</td></tr><tr><th>Plaka</th><td>${x.plate||''}</td></tr><tr><th>Şoför</th><td>${x.driver_name||''}</td></tr>
   <tr><th>exit_cash</th><td><b>${money(x.exit_cash||0)} IQD</b></td></tr><tr><th>exit_done</th><td>${x.exit_done}</td></tr>
   <tr><th>exit_at</th><td>${fmtDateTime(x.exit_at)}</td></tr><tr><th>trip_date</th><td>${x.trip_date||''}</td></tr><tr><th>created_at</th><td>${fmtDateTime(x.created_at)}</td></tr>
   <tr><th>Efektif Çıkış Tarihi</th><td>${fmtDateTime(x.effective_exit_date)}</td></tr><tr><th>Silinmiş</th><td>${x.is_deleted||0}</td></tr>
   <tr><th>Harcırah</th><td>${money(x.exit_allowance||0)}</td></tr><tr><th>Prim</th><td>${money(x.exit_premium||0)}</td></tr><tr><th>Other</th><td>${money(x.exit_other||0)}</td></tr>
   <tr><th>Kasa Çıkışına Uygun</th><td><b>${x.cash_out_eligible?'EVET':'HAYIR'}</b></td></tr><tr><th>Engel Nedeni</th><td>${reasons}</td></tr>
   </tbody></table></div>`,null);
 }).catch(e=>alert(e.message));
}

function openShipmentCashOutDetails(){
 const x=cashSummary||{},rows=x.shipment_cash_out_entries||[],cur=x.currency||cashCur.value;
 const body=rows.length?rows.map(r=>`<tr><td><b>${r.scna||''}</b></td><td>${r.plate||''}</td><td>${r.driver_name||''}</td><td>${money(r.amount)} ${cur}</td><td>${fmtDateTime(r.exit_date)}</td></tr>`).join(''):`<tr><td colspan="5" class="muted">Bu tarihte sevkiyat cikisinda sofore verilen para yok.</td></tr>`;
 openM('Sevkiyat Çıkışı / Şoföre Verilen',`<div class="section-note">Bu liste sevkiyat çıkışındaki exit_cash alanını salt okunur gösterir. SCNA kaydını değiştirmez.</div><div class="table"><table><thead><tr><th>SCNA</th><th>Plaka</th><th>Şoför</th><th>Verilen</th><th>Çıkış Tarihi</th></tr></thead><tbody>${body}</tbody></table></div>`,null);
}
function openShipmentCashDetails(){
 const x=cashSummary||{},rows=x.shipment_cash_entries||[],cur=x.currency||cashCur.value;
 const body=rows.length?rows.map(r=>`<tr><td><b>${r.scna||''}</b></td><td>${r.plate||''}</td><td>${money(r.amount)} ${cur}</td><td>${fmtDateTime(r.entry_date)}</td></tr>`).join(''):`<tr><td colspan="4" class="empty-state">Bu tarih ve para biriminde kasaya teslim edilmiş sevkiyat tahsilatı yok.</td></tr>`;
 openM('Sevkiyattan Gelen — '+(x.business_date||cashDay.value),`<div class="section-note">Yalnızca girişi tamamlanmış sevkiyatlardaki “Şoförün Teslim Ettiği Para” okunur. Navlun, müşteriden alınan toplam ve müşteriden kalan tutar dahil değildir.</div><div class="table"><table><thead><tr><th>SCNA</th><th>PLAKA</th><th>TAHSİLAT / KASAYA TESLİM</th><th>GİRİŞ TARİHİ</th></tr></thead><tbody>${body}</tbody></table></div>`,()=>closeM());
}
function cashAddExpenseRow(){
 cashExpenseEntry.insertAdjacentHTML('beforeend',`<tr><td><input class="ceDate" type="date" title="Boşsa bugün"></td><td><input class="ceDoc" placeholder="Fiş/Fatura No *"></td><td><input class="ceAmt" type="number" min="0" step="1" placeholder="Tutar *"></td><td><input class="ceNote" placeholder="Açıklama (opsiyonel)"></td><td><button class="btn secondary" onclick="this.closest('tr').remove()">Sil</button></td></tr>`);
}
async function saveCashExpenses(){
 const rows=[...cashExpenseEntry.querySelectorAll('tr')]; if(!rows.length)return alert('Harcama satırı girin.');
 try{for(let i=0;i<rows.length;i++){const tr=rows[i],doc=tr.querySelector('.ceDoc').value.trim(),amt=Number(tr.querySelector('.ceAmt').value||0);if(!doc)throw new Error((i+1)+'. satırda fiş/fatura numarası zorunlu.');if(amt<=0)throw new Error((i+1)+'. satırda tutar zorunlu.');await api('/api/cash-control/expense',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({document_no:doc,amount:amt,currency:cashCur.value,note:tr.querySelector('.ceNote').value,expense_date:tr.querySelector('.ceDate').value})});} cashExpenseEntry.innerHTML=''; await loadCashControl(); alert(rows.length+' harcama kaydedildi.');}catch(e){alert(e.message)}
}
function openCashCount(){
 const x=cashSummary||{opening_amount:0,other_cash_in:0,actual_amount:0}; openM('Kasa Tutarı / Gün Sonu Kontrolü',`<div class="grid"><div class="field"><label>Tarih</label><input id="ccDay" type="date" value="${cashDay.value||cashToday()}"></div><div class="field"><label>Para Birimi</label><input value="${cashCur.value}" disabled></div><div class="field"><label>Gün Başı Kasa</label><input id="ccOpen" type="number" value="${Number(x.opening_amount||0)}"></div><div class="field"><label>Diğer Kasa Girişi</label><input id="ccIn" type="number" value="${Number(x.other_cash_in||0)}"></div><div class="field wide"><label>Fiilen Saydığım Kasa</label><input id="ccActual" type="number" value="${Number(x.actual_amount||0)}"></div><div class="field wide"><label>Not</label><input id="ccNote"></div></div>`,async()=>{try{await api('/api/cash-control/count',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({business_date:ccDay.value,currency:cashCur.value,opening_amount:Number(ccOpen.value||0),cash_in:Number(ccIn.value||0),actual_amount:Number(ccActual.value||0),note:ccNote.value})});cashDay.value=ccDay.value;closeM();await loadCashControl();}catch(e){alert(e.message)}});
}

let advanceRows=[];

window._advanceDriverSelected=false;
window._advanceDriverTimer=null;
function advanceRecipientTypeChanged(){
 window._advanceDriverSelected=false;
 window._advanceRecipientSelected=false;
 const hint=document.getElementById('aDriverHint'),lst=document.getElementById('aDriverList'),name=document.getElementById('aName');
 if(lst){lst.innerHTML='';lst.style.display='none';}
 if(name){name.value='';name.focus();}
 const typ=document.getElementById('aType')?.value||'';
 if(hint)hint.textContent=typ==='SOFOR'?'Kayıtlı şoförü seçin. 3 harften sonra arama başlar.':'Kayıtlı '+(typ==='USTA'?'usta':typ==='PERSONEL'?'personel':'kişi')+' için 3 harf yazın.';
}
// SAMA_ADVANCE_AUTOCOMPLETE_V3
function advanceDriverSearch(el){
 window._advanceDriverSelected=false;
 window._advanceRecipientSelected=false;
 const typ=document.getElementById('aType')?.value||'',lst=document.getElementById('aDriverList'),hint=document.getElementById('aDriverHint');
 if(!lst)return;
 const q=(el.value||'').trim(); clearTimeout(window._advanceDriverTimer);
 if(q.length<3){lst.style.display='none';lst.innerHTML='';if(hint)hint.textContent='En az 3 harf yazın.';return;}
 if(hint)hint.textContent='Aranıyor...';
 lst.innerHTML='<div class="autocomplete-item"><div class="autocomplete-sub">Aranıyor...</div></div>';lst.style.display='block';
 window._advanceDriverTimer=setTimeout(async()=>{
   let rows=[];
   // Önce Avans Takip içindeki mevcut cari/kişi kayıtlarını ara.
   try{
     const rec=await api('/api/advances/recipient-search?q='+encodeURIComponent(q)+'&recipient_type='+encodeURIComponent(typ));
     if(Array.isArray(rec)) rows=rec.map(x=>({kind:'RECIPIENT',name:x.name||'',open_count:Number(x.open_count||0),open_balance:Number(x.open_balance||0),record_count:Number(x.record_count||0)}));
   }catch(e){}

   // Şoförde ayrıca ana şoför listesini de tara; mevcut avans kaydı olmasa bile seçilebilsin.
   if(typ==='SOFOR'){
     let drv=[];
     try{drv=await api('/api/advances/driver-search?q='+encodeURIComponent(q));}catch(e){}
     if(!Array.isArray(drv)||!drv.length){
       try{
         const r=await api('/api/fleet/drivers');
         const all=Array.isArray(r)?r:(r?.rows||[]),Q=q.toLocaleUpperCase('tr-TR');
         drv=all.filter(x=>Number(x.is_active)!==0 && [x.name,x.d_no,x.phone,x.last_plate].join(' ').toLocaleUpperCase('tr-TR').includes(Q))
           .slice(0,30).map(x=>({name:x.name||'',d_no:x.d_no||'',phone:x.phone||'',plate:x.last_plate||''}));
       }catch(e){drv=[];}
     }
     const seen=new Set(rows.map(x=>String(x.name||'').trim().toLocaleUpperCase('tr-TR')));
     for(const x of (drv||[])){
       const k=String(x.name||'').trim().toLocaleUpperCase('tr-TR');
       if(k&&!seen.has(k)){seen.add(k);rows.push({kind:'DRIVER',name:x.name||'',d_no:x.d_no||'',phone:x.phone||'',plate:x.plate||''});}
     }
   }

   if(!rows.length){
     lst.innerHTML='<div class="autocomplete-item"><div class="autocomplete-sub">Kayıtlı kişi/cari bulunamadı</div></div>';lst.style.display='block';
     if(hint)hint.textContent=typ==='SOFOR'?'⚠ Kayıtlı şoför bulunamadı. Kaydederken yeni şoför oluşturma onayı istenecek.':'Kayıtlı cari bulunamadı; yeni isim olarak kaydedilebilir.';
     return;
   }
   window._advanceDriverMatches=rows;
   if(hint)hint.textContent=rows.length+' kayıt bulundu.';
   lst.innerHTML=rows.slice(0,30).map((x,i)=>{
     const sub=x.kind==='RECIPIENT'
       ? ((x.open_count>0?'AÇIK CARİ • '+x.open_count+' kayıt • Bakiye '+money(x.open_balance):'Geçmiş cari • '+x.record_count+' kayıt'))
       : ([x.d_no,x.phone,x.plate].filter(Boolean).join(' • ')||'Kayıtlı şoför');
     return `<div class="autocomplete-item" onmousedown="event.preventDefault();selectAdvanceDriver(${i})"><div class="autocomplete-main">${advEsc?advEsc(x.name||''):x.name||''}</div><div class="autocomplete-sub">${sub}</div></div>`;
   }).join('');
   lst.style.display='block';
 },180);
}
function selectAdvanceDriver(i){
 const x=(window._advanceDriverMatches||[])[i];if(!x)return;
 const typ=document.getElementById('aType')?.value||'';
 document.getElementById('aName').value=x.name||'';
 if(typ==='SOFOR'&&x.plate)document.getElementById('aPlate').value=x.plate;
 window._advanceRecipientSelected=true;
 window._advanceDriverSelected=(typ!=='SOFOR'||x.kind==='DRIVER'||x.kind==='RECIPIENT');
 const lst=document.getElementById('aDriverList');if(lst)lst.style.display='none';
 const hint=document.getElementById('aDriverHint');if(hint)hint.textContent=x.open_count>0?'✓ Açık cari seçildi.':'✓ Kayıtlı kişi seçildi.';
}


function advTypeLabel(t){return {USTA:'USTA',SOFOR:'ŞOFÖR',PERSONEL:'PERSONEL',DIGER:'DİĞER'}[t]||t||'';}
async function loadAdvances(){advanceRows=await api('/api/advances');renderAdvances();}
function renderAdvances(){
 const q=(document.getElementById('advQ')?.value||'').toLocaleUpperCase('tr-TR'),typ=document.getElementById('advType')?.value||'',st=document.getElementById('advStatus')?.value||'';
 const rows=advanceRows.filter(x=>(!typ||x.recipient_type===typ)&&(!st||x.account_status===st)&&(!q||[x.recipient_name,x.plate,x.scna,x.purpose,x.note].join(' ').toLocaleUpperCase('tr-TR').includes(q)));
 // SAMA_ADVANCE_PERSON_SUMMARY_V1
 const sums={};
 rows.forEach(x=>{
   const cur=(x.currency||'').toUpperCase();
   if(!sums[cur]) sums[cur]={total:0,settled:0,remaining:0};
   sums[cur].total+=Number(x.amount||0);
   sums[cur].settled+=Number(x.settled_amount||0);
   sums[cur].remaining+=Number(x.remaining||0);
 });
 const fmtS=(k)=>Object.entries(sums).map(([cur,v])=>`${money(v[k])} ${cur}`).join(' + ')||'0';
 const names=[...new Set(rows.map(x=>(x.recipient_name||'').trim()).filter(Boolean))];
 const personTitle=(q&&names.length===1)?names[0]:(q?`${rows.length} FİLTRELİ İŞLEM`:'TÜM AVANS HESAPLARI');
 const summary=document.getElementById('advPersonSummary');
 if(summary){
   summary.style.display=(q||typ||st)?'block':'none';
   summary.innerHTML=`<div style="font-weight:800;font-size:17px;margin-bottom:8px">${personTitle}</div><div class="cards"><div class="card">İşlem Sayısı<b>${rows.length}</b></div><div class="card">Toplam Alınan<b>${fmtS('total')}</b></div><div class="card">Fatura / Fiş / İade<b>${fmtS('settled')}</b></div><div class="card">Güncel Bakiye / Borç<b>${fmtS('remaining')}</b></div></div>`;
 }
 advRows.innerHTML=rows.map(x=>`<tr><td>${fmtDateTime(x.created_at)}</td><td>${advTypeLabel(x.recipient_type)}</td><td><b>${x.recipient_name||''}</b></td><td>${x.plate||''}</td><td>${x.scna||''}</td><td>${x.purpose||''}</td><td>${money(x.amount)}</td><td>${money(x.settled_amount)}</td><td><b>${money(x.remaining)}</b></td><td>${x.currency||''}</td><td>${x.note||''}</td><td><b>${x.account_status||x.status||''}</b><div class="small">Kayıt: ${x.record_status||x.status||''}</div></td><td><div class="compact-actions"><button class="btn green" onclick="openAdvanceSettlement(${x.id})">FATURA / FİŞ / İADE</button><button class="btn secondary" onclick="openAdvanceHistory(${x.id})">GEÇMİŞ</button><button class="btn secondary" onclick="editAdvance(${x.id})">DÜZENLE</button></div></td></tr>`).join('');
 const open=rows.filter(x=>x.account_status!=='KAPANDI'); advOpen.innerText=open.length; advTotal.innerText=money(rows.reduce((a,x)=>a+Number(x.amount||0),0)); advSettled.innerText=money(rows.reduce((a,x)=>a+Number(x.settled_amount||0),0)); advRemaining.innerText=money(open.reduce((a,x)=>a+Number(x.remaining||0),0));
 if(typeof v54InitTables==='function')v54InitTables();
}
// SAMA_ADVANCE_PERSON_ACCOUNTS_UI_V1
async function openAdvancePersonAccounts(){
 try{const rows=await api('/api/advances/person-accounts');const body=rows.length?rows.map((x,i)=>`<tr><td>${advEsc(x.recipient_type)}</td><td><b>${advEsc(x.recipient_name)}</b></td><td>${advEsc(x.currency)}</td><td>${money(x.total_advance)}</td><td>${money(x.total_settled)}</td><td><b>${money(x.open_balance)}</b></td><td>${x.open_count||0}</td><td><button class="btn secondary" onclick="openAdvancePersonDetail(${i})">HESAP</button></td></tr>`).join(''):'<tr><td colspan="8">Kişi hesabı bulunamadı.</td></tr>';window._advancePersonRows=rows;openM('KİŞİ / CARİ HESAPLARI',`<div class="section-note">Avans, fatura/fiş, iade ve mahsuplar kişi bazında tek hesapta gösterilir.</div><div class="table"><table><thead><tr><th>TİP</th><th>KİŞİ</th><th>PB</th><th>VERİLEN AVANS</th><th>KAPANAN</th><th>CARİ BAKİYE</th><th>AÇIK KAYIT</th><th>İŞLEM</th></tr></thead><tbody>${body}</tbody></table></div>`,()=>closeM());}catch(e){alert(e.message)}
}
async function openAdvancePersonDetail(i){const p=(window._advancePersonRows||[])[i];if(!p)return;try{const x=await api('/api/advances/person-account?recipient_type='+encodeURIComponent(p.recipient_type)+'&recipient_name='+encodeURIComponent(p.recipient_name)+'&currency='+encodeURIComponent(p.currency));window._advancePersonCurrent=x;const body=(x.movements||[]).map(m=>`<tr><td>${fmtDateTime(m.date)}</td><td>${advEsc(m.kind)}</td><td>${m.amount>=0?money(m.amount):''}</td><td>${m.amount<0?money(-m.amount):''}</td><td>${advEsc(m.document_no||'')}</td><td>${advEsc(m.note||'')}</td></tr>`).join('');openM(`${advEsc(x.recipient_name)} — Kişi Hesabı`,`<div class="calc"><b>Toplam Avans:</b> ${money(x.total_advance)} ${x.currency} &nbsp; <b>Kapanan:</b> ${money(x.total_settled)} &nbsp; <b>Cari Bakiye:</b> ${money(x.open_balance)} ${x.currency} &nbsp; <b>Durum:</b> ${x.account_status||''}<br><span class="small">Pozitif: kişi şirkete borçlu. Negatif: şirket kişiye borçlu.</span></div><div style="margin:10px 0;display:flex;gap:8px;flex-wrap:wrap"><button class="btn green" onclick="openAdvancePersonOffset()">YENİ AVANS / MAHSUP</button><button class="btn primary" onclick="openAdvancePersonDocument()">CARİYE FATURA / FİŞ</button><button class="btn secondary" onclick="reconcileAdvancePerson()">HESABI MUTABAKAT ET</button></div><div class="table"><table><thead><tr><th>TARİH</th><th>İŞLEM</th><th>AVANS</th><th>FATURA/İADE/MAHSUP</th><th>BELGE</th><th>AÇIKLAMA</th></tr></thead><tbody>${body}</tbody></table></div>`,()=>closeM());}catch(e){alert(e.message)}}

async function reconcileAdvancePerson(){
 const x=window._advancePersonCurrent;if(!x)return;
 try{const r=await api('/api/advances/person-account/reconcile?recipient_type='+encodeURIComponent(x.recipient_type)+'&recipient_name='+encodeURIComponent(x.recipient_name)+'&currency='+encodeURIComponent(x.currency),{method:'POST'});alert('Mutabakat tamamlandı. Cari bakiye: '+money(r.balance)+' '+x.currency+' / '+r.account_status);closeM();await openAdvancePersonAccounts();await loadAdvances();}catch(e){alert(e.message)}
}
function openAdvancePersonDocument(){
 const x=window._advancePersonCurrent;if(!x)return;
 openM(`${advEsc(x.recipient_name)} — Cariye Fatura / Fiş`,`<div class="section-note">Belge kişi hesabına girilir. Sistem eski açık avanslara FIFO dağıtır; fazlası kişi alacağı olarak eksi cari bakiyede kalır.</div><div class="grid"><div class="field"><label>Belge Türü</label><select id="apdType"><option value="FATURA">FATURA</option><option value="FIS">FİŞ</option></select></div><div class="field"><label>Belge No *</label><input id="apdDoc"></div><div class="field"><label>Tutar *</label><input id="apdAmount" type="number" min="0" step="1"></div><div class="field"><label>Para Birimi</label><input value="${x.currency}" disabled></div><div class="field wide"><label>Açıklama</label><input id="apdNote"></div></div>`,async()=>{try{const amt=Number(apdAmount.value||0),doc=apdDoc.value.trim();if(amt<=0)throw new Error('Tutar 0 dan büyük olmalı.');if(!doc)throw new Error('Belge numarası zorunlu.');const r=await api('/api/advances/person-account/settlement?recipient_type='+encodeURIComponent(x.recipient_type)+'&recipient_name='+encodeURIComponent(x.recipient_name),{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({settlement_type:apdType.value,amount:amt,document_no:doc,note:apdNote.value,currency:x.currency})});alert('Belge işlendi. Cari bakiye: '+money(r.balance)+' '+x.currency+' / '+r.account_status);closeM();await loadAdvances();await openAdvancePersonAccounts();}catch(e){alert(e.message)}});
}
function openAdvancePersonOffset(){const x=window._advancePersonCurrent;if(!x)return;openM(`${advEsc(x.recipient_name)} — Yeni Avans / Mahsup`,`<div class="calc"><b>Mevcut cari bakiye:</b> ${money(x.open_balance)} ${x.currency}<br><span class="small">Pozitif: kişinin elinde şirket parası. Negatif: şirket kişiye borçlu.</span><br><span class="small">Yeni ihtiyaçtan mevcut bakiye otomatik mahsup edilir. Sadece gerçekten verilen nakit yeni avans/kasa çıkışı olur.</span></div><div class="grid" style="margin-top:12px"><div class="field"><label>Yeni Avans İhtiyacı</label><input id="apoNeed" type="number" min="0" step="1" oninput="calcAdvanceOffset()"></div><div class="field"><label>Eski Bakiyeden Mahsup</label><input id="apoOffset" disabled></div><div class="field"><label>Kasadan Verilecek</label><input id="apoCash" disabled></div><div class="field wide"><label>Yeni İş / Avans Nedeni</label><input id="apoPurpose"></div><div class="field"><label>Plaka</label><input id="apoPlate"></div><div class="field"><label>SCNA</label><input id="apoScna"></div><div class="field wide"><label>Not</label><input id="apoNote"></div></div>`,async()=>{try{const need=Number(apoNeed.value||0),bal=Number(x.open_balance||0),off=Math.min(need,Math.max(0,bal)),debt=Math.max(0,-bal),cash=Math.max(0,need-off+debt);if(!confirm(`Yeni ihtiyaç: ${money(need)} ${x.currency}\nMahsup: ${money(off)}\nKasadan verilecek: ${money(cash)}\n\nİşlem kaydedilsin mi?`))return;await api('/api/advances/person-account/offset?recipient_type='+encodeURIComponent(x.recipient_type)+'&recipient_name='+encodeURIComponent(x.recipient_name),{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({new_need:need,purpose:apoPurpose.value,plate:apoPlate.value,scna:apoScna.value,note:apoNote.value,currency:x.currency})});closeM();await loadAdvances();alert('Mahsup ve yeni avans işlemi kaydedildi.');}catch(e){alert(e.message)}});calcAdvanceOffset();}
function calcAdvanceOffset(){const x=window._advancePersonCurrent||{},need=Number(document.getElementById('apoNeed')?.value||0),bal=Number(x.open_balance||0),off=Math.min(need,Math.max(0,bal)),debt=Math.max(0,-bal),cash=Math.max(0,need-off+debt);if(document.getElementById('apoOffset'))apoOffset.value=off;if(document.getElementById('apoCash'))apoCash.value=cash;}

function openAdvanceNew(){openM('Yeni Avans',`<div class="grid"><div class="field"><label>Alıcı Tipi</label><select id="aType" onchange="advanceRecipientTypeChanged()"><option value="SOFOR">Şoför</option><option value="USTA">Usta</option><option value="PERSONEL">Personel</option><option value="DIGER">Diğer</option></select></div><div class="field"><label>Ad Soyad / Kişi</label><div class="autocomplete-wrap"><input id="aName" autocomplete="off" oninput="advanceDriverSearch(this)" placeholder="3 harften sonra kayıtlı cari/kişi ara"><div id="aDriverList" class="autocomplete-list"></div></div><div id="aDriverHint" class="small"></div></div><div class="field"><label>Plaka (opsiyonel)</label><input id="aPlate"></div><div class="field"><label>SCNA (opsiyonel)</label><input id="aScna"></div><div class="field"><label>Tutar</label><input id="aAmount" type="number" min="0" step="1"></div><div class="field"><label>Para Birimi</label><select id="aCurrency"><option>IQD</option><option>USD</option><option>TRY</option></select></div><div class="field wide"><label>İş / Avans Nedeni</label><input id="aPurpose" placeholder="Örn: 22L29155 fren tamiri için parça ve işçilik"></div><div class="field wide"><label>Not</label><textarea id="aNote"></textarea></div></div>`,async()=>{try{if(aType.value==='SOFOR'&&!window._advanceDriverSelected){const nm=aName.value.trim();if(!nm)throw new Error('Şoför adı zorunlu.');const ok=confirm('⚠ Şoför kayıtlarında bu isim seçilmedi veya bulunamadı.\n\nYeni şoför kaydı oluşturulacak: '+nm+'\n\nDevam etmek istiyor musunuz?');if(!ok)return;const cr=await api('/api/advances/driver-create',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({name:nm,plate:aPlate.value})});aName.value=cr.name||nm;window._advanceDriverSelected=true;}await api('/api/advances',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({recipient_type:aType.value,recipient_name:aName.value,plate:aPlate.value,scna:aScna.value,purpose:aPurpose.value,amount:Number(aAmount.value||0),currency:aCurrency.value,note:aNote.value})});closeM();await loadAdvances();alert('Avans kaydedildi.');}catch(e){alert(e.message);}});}
function openAdvanceSettlement(id){const x=advanceRows.find(r=>Number(r.id)===Number(id));if(!x)return;openM(`${x.recipient_name} — Avans Kapatma`,`<div class="calc"><b>Verilen:</b> ${money(x.amount)} ${x.currency}<br><b>Belgelenen/İade:</b> ${money(x.settled_amount)} ${x.currency}<br><b>Kalan:</b> ${money(x.remaining)} ${x.currency}</div><div class="grid" style="margin-top:12px"><div class="field"><label>İşlem Türü</label><select id="asType"><option value="FATURA">Fatura</option><option value="FIS">Fiş</option><option value="NAKIT_IADE">Nakit İade</option><option value="MAHSUP">Mahsup</option></select></div><div class="field"><label>Tutar</label><input id="asAmount" type="number" min="0" step="1"></div><div class="field"><label>Belge No</label><input id="asDoc"></div><div class="field wide"><label>Açıklama</label><textarea id="asNote"></textarea></div></div>`,async()=>{try{await api('/api/advances/'+id+'/settlement',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({settlement_type:asType.value,amount:Number(asAmount.value||0),document_no:asDoc.value,note:asNote.value})});closeM();await loadAdvances();alert('Avans hareketi işlendi.');}catch(e){alert(e.message);}});}
function advEsc(v){return String(v??'').replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;').replace(/"/g,'&quot;').replace(/'/g,'&#39;');}
async function openAdvanceHistory(id){const rows=await api('/api/advances/'+id+'/settlements');const x=advanceRows.find(r=>Number(r.id)===Number(id));window._advanceHistoryRows=rows;openM(`${x?.recipient_name||''} — Avans Geçmişi`,`<div class="table"><table><thead><tr><th>TARİH</th><th>TÜR</th><th>TUTAR</th><th>BELGE NO</th><th>AÇIKLAMA</th><th>GİREN</th><th>İŞLEM</th></tr></thead><tbody>${rows.map(r=>`<tr><td>${fmtDateTime(r.created_at)}</td><td>${advEsc(r.settlement_type)}</td><td>${money(r.amount)} ${advEsc(x?.currency||'')}</td><td>${advEsc(r.document_no||'')}</td><td>${advEsc(r.note||'')}</td><td>${advEsc(r.created_by||'')}</td><td><button class="btn secondary" onclick="editAdvanceSettlement(${id},${r.id})">DÜZELT</button></td></tr>`).join('')}</tbody></table></div>`,()=>closeM());}
function editAdvanceSettlement(advanceId,settlementId){const r=(window._advanceHistoryRows||[]).find(z=>Number(z.id)===Number(settlementId));const x=advanceRows.find(z=>Number(z.id)===Number(advanceId));if(!r)return;const day=String(r.created_at||'').slice(0,10);openM(`${x?.recipient_name||''} — Hareketi Düzelt`,`<div class="calc"><b>Bu işlem mevcut hareketi düzeltir; yeni hareket oluşturmaz.</b><br><span class="small">Değişiklik işlem geçmişine kaydedilir.</span></div><div class="grid" style="margin-top:12px"><div class="field"><label>İşlem Türü</label><select id="aseType"><option value="FATURA">Fatura</option><option value="FIS">Fiş</option><option value="NAKIT_IADE">Nakit İade</option><option value="MAHSUP">Mahsup</option></select></div><div class="field"><label>Tutar</label><input id="aseAmount" type="number" min="0" step="1" value="${Number(r.amount||0)}"></div><div class="field"><label>Belge No</label><input id="aseDoc" value="${advEsc(r.document_no||'')}"></div><div class="field"><label>İşlem Tarihi</label><input id="aseDate" type="date" value="${advEsc(day)}"></div><div class="field wide"><label>Açıklama</label><textarea id="aseNote">${advEsc(r.note||'')}</textarea></div></div>`,async()=>{try{await api('/api/advances/'+advanceId+'/settlements/'+settlementId,{method:'PATCH',headers:{'Content-Type':'application/json'},body:JSON.stringify({settlement_type:aseType.value,amount:Number(aseAmount.value||0),document_no:aseDoc.value,note:aseNote.value,settlement_date:aseDate.value})});closeM();await loadAdvances();await openAdvanceHistory(advanceId);alert('Avans hareketi düzeltildi.');}catch(e){alert(e.message);}});setTimeout(()=>{aseType.value=r.settlement_type||'FATURA';},0);}
function editAdvance(id){const x=advanceRows.find(r=>Number(r.id)===Number(id));if(!x)return;openM(`${x.recipient_name} — Avansı Düzelt`,`<div class="grid"><div class="field"><label>Alıcı Tipi</label><select id="aeType"><option value="USTA">Usta</option><option value="SOFOR">Şoför</option><option value="PERSONEL">Personel</option><option value="DIGER">Diğer</option></select></div><div class="field"><label>Ad Soyad / Kişi</label><input id="aeName" value="${x.recipient_name||''}"></div><div class="field"><label>Plaka</label><input id="aePlate" value="${x.plate||''}"></div><div class="field"><label>SCNA</label><input id="aeScna" value="${x.scna||''}"></div><div class="field"><label>Tutar</label><input id="aeAmount" type="number" value="${Number(x.amount||0)}"></div><div class="field"><label>Para Birimi</label><select id="aeCurrency"><option>IQD</option><option>USD</option><option>TRY</option></select></div><div class="field wide"><label>İş / Avans Nedeni</label><input id="aePurpose" value="${String(x.purpose||'').replace(/\"/g,'&quot;')}"></div><div class="field wide"><label>Not</label><textarea id="aeNote">${x.note||''}</textarea></div></div>`,async()=>{try{await api('/api/advances/'+id+'/update',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({recipient_type:aeType.value,recipient_name:aeName.value,plate:aePlate.value,scna:aeScna.value,purpose:aePurpose.value,amount:Number(aeAmount.value||0),currency:aeCurrency.value,note:aeNote.value})});closeM();await loadAdvances();alert('Avans kaydı güncellendi.');}catch(e){alert(e.message);}});setTimeout(()=>{aeType.value=x.recipient_type||'USTA';aeCurrency.value=x.currency||'IQD';},0);}

async function loadDeletedTrips(){
  if(!AUTH_USER||AUTH_USER.role!=='ADMIN')return;
  const d=await api('/api/trips-deleted'); deletedRows.innerHTML=d.map(x=>`<tr><td><b>${x.scna}</b></td><td>${x.plate||''}</td><td>${fmtDateTime(x.deleted_at)}</td><td>${x.deleted_by||''}</td><td>${Math.max(0,Number(x.days_left||0))}</td><td>${Number(x.days_left||0)>=0?`<button class="btn green" onclick="restoreTrip('${x.scna}')">Geri Al</button>`:'Süre Doldu'}</td></tr>`).join('');
}
async function restoreTrip(scna){await api('/api/trips/'+encodeURIComponent(scna)+'/restore',{method:'POST'});await loadDeletedTrips();}
async function printTrip(scna,mode){
  const d=await api('/api/scna-detail/'+encodeURIComponent(scna)); const x=d.trip||{};
  const w=window.open('','_blank'); if(!w)return;
  const rows=[['SCNA',x.scna],['PLAKA',x.plate],['ŞOFÖR',x.driver_name||''],['MÜŞTERİ',x.customer_name||''],['BÖLGE',x.area_name||''],['KG',money(x.net_kg)],['NAVLUN',money(x.freight_total)],['ÇIKIŞ',fmtDateTime(x.exit_at)],['GİRİŞ',fmtDateTime(x.entry_display_at||x.entry_at)],['OTHER',money(x.exit_other)],['OTHER AÇIKLAMA',x.exit_other_note||'']];
  w.document.write(`<html><head><title>${mode} ${x.scna}</title><style>body{font-family:Arial;padding:28px}h2{margin-bottom:20px}table{border-collapse:collapse;width:100%;max-width:800px}td{border:1px solid #333;padding:10px}td:first-child{font-weight:bold;width:230px}@media print{button{display:none}}</style></head><body><h2>SAMA TRACK - ${mode} FORMU</h2><table>${rows.map(r=>`<tr><td>${r[0]}</td><td>${r[1]??''}</td></tr>`).join('')}</table><br><button onclick="window.print()">YAZDIR / PDF</button></body></html>`); w.document.close();
}

</script>
</body>
</html>
"""

@app.get("/", response_class=HTMLResponse)
def home():
    return HTML



if __name__ == "__main__":
    import uvicorn
    print("")
    print("=" * 60)
    print("SAMA TRACK BASLATILIYOR")
    print("Yerel adres : http://127.0.0.1:8000")
    print("Network     : http://BILGISAYAR_IP:8000")
    print("Otomatik yeniden yukleme: AKTIF")
    print("=" * 60)
    print("")
    uvicorn.run(
        "app:app",
        host="0.0.0.0",
        port=8000,
        reload=True,
        reload_includes=["*.py"]
    )
